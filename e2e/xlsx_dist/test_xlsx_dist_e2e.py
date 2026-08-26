"""Release-artifact E2E coverage for the XLSX public process boundary."""

from __future__ import annotations

import json
import os
from pathlib import Path
import warnings

import pytest

from support import (
    DistXlsx,
    XLSX_OPERATIONS,
    as_xltx,
    canonical_document,
    rich_workbook,
    sha256,
    write_request,
)


@pytest.fixture(scope="session")
def dist_xlsx() -> DistXlsx:
    root = os.environ.get("DOCUMENT_SKILLS_DIST_ROOT")
    assert root, "run through `npm run test:xlsx:dist-e2e` so dist is rebuilt first"
    marker = os.environ.get("DOCUMENT_SKILLS_DIST_BUILD_SHA256")
    assert marker and len(marker) == 64, "fresh dist build marker is missing"
    return DistXlsx(Path(root))


def test_dist_capabilities_are_the_exact_eleven_operation_contract(
    dist_xlsx: DistXlsx,
) -> None:
    report = dist_xlsx.invoke("capabilities", "--json")
    operations = {item["operation"]: item for item in report["operations"]}

    assert report["format"] == "xlsx"
    assert set(operations) == XLSX_OPERATIONS
    assert len(operations) == 11
    for name in XLSX_OPERATIONS - {"xlsx.render", "xlsx.validate.schema"}:
        assert operations[name]["available"] is True
        assert operations[name]["providers"] == ["core-python"]
    assert operations["xlsx.render"]["providers"] in ([], ["libreoffice"])
    assert operations["xlsx.validate.schema"]["providers"] in ([], ["dotnet-openxml"])

    truth_table = json.loads(
        (
            dist_xlsx.root
            / "skills"
            / "document-xlsx"
            / "references"
            / "feature-truth-table.json"
        ).read_text(encoding="utf-8")
    )
    assert set(truth_table["operations"]) == XLSX_OPERATIONS


def test_dist_feature_truth_table_has_auditable_execution_nodeids(
    dist_xlsx: DistXlsx,
    request: pytest.FixtureRequest,
) -> None:
    truth_table = json.loads(
        (
            dist_xlsx.root
            / "skills"
            / "document-xlsx"
            / "references"
            / "feature-truth-table.json"
        ).read_text(encoding="utf-8")
    )
    coverage_path = Path(__file__).with_name("feature-nodeids.json")
    coverage = json.loads(coverage_path.read_text(encoding="utf-8"))
    available = {
        f"{operation}/{feature}"
        for operation, record in truth_table["operations"].items()
        for feature in record.get("available", [])
    }
    provider_features = set(coverage["provider_feature_ids"])
    assert provider_features <= available
    assert set(coverage["core_operation_nodeids"]) == XLSX_OPERATIONS
    assert set(coverage["provider_operation_nodeids"]) == {
        "xlsx.convert",
        "xlsx.recalculate",
        "xlsx.render",
        "xlsx.validate.schema",
    }

    resolved: dict[str, dict[str, object]] = {}
    for feature_id in sorted(available):
        operation = feature_id.rsplit("/", 1)[0]
        if feature_id in provider_features:
            profile = "provider-dist"
            nodeids = coverage["provider_operation_nodeids"][operation]
        else:
            profile = "core-dist"
            nodeids = coverage["core_operation_nodeids"][operation]
        assert nodeids, feature_id
        resolved[feature_id] = {"profile": profile, "nodeids": nodeids}
    assert set(resolved) == available

    collected = {
        item.nodeid.replace("\\", "/")
        for item in request.session.items
        if item.nodeid.replace("\\", "/").startswith("e2e/xlsx_dist/")
    }
    assert set(coverage["core_dist_nodeids"]) == collected
    for record in resolved.values():
        allowed = set(coverage[f"{record['profile'].replace('-', '_')}_nodeids"])
        assert set(record["nodeids"]) <= allowed

    for nodeid in coverage["provider_dist_nodeids"]:
        relative, function_name = nodeid.split("::", 1)
        test_path = dist_xlsx.root.parent.parent / relative
        assert test_path.is_file(), nodeid
        assert f"def {function_name}(" in test_path.read_text(encoding="utf-8")


def test_dist_rich_workbook_survives_public_create_read_inspect_edit_and_validate(
    dist_xlsx: DistXlsx,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "rich-source.xlsx"
    create_request = write_request(
        tmp_path / "rich-create.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(source),
            "arguments": {"workbook": rich_workbook()},
        },
    )
    created = dist_xlsx.run(create_request)
    assert created["status"] in {"success", "degraded"}
    assert created["provider_chain"] == ["core-python"]
    source_hash = sha256(source)

    read_request = write_request(
        tmp_path / "rich-read.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(source),
            "arguments": {},
        },
    )
    read = dist_xlsx.run(read_request)
    projection = read["diagnostics"]["operation_result"]
    assert [item["name"] for item in projection["charts"]] == ["RevenueChart"]
    assert [item["location"] for item in projection["sparklines"]] == ["G2", "G3"]
    assert projection["hyperlinks"][0]["location"] == "Details!A1"
    assert projection["comments"][0]["text"] == "Check margin"
    assert projection["workbook_properties"]["title"] == "XLSX dist E2E"

    inspect_request = write_request(
        tmp_path / "rich-inspect.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.inspect.structure",
            "input": str(source),
            "arguments": {},
        },
    )
    inspection = dist_xlsx.run(inspect_request)["diagnostics"]["operation_result"]
    assert inspection["mutation_authorized"] is False
    assert inspection["worksheet_count"] == 2
    assert len(inspection["charts"]) == 1
    assert sha256(source) == source_hash

    edited = tmp_path / "rich-edited.xlsx"
    replacement_chart = rich_workbook()["charts"][0]
    replacement_chart.update(
        {
            "type": "line",
            "title": "Revenue trend",
            "anchor": "H3:O17",
        }
    )
    edits = [
        {"sheet": "Data", "type": "cell_value", "ref": "B2", "value": "15"},
        {
            "sheet": "Data",
            "type": "comment_update",
            "ref": "E2",
            "comment": {
                "ref": "E2",
                "text": "Margin changed by E2E",
                "author": "Dist E2E",
            },
        },
        {
            "sheet": "Data",
            "type": "sparkline_update",
            "ref": "G2",
            "sparkline": {
                "location": "G2",
                "data": "Data!B2:C2",
                "type": "column",
                "markers": True,
                "color": "#ED7D31",
            },
        },
        {
            "sheet": "Data",
            "type": "chart_update",
            "name": "RevenueChart",
            "chart": replacement_chart,
        },
    ]
    edit_request = write_request(
        tmp_path / "rich-edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(edited),
            "arguments": {"edits": edits, "expected_edits": len(edits)},
        },
    )
    edit_result = dist_xlsx.run(edit_request)
    assert edit_result["status"] == "degraded"
    assert sha256(source) == source_hash

    validation = dist_xlsx.invoke("validate", "--input", str(edited), "--json")
    assert validation["status"] == "pass"
    outcomes = {item["id"]: item["outcome"] for item in validation["gates"]}
    assert outcomes["provider.reopen"] == "pass"

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        reopened = load_workbook(edited, data_only=False)
    assert reopened.sheetnames == ["Data", "Details"]
    assert reopened["Data"]["B2"].value == 15
    assert reopened["Data"]["E2"].value == "=B2-C2"
    assert reopened["Data"]["E2"].comment.text == "Margin changed by E2E"
    assert reopened["Data"].tables["SalesTable"].ref == "A1:C5"
    assert type(reopened["Data"]._charts[0]).__name__ == "LineChart"
    assert reopened["Data"]["A2"].hyperlink.location == "Details!A1"
    reopened.close()


def test_dist_template_summary_pivot_and_transactional_followup_edit(
    dist_xlsx: DistXlsx,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    created_xlsx = tmp_path / "workflow-created.xlsx"
    created = dist_xlsx.run(
        write_request(
            tmp_path / "workflow-create.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.create",
                "output": str(created_xlsx),
                "arguments": {"workbook": rich_workbook()},
            },
        )
    )
    assert created["status"] in {"success", "degraded"}

    template = as_xltx(created_xlsx, tmp_path / "workflow-template.xltx")
    template_hash = sha256(template)
    instantiated = tmp_path / "workflow-instantiated.xlsx"
    instantiate = dist_xlsx.run(
        write_request(
            tmp_path / "workflow-instantiate.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.template.instantiate",
                "input": str(template),
                "output": str(instantiated),
                "arguments": {
                    "recalculation": "skip",
                    "edits": [
                        {
                            "sheet": "Details",
                            "type": "cell_value",
                            "ref": "A1",
                            "value": "Instantiated through dist",
                        }
                    ],
                },
            },
        )
    )
    assert instantiate["status"] in {"success", "degraded"}
    assert sha256(template) == template_hash

    summary_output = tmp_path / "workflow-summary.xlsx"
    summary = dist_xlsx.run(
        write_request(
            tmp_path / "workflow-summary.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.summary.aggregate",
                "input": str(instantiated),
                "output": str(summary_output),
                "arguments": {
                    "source": {"sheet": "Data", "range": "A1:D5"},
                    "group_by": ["Region"],
                    "aggregates": [
                        {
                            "column": "Revenue",
                            "function": "sum",
                            "as": "Revenue Total",
                        },
                        {"function": "count", "as": "Rows"},
                    ],
                    "sort": [
                        {"column": "Revenue Total", "direction": "desc"},
                        {"column": "Region", "direction": "asc"},
                    ],
                    "target": {
                        "sheet": "Summary",
                        "start_cell": "A1",
                        "table_name": "DistSummary",
                    },
                },
            },
        )
    )
    assert summary["status"] in {"success", "degraded"}
    summary_details = summary["diagnostics"]["operation_result"]["summary"]
    assert summary_details["kind"] == "ordinary_table"
    assert summary_details["native_pivot"] is False

    pivot_output = tmp_path / "workflow-pivot.xlsx"
    pivot = dist_xlsx.run(
        write_request(
            tmp_path / "workflow-pivot.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.pivot.create",
                "input": str(summary_output),
                "output": str(pivot_output),
                "arguments": {
                    "source": {"sheet": "Data", "range": "A1:D5"},
                    "rows": [{"column": "Region", "sort": "asc"}],
                    "columns": [],
                    "values": [
                        {
                            "column": "Revenue",
                            "function": "sum",
                            "as": "Pivot Revenue",
                        }
                    ],
                    "filters": [{"column": "Approved", "value": None}],
                    "target": {
                        "sheet": "Pivot",
                        "start_cell": "A1",
                        "name": "DistPivot",
                        "style": "PivotStyleMedium9",
                    },
                    "recalculation": "skip",
                },
            },
        )
    )
    assert pivot["status"] in {"success", "degraded"}
    assert pivot["diagnostics"]["operation_result"]["pivot"]["native"] is True
    pivot_hash = sha256(pivot_output)

    safely_edited = tmp_path / "workflow-safe-edit.xlsx"
    safe_edit = dist_xlsx.run(
        write_request(
            tmp_path / "workflow-safe-edit.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.edit",
                "input": str(pivot_output),
                "output": str(safely_edited),
                "arguments": {
                    "edits": [
                        {
                            "sheet": "Data",
                            "type": "cell_value",
                            "ref": "B2",
                            "value": "16",
                        }
                    ],
                    "expected_edits": 1,
                },
            },
        )
    )
    assert safe_edit["status"] == "degraded"
    assert sha256(pivot_output) == pivot_hash

    read = dist_xlsx.run(
        write_request(
            tmp_path / "workflow-read.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.read",
                "input": str(safely_edited),
                "arguments": {},
            },
        )
    )
    projected_pivots = read["diagnostics"]["operation_result"]["pivot_tables"]
    assert [item["name"] for item in projected_pivots] == ["DistPivot"]

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        reopened = load_workbook(safely_edited, data_only=False)
    assert reopened["Details"]["A1"].value == "Instantiated through dist"
    assert reopened["Summary"].tables["DistSummary"].ref == "A1:C3"
    assert reopened["Data"]["B2"].value == 16
    assert len(reopened["Pivot"]._pivots) == 1
    assert reopened["Pivot"]._pivots[0].name == "DistPivot"
    reopened.close()

    safe_hash = sha256(safely_edited)
    forbidden_output = tmp_path / "existing-structural-output.xlsx"
    sentinel = b"existing destination must survive rejected structural edit"
    forbidden_output.write_bytes(sentinel)
    forbidden = dist_xlsx.run(
        write_request(
            tmp_path / "workflow-forbidden-structural-edit.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.edit",
                "input": str(safely_edited),
                "output": str(forbidden_output),
                "arguments": {
                    "edits": [
                        {
                            "sheet": "Data",
                            "type": "row_insert",
                            "ref": "2",
                            "count": 1,
                        }
                    ],
                    "expected_edits": 1,
                },
            },
        ),
        check=False,
    )
    assert forbidden["status"] == "enhancement_required"
    assert forbidden["errors"][0]["code"] == "DS_ENHANCEMENT_REQUIRED"
    assert forbidden_output.read_bytes() == sentinel
    assert sha256(safely_edited) == safe_hash


def test_dist_canonical_json_xlsx_read_json_and_csv_typed_roundtrip(
    dist_xlsx: DistXlsx,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    canonical = tmp_path / "typed-source.json"
    canonical.write_text(
        json.dumps(canonical_document(), ensure_ascii=False),
        encoding="utf-8",
        newline="\n",
    )
    converted_xlsx = tmp_path / "typed-output.xlsx"
    to_xlsx = dist_xlsx.run(
        write_request(
            tmp_path / "typed-to-xlsx.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.convert",
                "input": str(canonical),
                "output": str(converted_xlsx),
                "arguments": {
                    "source_format": "json",
                    "target_format": "xlsx",
                    "values": {"formula_policy": "evaluated"},
                },
            },
        )
    )
    assert to_xlsx["status"] == "degraded"
    assert _loss_codes(to_xlsx) == [
        "formulas-replaced-with-cached-values",
        "null-values-mapped-to-blank-cells",
        "timezone-values-preserved-as-text",
    ]

    read = dist_xlsx.run(
        write_request(
            tmp_path / "typed-read.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.read",
                "input": str(converted_xlsx),
                "arguments": {},
            },
        )
    )
    projection = read["diagnostics"]["operation_result"]
    cells = {
        cell["ref"]: cell
        for row in projection["sheets"][0]["rows"]
        for cell in row["cells"]
    }
    assert cells["A2"]["value"] == "00123"
    assert cells["B2"]["value"] == "12.50"
    assert cells["C2"]["type"] == "b"
    assert cells["D2"]["number_format"] == "yyyy-mm-dd"
    assert cells["E2"]["number_format"] == "hh:mm:ss"
    assert cells["F2"]["number_format"] == 'yyyy-mm-dd"T"hh:mm:ss'
    assert cells["G2"]["value"] == "2026-08-24T12:30:15+08:00"
    assert "H2" not in cells
    assert cells["I2"]["value"] == ""
    assert cells["J2"]["value"] == "=1+1"
    assert cells["K2"]["value"] == "25"
    assert projection["formula_state"]["cells"] == {}

    reopened = load_workbook(converted_xlsx, data_only=False)
    assert reopened["Typed"]["A2"].value == "00123"
    assert reopened["Typed"]["C2"].value is True
    assert reopened["Typed"]["G2"].value == "2026-08-24T12:30:15+08:00"
    assert reopened["Typed"]["K2"].value == 25
    reopened.close()

    roundtrip_json = tmp_path / "typed-roundtrip.json"
    to_json = dist_xlsx.run(
        write_request(
            tmp_path / "typed-to-json.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.convert",
                "input": str(converted_xlsx),
                "output": str(roundtrip_json),
                "arguments": {"source_format": "xlsx", "target_format": "json"},
            },
        )
    )
    assert to_json["status"] == "degraded"
    assert _loss_codes(to_json) == ["xlsx-formatting-and-objects-dropped"]
    roundtrip = json.loads(roundtrip_json.read_text(encoding="utf-8"))
    values = roundtrip["sheets"][0]["rows"][1]
    assert values[0] == {"type": "string", "value": "00123"}
    assert values[1] == {"type": "number", "value": "12.50"}
    assert values[2] == {"type": "boolean", "value": True}
    assert values[3] == {"type": "date", "value": "2026-08-24"}
    assert values[6] == {
        "type": "string",
        "value": "2026-08-24T12:30:15+08:00",
    }
    assert values[7] == {"type": "null", "value": None}
    assert values[8] == {"type": "empty", "value": ""}
    assert values[10] == {"type": "number", "value": "25"}

    roundtrip_csv = tmp_path / "typed-roundtrip.csv"
    to_csv = dist_xlsx.run(
        write_request(
            tmp_path / "typed-to-csv.json",
            {
                "schema_version": "1.0",
                "operation": "xlsx.convert",
                "input": str(converted_xlsx),
                "output": str(roundtrip_csv),
                "arguments": {"source_format": "xlsx", "target_format": "csv"},
            },
        )
    )
    assert to_csv["status"] == "degraded"
    assert _loss_codes(to_csv) == [
        "xlsx-formatting-and-objects-dropped",
        "typed-values-rendered-as-text",
        "csv-injection-escaped",
    ]
    csv_lines = roundtrip_csv.read_text(encoding="utf-8").splitlines()
    assert csv_lines[1].startswith("00123,12.50,TRUE,2026-08-24,12:30:15")
    assert "'=1+1" in csv_lines[1]


def _loss_codes(result: dict[str, object]) -> list[str]:
    operation_result = result["diagnostics"]["operation_result"]
    return [item["code"] for item in operation_result["semantic_losses"]]
