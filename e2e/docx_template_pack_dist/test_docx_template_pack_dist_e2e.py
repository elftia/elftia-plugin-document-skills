"""Distribution-only DOCX template-pack catalog, import, and reuse coverage."""

from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
import json
import os
from pathlib import Path
import shutil
import subprocess
from xml.etree import ElementTree

from docx import Document
import pytest

import document_skills_core
import document_skills_core.formats.docx.template_pack_operation as pack_operation
from document_skills_core.core.io.paths import sha256_file
from document_skills_core.formats.docx.constants import qn
from document_skills_core.formats.docx.create import create_docx
from document_skills_core.formats.docx.package import OpcPackage, write_deterministic_zip
from document_skills_core.formats.docx.service import DocxService


@pytest.fixture(scope="session")
def dist_root() -> Path:
    value = os.environ.get("DOCUMENT_SKILLS_DIST_ROOT")
    assert value, "run through `npm run test:docx:dist-e2e` so dist is rebuilt first"
    root = Path(value).resolve()
    marker = os.environ.get("DOCUMENT_SKILLS_DIST_BUILD_SHA256")
    assert marker and len(marker) == 64, "fresh dist build marker is missing"
    runtime_module = Path(document_skills_core.__file__).resolve()
    assert runtime_module.is_relative_to(root), (
        f"DOCX runtime escaped the built artifact: {runtime_module}"
    )
    return root


def test_dist_catalog_import_read_and_instantiate_without_producer_source(
    dist_root: Path,
    tmp_path: Path,
) -> None:
    service = DocxService(dist_root)
    listed = service.execute("docx.template.pack.list", _request("docx.template.pack.list"))
    assert listed["status"] == "success", listed
    assert [
        (item["id"], item["version"])
        for item in listed["diagnostics"]["operation_result"]["packs"]
    ] == [("general-academic-paper", "1.0.0")]

    catalog_path = dist_root / "skills/document-docx/assets/template-packs/catalog.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    builtin = {
        "kind": "builtin",
        "id": "general-academic-paper",
        "version": "1.0.0",
    }
    read = service.execute(
        "docx.template.pack.read",
        _request("docx.template.pack.read", arguments={"pack": builtin}),
    )
    assert read["status"] == "success", read
    builtin_pack = read["diagnostics"]["operation_result"]["pack"]
    assert builtin_pack["manifest_sha256"] == catalog["packs"][0][
        "expected_manifest_sha256"
    ]
    for locale in ("en-US", "zh-CN"):
        academic_output = tmp_path / f"builtin-academic-{locale}.docx"
        generated = service.execute(
            "docx.template.pack.instantiate",
            _request(
                "docx.template.pack.instantiate",
                output=str(academic_output),
                arguments={
                    "pack": builtin,
                    "mode": "document-spec",
                    "document_spec": _academic_spec(locale, include_table=True),
                },
            ),
        )
        assert generated["status"] == "success", json.dumps(
            generated, ensure_ascii=False, indent=2
        )
        _strict_traverse(academic_output)
        _assert_academic_geometry_and_fields(academic_output)
        _libreoffice_open_and_render(academic_output, tmp_path / f"lo-{locale}")

    source = (
        dist_root
        / "skills/document-docx/assets/template-packs/general-academic-paper/1.0.0/template.docx"
    )
    source_before = source.read_bytes()
    destination = tmp_path / "user-pack"
    imported = service.execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(source),
            output=str(destination),
            arguments={
                "expected_source_sha256": sha256_file(source),
                "manifest": {
                    "id": "dist-private-academic",
                    "version": "1.0.0",
                    "display_name": "Dist Private Academic",
                    "description": "Imported only by the distribution E2E.",
                    "modes": ["document-spec"],
                    "role_styles": builtin_pack["capabilities"]["role_styles"],
                    "authoring_format": builtin_pack["capabilities"]["authoring_format"],
                    "compatibility": {
                        "document_spec_versions": ["1.0"],
                        "domain_profiles": [{"id": "academic-paper", "version": "1.0"}],
                        "template_engine_versions": [],
                        "style_profile_versions": ["template-mapped/1.0"],
                    },
                },
                "provenance": {
                    "source_kind": "local",
                    "imported_at": "2026-08-29T00:00:00Z",
                    "redistributable": False,
                },
            },
        ),
    )
    assert imported["status"] == "success", json.dumps(imported, ensure_ascii=False, indent=2)
    assert source.read_bytes() == source_before
    pack = imported["diagnostics"]["operation_result"]["pack"]
    local = {
        "kind": "local",
        "path": str(destination),
        "expected_manifest_sha256": pack["manifest_sha256"],
    }
    local_read = service.execute(
        "docx.template.pack.read",
        _request("docx.template.pack.read", arguments={"pack": local}),
    )
    assert local_read["status"] == "success", local_read

    output = tmp_path / "instantiated.docx"
    instantiated = service.execute(
        "docx.template.pack.instantiate",
        _request(
            "docx.template.pack.instantiate",
            output=str(output),
            arguments={
                "pack": local,
                "mode": "document-spec",
                "document_spec": _academic_spec(),
            },
        ),
    )
    assert instantiated["status"] == "success", json.dumps(
        instantiated, ensure_ascii=False, indent=2
    )
    assert output.is_file()
    assert sha256(output.read_bytes()).hexdigest() == instantiated["artifacts"][0]["sha256"]

    provenance_text = (dist_root / "provenance/modules.json").read_text(encoding="utf-8")
    allowlist_text = (dist_root / "provenance/runtime-source-allowlist.json").read_text(
        encoding="utf-8"
    )
    assert str(destination) not in provenance_text
    assert str(destination) not in allowlist_text


def test_dist_footer_inventory_overflow_is_fail_closed_per_mode(
    dist_root: Path,
    tmp_path: Path,
) -> None:
    service = DocxService(dist_root)
    source = tmp_path / "dist-footer-overflow.docx"
    builtin_pack = (
        dist_root
        / "skills/document-docx/assets/template-packs/general-academic-paper/1.0.0"
    )
    shutil.copy2(builtin_pack / "template.docx", source)
    _replace_footer_paragraphs(source, count=17)
    builtin_manifest = json.loads(
        (builtin_pack / "manifest.json").read_text(encoding="utf-8")
    )
    bounded_profile = deepcopy(builtin_manifest["capabilities"]["authoring_format"])
    bounded_profile["default_footer"] = {
        "paragraphs": [
            {"alignment": "center", "runs": [{"text": "Synthetic footer"}]}
            for _paragraph_index in range(16)
        ]
    }

    inspected = service.execute(
        "docx.template.import.inspect",
        _request(
            "docx.template.import.inspect",
            input=str(source),
            arguments={"expected_source_sha256": sha256_file(source)},
        ),
    )
    assert inspected["status"] == "success", inspected
    operation = inspected["diagnostics"]["operation_result"]
    document_spec = operation["mode_compatibility"]["document-spec"]
    assert document_spec["status"] == "incompatible"
    assert "footer-paragraph-inventory-overflow" in document_spec["reasons"]

    template_source = tmp_path / "dist-footer-template-source.docx"
    create_docx(
        template_source,
        {**_raw_report("Template base copy"), "footer": "Synthetic footer"},
    )
    _replace_footer_paragraphs(template_source, count=17)
    template_inspection = service.execute(
        "docx.template.import.inspect",
        _request(
            "docx.template.import.inspect",
            input=str(template_source),
            arguments={"expected_source_sha256": sha256_file(template_source)},
        ),
    )
    assert template_inspection["status"] == "success", template_inspection
    assert template_inspection["diagnostics"]["operation_result"][
        "mode_compatibility"
    ]["template"]["status"] == "compatible"
    template_pack = tmp_path / "dist-footer-template-pack"
    imported_template = service.execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(template_source),
            output=str(template_pack),
            arguments={
                "expected_source_sha256": sha256_file(template_source),
                "manifest": {
                    "id": "dist-footer-template",
                    "version": "1.0.0",
                    "display_name": "Dist Footer Template",
                    "description": "Distribution base-copy fidelity regression.",
                    "modes": ["template"],
                    "variables": [],
                    "regions": [],
                    "role_styles": {},
                    "compatibility": {
                        "document_spec_versions": [],
                        "domain_profiles": [],
                        "template_engine_versions": ["1.0"],
                        "style_profile_versions": [],
                    },
                },
                "provenance": {
                    "source_kind": "local",
                    "imported_at": "2026-08-30T00:00:00Z",
                    "redistributable": False,
                },
            },
        ),
    )
    assert imported_template["status"] == "success", imported_template
    template_output = tmp_path / "dist-footer-template.docx"
    instantiated = service.execute(
        "docx.template.pack.instantiate",
        _request(
            "docx.template.pack.instantiate",
            output=str(template_output),
            arguments={
                "pack": {
                    "kind": "local",
                    "path": str(template_pack),
                    "expected_manifest_sha256": imported_template["diagnostics"][
                        "operation_result"
                    ]["manifest_sha256"],
                },
                "mode": "template",
                "variables": {},
            },
        ),
    )
    assert instantiated["status"] == "success", json.dumps(
        instantiated,
        ensure_ascii=False,
        indent=2,
    )
    assert len(OpcPackage.open(template_output).xml("word/footer1.xml").findall(qn("w", "p"))) == 17

    document_spec_pack = tmp_path / "dist-footer-document-spec-pack"
    imported_document_spec = service.execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(source),
            output=str(document_spec_pack),
            arguments={
                "expected_source_sha256": sha256_file(source),
                "manifest": {
                    "id": "dist-footer-document-spec",
                    "version": "1.0.0",
                    "display_name": "Dist Footer Document Spec",
                    "description": "Distribution reconstruction fidelity regression.",
                    "modes": ["document-spec"],
                    "role_styles": builtin_manifest["capabilities"]["role_styles"],
                    "authoring_format": bounded_profile,
                    "compatibility": {
                        "document_spec_versions": ["1.0"],
                        "domain_profiles": [{"id": "academic-paper", "version": "1.0"}],
                        "template_engine_versions": [],
                        "style_profile_versions": ["template-mapped/1.0"],
                    },
                },
                "provenance": {
                    "source_kind": "local",
                    "imported_at": "2026-08-30T00:00:00Z",
                    "redistributable": False,
                },
            },
        ),
    )
    assert imported_document_spec["status"] == "enhancement_required"
    details = imported_document_spec["errors"][0]["details"]
    assert details["diagnostic"] == "DS_DOCX_TEMPLATE_FORMAT_INVENTORY_OVERFLOW"
    assert "footer-paragraph-inventory-overflow" in details["reasons"]
    assert not document_spec_pack.exists()


@pytest.mark.parametrize("mode", ["template", "document-spec"])
def test_dist_instantiation_uses_verified_snapshot_after_original_payload_mutates(
    dist_root: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    mode: str,
) -> None:
    for dependency in ("docxtemplater", "pizzip"):
        assert (dist_root / "node_modules" / dependency / "package.json").is_file(), (
            f"declared production dependency is unavailable in dist: {dependency}"
        )
    service = DocxService(dist_root)
    if mode == "template":
        source = tmp_path / "template-source.docx"
        create_docx(source, _raw_report("DIST ORIGINAL {name}"))
        pack_root = tmp_path / "dist-template-pack"
        imported = service.execute(
            "docx.template.import.create",
            _request(
                "docx.template.import.create",
                input=str(source),
                output=str(pack_root),
                arguments={
                    "expected_source_sha256": sha256_file(source),
                    "manifest": {
                        "id": "dist-snapshot-template",
                        "version": "1.0.0",
                        "display_name": "Dist Snapshot Template",
                        "description": "Distribution template snapshot regression.",
                        "modes": ["template"],
                        "variables": [{"name": "name", "required": True}],
                        "regions": [],
                        "role_styles": {},
                        "compatibility": {
                            "document_spec_versions": [],
                            "domain_profiles": [],
                            "template_engine_versions": ["1.0"],
                            "style_profile_versions": [],
                        },
                    },
                    "provenance": {
                        "source_kind": "local",
                        "imported_at": "2026-08-30T00:00:00Z",
                        "redistributable": False,
                    },
                },
            ),
        )
        assert imported["status"] == "success", imported
        manifest_sha256 = imported["diagnostics"]["operation_result"][
            "manifest_sha256"
        ]
        instantiate_arguments = {
            "mode": "template",
            "variables": {"name": "Agent"},
        }
        substitute_text = "DIST MUTATED {name}"
    else:
        builtin = (
            dist_root
            / "skills/document-docx/assets/template-packs/general-academic-paper/1.0.0"
        )
        pack_root = tmp_path / "dist-document-spec-pack"
        shutil.copytree(builtin, pack_root)
        manifest_sha256 = sha256_file(pack_root / "manifest.json")
        instantiate_arguments = {
            "mode": "document-spec",
            "document_spec": _academic_spec(),
        }
        substitute_text = "DIST MUTATED DOCUMENT SPEC SOURCE"
    payload = pack_root / "template.docx"
    original_payload_sha256 = sha256_file(payload)
    substitute = tmp_path / f"{mode}-substitute.docx"
    create_docx(substitute, _raw_report(substitute_text))
    assert sha256_file(substitute) != original_payload_sha256
    original_snapshot = pack_operation.snapshot_pack_reference

    def snapshot_then_mutate(*args: object, **kwargs: object):
        snapshot = original_snapshot(*args, **kwargs)
        shutil.copyfile(substitute, payload)
        return snapshot

    monkeypatch.setattr(
        pack_operation,
        "snapshot_pack_reference",
        snapshot_then_mutate,
    )
    output = tmp_path / f"dist-{mode}-snapshot.docx"
    result = service.execute(
        "docx.template.pack.instantiate",
        _request(
            "docx.template.pack.instantiate",
            output=str(output),
            arguments={
                "pack": {
                    "kind": "local",
                    "path": str(pack_root),
                    "expected_manifest_sha256": manifest_sha256,
                },
                **instantiate_arguments,
            },
        ),
    )
    assert result["status"] == "success", json.dumps(result, ensure_ascii=False)
    if mode == "template":
        document_xml = OpcPackage.open(output).parts["word/document.xml"]
        assert b"DIST ORIGINAL Agent" in document_xml
        assert b"DIST MUTATED Agent" not in document_xml
    else:
        styles = OpcPackage.open(output).xml("word/styles.xml")
        body = next(
            item
            for item in styles.findall(qn("w", "style"))
            if item.attrib.get(qn("w", "styleId")) == "ElftiaBody"
        )
        fonts = body.find(f"./{qn('w', 'rPr')}/{qn('w', 'rFonts')}")
        assert fonts is not None
        assert fonts.attrib[qn("w", "ascii")] == "Times New Roman"
        assert fonts.attrib[qn("w", "eastAsia")] == "SimSun"
    pack_result = result["diagnostics"]["operation_result"]["pack"]
    assert pack_result["manifest_sha256"] == manifest_sha256
    assert pack_result["payload_sha256"] == original_payload_sha256
    assert pack_result["authoring_source"] == "operation-owned-verified-snapshot"
    assert sha256_file(payload) != original_payload_sha256


def _request(operation: str, **values: object) -> dict[str, object]:
    request: dict[str, object] = {
        "schema_version": "1.0",
        "operation": operation,
        "arguments": {},
    }
    request.update(values)
    return request


def _raw_report(text: str) -> dict[str, object]:
    return {
        "blocks": [{"type": "paragraph", "text": text, "style": "Normal"}],
        "image": None,
        "header": None,
        "footer": None,
        "sections": [{"orientation": "portrait", "title": None}],
        "metadata": {
            "title": "Dist snapshot fixture",
            "subject": "",
            "creator": "",
            "keywords": "",
        },
        "style_profile": None,
    }


def _academic_spec(
    locale: str = "en-US",
    *,
    include_table: bool = False,
) -> dict[str, object]:
    chinese = locale == "zh-CN"
    nodes: list[dict[str, object]] = [
        {"id": "title", "type": "title", "text": "分发学术模板" if chinese else "Distribution Template Pack"},
        {"id": "authors", "type": "authors", "items": ["Ada Lovelace"]},
        {"id": "abstract", "type": "abstract", "text": "仅使用分发产物。" if chinese else "Built from dist only."},
        {"id": "keywords", "type": "keywords", "items": ["DOCX", "dist"]},
        {"id": "heading", "type": "heading", "level": 1, "text": "引言" if chinese else "Introduction"},
        {"id": "body", "type": "paragraph", "text": "正文保持原样。" if chinese else "No producer import path."},
    ]
    if include_table:
        nodes.append(
            {
                "id": "table",
                "type": "table",
                "rows": [["指标" if chinese else "Metric", "Value"], ["A", "1"]],
            }
        )
    return {
        "version": "1.0",
        "domain_profile": {
            "id": "academic-paper",
            "version": "1.0",
            "locale": locale,
        },
        "nodes": nodes,
    }


def _strict_traverse(path: Path) -> None:
    document = Document(path)
    _ = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                _ = [paragraph.text for paragraph in cell.paragraphs]
    for section in document.sections:
        _ = [paragraph.text for paragraph in section.header.paragraphs]
        _ = [paragraph.text for paragraph in section.footer.paragraphs]


def _replace_footer_paragraphs(path: Path, *, count: int) -> None:
    package = OpcPackage.open(path)
    footer = package.xml("word/footer1.xml")
    for child in list(footer):
        footer.remove(child)
    for _paragraph_index in range(count):
        paragraph = ElementTree.SubElement(footer, qn("w", "p"))
        properties = ElementTree.SubElement(paragraph, qn("w", "pPr"))
        ElementTree.SubElement(
            properties,
            qn("w", "jc"),
            {qn("w", "val"): "center"},
        )
        run = ElementTree.SubElement(paragraph, qn("w", "r"))
        ElementTree.SubElement(run, qn("w", "t")).text = "Synthetic footer"
    parts = dict(package.parts)
    from document_skills_core.formats.docx.xml_utils import xml_bytes

    parts["word/footer1.xml"] = xml_bytes(footer)
    write_deterministic_zip(path, parts)


def _assert_academic_geometry_and_fields(path: Path) -> None:
    import zipfile

    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as archive:
        document = ElementTree.fromstring(archive.read("word/document.xml"))
        footer = ElementTree.fromstring(archive.read("word/footer1.xml"))
    table = next(document.iter(namespace + "tbl"))
    grid = table.find(namespace + "tblGrid")
    assert grid is not None
    assert len(grid.findall(namespace + "gridCol")) == 2
    assert all(int(item.attrib[namespace + "w"]) > 0 for item in grid)
    assert [" ".join((item.text or "").split()) for item in footer.iter(namespace + "instrText")] == [
        "PAGE",
        "NUMPAGES",
    ]


def _libreoffice_open_and_render(path: Path, output: Path) -> None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None and os.name == "nt":
        known = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
        executable = str(known) if known.is_file() else None
    assert executable is not None, "LibreOffice is mandatory for the DOCX pack release gate"
    output.mkdir()
    profile = output / "profile"
    completed = subprocess.run(
        [
            executable,
            "--headless",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output),
            str(path),
        ],
        check=False,
        capture_output=True,
        timeout=120,
    )
    assert completed.returncode == 0, completed.stderr.decode(errors="replace")[-2000:]
    rendered = output / f"{path.stem}.pdf"
    assert rendered.is_file() and rendered.stat().st_size > 0
