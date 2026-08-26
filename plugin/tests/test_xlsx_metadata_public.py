"""Public run.py coverage for worksheet metadata, annotations, and properties."""

from hashlib import sha256
from pathlib import Path

from test_xlsx_annotations import _workbook
from test_xlsx_public import _public, _request


def _page_setup(*, fit: bool) -> dict:
    result = {
        "orientation": "landscape",
        "paper_size": "a4",
        "margins": {
            "left": 0.5,
            "right": 0.5,
            "top": 0.6,
            "bottom": 0.6,
            "header": 0.2,
            "footer": 0.2,
        },
        "horizontal_centered": True,
        "vertical_centered": False,
    }
    if fit:
        result.update({"fit_to_width": 1, "fit_to_height": 0})
    else:
        result["scale"] = 90
    return result


def _header_footer(label: str) -> dict:
    return {
        "odd_header": f"&C{label}",
        "odd_footer": "&RPage &P of &N",
        "even_header": "",
        "even_footer": "",
        "first_header": "",
        "first_footer": "",
        "different_first": False,
        "different_odd_even": False,
        "scale_with_doc": True,
        "align_with_margins": True,
    }


def test_public_metadata_and_annotation_create_edit_delete_reopen(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "public-metadata-source.xlsx"
    edited = tmp_path / "public-metadata-edited.xlsx"
    deleted = tmp_path / "public-metadata-deleted.xlsx"
    workbook = _workbook()
    workbook["sheets"][0].update(
        {
            "view": {
                "show_grid_lines": False,
                "zoom_scale": 125,
                "selected_cell": "B2",
            },
            "page_setup": _page_setup(fit=True),
            "header_footer": _header_footer("Created"),
            "print_area": "A1:C20",
            "print_titles": {"rows": "1:2", "columns": "A:B"},
        }
    )
    create_request = _request(
        tmp_path,
        "public-metadata-create.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(source),
            "arguments": {"workbook": workbook},
        },
    )
    create_result = _public(project_root, "run", "--request", str(create_request))
    assert create_result["status"] == "success", create_result
    reopened = load_workbook(source)
    assert reopened["Report"].sheet_view.showGridLines is False
    assert reopened["Report"].page_setup.fitToWidth == 1
    assert reopened["Report"].oddHeader.center.text == "Created"
    assert reopened["Report"].print_title_rows == "$1:$2"
    assert reopened["Report"]["A1"].hyperlink.location == "Details!A1"
    assert reopened["Report"]["B2"].comment.text == "Initial note"
    assert reopened.properties.title == "Annotated workbook"

    source_hash = sha256(source.read_bytes()).hexdigest()
    edits = [
        {"sheet": "Report", "type": "page_setup", "page_setup": _page_setup(fit=False)},
        {
            "sheet": "Report",
            "type": "header_footer",
            "header_footer": _header_footer("Edited"),
        },
        {
            "sheet": "Report",
            "type": "sheet_view",
            "view": {
                "show_grid_lines": True,
                "zoom_scale": 140,
                "selected_cell": "C2",
            },
        },
        {
            "sheet": "Report",
            "type": "print_titles",
            "print_titles": {"rows": "1:1", "columns": "A:A"},
        },
        {
            "sheet": "Report",
            "type": "hyperlink_update",
            "ref": "A1",
            "hyperlink": {
                "ref": "A1",
                "location": "Details!A1",
                "display": "Updated internal link",
                "tooltip": None,
            },
        },
        {
            "sheet": "Report",
            "type": "hyperlink_add",
            "hyperlink": {
                "ref": "C1",
                "location": "Details!A1",
                "display": None,
                "tooltip": "Added internal link",
            },
        },
        {
            "sheet": "Report",
            "type": "comment_update",
            "ref": "B2",
            "comment": {"ref": "B2", "text": "Updated note", "author": "Bob"},
        },
        {
            "sheet": "Report",
            "type": "comment_add",
            "comment": {"ref": "C2", "text": "Added note", "author": "Alice"},
        },
        {
            "sheet": "",
            "type": "workbook_properties",
            "properties": {"title": "Public edited title", "company": "Public Company"},
        },
    ]
    edit_request = _request(
        tmp_path,
        "public-metadata-edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(edited),
            "arguments": {"edits": edits, "expected_edits": len(edits)},
        },
    )
    edit_result = _public(project_root, "run", "--request", str(edit_request))
    assert edit_result["status"] == "success", edit_result
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    reopened = load_workbook(edited)
    assert reopened["Report"].page_setup.scale == 90
    assert reopened["Report"].oddHeader.center.text == "Edited"
    assert reopened["Report"].sheet_view.zoomScale == 140
    assert reopened["Report"].print_title_rows == "$1:$1"
    assert reopened["Report"]["A1"].hyperlink.display == "Updated internal link"
    assert reopened["Report"]["C1"].hyperlink.location == "Details!A1"
    assert reopened["Report"]["B2"].comment.text == "Updated note"
    assert reopened["Report"]["C2"].comment.text == "Added note"
    assert reopened.properties.title == "Public edited title"

    delete_request = _request(
        tmp_path,
        "public-metadata-delete.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(edited),
            "output": str(deleted),
            "arguments": {
                "edits": [
                    {"sheet": "Report", "type": "hyperlink_delete", "ref": "A1"},
                    {"sheet": "Report", "type": "hyperlink_delete", "ref": "C1"},
                    {"sheet": "Report", "type": "comment_delete", "ref": "B2"},
                    {"sheet": "Report", "type": "comment_delete", "ref": "C2"},
                    {"sheet": "Report", "type": "print_titles_clear"},
                ],
                "expected_edits": 5,
            },
        },
    )
    delete_result = _public(project_root, "run", "--request", str(delete_request))
    assert delete_result["status"] == "success", delete_result
    reopened = load_workbook(deleted)
    assert reopened["Report"]["A1"].hyperlink is None
    assert reopened["Report"]["C1"].hyperlink is None
    assert reopened["Report"]["B2"].comment is None
    assert reopened["Report"]["C2"].comment is None
    assert reopened["Report"].print_title_rows is None
