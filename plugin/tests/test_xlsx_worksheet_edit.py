"""Worksheet metadata, range, dimension, and defined-name edit tests."""

from pathlib import Path

from document_skills_core.formats.xlsx.service import XlsxService


def _source(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount", "Category", "Derived", "Merge me"])
    sheet.append(["One", 10, "A", "=B2*2"])
    sheet.append(["Two", 20, "B", "=B3*2"])
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 16
    workbook.save(path)


def test_worksheet_metadata_and_range_edits_reopen_in_consumer(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "edited.xlsx"
    _source(source)
    edits = [
        {"sheet": "Data", "type": "row_height", "ref": "2:3", "height": 28},
        {"sheet": "Data", "type": "row_hidden", "ref": "3", "hidden": True},
        {"sheet": "Data", "type": "column_width", "ref": "B:C", "width": 22},
        {"sheet": "Data", "type": "column_hidden", "ref": "C", "hidden": True},
        {"sheet": "Data", "type": "cells_merge", "ref": "E1:F1"},
        {"sheet": "Data", "type": "range_clear", "ref": "B2", "clear": "contents"},
        {"sheet": "Data", "type": "freeze_panes", "ref": "C3"},
        {"sheet": "Data", "type": "auto_filter", "ref": "A1:C3"},
        {"sheet": "Data", "type": "print_area", "ref": "A1:F5"},
        {"sheet": "Data", "type": "row_page_break", "ref": "4", "enabled": True},
        {"sheet": "Data", "type": "column_page_break", "ref": "D", "enabled": True},
        {
            "sheet": "Data",
            "type": "defined_name_add",
            "name": "InputArea",
            "ref": "Data!$A$1:$C$3",
            "scope": "workbook",
        },
    ]
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {"edits": edits, "expected_edits": len(edits)},
        },
    )

    assert result["status"] == "degraded"
    reopened = load_workbook(output, data_only=False)
    sheet = reopened["Data"]
    assert sheet.row_dimensions[2].height == 28
    assert sheet.row_dimensions[3].height == 28
    assert sheet.row_dimensions[3].hidden is True
    assert sheet.column_dimensions["B"].width == 22
    assert sheet.column_dimensions["C"].width == 22
    assert sheet.column_dimensions["C"].hidden is True
    assert [str(item) for item in sheet.merged_cells.ranges] == ["E1:F1"]
    assert sheet["B2"].value is None
    assert sheet["D2"].value == "=B2*2"
    assert sheet.freeze_panes == "C3"
    assert sheet.auto_filter.ref == "A1:C3"
    assert sheet.print_area == "'Data'!$A$1:$F$5"
    assert [item.id for item in sheet.row_breaks.brk] == [4]
    assert [item.id for item in sheet.col_breaks.brk] == [4]
    assert reopened.defined_names["InputArea"].attr_text == "Data!$A$1:$C$3"
    formula_cells = result["diagnostics"]["operation_result"]["formula_state"]["cells"]
    assert formula_cells["Data!D2"]["state"] == "recalculation_required"


def test_worksheet_metadata_clear_and_defined_name_crud(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    first = tmp_path / "first.xlsx"
    updated = tmp_path / "updated.xlsx"
    output = tmp_path / "deleted.xlsx"
    _source(source)
    first_edits = [
        {"sheet": "Data", "type": "cells_merge", "ref": "E1:F1"},
        {"sheet": "Data", "type": "auto_filter", "ref": "A1:C3"},
        {"sheet": "Data", "type": "print_area", "ref": "A1:F5"},
        {"sheet": "Data", "type": "freeze_panes", "ref": "C3"},
        {"sheet": "Data", "type": "row_page_break", "ref": "4"},
        {"sheet": "Data", "type": "column_page_break", "ref": "D"},
        {
            "sheet": "Data",
            "type": "defined_name_add",
            "name": "InputArea",
            "ref": "Data!$A$1:$C$3",
            "scope": "workbook",
        },
    ]
    first_result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(first),
            "arguments": {"edits": first_edits},
        },
    )
    assert first_result["status"] == "degraded"
    second_edits = [
        {"sheet": "Data", "type": "cells_unmerge", "ref": "E1:F1"},
        {"sheet": "Data", "type": "auto_filter_clear"},
        {"sheet": "Data", "type": "print_area_clear"},
        {"sheet": "Data", "type": "freeze_panes", "ref": "A1"},
        {"sheet": "Data", "type": "row_page_break", "ref": "4", "enabled": False},
        {"sheet": "Data", "type": "column_page_break", "ref": "D", "enabled": False},
        {
            "sheet": "Data",
            "type": "defined_name_update",
            "name": "InputArea",
            "ref": "Data!$A$1:$B$2",
            "scope": "workbook",
        },
    ]
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(first),
            "output": str(updated),
            "arguments": {"edits": second_edits},
        },
    )

    assert result["status"] == "degraded"
    reopened = load_workbook(updated, data_only=False)
    sheet = reopened["Data"]
    assert not sheet.merged_cells.ranges
    assert sheet.auto_filter.ref is None
    assert sheet.print_area == ""
    assert sheet.freeze_panes is None
    assert not sheet.row_breaks.brk
    assert not sheet.col_breaks.brk
    assert reopened.defined_names["InputArea"].attr_text == "Data!$A$1:$B$2"

    delete_result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(updated),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "defined_name_delete",
                        "name": "InputArea",
                        "scope": "workbook",
                    }
                ]
            },
        },
    )
    assert delete_result["status"] == "degraded"
    assert "InputArea" not in load_workbook(output).defined_names
