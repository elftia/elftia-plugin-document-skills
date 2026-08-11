"""Public committed-output truth across all four format service seams."""

from __future__ import annotations

import errno
import hashlib
import importlib
from pathlib import Path
from typing import Any

from openpyxl import Workbook
import pytest
from pptx import Presentation

import document_skills_core.core.io.paths as paths_module
from document_skills_core.cli import execute_request
from document_skills_core.core.contracts.schemas import SchemaCatalog
from document_skills_core.formats.pdf.create import create_pdf


_EDIT_FORMATS = ("docx", "xlsx", "pptx", "pdf")


def test_public_supported_docx_create_with_existing_output_returns_committed_result(
    project_root: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "existing.docx"
    output.write_bytes(b"prior-output")
    request = _create_request("docx", output)

    result = _execute(project_root, request)

    _assert_committed_existing_result(result, output)


@pytest.mark.parametrize("format_id", _EDIT_FORMATS)
def test_public_edit_with_existing_output_returns_committed_result(
    project_root: Path,
    tmp_path: Path,
    format_id: str,
) -> None:
    source, request = _edit_request(project_root, tmp_path, format_id)
    output = Path(str(request["output"]))
    output.write_bytes(b"prior-output")
    source_before = source.read_bytes()

    result = _execute(project_root, request)

    _assert_committed_existing_result(result, output, expected_artifacts=2)
    assert source.read_bytes() == source_before


@pytest.mark.parametrize("format_id", _EDIT_FORMATS)
def test_public_edit_source_race_merges_with_committed_output(
    project_root: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    format_id: str,
) -> None:
    source, request = _edit_request(project_root, tmp_path, format_id)
    output = Path(str(request["output"]))
    output.write_bytes(b"prior-output")
    transaction = importlib.import_module(
        f"document_skills_core.formats.{format_id}.transaction"
    )
    real_assert = transaction.assert_source_preserved
    checks = 0

    def change_source_after_commit(path: str | Path, expected: str) -> None:
        nonlocal checks
        checks += 1
        if checks == 2:
            Path(path).write_bytes(b"concurrent-source-writer")
        real_assert(path, expected)

    monkeypatch.setattr(
        transaction,
        "assert_source_preserved",
        change_source_after_commit,
    )
    result = _execute(project_root, request)

    _assert_committed_existing_result(result, output, expected_artifacts=2)
    assert checks == 2
    assert source.read_bytes() == b"concurrent-source-writer"
    assert {warning["code"] for warning in result["warnings"]} == {
        "DS_PROMOTION_RESIDUE_PRESERVED",
        "DS_SOURCE_CHANGED_AFTER_COMMIT",
    }
    promotion = result["diagnostics"]["promotion"]
    assert promotion["state"] == "committed_with_warnings"
    assert promotion["filesystem_state"] == "committed_with_residue"
    assert promotion["source_preservation"]["status"] == "fail"
    assert (
        promotion["source_preservation"]["error"]["code"]
        == "DS_VALIDATION_FAILED"
    )


@pytest.mark.parametrize("format_id", _EDIT_FORMATS)
def test_public_edit_failure_preserves_promotion_and_source_race_truth(
    project_root: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    format_id: str,
) -> None:
    source, request = _edit_request(project_root, tmp_path, format_id)
    output = Path(str(request["output"]))
    output.write_bytes(b"prior-output")
    real_rename = paths_module._rename_no_replace
    failures: list[str] = []

    def fail_install_and_restore(
        rename_source: Path,
        rename_destination: Path,
        *,
        state: Any,
    ) -> None:
        if rename_destination.name == output.name and rename_source.name.startswith(
            ".document-skills-stage-"
        ):
            source.write_bytes(b"concurrent-source-writer")
            failures.append("candidate_install")
            raise OSError(errno.EIO, "injected candidate install failure")
        if rename_destination.name == output.name and rename_source.name.startswith(
            ".document-skills-capture-"
        ):
            failures.append("capture_restore")
            raise OSError(errno.EIO, "injected capture restore failure")
        real_rename(rename_source, rename_destination, state=state)

    monkeypatch.setattr(
        paths_module,
        "_rename_no_replace",
        fail_install_and_restore,
    )
    result = _execute(project_root, request)

    SchemaCatalog(project_root).validate("operation-result", result)
    assert failures == ["candidate_install", "capture_restore"]
    assert result["status"] == "failed"
    assert result["artifacts"] == []
    assert len(result["errors"]) == 1
    error = result["errors"][0]
    assert error["code"] == "DS_VALIDATION_FAILED"
    assert error["message"] == (
        "Atomic no-replace promotion is unavailable for this destination."
    )
    details = error["details"]
    assert details["atomic_no_replace_unavailable"] is True
    assert details["errno"] == errno.EIO
    assert details["rollback_complete"] is False
    assert details["rollback_reason"] == "OSError"
    assert details["destination_capture_preserved"] is True
    assert details["residue_observation_stable"] is True
    source_preservation = details["source_preservation"]
    assert source_preservation["status"] == "fail"
    assert source_preservation["error"]["code"] == "DS_VALIDATION_FAILED"
    assert source_preservation["error"]["message"] == (
        "Source artifact changed during a non-destructive operation."
    )
    assert source_preservation["error"]["details"]["actual_sha256"] == (
        hashlib.sha256(b"concurrent-source-writer").hexdigest()
    )
    assert source.read_bytes() == b"concurrent-source-writer"
    assert not output.exists()

    residues = details["transaction_residues"]
    residue_paths = sorted(Path(path) for path in details["transaction_residue_paths"])
    assert residue_paths == sorted(tmp_path.glob(".document-skills-*"))
    assert sorted(Path(item["path"]) for item in residues) == residue_paths
    assert {item["role"] for item in residues} == {
        "destination_capture",
        "stage",
    }
    for item in residues:
        path = Path(str(item["path"]))
        assert item["stable"] is True
        assert item["bytes"] == path.stat().st_size
        assert item["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    recovery_bytes = {
        item["role"]: Path(str(item["path"])).read_bytes() for item in residues
    }
    assert recovery_bytes["destination_capture"] == b"prior-output"
    assert recovery_bytes["stage"]


def _execute(project_root: Path, request: dict[str, Any]) -> dict[str, Any]:
    return execute_request(request, project_root, SchemaCatalog(project_root))


def _create_request(format_id: str, output: Path) -> dict[str, Any]:
    if format_id == "docx":
        arguments = {
            "report": {
                "blocks": [
                    {"type": "heading", "text": "Committed", "level": 1},
                    {"type": "paragraph", "text": "Bounded public output."},
                ]
            }
        }
    else:
        arguments = {"document": _pdf_document()}
    return {
        "schema_version": "1.0",
        "operation": f"{format_id}.create",
        "output": str(output),
        "arguments": arguments,
    }


def _edit_request(
    project_root: Path,
    root: Path,
    format_id: str,
) -> tuple[Path, dict[str, Any]]:
    source = root / f"source.{format_id}"
    output = root / f"existing-edit.{format_id}"
    if format_id == "docx":
        created = _execute(
            project_root,
            {
                "schema_version": "1.0",
                "operation": "docx.create",
                "output": str(source),
                "arguments": {
                    "report": {
                        "blocks": [
                            {"type": "paragraph", "text": "source marker"}
                        ]
                    }
                },
            },
        )
        assert created["status"] == "success", created
        arguments = {
            "replacements": [
                {
                    "search": "source marker",
                    "replace": "edited marker",
                    "expected_matches": 1,
                }
            ]
        }
        operation = "docx.edit.replace-text"
    elif format_id == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet["A1"] = "Name"
        sheet["B1"] = "Value"
        sheet["A2"] = "Alpha"
        sheet["B2"] = 10
        workbook.save(source)
        arguments = {
            "edits": [
                {
                    "sheet": "Data",
                    "type": "cell_value",
                    "ref": "B2",
                    "value": "99",
                }
            ]
        }
        operation = "xlsx.edit"
    elif format_id == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Source title"
        presentation.save(source)
        arguments = {
            "edits": [
                {
                    "slide": 1,
                    "type": "slide_text",
                    "ref": "",
                    "value": "Edited title",
                }
            ]
        }
        operation = "pptx.edit"
    else:
        create_pdf(source, _pdf_document())
        arguments = {
            "primitives": [
                {"type": "rotate", "pages": [1], "degrees": 90}
            ]
        }
        operation = "pdf.edit"
    return source, {
        "schema_version": "1.0",
        "operation": operation,
        "input": str(source),
        "output": str(output),
        "arguments": arguments,
    }


def _pdf_document() -> dict[str, Any]:
    return {
        "metadata": {"title": "Committed PDF", "author": "Elftia"},
        "page_size": "A4",
        "pages": [
            {
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": "Page one",
                        "style": None,
                        "table": None,
                        "image": None,
                        "shape": None,
                    }
                ],
                "metadata": None,
            },
            {
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": "Page two",
                        "style": None,
                        "table": None,
                        "image": None,
                        "shape": None,
                    }
                ],
                "metadata": None,
            },
        ],
    }


def _assert_committed_existing_result(
    result: dict[str, Any],
    output: Path,
    *,
    expected_artifacts: int = 1,
) -> None:
    assert result["status"] in {"success", "degraded"}, result
    assert result["achieved_fidelity"] != "none"
    assert result["validation"]["status"] == "pass"
    assert result["errors"] == []
    assert len(result["artifacts"]) == expected_artifacts
    output_artifact = result["artifacts"][-1]
    assert output_artifact["path"] == str(output.resolve())
    assert output_artifact["sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert output_artifact["bytes"] == output.stat().st_size
    promotion = result["diagnostics"]["promotion"]
    assert promotion["promotion_committed"] is True
    assert promotion["filesystem_state"] == "committed_with_residue"
    assert promotion["destination_capture_preserved"] is True
    assert promotion["residue_observation_stable"] is True
    assert len(promotion["transaction_residues"]) == 1
    capture = promotion["transaction_residues"][0]
    assert capture["role"] == "destination_capture"
    assert capture["stable"] is True
    assert capture["identity_matches_expected"] is True
    assert Path(str(capture["path"])).read_bytes() == b"prior-output"
    assert "DS_PROMOTION_RESIDUE_PRESERVED" in {
        warning["code"] for warning in result["warnings"]
    }
