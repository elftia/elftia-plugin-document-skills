"""XLSX operation tests — create, read, inspect, edit through the real service."""

import hashlib
from pathlib import Path
import zipfile

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.core.contracts.schemas import SchemaCatalog
from document_skills_core.formats.xlsx.constants import (
    FORMULA_STATE_RECALCULATION_REQUIRED,
    FORMULA_STATE_STALE,
)
from document_skills_core.formats.xlsx.create import create_xlsx
from document_skills_core.formats.xlsx.mapping import map_workbook
from document_skills_core.formats.xlsx.package import OpcPackage
from document_skills_core.formats.xlsx.service import XlsxService


def _service(project_root: Path) -> XlsxService:
    return XlsxService(project_root)


def _strip_style_children(path: Path) -> None:
    """Replace styles.xml with an empty-container form that fails the style gate."""

    with zipfile.ZipFile(path, "r") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/styles.xml"] = (
        b'<?xml version="1.0" encoding="UTF-8"?>\n'
        b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<fonts count="1"/><fills count="2"/><borders count="1"/>'
        b'<cellStyleXfs count="1"/><cellXfs count="1"/>'
        b'<cellStyles count="1"/><dxfs count="0"/></styleSheet>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in sorted(parts.items()):
            archive.writestr(name, data)


def _sample_workbook() -> dict:
    return {
        "metadata": {"title": "Test Workbook", "creator": "Test", "subject": "Testing"},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Name", "type": "s"},
                            {"ref": "B1", "value": "Value", "type": "s"},
                            {"ref": "A2", "value": "Alpha", "type": "s"},
                            {"ref": "B2", "value": "10", "type": "n"},
                            {"ref": "A3", "value": "Beta", "type": "s"},
                            {"ref": "B3", "value": "20", "type": "n"},
                        ],
                    },
                    {
                        "cells": [
                            {"ref": "A5", "value": "Total", "type": "s"},
                            {"ref": "B5", "formula": "SUM(B2:B3)", "type": "n"},
                        ],
                    },
                ],
                "number_formats": [],
            },
            {
                "name": "Summary",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Result", "type": "s"},
                            {"ref": "B1", "formula": "Data!B5*2", "type": "n"},
                        ],
                    }
                ],
                "number_formats": [],
            },
        ],
        "defined_names": [
            {"name": "TotalVal", "ref": "Data!$B$5", "scope": "workbook"},
        ],
        "tables": [
            {"name": "DataTable", "ref": "A1:B3", "sheet": "Data", "style": "TableStyleMedium2"},
        ],
        "chart_reference": {
            "title": "Values Chart",
            "data_ref": "Data!B2:B3",
            "sheet": "Data",
        },
        "page_setup": {
            "orientation": "portrait",
            "header": "Test Header",
            "footer": "Test Footer",
        },
    }


def _bounded_workbook() -> dict:
    return {
        "metadata": {"title": "Bounded", "creator": "Test", "subject": ""},
        "sheets": [
            {
                "name": "Sheet1",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Name", "type": "s"},
                            {"ref": "B1", "value": "10", "type": "n"},
                        ]
                    }
                ],
                "number_formats": [],
            }
        ],
        "defined_names": [],
        "tables": [],
        "chart_reference": None,
        "page_setup": None,
    }


def _styled_workbook() -> dict:
    workbook = _bounded_workbook()
    sheet = workbook["sheets"][0]
    sheet["columns"] = [
        {"ref": "A:B", "width": 16, "hidden": False, "style": None},
    ]
    sheet["rows"][0]["height"] = 22
    sheet["rows"][0]["style"] = {
        "font": {"bold": True},
        "fill": {"pattern": "solid", "color": "#DDEEFF"},
    }
    sheet["rows"][0]["cells"][0]["style"] = {
        "alignment": {"horizontal": "center", "wrap": True},
    }
    repeated_style = {
        "number_format": {"id": 165},
        "protection": {"locked": False},
    }
    sheet["rows"][0]["cells"][1]["style"] = repeated_style
    sheet["rows"][0]["cells"].append(
        {"ref": "C1", "value": "20", "type": "n", "style": repeated_style}
    )
    sheet["number_formats"] = [{"id": 165, "code": "0.000"}]
    return workbook


@pytest.fixture
def created_xlsx(project_root: Path, tmp_path: Path) -> Path:
    from openpyxl import Workbook

    output = tmp_path / "created.xlsx"
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data["A1"] = "Name"
    data["B1"] = "Value"
    data["A2"] = "Alpha"
    data["B2"] = 10
    data["A3"] = "Beta"
    data["B3"] = 20
    data["A5"] = "Total"
    data["B5"] = "=SUM(B2:B3)"
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Result"
    summary["B1"] = "=Data!B5*2"
    workbook.save(output)
    return output


class TestCreateOperation:
    def test_disconnected_create_request_fails_before_output(self, project_root: Path, tmp_path: Path):
        output = tmp_path / "test.xlsx"
        service = _service(project_root)
        result = service.execute("xlsx.create", {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": _sample_workbook()},
        })
        assert result["status"] == "enhancement_required"
        assert result["provider_chain"] == []
        assert not output.exists()
        assert not result["artifacts"]

    def test_disconnected_create_does_not_report_formula_success(self, project_root: Path, tmp_path: Path):
        output = tmp_path / "formulas.xlsx"
        service = _service(project_root)
        result = service.execute("xlsx.create", {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": _sample_workbook()},
        })
        assert result["status"] == "enhancement_required"
        assert "operation_result" not in result["diagnostics"]
        assert not output.exists()

    def test_create_with_disconnected_features_is_not_degraded(self, project_root: Path, tmp_path: Path):
        output = tmp_path / "degraded.xlsx"
        service = _service(project_root)
        result = service.execute("xlsx.create", {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": _sample_workbook()},
        })
        assert result["status"] == "enhancement_required"
        assert result["degraded"] is False
        assert not result["degradations"]
        assert not output.exists()

    def test_create_styles_are_deduplicated_and_deterministic(
        self,
        project_root: Path,
        tmp_path: Path,
    ) -> None:
        service = _service(project_root)
        outputs = [tmp_path / "styled-a.xlsx", tmp_path / "styled-b.xlsx"]
        results = [
            service.execute(
                "xlsx.create",
                {
                    "schema_version": "1.0",
                    "operation": "xlsx.create",
                    "output": str(output),
                    "arguments": {"workbook": _styled_workbook()},
                },
            )
            for output in outputs
        ]

        assert all(result["status"] == "success" for result in results)
        assert outputs[0].read_bytes() == outputs[1].read_bytes()
        first_styles = results[0]["diagnostics"]["operation_result"]["creation"]["styles"]
        assignments = first_styles["assignments"]["cells"]
        assert assignments["Sheet1!B1"] == assignments["Sheet1!C1"]
        assert first_styles["cell_xfs"] < len(assignments) + 3

    def test_invalid_style_preserves_existing_destination(
        self,
        project_root: Path,
        tmp_path: Path,
    ) -> None:
        output = tmp_path / "existing.xlsx"
        output.write_bytes(b"existing-style-destination")
        workbook = _bounded_workbook()
        workbook["sheets"][0]["rows"][0]["cells"][0]["style"] = {
            "font": {"color": "invalid"}
        }

        result = _service(project_root).execute(
            "xlsx.create",
            {
                "schema_version": "1.0",
                "operation": "xlsx.create",
                "output": str(output),
                "arguments": {"workbook": workbook},
            },
        )

        assert result["status"] == "invalid_request"
        assert output.read_bytes() == b"existing-style-destination"


class TestReadOperation:
    def test_read_returns_structure(self, project_root: Path, created_xlsx: Path):
        service = _service(project_root)
        result = service.execute("xlsx.read", {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(created_xlsx),
            "arguments": {},
        })
        op_result = result["diagnostics"]["operation_result"]
        assert result["status"] in {"success", "degraded"}
        assert op_result["sheet_count"] == 2
        sheets = op_result["sheets"]
        assert sheets[0]["name"] == "Data"
        assert sheets[1]["name"] == "Summary"

    def test_read_reports_formula_state(self, project_root: Path, created_xlsx: Path):
        service = _service(project_root)
        result = service.execute("xlsx.read", {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(created_xlsx),
            "arguments": {},
        })
        formula_state = result["diagnostics"]["operation_result"]["formula_state"]
        # There should be formula cells (SUM and cross-sheet ref)
        assert len(formula_state["cells"]) > 0
        # No formula should report recalculated without a provider
        for ref, cell in formula_state["cells"].items():
            assert cell["state"] != "recalculated"
        assert formula_state["summary"]["no_unverified_claimed_recalculated"] is True

    def test_read_preserves_source(self, project_root: Path, created_xlsx: Path):
        import hashlib
        original_hash = hashlib.sha256(created_xlsx.read_bytes()).hexdigest()
        service = _service(project_root)
        service.execute("xlsx.read", {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(created_xlsx),
            "arguments": {},
        })
        assert hashlib.sha256(created_xlsx.read_bytes()).hexdigest() == original_hash


class TestInspectOperation:
    def test_inspect_returns_inventory(self, project_root: Path, created_xlsx: Path):
        service = _service(project_root)
        result = service.execute("xlsx.inspect.structure", {
            "schema_version": "1.0",
            "operation": "xlsx.inspect.structure",
            "input": str(created_xlsx),
            "arguments": {},
        })
        op_result = result["diagnostics"]["operation_result"]
        assert result["status"] == "success"
        assert op_result["mutation_authorized"] is False
        assert op_result["part_count"] > 0
        assert op_result["worksheet_count"] == 2

    def test_inspection_does_not_authorize_mutation(self, project_root: Path, created_xlsx: Path):
        service = _service(project_root)
        result = service.execute("xlsx.inspect.structure", {
            "schema_version": "1.0",
            "operation": "xlsx.inspect.structure",
            "input": str(created_xlsx),
            "arguments": {},
        })
        assert result["diagnostics"]["operation_result"]["mutation_authorized"] is False


class TestEditOperation:
    def test_edit_cell_value(self, project_root: Path, created_xlsx: Path, tmp_path: Path):
        output = tmp_path / "edited.xlsx"
        service = _service(project_root)
        result = service.execute("xlsx.edit", {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(created_xlsx),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "cell_value", "ref": "B2", "value": "99"},
                ],
            },
        })
        assert result["status"] in {"success", "degraded"}
        assert output.is_file()
        assert created_xlsx.is_file()  # source preserved
        assert any(
            gate["id"] == "operation.consumer-package-conformance"
            and gate["outcome"] == "pass"
            for gate in result["validation"]["gates"]
        )

    def test_edit_rejects_consumer_invalid_package_without_promotion(
        self,
        project_root: Path,
        tmp_path: Path,
    ) -> None:
        source = tmp_path / "core-invalid.xlsx"
        create_xlsx(source, _bounded_workbook())
        _strip_style_children(source)
        source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
        output = tmp_path / "existing.xlsx"
        existing = b"existing-destination-must-survive"
        output.write_bytes(existing)
        service = _service(project_root)

        result = service.execute(
            "xlsx.edit",
            {
                "schema_version": "1.0",
                "operation": "xlsx.edit",
                "input": str(source),
                "output": str(output),
                "arguments": {
                    "edits": [
                        {
                            "sheet": "Sheet1",
                            "type": "cell_value",
                            "ref": "B1",
                            "value": "42",
                        }
                    ]
                },
            },
        )

        SchemaCatalog(project_root).validate("operation-result", result)
        assert result["status"] == "failed"
        assert result["validation"]["status"] == "fail"
        assert any(
            gate["id"] == "operation.consumer-package-conformance"
            and gate["required"] is True
            and gate["outcome"] == "fail"
            for gate in result["validation"]["gates"]
        )
        assert result["artifacts"] == []
        assert output.read_bytes() == existing
        assert hashlib.sha256(source.read_bytes()).hexdigest() == source_sha256

    def test_edit_preserves_source(self, project_root: Path, created_xlsx: Path, tmp_path: Path):
        import hashlib
        original_hash = hashlib.sha256(created_xlsx.read_bytes()).hexdigest()
        output = tmp_path / "edited.xlsx"
        service = _service(project_root)
        service.execute("xlsx.edit", {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(created_xlsx),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "cell_value", "ref": "B3", "value": "21"},
                ],
            },
        })
        assert hashlib.sha256(created_xlsx.read_bytes()).hexdigest() == original_hash

    def test_edit_formula_invalidates_dependents(self, project_root: Path, created_xlsx: Path, tmp_path: Path):
        """Editing a precedent cell should mark dependents as recalculation_required."""
        output = tmp_path / "invalidated.xlsx"
        service = _service(project_root)
        result = service.execute("xlsx.edit", {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(created_xlsx),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "cell_value", "ref": "B2", "value": "100"},
                ],
            },
        })
        formula_state = result["diagnostics"]["operation_result"].get("formula_state", {})
        cells = formula_state.get("cells", {})
        # At least one formula should be recalculation_required
        recalc_required = [
            ref for ref, c in cells.items()
            if c.get("state") == FORMULA_STATE_RECALCULATION_REQUIRED
        ]
        assert len(recalc_required) > 0, "Expected at least one dependent invalidated"

    def test_edit_output_equals_input_rejected(self, project_root: Path, created_xlsx: Path):
        service = _service(project_root)
        result = service.execute("xlsx.edit", {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(created_xlsx),
            "output": str(created_xlsx),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "cell_value", "ref": "A1", "value": "X"},
                ],
            },
        })
        assert result["status"] == "invalid_request"
        assert result["errors"][0]["code"] == "DS_OUTPUT_EQUALS_INPUT"

    def test_style_edit_appends_records_and_preserves_existing_style_payloads(
        self,
        project_root: Path,
        tmp_path: Path,
    ) -> None:
        from defusedxml.ElementTree import fromstring
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Color, Font, NamedStyle
        from xml.etree.ElementTree import tostring

        source = tmp_path / "styled-source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet["A1"] = "Theme font"
        sheet["B1"] = 3.5
        sheet["A1"].font = Font(name="Aptos", bold=True, color=Color(theme=1))
        unused = NamedStyle(name="PreservedNamedStyle", number_format="0.00%")
        workbook.add_named_style(unused)
        workbook.save(source)
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        before = _style_record_payloads(source, fromstring, tostring)
        output = tmp_path / "styled-output.xlsx"

        result = _service(project_root).execute(
            "xlsx.edit",
            {
                "schema_version": "1.0",
                "operation": "xlsx.edit",
                "input": str(source),
                "output": str(output),
                "arguments": {
                    "edits": [
                        {
                            "sheet": "Sheet1",
                            "type": "cell_style",
                            "ref": "A1",
                            "style": {
                                "font": {"italic": True},
                                "fill": {"color": "#DDEEFF"},
                            },
                        },
                        {
                            "sheet": "Sheet1",
                            "type": "row_style",
                            "ref": "2",
                            "style": {"font": {"bold": True}},
                        },
                        {
                            "sheet": "Sheet1",
                            "type": "column_style",
                            "ref": "B",
                            "style": {"number_format": {"code": "0.0000"}},
                        },
                    ],
                    "expected_edits": 3,
                },
            },
        )

        assert result["status"] == "success"
        assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
        after = _style_record_payloads(output, fromstring, tostring)
        for container, original_records in before.items():
            assert after[container][: len(original_records)] == original_records
        reopened = load_workbook(output)
        edited = reopened["Sheet1"]["A1"]
        assert edited.font.bold is True
        assert edited.font.italic is True
        assert edited.font.color.type == "theme"
        assert edited.font.color.theme == 1
        assert edited.fill.fgColor.rgb == "FFDDEEFF"
        assert reopened["Sheet1"].row_dimensions[2].style_id > 0
        assert reopened["Sheet1"].column_dimensions["B"].number_format == "0.0000"
        mapped = map_workbook(OpcPackage.open(output))
        mapped_cells = {
            cell["ref"]: cell
            for row in mapped["sheets"][0]["rows"]
            for cell in row["cells"]
        }
        assert mapped_cells["B1"]["style_source"] == "column"
        assert mapped_cells["B1"]["number_format"] == "0.0000"

    def test_missing_custom_number_format_id_preserves_existing_destination(
        self,
        project_root: Path,
        created_xlsx: Path,
        tmp_path: Path,
    ) -> None:
        output = tmp_path / "existing-style-output.xlsx"
        output.write_bytes(b"existing-style-output")
        source_hash = hashlib.sha256(created_xlsx.read_bytes()).hexdigest()

        result = _service(project_root).execute(
            "xlsx.edit",
            {
                "schema_version": "1.0",
                "operation": "xlsx.edit",
                "input": str(created_xlsx),
                "output": str(output),
                "arguments": {
                    "edits": [
                        {
                            "sheet": "Data",
                            "type": "cell_style",
                            "ref": "A1",
                            "style": {"number_format": {"id": 165}},
                        }
                    ]
                },
            },
        )

        assert result["status"] == "invalid_request"
        assert output.read_bytes() == b"existing-style-output"
        assert hashlib.sha256(created_xlsx.read_bytes()).hexdigest() == source_hash


def _style_record_payloads(path: Path, fromstring, serialize) -> dict[str, list[bytes]]:
    with zipfile.ZipFile(path, "r") as archive:
        root = fromstring(archive.read("xl/styles.xml"))
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    result: dict[str, list[bytes]] = {}
    for container_name in ("fonts", "fills", "borders", "cellXfs"):
        container = root.find(f"{{{namespace}}}{container_name}")
        assert container is not None
        result[container_name] = [serialize(item, encoding="UTF-8") for item in container]
    return result
