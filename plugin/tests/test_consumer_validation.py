"""Independent portable and conditional Office consumer-gate contracts."""

from __future__ import annotations

import base64
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import time
from typing import Any
import xml.etree.ElementTree as ElementTree
import zipfile

import pytest

from consumer_validation.contracts import validate_consumer_report
from consumer_validation.harness import qualify_artifact
from consumer_validation.office import _parse_probe_output, detect_office, open_with_office
import consumer_validation.harness as harness_module
import consumer_validation.office as office_module
import consumer_validation.pdf as pdf_module


_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAABAAAAAQCAYAAAAf8/9hAAAAI0lEQVR42mP4sEXjv0LCg"
    "f/k0gyUaAbRDKMuGHXBqAsGiQsA0UOFH7DDlt4AAAAASUVORK5CYII="
)


@pytest.mark.parametrize("format_id", ["docx", "xlsx", "pptx", "pdf"])
def test_independently_produced_fixture_passes(
    tmp_path: Path,
    format_id: str,
) -> None:
    artifact, expectations = _known_good(tmp_path, format_id)
    report = qualify_artifact(
        format_id=format_id,
        operation=f"{format_id}.create",
        artifact=artifact,
        expectations=expectations,
        office_policy="off",
    )

    validate_consumer_report(report)
    assert report["portable"]["outcome"] == "pass"
    assert report["artifact"]["sha256"] == _sha256(artifact)
    assert report["artifact"]["bytes"] == artifact.stat().st_size
    assert report["office_acceptance"] is False


@pytest.mark.parametrize("format_id", ["docx", "xlsx", "pptx", "pdf"])
def test_deliberately_corrupted_fixture_fails_independently(
    tmp_path: Path,
    format_id: str,
) -> None:
    artifact, expectations = _known_good(tmp_path, format_id)
    corrupted = _corrupt(tmp_path, artifact, format_id)

    report = qualify_artifact(
        format_id=format_id,
        operation=f"{format_id}.create",
        artifact=corrupted,
        expectations=expectations,
        office_policy="off",
    )

    validate_consumer_report(report)
    assert report["status"] == "fail"
    assert report["portable"]["outcome"] == "fail"
    assert any(
        assertion["outcome"] == "fail"
        for assertion in report["portable"]["assertions"]
    )


@pytest.mark.parametrize("format_id", ["docx", "xlsx", "pptx"])
def test_external_relationships_are_rejected_for_every_ooxml_format(
    tmp_path: Path,
    format_id: str,
) -> None:
    artifact, expectations = _known_good(tmp_path, format_id)
    relationship_part = {
        "docx": "word/_rels/document.xml.rels",
        "xlsx": "xl/_rels/workbook.xml.rels",
        "pptx": "ppt/_rels/presentation.xml.rels",
    }[format_id]
    mutated = _mutate_relationships(
        tmp_path,
        artifact,
        relationship_part,
        lambda root: ElementTree.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            {
                "Id": "rIdExternalRegression",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate",
                "Target": "https://example.invalid/never-fetch",
                "TargetMode": "External",
            },
        ),
    )

    report = qualify_artifact(
        format_id=format_id,
        operation=f"{format_id}.create",
        artifact=mutated,
        expectations=expectations,
        office_policy="off",
    )

    assertion = _assertions(report)["ooxml.relationship-graph"]
    assert assertion["outcome"] == "fail"
    assert assertion["evidence"]["issues"][0]["category"] == "external-relationship"


def test_pptx_slide_layout_master_closure_is_required(tmp_path: Path) -> None:
    artifact, expectations = _known_good(tmp_path, "pptx")

    def remove_slide_layout(root: ElementTree.Element) -> None:
        for relationship in list(root):
            if relationship.attrib.get("Type", "").endswith("/slideLayout"):
                root.remove(relationship)

    mutated = _mutate_relationships(
        tmp_path,
        artifact,
        "ppt/slides/_rels/slide1.xml.rels",
        remove_slide_layout,
    )
    report = qualify_artifact(
        format_id="pptx",
        operation="pptx.create",
        artifact=mutated,
        expectations=expectations,
        office_policy="off",
    )

    assert report["status"] == "fail"
    assert _assertions(report)["ooxml.critical-relationship-closure"]["outcome"] == "fail"


@pytest.mark.parametrize(
    ("mutation", "category"),
    [
        ("dangling", "dangling-relationship-target"),
        ("unsafe", "unsafe-relationship-target"),
        ("duplicate", "duplicate-relationship-id"),
        ("empty", "empty-relationship-field"),
    ],
)
def test_relationship_graph_rejects_invalid_internal_records(
    tmp_path: Path,
    mutation: str,
    category: str,
) -> None:
    artifact, expectations = _known_good(tmp_path, "docx")

    def mutate(root: ElementTree.Element) -> None:
        first = next(iter(root))
        attributes = {
            "Id": "rIdRelationshipRegression",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
            "Target": "media/missing.png",
        }
        if mutation == "unsafe":
            attributes["Target"] = "../../../outside.xml"
        elif mutation == "duplicate":
            attributes.update(
                {
                    "Id": first.attrib["Id"],
                    "Type": first.attrib["Type"],
                    "Target": first.attrib["Target"],
                }
            )
        elif mutation == "empty":
            attributes["Target"] = ""
        ElementTree.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            attributes,
        )

    mutated = _mutate_relationships(
        tmp_path,
        artifact,
        "word/_rels/document.xml.rels",
        mutate,
    )
    report = qualify_artifact(
        format_id="docx",
        operation="docx.create",
        artifact=mutated,
        expectations=expectations,
        office_policy="off",
    )

    issues = _assertions(report)["ooxml.relationship-graph"]["evidence"]["issues"]
    assert any(issue["category"] == category for issue in issues)


def test_pdf_text_delta_and_placeholder_only_image_fail(tmp_path: Path) -> None:
    import fitz

    artifact = tmp_path / "placeholder.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "????")
    page.draw_rect(fitz.Rect(72, 100, 172, 200))
    page.insert_text((80, 150), "[Image: placeholder.png]")
    document.save(artifact)
    document.close()
    reference = tmp_path / "expected.pdf"
    _write_pdf(reference, "Expected PDF", include_image=True)

    report = qualify_artifact(
        format_id="pdf",
        operation="pdf.create",
        artifact=artifact,
        expectations={
            "text": ["中文"],
            "real_images": 1,
            "render": {
                "reference": str(reference),
                "reference_kind": "expected",
                "mode": "match",
            },
        },
        office_policy="off",
    )

    assert report["status"] == "fail"
    assertions = {item["id"]: item for item in report["portable"]["assertions"]}
    assert assertions["pdf.requested-text"]["outcome"] == "fail"
    assert assertions["pdf.real-images"]["outcome"] == "fail"
    assert assertions["pdf.render-delta"]["outcome"] == "fail"


def test_pdf_visible_mutation_fails_source_render_delta_with_same_page_count(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.pdf"
    candidate = tmp_path / "candidate.pdf"
    _write_pdf(source, "Stable visible output")
    _write_pdf(candidate, "Changed visible output")

    report = qualify_artifact(
        format_id="pdf",
        operation="pdf.edit",
        artifact=candidate,
        expectations={
            "text": ["Changed visible output"],
            "render": {
                "reference": str(source),
                "reference_kind": "source",
                "mode": "match",
                "max_changed_sample_ratio": 0.0,
            },
        },
        office_policy="off",
    )

    assertion = _assertions(report)["pdf.render-delta"]
    assert report["status"] == "fail"
    assert assertion["outcome"] == "fail"
    assert assertion["evidence"]["pages"] == 1
    assert assertion["evidence"]["changed_samples"] > 0


def test_pdf_no_render_page_limit_fails_closed_with_bounded_evidence(
    tmp_path: Path,
) -> None:
    import fitz

    artifact = tmp_path / "thirty-three-pages.pdf"
    document = fitz.open()
    for _index in range(33):
        document.new_page()
    document.save(artifact)
    document.close()

    report = qualify_artifact(
        format_id="pdf",
        operation="pdf.create",
        artifact=artifact,
        expectations={},
        office_policy="off",
    )

    validate_consumer_report(report)
    assertion = _assertions(report)["pdf.resource-bounds"]
    assert report["status"] == "fail"
    assert assertion["outcome"] == "fail"
    assert assertion["evidence"] == {
        "actual": 33,
        "category": "artifact-page-limit",
        "maximum": 32,
    }


def test_pdf_oversized_artifact_fails_before_full_read(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    artifact = tmp_path / "oversized-relative-to-policy.pdf"
    _write_pdf(artifact, "Bounded artifact")
    policy_limit = artifact.stat().st_size - 1
    monkeypatch.setattr(
        pdf_module,
        "_MAX_ARTIFACT_BYTES",
        policy_limit,
        raising=False,
    )
    hash_calls = 0

    def reject_candidate_hash(path: Path) -> str:
        nonlocal hash_calls
        hash_calls += 1
        raise AssertionError(
            f"oversized candidate must be rejected before aggregate hash: {path}"
        )

    monkeypatch.setattr(harness_module, "_sha256", reject_candidate_hash)

    report = qualify_artifact(
        format_id="pdf",
        operation="pdf.create",
        artifact=artifact,
        expectations={},
        office_policy="off",
    )

    validate_consumer_report(report)
    assertion = _assertions(report)["pdf.resource-bounds"]
    assert report["status"] == "fail"
    assert assertion["outcome"] == "fail"
    assert assertion["evidence"] == {
        "actual": artifact.stat().st_size,
        "category": "artifact-byte-limit",
        "maximum": policy_limit,
    }
    assert report["artifact"] == {
        "path": str(artifact),
        "sha256": None,
        "sha256_status": "not-computed-resource-limit",
        "bytes": artifact.stat().st_size,
    }
    assert hash_calls == 0


def test_pdf_oversized_render_reference_fails_before_reference_read(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    artifact = tmp_path / "candidate.pdf"
    reference = tmp_path / "oversized-reference.pdf"
    _write_pdf(artifact, "Stable render")
    _write_pdf(reference, "Stable render")
    reference.write_bytes(reference.read_bytes() + b" " * 128)
    policy_limit = max(artifact.stat().st_size, reference.stat().st_size - 1)
    monkeypatch.setattr(
        pdf_module,
        "_MAX_ARTIFACT_BYTES",
        policy_limit,
        raising=False,
    )

    report = qualify_artifact(
        format_id="pdf",
        operation="pdf.create",
        artifact=artifact,
        expectations={
            "text": ["Stable render"],
            "render": {
                "reference": str(reference),
                "reference_kind": "expected",
                "mode": "match",
            },
        },
        office_policy="off",
    )

    validate_consumer_report(report)
    assertion = _assertions(report)["pdf.render-delta"]
    assert report["status"] == "fail"
    assert assertion["outcome"] == "fail"
    assert assertion["evidence"] == {
        "actual": reference.stat().st_size,
        "category": "reference-artifact-byte-limit",
        "maximum": policy_limit,
    }


def test_pdf_over_page_render_reference_fails_before_render_or_hash(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import fitz

    artifact = tmp_path / "candidate.pdf"
    reference = tmp_path / "thirty-three-page-reference.pdf"
    _write_pdf(artifact, "Stable render")
    document = fitz.open()
    for _index in range(33):
        document.new_page()
    document.save(reference)
    document.close()

    def reject_reference_hash(_path: Path) -> str:
        raise AssertionError("over-page reference must not be hashed")

    def reject_reference_render(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        raise AssertionError("over-page reference must not be rendered")

    monkeypatch.setattr(pdf_module, "_sha256_file", reject_reference_hash)
    monkeypatch.setattr(pdf_module, "_compare_rendered_pages", reject_reference_render)

    report = qualify_artifact(
        format_id="pdf",
        operation="pdf.create",
        artifact=artifact,
        expectations={
            "render": {
                "reference": str(reference),
                "reference_kind": "expected",
                "mode": "match",
            },
        },
        office_policy="off",
    )

    validate_consumer_report(report)
    assertion = _assertions(report)["pdf.render-delta"]
    assert report["status"] == "fail"
    assert assertion["outcome"] == "fail"
    assert assertion["evidence"] == {
        "actual": 33,
        "category": "reference-artifact-page-limit",
        "maximum": 32,
    }


def test_absent_office_is_unavailable_not_pass(tmp_path: Path) -> None:
    artifact, expectations = _known_good(tmp_path, "docx")
    report = qualify_artifact(
        format_id="docx",
        operation="docx.create",
        artifact=artifact,
        expectations=expectations,
        office_policy="auto",
        office_detector=lambda _application: {
            "available": False,
            "application": "word",
            "reason": "not-installed",
        },
    )

    assert report["portable"]["outcome"] == "pass"
    assert report["office"]["outcome"] == "unavailable"
    assert report["office_acceptance"] is False
    assert report["status"] == "unavailable"


def test_installed_office_cannot_silently_skip_failed_open(tmp_path: Path) -> None:
    artifact, expectations = _known_good(tmp_path, "xlsx")
    report = qualify_artifact(
        format_id="xlsx",
        operation="xlsx.create",
        artifact=artifact,
        expectations=expectations,
        office_policy="auto",
        office_detector=lambda _application: {
            "available": True,
            "application": "excel",
            "version": "test-version",
        },
        office_runner=lambda _application, _artifact, _timeout: {
            "outcome": "fail",
            "category": "open-rejected",
        },
    )

    assert report["office"]["outcome"] == "fail"
    assert report["status"] == "fail"
    assert report["office_acceptance"] is False
    assert report["office"]["evidence"]["application"] == "excel"
    assert report["office"]["evidence"]["version"] == "test-version"


def test_office_timeout_preserves_artifact_and_reports_cleanup(tmp_path: Path) -> None:
    artifact, expectations = _known_good(tmp_path, "pptx")
    before = _sha256(artifact)

    def timeout_runner(_application: str, _artifact: Path, _timeout: float) -> dict[str, Any]:
        return {
            "outcome": "fail",
            "category": "timeout",
            "descendants_cleaned": True,
        }

    report = qualify_artifact(
        format_id="pptx",
        operation="pptx.create",
        artifact=artifact,
        expectations=expectations,
        office_policy="auto",
        office_detector=lambda _application: {
            "available": True,
            "application": "powerpoint",
            "version": "test-version",
        },
        office_runner=timeout_runner,
        timeout_seconds=0.01,
    )

    assert report["office"]["outcome"] == "fail"
    assert report["office"]["evidence"]["category"] == "timeout"
    assert report["office"]["evidence"]["descendants_cleaned"] is True
    assert report["office"]["evidence"]["application"] == "powerpoint"
    assert report["office"]["evidence"]["version"] == "test-version"
    assert _sha256(artifact) == before


@pytest.mark.skipif(os.name != "nt", reason="mocked COM process timeout is Windows-only")
def test_office_second_timeout_returns_identified_schema_valid_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    artifact, expectations = _known_good(tmp_path, "docx")
    identity = b'{"event":"identity","application":"word","version":"16.0"}\n'
    child_json = (
        b'{"raw_output":"TOPSECRET","document_text":"CONFIDENTIAL"}\n'
    )
    timeout_output = identity + child_json

    class UnresponsiveProcess:
        pid = 4242
        returncode = None

        def __init__(self) -> None:
            self.calls = 0

        def communicate(self, timeout: float) -> tuple[bytes, bytes]:
            self.calls += 1
            raise subprocess.TimeoutExpired(
                cmd="powershell.exe",
                timeout=timeout,
                output=timeout_output,
                stderr=b"x" * 5000,
            )

    process = UnresponsiveProcess()
    monkeypatch.setattr(office_module.subprocess, "Popen", lambda *args, **kwargs: process)
    monkeypatch.setattr(
        office_module,
        "detect_office",
        lambda _application: {
            "available": True,
            "application": "word",
            "version": "16.0",
        },
    )
    monkeypatch.setattr(
        office_module,
        "_terminate_tree",
        lambda _pid: {
            "descendants_cleaned": False,
            "cleanup_category": "taskkill-incomplete",
        },
    )

    report = qualify_artifact(
        format_id="docx",
        operation="docx.create",
        artifact=artifact,
        expectations=expectations,
        office_policy="auto",
        office_detector=lambda _application: {
            "available": True,
            "application": "word",
            "version": "16.0",
        },
        office_runner=office_module.open_with_office,
        timeout_seconds=0.01,
    )

    validate_consumer_report(report)
    evidence = report["office"]["evidence"]
    assert report["status"] == "fail"
    assert report["office"]["outcome"] == "fail"
    assert evidence["application"] == "word"
    assert evidence["version"] == "16.0"
    assert evidence["category"] == "timeout"
    assert evidence["descendants_cleaned"] is False
    assert evidence["cleanup_category"] == "taskkill-incomplete"
    assert evidence["post_termination_drain"] == "timeout"
    assert evidence["stdout_bytes"] == len(timeout_output)
    assert evidence["stdout_truncated"] is False
    assert evidence["stderr_bytes"] == 4096
    assert evidence["stderr_truncated"] is True
    assert "stdout" not in evidence
    assert "stderr" not in evidence
    serialized_office = json.dumps(report["office"], sort_keys=True)
    assert "raw_output" not in serialized_office
    assert "TOPSECRET" not in serialized_office
    assert "document_text" not in serialized_office
    assert "CONFIDENTIAL" not in serialized_office
    assert process.calls == 2


def test_nonzero_office_child_json_retains_identity() -> None:
    payload = (
        b'{"event":"identity","application":"word","version":"16.0"}\r\n'
        b'{"outcome":"fail","application":"word","version":"16.0",'
        b'"category":"open-rejected","exception":"COMException"}\r\n'
    )

    evidence = _parse_probe_output(payload, require_result=True)

    assert evidence["outcome"] == "fail"
    assert evidence["application"] == "word"
    assert evidence["version"] == "16.0"
    assert evidence["category"] == "open-rejected"


def test_consumer_module_does_not_import_core_format_or_validation_modules(
    project_root: Path,
) -> None:
    forbidden = (
        "document_skills_core.formats",
        "document_skills_core.core.validation",
        "provider.reopen",
    )
    root = project_root / "consumer_validation"
    sources = [path.read_text(encoding="utf-8") for path in root.glob("*.py")]
    assert sources
    assert not any(token in source for source in sources for token in forbidden)


def test_docx_unknown_expectation_key_fails_closed(tmp_path: Path) -> None:
    artifact, expectations = _known_good(tmp_path, "docx")
    report = qualify_artifact(
        format_id="docx",
        operation="docx.create",
        artifact=artifact,
        expectations={**expectations, "unknown_semantic": True},
        office_policy="off",
    )

    validate_consumer_report(report)
    assertion = _assertions(report)["docx.expectations-schema"]
    assert report["status"] == "fail"
    assert assertion["outcome"] == "fail"
    assert assertion["evidence"] == {
        "category": "unknown-expectation-keys",
        "unknown": ["unknown_semantic"],
    }


def test_docx_missing_requested_image_fails_with_other_semantics_intact(
    tmp_path: Path,
) -> None:
    artifact, expectations = _known_good(tmp_path, "docx")

    def remove_image(root: ElementTree.Element) -> None:
        body = root.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
        assert body is not None
        for paragraph in list(body):
            if paragraph.find(
                ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"
            ) is not None:
                body.remove(paragraph)

    mutated = _mutate_xml_part(
        tmp_path,
        artifact,
        "word/document.xml",
        remove_image,
        "without-image",
    )
    report = qualify_artifact(
        format_id="docx",
        operation="docx.create",
        artifact=mutated,
        expectations=expectations,
        office_policy="off",
    )

    validate_consumer_report(report)
    assertions = _assertions(report)
    assert report["status"] == "fail"
    assert assertions["ooxml.relationship-graph"]["outcome"] == "pass"
    assert assertions["docx.requested-text"]["outcome"] == "pass"
    assert assertions["docx.tables"]["outcome"] == "pass"
    assert assertions["docx.paragraph-styles"]["outcome"] == "pass"
    assert assertions["docx.images"]["outcome"] == "fail"


def test_docx_replaced_requested_registered_style_fails_with_content_intact(
    tmp_path: Path,
) -> None:
    artifact, expectations = _known_good(tmp_path, "docx")
    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def replace_heading_style(root: ElementTree.Element) -> None:
        style = root.find(f".//{{{word_namespace}}}pStyle")
        assert style is not None
        style.set(f"{{{word_namespace}}}val", "Heading2")

    mutated = _mutate_xml_part(
        tmp_path,
        artifact,
        "word/document.xml",
        replace_heading_style,
        "wrong-style",
    )
    report = qualify_artifact(
        format_id="docx",
        operation="docx.create",
        artifact=mutated,
        expectations=expectations,
        office_policy="off",
    )

    validate_consumer_report(report)
    assertions = _assertions(report)
    assert report["status"] == "fail"
    assert assertions["ooxml.relationship-graph"]["outcome"] == "pass"
    assert assertions["docx.requested-text"]["outcome"] == "pass"
    assert assertions["docx.tables"]["outcome"] == "pass"
    assert assertions["docx.images"]["outcome"] == "pass"
    assert assertions["docx.paragraph-styles"]["outcome"] == "fail"


def test_consumer_cli_emits_schema_valid_machine_report(
    project_root: Path,
    tmp_path: Path,
) -> None:
    artifact, expectations = _known_good(tmp_path, "docx")
    expectation_path = tmp_path / "expectations.json"
    expectation_path.write_text(
        json.dumps(expectations, ensure_ascii=False),
        encoding="utf-8",
        newline="\n",
    )

    process = subprocess.run(
        [
            sys.executable,
            "-m",
            "consumer_validation",
            "--format",
            "docx",
            "--operation",
            "docx.create",
            "--artifact",
            str(artifact),
            "--expectations",
            str(expectation_path),
            "--office-policy",
            "off",
        ],
        cwd=project_root,
        check=False,
        capture_output=True,
        text=False,
        shell=False,
        timeout=30,
    )

    assert process.returncode == 0
    assert process.stderr == b""
    report = json.loads(process.stdout.decode("utf-8", errors="strict"))
    validate_consumer_report(report)
    assert report["artifact"]["sha256"] == _sha256(artifact)


@pytest.mark.skipif(os.name != "nt", reason="Microsoft Office COM is Windows-only")
@pytest.mark.parametrize(
    ("format_id", "application"),
    [("docx", "word"), ("xlsx", "excel"), ("pptx", "powerpoint")],
)
def test_installed_office_real_safe_open_is_mandatory(
    tmp_path: Path,
    format_id: str,
    application: str,
) -> None:
    detection = detect_office(application)
    if not detection["available"]:
        pytest.skip(f"{application} is not installed")
    artifact, expectations = _known_good(tmp_path, format_id)
    before = _sha256(artifact)

    report = qualify_artifact(
        format_id=format_id,
        operation=f"{format_id}.create",
        artifact=artifact,
        expectations=expectations,
        office_policy="auto",
        # Cold Word DCOM activation plus bounded teardown measured 62.794s on
        # the loaded Windows gate host; mocked 0.01s tests keep timeout failure bounded.
        timeout_seconds=90,
    )

    assert report["office"]["outcome"] == "pass", report["office"]
    assert report["office"]["evidence"]["application"] == application
    assert report["office"]["evidence"]["version"]
    assert report["office"]["evidence"]["saved"] is False
    assert report["office_acceptance"] is True
    assert _sha256(artifact) == before


@pytest.mark.skipif(os.name != "nt", reason="bounded process-tree cleanup is Windows-only")
def test_real_consumer_timeout_kills_descendant_and_preserves_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    artifact, _expectations = _known_good(tmp_path, "docx")
    child_pid_path = tmp_path / "child.pid"
    before = _sha256(artifact)
    monkeypatch.setenv("DS_CHILD_PID", str(child_pid_path))
    monkeypatch.setitem(
        office_module._POWERSHELL,
        "word",
        r"""
$child = Start-Process powershell.exe -ArgumentList @('-NoProfile', '-NonInteractive', '-Command', 'Start-Sleep -Seconds 30') -WindowStyle Hidden -PassThru
[IO.File]::WriteAllText($env:DS_CHILD_PID, [string]$child.Id)
Start-Sleep -Seconds 30
""",
    )

    started = time.monotonic()
    result = open_with_office("word", artifact, 4.0)
    elapsed = time.monotonic() - started

    assert result["outcome"] == "fail"
    assert result["category"] == "timeout"
    assert result["descendants_cleaned"] is True
    assert result["application"] == "word"
    assert result["version"]
    assert elapsed < 10
    assert _sha256(artifact) == before
    child_pid = int(child_pid_path.read_text(encoding="utf-8"))
    probe = subprocess.run(
        [
            "powershell.exe",
            "-NoProfile",
            "-NonInteractive",
            "-Command",
            f"if (Get-Process -Id {child_pid} -ErrorAction SilentlyContinue) {{ exit 1 }}",
        ],
        stdin=subprocess.DEVNULL,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
        shell=False,
        timeout=10,
    )
    assert probe.returncode == 0


@pytest.mark.parametrize("format_id", ["xlsx", "pptx"])
def test_independently_valid_office_fixture_passes_runtime_package_gate(
    tmp_path: Path,
    format_id: str,
) -> None:
    artifact, _expectations = _known_good(tmp_path, format_id)
    if format_id == "xlsx":
        from document_skills_core.formats.xlsx.validation import validate_created

        model = {
            "metadata": {"title": "", "creator": "", "subject": ""},
            "sheets": [
                {
                    "name": "Data",
                    "rows": [
                        {
                            "cells": [
                                {"ref": "A1", "value": "Independent XLSX", "type": "s"},
                                {"ref": "B1", "value": "10", "type": "n"},
                            ]
                        }
                    ],
                    "number_formats": [],
                }
            ],
            "defined_names": [],
            "tables": [],
            "chart_reference": None,
            "page_setup": None,
        }
        report = validate_created(artifact, model)
    else:
        from document_skills_core.formats.pptx.validation import validate_created

        model = {
            "metadata": {"title": "", "creator": "", "subject": ""},
            "slide_size": {"cx": "9144000", "cy": "6858000", "type": "screen4x3"},
            "slides": [
                {
                    "layout": "content",
                    "title": "Independent PPTX",
                    "shapes": [],
                    "table": None,
                    "chart_reference": None,
                    "image_reference": None,
                    "notes": None,
                }
            ],
        }
        report = validate_created(artifact, model)

    gate = next(
        item
        for item in report["gates"]
        if item["id"] == "operation.consumer-package-conformance"
    )
    assert gate["outcome"] == "pass"


def _known_good(tmp_path: Path, format_id: str) -> tuple[Path, dict[str, Any]]:
    if format_id == "docx":
        from docx import Document

        artifact = tmp_path / "independent.docx"
        document = Document()
        document.add_heading("Independent DOCX", level=1)
        document.add_paragraph("Portable consumer text")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        image_path = tmp_path / "independent.png"
        image_path.write_bytes(_PNG)
        inline_shape = document.add_picture(str(image_path))
        inline_shape._inline.docPr.set("descr", "Independent pixel")
        document.save(artifact)
        return artifact, {
            "text": ["Independent DOCX"],
            "tables": 1,
            "images": [
                {
                    "alt_text": "Independent pixel",
                    "sha256": hashlib.sha256(_PNG).hexdigest(),
                }
            ],
            "paragraph_styles": [
                {"text": "Independent DOCX", "style": "Heading1"}
            ],
        }
    if format_id == "xlsx":
        from openpyxl import Workbook

        artifact = tmp_path / "independent.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet["A1"] = "Independent XLSX"
        sheet["B1"] = 10
        workbook.properties.title = ""
        workbook.properties.creator = ""
        workbook.properties.subject = ""
        workbook.save(artifact)
        return artifact, {
            "sheets": ["Data"],
            "cells": {"Data!A1": "Independent XLSX"},
        }
    if format_id == "pptx":
        from pptx import Presentation

        artifact = tmp_path / "independent.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Independent PPTX"
        slide.placeholders[1].text = "Portable consumer text"
        presentation.save(artifact)
        return artifact, {"slide_count": 1, "text": ["Independent PPTX"]}
    artifact = tmp_path / "independent.pdf"
    reference = tmp_path / "independent.expected-render.pdf"
    _write_pdf(artifact, "Independent PDF", include_image=True)
    _write_pdf(reference, "Independent PDF", include_image=True)
    return artifact, {
        "text": ["Independent PDF"],
        "real_images": 1,
        "render": {
            "reference": str(reference),
            "reference_kind": "expected",
            "mode": "match",
        },
    }


def _write_pdf(path: Path, text: str, *, include_image: bool = False) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    if include_image:
        page.insert_image(fitz.Rect(72, 100, 136, 164), stream=_PNG)
    document.save(path)
    document.close()


def _mutate_relationships(
    tmp_path: Path,
    artifact: Path,
    relationship_part: str,
    mutation: Any,
) -> Path:
    mutated = tmp_path / f"{artifact.stem}-{hashlib.sha256(relationship_part.encode()).hexdigest()[:8]}-mutated{artifact.suffix}"
    with zipfile.ZipFile(artifact) as source, zipfile.ZipFile(
        mutated,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == relationship_part:
                root = ElementTree.fromstring(payload)
                mutation(root)
                payload = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, payload)
    return mutated


def _mutate_xml_part(
    tmp_path: Path,
    artifact: Path,
    part: str,
    mutation: Any,
    suffix: str,
) -> Path:
    mutated = tmp_path / f"{artifact.stem}-{suffix}{artifact.suffix}"
    with zipfile.ZipFile(artifact) as source, zipfile.ZipFile(
        mutated,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == part:
                root = ElementTree.fromstring(payload)
                mutation(root)
                payload = ElementTree.tostring(
                    root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            target.writestr(info, payload)
    return mutated


def _assertions(report: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {item["id"]: item for item in report["portable"]["assertions"]}


def _corrupt(
    tmp_path: Path,
    artifact: Path,
    format_id: str,
) -> Path:
    if format_id == "pdf":
        corrupted = tmp_path / "corrupted.pdf"
        payload = artifact.read_bytes()
        corrupted.write_bytes(payload[: max(16, len(payload) // 2)])
        return corrupted
    removed = {
        "docx": "word/document.xml",
        "xlsx": "xl/workbook.xml",
        "pptx": "ppt/theme/theme1.xml",
    }[format_id]
    corrupted = tmp_path / f"corrupted.{format_id}"
    with zipfile.ZipFile(artifact) as source, zipfile.ZipFile(
        corrupted,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        for info in source.infolist():
            if info.filename != removed:
                target.writestr(info, source.read(info.filename))
    return corrupted


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
