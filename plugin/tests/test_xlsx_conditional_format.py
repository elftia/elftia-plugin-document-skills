"""Native XLSX conditional-format creation, editing, and readback tests."""

from hashlib import sha256
from pathlib import Path

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.service import XlsxService


def _rules() -> list[dict]:
    return [
        {
            "ref": "B2:B6",
            "type": "cellIs",
            "operator": "greaterThan",
            "formulas": ["10"],
            "style": {"fill": {"pattern": "solid", "color": "FFF8696B"}},
            "stop_if_true": True,
        },
        {
            "ref": "C2:C6",
            "type": "expression",
            "formulas": ["MOD(C2,2)=0"],
            "style": {"font": {"bold": True, "color": "FF0000FF"}},
        },
        {
            "ref": "D2:D6",
            "type": "colorScale",
            "thresholds": [
                {"type": "min"},
                {"type": "percentile", "value": "50"},
                {"type": "max"},
            ],
            "colors": ["FFF8696B", "FFFFEB84", "FF63BE7B"],
        },
        {
            "ref": "E2:E6",
            "type": "dataBar",
            "thresholds": [{"type": "min"}, {"type": "max"}],
            "color": "FF638EC6",
            "show_value": False,
        },
        {
            "ref": "F2:F6",
            "type": "iconSet",
            "thresholds": [
                {"type": "percent", "value": "0"},
                {"type": "percent", "value": "33"},
                {"type": "percent", "value": "67"},
            ],
            "icon_set": "3TrafficLights1",
            "show_value": True,
            "reverse": True,
        },
    ]


def _workbook() -> dict:
    return {
        "metadata": {"title": "Conditional formats", "creator": "Test", "subject": ""},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": f"{column}{row}", "value": str(row), "type": "n"}
                            for column in "BCDEF"
                            for row in range(2, 7)
                        ]
                    }
                ],
                "columns": [],
                "number_formats": [],
                "data_validations": [],
                "conditional_formats": _rules(),
            }
        ],
        "defined_names": [],
        "tables": [],
        "chart_reference": None,
        "page_setup": None,
    }


def _request(operation: str, **paths: Path) -> dict:
    return {
        "schema_version": "1.0",
        "operation": operation,
        **{name: str(path) for name, path in paths.items()},
    }


def test_conditional_format_create_supports_five_rule_types_and_readback(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "conditional-formats.xlsx"
    request = _request("xlsx.create", output=output)
    request["arguments"] = {"workbook": _workbook()}
    result = XlsxService(project_root).execute("xlsx.create", request)

    assert result["status"] == "success"
    reopened = load_workbook(output)
    worksheet = reopened["Data"]
    rules = [
        rule
        for conditional_format in worksheet.conditional_formatting
        for rule in worksheet.conditional_formatting[conditional_format]
    ]
    assert [rule.type for rule in rules] == [
        "cellIs",
        "expression",
        "colorScale",
        "dataBar",
        "iconSet",
    ]
    assert [rule.priority for rule in rules] == [1, 2, 3, 4, 5]
    assert len(reopened._differential_styles.styles) == 2
    assert rules[0].dxfId == 0
    assert rules[1].dxfId == 1
    assert rules[2].colorScale.color[2].rgb == "FF63BE7B"
    assert rules[3].dataBar.showValue is False
    assert rules[4].iconSet.iconSet == "3TrafficLights1"

    read_request = _request("xlsx.read", input=output)
    read_request["arguments"] = {}
    read_result = XlsxService(project_root).execute("xlsx.read", read_request)
    projected = read_result["diagnostics"]["operation_result"]["conditional_formats"]
    assert len(projected) == 5
    assert all(rule["dxf_valid"] for rule in projected)
    assert projected[0]["style"] == {
        "fill": {"pattern": "solid", "color": "FFF8696B"}
    }
    assert projected[2]["thresholds"][1] == {
        "type": "percentile",
        "value": "50",
        "gte": True,
    }
    assert projected[4]["reverse"] is True


def test_conditional_format_create_is_deterministic(
    project_root: Path,
    tmp_path: Path,
) -> None:
    outputs = [tmp_path / "first.xlsx", tmp_path / "second.xlsx"]
    for output in outputs:
        request = _request("xlsx.create", output=output)
        request["arguments"] = {"workbook": _workbook()}
        assert XlsxService(project_root).execute("xlsx.create", request)["status"] == "success"
    assert outputs[0].read_bytes() == outputs[1].read_bytes()


@pytest.mark.parametrize(
    "rule",
    [
        {
            "ref": "A1:A5",
            "type": "cellIs",
            "operator": "between",
            "formulas": ["1"],
            "style": {"font": {"bold": True}},
        },
        {"ref": "A1:A5", "type": "expression", "formulas": ["A1>0"]},
        {
            "ref": "A1:A5",
            "type": "colorScale",
            "thresholds": [{"type": "min"}, {"type": "max"}],
            "colors": ["FFFF0000"],
        },
        {
            "ref": "A1:A5",
            "type": "dataBar",
            "thresholds": [{"type": "min"}],
            "color": "FF0000FF",
        },
        {
            "ref": "A1:A5",
            "type": "iconSet",
            "thresholds": [{"type": "percent", "value": "0"}],
            "icon_set": "3Arrows",
        },
        {
            "ref": "A1:A5",
            "type": "expression",
            "formulas": ["[Book.xlsx]Data!A1>0"],
            "style": {"font": {"bold": True}},
        },
    ],
)
def test_conditional_format_contract_rejects_invalid_boundaries(rule: dict) -> None:
    workbook = _workbook()
    workbook["sheets"][0]["conditional_formats"] = [rule]
    with pytest.raises(DocumentSkillsError) as exc:
        parse_xlsx_request(
            {
                "operation": "xlsx.create",
                "output": "conditional-format.xlsx",
                "arguments": {"workbook": workbook},
            }
        )
    assert exc.value.code.value == "DS_REQUEST_INVALID"


def _source(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in range(1, 8):
        sheet.cell(row=row, column=1, value=row)
        sheet.cell(row=row, column=2, value=row)
    red = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    sheet.conditional_formatting.add(
        "A1:A7",
        CellIsRule(operator="greaterThan", formula=["3"], fill=red),
    )
    sheet.conditional_formatting.add(
        "B1:B7",
        FormulaRule(formula=["MOD(B1,2)=0"], fill=red),
    )
    workbook.save(path)


def test_conditional_format_edit_add_update_delete_preserves_existing_dxfs(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "edited.xlsx"
    _source(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    edits = [
        {
            "sheet": "Data",
            "type": "conditional_format_update",
            "ref": "A1:A7",
            "priority": 1,
            "rule": {
                "ref": "A1:A7",
                "type": "cellIs",
                "operator": "lessThan",
                "formulas": ["6"],
                "style": {"fill": {"pattern": "solid", "color": "FF00FF00"}},
            },
        },
        {
            "sheet": "Data",
            "type": "conditional_format_delete",
            "ref": "B1:B7",
            "priority": 2,
        },
        {
            "sheet": "Data",
            "type": "conditional_format_add",
            "rule": {
                "ref": "C1:C7",
                "type": "dataBar",
                "thresholds": [{"type": "min"}, {"type": "max"}],
                "color": "FF638EC6",
            },
        },
    ]
    request = _request("xlsx.edit", input=source, output=output)
    request["arguments"] = {"edits": edits, "expected_edits": 3}
    result = XlsxService(project_root).execute("xlsx.edit", request)

    assert result["status"] == "success"
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    reopened = load_workbook(output)
    rules = [
        rule
        for conditional_format in reopened["Data"].conditional_formatting
        for rule in reopened["Data"].conditional_formatting[conditional_format]
    ]
    assert {(rule.type, rule.priority) for rule in rules} == {
        ("cellIs", 1),
        ("dataBar", 2),
    }
    assert len(reopened._differential_styles.styles) >= 2


def test_conditional_format_edit_failure_preserves_existing_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "existing.xlsx"
    _source(source)
    output.write_bytes(b"existing-conditional-format-destination")
    request = _request("xlsx.edit", input=source, output=output)
    request["arguments"] = {
        "edits": [
            {
                "sheet": "Data",
                "type": "conditional_format_delete",
                "ref": "A1:A7",
                "priority": 99,
            }
        ]
    }
    result = XlsxService(project_root).execute("xlsx.edit", request)

    assert result["status"] == "invalid_request"
    assert output.read_bytes() == b"existing-conditional-format-destination"
