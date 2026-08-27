"""Deck-level PPTX slide-size editing and object-boundary validation."""

from hashlib import sha256
from pathlib import Path
from typing import Any

from pptx import Presentation
import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError, ErrorCode
from document_skills_core.formats.pptx.contracts import parse_pptx_request
from document_skills_core.formats.pptx.create import create_pptx
from document_skills_core.formats.pptx.edit import edit_pptx
from document_skills_core.formats.pptx.package import OpcPackage
from document_skills_core.formats.pptx.service import PptxService

_WIDE_SIZE = {"cx": 12_192_000, "cy": 6_858_000, "type": "screen16x9"}


def _deck() -> dict[str, Any]:
    return {
        "metadata": {"title": "Slide size", "creator": "Test", "subject": ""},
        "slide_size": {"cx": "9144000", "cy": "6858000", "type": "screen4x3"},
        "slides": [{
            "layout": "content",
            "title": "Bounded content",
            "shapes": [{"text": "Inside the original canvas", "runs": []}],
            "table": None,
            "chart_reference": None,
            "image_reference": None,
            "notes": None,
        }],
    }


def _created(tmp_path: Path) -> Path:
    source = tmp_path / "slide-size-source.pptx"
    create_pptx(source, _deck())
    return source


def _request(source: Path, output: Path, edits: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "schema_version": "1.0",
        "operation": "pptx.edit",
        "input": str(source),
        "output": str(output),
        "arguments": {"edits": edits},
    }


def test_contract_accepts_only_explicit_deck_level_slide_size(tmp_path: Path) -> None:
    parsed = parse_pptx_request(_request(
        tmp_path / "source.pptx",
        tmp_path / "output.pptx",
        [{"type": "slide_size", "size": _WIDE_SIZE}],
    ))
    edit = parsed.arguments["edits"][0]
    assert edit == {"type": "slide_size", "size": _WIDE_SIZE}
    assert "slide" not in edit

    with pytest.raises(DocumentSkillsError) as captured:
        parse_pptx_request(_request(
            tmp_path / "source.pptx",
            tmp_path / "output.pptx",
            [{"type": "slide_size", "slide": 1, "size": _WIDE_SIZE}],
        ))
    assert captured.value.code == ErrorCode.REQUEST_INVALID


def test_contract_rejects_ambiguous_or_unbounded_slide_sizes(tmp_path: Path) -> None:
    for edits in (
        [
            {"type": "slide_size", "size": _WIDE_SIZE},
            {"type": "slide_size", "size": _WIDE_SIZE},
        ],
        [{
            "type": "slide_size",
            "size": {"cx": 914_399, "cy": 6_858_000, "type": "custom"},
        }],
        [{
            "type": "slide_size",
            "size": {"cx": 12_192_000, "cy": 6_858_000, "type": "invented"},
        }],
    ):
        with pytest.raises(DocumentSkillsError) as captured:
            parse_pptx_request(_request(
                tmp_path / "source.pptx",
                tmp_path / "output.pptx",
                edits,
            ))
        assert captured.value.code == ErrorCode.REQUEST_INVALID


def test_slide_size_edit_changes_only_deck_size_and_reports_bounds(tmp_path: Path) -> None:
    source = _created(tmp_path)
    source_hash = sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "wide.pptx"
    result, manifest = edit_pptx(
        source,
        output,
        {"edits": [{"type": "slide_size", "size": _WIDE_SIZE}]},
    )

    assert sha256(source.read_bytes()).hexdigest() == source_hash
    assert manifest.changed == ("ppt/presentation.xml",)
    assert not manifest.added
    assert not manifest.removed
    evidence = result["slide_size"]
    assert evidence["scope"] == "deck"
    assert evidence["before"] == {
        "cx": 9_144_000,
        "cy": 6_858_000,
        "type": "screen4x3",
    }
    assert evidence["after"] == _WIDE_SIZE
    assert evidence["objects_checked"] >= 2
    assert evidence["object_bounds_valid"] is True
    presentation = Presentation(output)
    assert presentation.slide_width == _WIDE_SIZE["cx"]
    assert presentation.slide_height == _WIDE_SIZE["cy"]


def test_slide_size_edit_rejects_out_of_bounds_objects_before_emit(tmp_path: Path) -> None:
    source = _created(tmp_path)
    output = tmp_path / "must-not-exist.pptx"
    with pytest.raises(DocumentSkillsError) as captured:
        edit_pptx(
            source,
            output,
            {"edits": [{
                "type": "slide_size",
                "size": {"cx": 7_000_000, "cy": 6_858_000, "type": "custom"},
            }]},
        )
    assert captured.value.code == ErrorCode.VALIDATION_FAILED
    assert captured.value.details["failures"][0]["code"] == "out-of-bounds"
    assert not output.exists()


def test_slide_size_resolves_inherited_placeholder_geometry(tmp_path: Path) -> None:
    source = tmp_path / "placeholder-source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Inherited title"
    slide.placeholders[1].text = "Inherited subtitle"
    presentation.save(source)

    output = tmp_path / "placeholder-wide.pptx"
    result, _manifest = edit_pptx(
        source,
        output,
        {"edits": [{"type": "slide_size", "size": _WIDE_SIZE}]},
    )

    assert output.is_file()
    assert result["slide_size"]["objects_checked"] == 2
    assert result["slide_size"]["inherited_bounds"] == 2


def test_service_validates_and_promotes_slide_size_edit(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _created(tmp_path)
    output = tmp_path / "service-wide.pptx"
    result = PptxService(project_root).execute(
        "pptx.edit",
        _request(
            source,
            output,
            [{"type": "slide_size", "size": _WIDE_SIZE}],
        ),
    )

    assert result["status"] == "success", result
    assert output.is_file()
    operation = result["diagnostics"]["operation_result"]
    assert operation["edit_counts"] == {"slide_size": 1}
    assert operation["slide_size"]["object_bounds_valid"] is True
    package = OpcPackage.open(output)
    size = package.xml("ppt/presentation.xml").find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}sldSz"
    )
    assert size is not None
    assert size.attrib == {
        "cx": str(_WIDE_SIZE["cx"]),
        "cy": str(_WIDE_SIZE["cy"]),
        "type": _WIDE_SIZE["type"],
    }
