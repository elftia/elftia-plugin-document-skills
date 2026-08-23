"""Native XLSX chart add/update/delete transaction tests."""

from hashlib import sha256
from pathlib import Path
import zipfile

from document_skills_core.formats.xlsx.service import XlsxService

from test_xlsx_chart_create import _chart, _workbook


def _source(project_root: Path, path: Path) -> None:
    workbook = _workbook()
    workbook["charts"] = [workbook["charts"][0]]
    result = XlsxService(project_root).execute(
        "xlsx.create",
        {
            "operation": "xlsx.create",
            "output": str(path),
            "arguments": {"workbook": workbook},
        },
    )
    assert result["status"] == "success"


def _edits() -> list[dict]:
    updated = _chart("UpdatedChart", "line", "F3:M17", value_column="C")
    updated["title"] = "Updated ranges"
    added = _chart("AddedPie", "pie", "N3:U17")
    return [
        {
            "sheet": "Data",
            "type": "chart_update",
            "name": "ColumnChart",
            "chart": updated,
        },
        {"sheet": "Data", "type": "chart_add", "chart": added},
    ]


def test_chart_add_and_update_reopen_with_safe_series_replacement(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "edited.xlsx"
    _source(project_root, source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    edits = _edits()
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {"edits": edits, "expected_edits": 2},
        },
    )

    assert result["status"] == "success", result
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    reopened = load_workbook(output, data_only=False)
    assert [type(chart).__name__ for chart in reopened["Data"]._charts] == [
        "LineChart",
        "PieChart",
    ]
    read_result = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(output), "arguments": {}},
    )
    charts = read_result["diagnostics"]["operation_result"]["charts"]
    assert [chart["name"] for chart in charts] == ["UpdatedChart", "AddedPie"]
    assert charts[0]["type"] == "line"
    assert charts[0]["anchor"] == "F3:M17"
    assert charts[0]["series"][0]["values"] == "'Data'!$C$2:$C$5"


def test_chart_delete_removes_declared_chart_and_empty_drawing_parts(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    populated = tmp_path / "populated.xlsx"
    output = tmp_path / "deleted.xlsx"
    _source(project_root, source)
    edits = _edits()
    first = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(populated),
            "arguments": {"edits": edits},
        },
    )
    assert first["status"] == "success"
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(populated),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "chart_delete", "name": "UpdatedChart"},
                    {"sheet": "Data", "type": "chart_delete", "name": "AddedPie"},
                ],
                "expected_edits": 2,
            },
        },
    )

    assert result["status"] == "success", result
    assert not load_workbook(output)["Data"]._charts
    with zipfile.ZipFile(output) as archive:
        names = set(archive.namelist())
    assert not any(name.startswith("xl/charts/") for name in names)
    assert not any(name.startswith("xl/drawings/") for name in names)
    removed = result["diagnostics"]["operation_result"]["preservation"]["removed_parts"]
    assert set(removed) >= {
        "xl/charts/chart1.xml",
        "xl/charts/chart2.xml",
        "xl/drawings/drawing1.xml",
        "xl/drawings/_rels/drawing1.xml.rels",
    }


def test_chart_edit_is_deterministic(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    outputs = [tmp_path / "first.xlsx", tmp_path / "second.xlsx"]
    _source(project_root, source)
    for output in outputs:
        result = XlsxService(project_root).execute(
            "xlsx.edit",
            {
                "operation": "xlsx.edit",
                "input": str(source),
                "output": str(output),
                "arguments": {"edits": _edits()},
            },
        )
        assert result["status"] == "success"
    assert outputs[0].read_bytes() == outputs[1].read_bytes()


def test_chart_edit_failure_preserves_existing_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "existing.xlsx"
    _source(project_root, source)
    output.write_bytes(b"existing-chart-destination")
    invalid = _chart("InvalidChart", "line", "F3:M17")
    invalid["series"][0]["values"] = "'Missing'!$B$2:$B$5"
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "chart_add", "chart": invalid}
                ]
            },
        },
    )

    assert result["status"] == "invalid_request"
    assert output.read_bytes() == b"existing-chart-destination"
