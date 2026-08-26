"""Native x14 sparkline create/edit/readback/reference-maintenance coverage."""

from __future__ import annotations

from pathlib import Path
import warnings
import zipfile

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.service import XlsxService


def _sparkline(location: str, data: str, sparkline_type: str = "line") -> dict:
    return {
        "location": location,
        "data": data,
        "type": sparkline_type,
        "empty_cells": "connect",
        "markers": True,
        "high_point": True,
        "low_point": True,
        "first_point": True,
        "last_point": True,
        "negative_points": True,
        "manual_min": -10,
        "manual_max": 100,
        "color": "#4472C4",
        "negative_color": "#C00000",
    }


def _workbook() -> dict:
    return {
        "metadata": {},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Name", "type": "s"},
                            *[
                                {"ref": f"{column}1", "value": column, "type": "s"}
                                for column in "BCD"
                            ],
                            *[
                                {"ref": f"A{row}", "value": f"R{row}", "type": "s"}
                                for row in range(2, 5)
                            ],
                            *[
                                {
                                    "ref": f"{column}{row}",
                                    "value": str(row * index),
                                    "type": "n",
                                }
                                for row in range(2, 5)
                                for index, column in enumerate("BCD", start=1)
                            ],
                        ]
                    }
                ],
                "number_formats": [],
                "sparklines": [
                    _sparkline("E2", "Data!B2:D2"),
                    _sparkline("E3", "Data!B3:D3", "column"),
                    _sparkline("E4", "Data!B4:D4", "win_loss"),
                ],
            }
        ],
        "defined_names": [],
        "tables": [],
    }


def _create(project_root: Path, output: Path) -> dict:
    return XlsxService(project_root).execute(
        "xlsx.create",
        {
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": _workbook()},
        },
    )


def test_sparkline_create_reopens_x14_extension_and_projects_details(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "sparklines.xlsx"
    result = _create(project_root, output)

    assert result["status"] == "success", result
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        reopened = load_workbook(output, data_only=False)
    assert reopened.sheetnames == ["Data"]
    reopened.close()
    with zipfile.ZipFile(output) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "x14:sparklineGroups" in worksheet
    assert worksheet.count("x14:sparklineGroup") >= 3
    read = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(output), "arguments": {}},
    )
    sparklines = read["diagnostics"]["operation_result"]["sparklines"]
    assert [item["type"] for item in sparklines] == ["line", "column", "win_loss"]
    assert sparklines[0]["location"] == "E2"
    assert sparklines[0]["data"] == "'Data'!$B$2:$D$2"
    assert sparklines[0]["manual_min"] == -10
    assert sparklines[0]["color"] == "FF4472C4"


def test_sparkline_crud_and_structural_references(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "sparkline-source.xlsx"
    assert _create(project_root, source)["status"] == "success"
    crud = tmp_path / "sparkline-crud.xlsx"
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(crud),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "sparkline_update",
                        "ref": "E2",
                        "sparkline": _sparkline("F2", "Data!B2:D2", "column"),
                    },
                    {"sheet": "Data", "type": "sparkline_delete", "ref": "E3"},
                    {
                        "sheet": "Data",
                        "type": "sparkline_add",
                        "sparkline": _sparkline("F3", "Data!B3:D3"),
                    },
                ],
                "expected_edits": 3,
            },
        },
    )
    assert result["status"] == "success", result
    read = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(crud), "arguments": {}},
    )
    assert [(item["location"], item["type"]) for item in read["diagnostics"]["operation_result"]["sparklines"]] == [
        ("F2", "column"),
        ("E4", "win_loss"),
        ("F3", "line"),
    ]

    shifted = tmp_path / "sparkline-shifted.xlsx"
    structural = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(crud),
            "output": str(shifted),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "row_insert", "ref": "2", "count": 1}
                ]
            },
        },
    )
    assert structural["status"] == "success", structural
    read = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(shifted), "arguments": {}},
    )
    projected = read["diagnostics"]["operation_result"]["sparklines"]
    assert projected[0]["location"] == "F3"
    assert projected[0]["data"] == "'Data'!$B$3:$D$3"

    renamed = tmp_path / "sparkline-renamed.xlsx"
    rename_result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(shifted),
            "output": str(renamed),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "sheet_rename", "value": "Renamed"}
                ]
            },
        },
    )
    assert rename_result["status"] == "success", rename_result
    read = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(renamed), "arguments": {}},
    )
    projected = read["diagnostics"]["operation_result"]["sparklines"]
    assert projected[0]["sheet"] == "Renamed"
    assert projected[0]["data"] == "'Renamed'!$B$3:$D$3"


@pytest.mark.parametrize(
    "mutate",
    [
        lambda item: item.update({"data": "Data!B2:C3"}),
        lambda item: item.update({"data": "Missing!B2:D2"}),
        lambda item: item.update({"location": "XFE1"}),
        lambda item: item.update({"manual_min": 100, "manual_max": 10}),
    ],
)
def test_sparkline_contract_rejects_invalid_boundaries(mutate) -> None:
    workbook = _workbook()
    workbook["sheets"][0]["sparklines"] = [workbook["sheets"][0]["sparklines"][0]]
    mutate(workbook["sheets"][0]["sparklines"][0])
    with pytest.raises(DocumentSkillsError) as exc:
        parse_xlsx_request(
            {
                "operation": "xlsx.create",
                "output": "sparklines.xlsx",
                "arguments": {"workbook": workbook},
            }
        )
    assert exc.value.code.value == "DS_REQUEST_INVALID"
