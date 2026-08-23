"""Native XLSX table creation and readback tests."""

from pathlib import Path

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.service import XlsxService


def _workbook() -> dict:
    return {
        "metadata": {"title": "Tables", "creator": "Test", "subject": ""},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Name", "type": "s"},
                            {"ref": "B1", "value": "Amount", "type": "s"},
                            {"ref": "A2", "value": "Alpha", "type": "s"},
                            {"ref": "B2", "value": "10", "type": "n"},
                            {"ref": "A3", "value": "Beta", "type": "s"},
                            {"ref": "B3", "value": "20", "type": "n"},
                        ]
                    }
                ],
                "number_formats": [],
            }
        ],
        "defined_names": [],
        "tables": [
            {
                "name": "DataTable",
                "ref": "A1:B3",
                "sheet": "Data",
                "style": "TableStyleMedium2",
            }
        ],
        "chart_reference": None,
        "page_setup": None,
    }


def test_native_table_create_reopens_and_reads_structured_details(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "table.xlsx"
    service = XlsxService(project_root)
    result = service.execute(
        "xlsx.create",
        {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": _workbook()},
        },
    )

    assert result["status"] == "success"
    reopened = load_workbook(output, data_only=False)
    table = reopened["Data"].tables["DataTable"]
    assert table.ref == "A1:B3"
    assert table.tableStyleInfo.name == "TableStyleMedium2"
    assert [column.name for column in table.tableColumns] == ["Name", "Amount"]

    read_result = service.execute(
        "xlsx.read",
        {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(output),
            "arguments": {},
        },
    )
    projected = read_result["diagnostics"]["operation_result"]["tables"][0]
    assert projected["name"] == "DataTable"
    assert projected["sheet"] == "Data"
    assert projected["ref"] == "A1:B3"
    assert projected["columns"] == ["Name", "Amount"]
    assert projected["auto_filter_ref"] == "A1:B3"
    assert projected["style_options"]["show_row_stripes"] is True


def test_native_table_create_is_deterministic(
    project_root: Path,
    tmp_path: Path,
) -> None:
    outputs = [tmp_path / "first.xlsx", tmp_path / "second.xlsx"]
    for output in outputs:
        result = XlsxService(project_root).execute(
            "xlsx.create",
            {
                "schema_version": "1.0",
                "operation": "xlsx.create",
                "output": str(output),
                "arguments": {"workbook": _workbook()},
            },
        )
        assert result["status"] == "success"
    assert outputs[0].read_bytes() == outputs[1].read_bytes()


@pytest.mark.parametrize("failure", ["duplicate-header", "overlap", "unknown-sheet"])
def test_native_table_contract_rejects_invalid_boundaries(failure: str) -> None:
    workbook = _workbook()
    if failure == "duplicate-header":
        workbook["sheets"][0]["rows"][0]["cells"][1]["value"] = "Name"
    elif failure == "overlap":
        workbook["tables"].append(
            {
                "name": "OtherTable",
                "ref": "B1:C3",
                "sheet": "Data",
                "style": "TableStyleLight1",
            }
        )
    else:
        workbook["tables"][0]["sheet"] = "Missing"

    with pytest.raises(DocumentSkillsError) as exc:
        parse_xlsx_request(
            {
                "schema_version": "1.0",
                "operation": "xlsx.create",
                "output": "table.xlsx",
                "arguments": {"workbook": workbook},
            }
        )
    assert exc.value.code.value == "DS_REQUEST_INVALID"
