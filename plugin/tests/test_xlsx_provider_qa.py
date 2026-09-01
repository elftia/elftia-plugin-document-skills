"""Provider-gated XLSX schema and PDF-render operation tests."""

import hashlib
import json
from pathlib import Path
from xml.etree.ElementTree import SubElement, fromstring, tostring
import zlib
import zipfile

import pytest

from document_skills_core.core.capabilities import (
    DetectionEvidence,
    ProviderCatalog,
)
from document_skills_core.core.capabilities.reports import build_capabilities
from document_skills_core.core.contracts.errors import DocumentSkillsError, ErrorCode
from document_skills_core.core.contracts.schemas import SchemaCatalog
from document_skills_core.core.process import ProcessResult
from document_skills_core.formats.pdf.create import create_pdf
from document_skills_core.formats.pdf.byte_preflight import decode_stream
from document_skills_core.formats.pdf.validation import reopen_pdf
from document_skills_core.formats.xlsx.create import create_xlsx
from document_skills_core.formats.xlsx.constants import (
    CONTENT_TYPES,
    MAX_PARTS,
    NS,
    PACKAGE_RELS,
    REL_TABLE,
    WORKBOOK_MAIN,
    WORKBOOK_RELS,
)
from document_skills_core.formats.xlsx.formula_security import (
    assert_provider_formula_safe,
)
from document_skills_core.formats.xlsx.render_operation import execute_render
from document_skills_core.providers.dotnet import build_dotnet_provider
from document_skills_core.providers.libreoffice import build_libreoffice_provider
from document_skills_core.public_cli.protocol import PublicCommand
from document_skills_core.public_cli.supervisor import PublicCommandSupervisor
from tests.support.xlsx_macro_fixture import create_package_fixture


class _CallableDetector:
    def __init__(self, path: str) -> None:
        self.path = path

    def detect(self) -> DetectionEvidence:
        return DetectionEvidence(True, version="test", path=self.path)


class _AbsentDetector:
    def detect(self) -> DetectionEvidence:
        return DetectionEvidence(False, reason="fixture provider unavailable")


class _DotnetRunner:
    def __init__(self, response: dict[str, object]) -> None:
        self.response = response
        self.calls: list[dict[str, object]] = []

    def set_executable(self, _executable: str) -> None:
        return None

    def run(self, subcommand: str, *, stdin_payload=None, **_kwargs) -> ProcessResult:
        input_path = Path(stdin_payload["input_path"])
        self.calls.append({
            "subcommand": subcommand,
            "payload": stdin_payload,
            "input_sha256": hashlib.sha256(input_path.read_bytes()).hexdigest(),
        })
        return ProcessResult(0, json.dumps(self.response), "", 5)


class _LibreOfficeRunner:
    def __init__(self, pdf_bytes: bytes) -> None:
        self.pdf_bytes = pdf_bytes
        self.calls: list[dict[str, object]] = []

    def set_executable(self, _executable: str) -> None:
        return None

    def convert(
        self,
        input_path: Path,
        target_format: str,
        output_dir: Path,
        *,
        timeout_seconds=None,
    ) -> Path:
        self.calls.append({
            "input": input_path,
            "format": target_format,
            "timeout": timeout_seconds,
        })
        output = output_dir / f"{input_path.stem}.{target_format}"
        output.write_bytes(self.pdf_bytes)
        return output


@pytest.fixture
def qa_xlsx(tmp_path: Path) -> Path:
    destination = tmp_path / "qa-source.xlsx"
    create_xlsx(
        destination,
        {
            "metadata": {},
            "sheets": [{
                "name": "Data",
                "rows": [{
                    "hidden": True,
                    "cells": [
                        {"ref": "A1", "value": "Wide column", "type": "s"},
                        {
                            "ref": "B1",
                            "value": "Long unwrapped text " * 8,
                            "type": "s",
                        },
                    ],
                }],
                "columns": [{
                    "min": 1,
                    "max": 1,
                    "ref": "A",
                    "width": 90.0,
                    "hidden": False,
                    "style": None,
                }],
                "number_formats": [],
            }],
            "defined_names": [],
            "tables": [],
            "charts": [],
        },
    )
    return destination


@pytest.fixture
def qa_pdf(tmp_path: Path) -> Path:
    destination = tmp_path / "qa-render.pdf"
    create_pdf(
        destination,
        {
            "metadata": {"title": "Render", "author": "Test", "subject": ""},
            "page_size": "A4",
            "pages": [{
                "blocks": [{
                    "type": "paragraph",
                    "text": "Rendered spreadsheet page",
                    "style": None,
                    "table": None,
                    "image": None,
                    "shape": None,
                }],
                "metadata": None,
            }],
        },
    )
    return destination


def test_dotnet_xlsx_schema_success_is_public_and_source_preserving(
    project_root: Path,
    qa_xlsx: Path,
) -> None:
    source_hash = hashlib.sha256(qa_xlsx.read_bytes()).hexdigest()
    runner = _DotnetRunner({
        "valid": True,
        "errors": [],
        "truncated": False,
        "file_format": "Microsoft365",
    })
    definition, _provider = build_dotnet_provider(
        project_root,
        detector=_CallableDetector("/fake/dotnet"),
        runner=runner,
    )
    registry = ProviderCatalog()
    registry.register_provider(definition)

    result = registry.execute({
        "schema_version": "1.0",
        "operation": "xlsx.validate.schema",
        "input": str(qa_xlsx),
        "arguments": {"max_errors": 25},
    })

    SchemaCatalog(project_root).validate("operation-result", result)
    assert result["status"] == "success"
    assert result["provider_chain"] == ["dotnet-openxml"]
    gates = {gate["id"]: gate for gate in result["validation"]["gates"]}
    assert gates["schema.full"]["outcome"] == "pass"
    assert gates["schema.full"]["validator"] == "dotnet-openxml"
    assert gates["schema.full"]["evidence"]["file_format"] == "Microsoft365"
    assert runner.calls[0]["subcommand"] == "--xlsx-schema-validate"
    assert runner.calls[0]["payload"]["max_errors"] == 25
    assert Path(runner.calls[0]["payload"]["input_path"]) != qa_xlsx
    assert hashlib.sha256(qa_xlsx.read_bytes()).hexdigest() == source_hash


def test_dotnet_xlsx_schema_invalid_retains_bounded_error_report(
    project_root: Path,
    qa_xlsx: Path,
) -> None:
    runner = _DotnetRunner({
        "valid": False,
        "errors": [
            {
                "part": "WorkbookPart",
                "path": "/x:workbook[1]",
                "description": "Fixture schema error",
                "error_type": "Schema",
            },
            {
                "part": "WorksheetPart",
                "path": "/x:worksheet[1]/x:sheetData[1]",
                "description": "é" * 200,
                "error_type": "Schema",
            },
        ],
        "truncated": True,
        "file_format": "Microsoft365",
    })
    definition, _provider = build_dotnet_provider(
        project_root,
        detector=_CallableDetector("/fake/dotnet"),
        runner=runner,
    )
    registry = ProviderCatalog()
    registry.register_provider(definition)

    result = registry.execute({
        "schema_version": "1.0",
        "operation": "xlsx.validate.schema",
        "input": str(qa_xlsx),
        "arguments": {"max_errors": 2},
    })

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_VALIDATION_FAILED"
    schema = result["diagnostics"]["operation_result"]["schema"]
    assert schema["error_count"] == 2
    assert schema["errors"][0] == {
        "part": "WorkbookPart",
        "path": "/x:workbook[1]",
        "description": "Fixture schema error",
        "error_type": "Schema",
    }
    assert schema["errors"][1]["part"] == "WorksheetPart"
    assert schema["errors"][1]["path"] == "/x:worksheet[1]/x:sheetData[1]"
    assert len(schema["errors"][1]["description"].encode("utf-8")) <= 256
    assert schema["errors"][1]["error_type"] == "Schema"
    assert schema["max_errors"] == 2
    assert schema["truncated"] is True
    assert schema["field_truncations"] == 1
    gates = {gate["id"]: gate for gate in result["validation"]["gates"]}
    assert gates["schema.full"]["outcome"] == "fail"
    assert gates["schema.full"]["evidence"]["error_count"] == 2
    assert gates["schema.full"]["evidence"]["truncated"] is True
    assert gates["schema.full"]["evidence"]["field_truncations"] == 1


def test_dotnet_xlsx_schema_signed_xlsm_is_inert_read_only_and_source_preserving(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = create_package_fixture(tmp_path / "signed.xlsm", "xlsm", signed=True)
    source_bytes = source.read_bytes()
    source_hash = hashlib.sha256(source_bytes).hexdigest()
    runner = _DotnetRunner({
        "valid": True,
        "errors": [],
        "truncated": False,
        "file_format": "Microsoft365",
    })
    definition, _provider = build_dotnet_provider(
        project_root,
        detector=_CallableDetector("/fake/dotnet"),
        runner=runner,
    )
    registry = ProviderCatalog()
    registry.register_provider(definition)

    result = registry.execute({
        "schema_version": "1.0",
        "operation": "xlsx.validate.schema",
        "input": str(source),
        "arguments": {"max_errors": 25},
    })

    assert result["status"] == "success"
    assert source.read_bytes() == source_bytes
    assert result["artifacts"] == [
        {
            "role": "input",
            "path": str(source.resolve()),
            "sha256": source_hash,
            "bytes": len(source_bytes),
        }
    ]
    private_input = Path(runner.calls[0]["payload"]["input_path"])
    assert private_input.suffix == ".xlsm"
    assert private_input != source
    assert runner.calls[0]["input_sha256"] == source_hash
    gates = {gate["id"]: gate for gate in result["validation"]["gates"]}
    package_evidence = gates["xlsx.package-security"]["evidence"]
    assert package_evidence["workbook_format"] == "xlsm"
    assert package_evidence["security"]["categories"]["vba"]
    assert gates["source.preservation"]["evidence"]["sha256"] == source_hash


def test_libreoffice_xlsx_render_promotes_reopened_pdf_with_scoped_evidence(
    project_root: Path,
    qa_xlsx: Path,
    qa_pdf: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "rendered.pdf"
    source_hash = hashlib.sha256(qa_xlsx.read_bytes()).hexdigest()
    runner = _LibreOfficeRunner(qa_pdf.read_bytes())
    definition, _provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector("/fake/soffice"),
        runner=runner,
    )
    registry = ProviderCatalog()
    registry.register_provider(definition)

    result = registry.execute({
        "schema_version": "1.0",
        "operation": "xlsx.render",
        "input": str(qa_xlsx),
        "output": str(output),
        "arguments": {"max_sheets": 5, "max_cells_per_sheet": 50},
    })

    SchemaCatalog(project_root).validate("operation-result", result)
    assert result["status"] == "success"
    assert result["provider_chain"] == ["libreoffice"]
    assert reopen_pdf(output)["pages"] == 1
    gates = {gate["id"]: gate for gate in result["validation"]["gates"]}
    assert gates["visual.render"]["outcome"] == "pass"
    assert gates["visual.render"]["evidence"]["claim_scope"] == (
        "render-produced-and-pdf-reopened"
    )
    assert gates["schema.full"]["outcome"] == "unavailable"
    sampling = result["diagnostics"]["operation_result"]["render"]["sampling"]
    assert sampling["risk_finding_counts"] == {
        "hidden-data": 1,
        "potential-text-truncation": 1,
        "wide-columns": 1,
    }
    assert any(
        warning["code"] == "DS_XLSX_RENDER_SCOPE_LIMITED"
        for warning in result["warnings"]
    )
    assert hashlib.sha256(qa_xlsx.read_bytes()).hexdigest() == source_hash


def test_libreoffice_xlsx_render_validation_failure_preserves_destination(
    project_root: Path,
    qa_xlsx: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "existing.pdf"
    output.write_bytes(b"sentinel")
    definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector("/fake/soffice"),
        runner=_LibreOfficeRunner(b"not a pdf"),
    )

    result = provider.execute("xlsx.render", {
        "schema_version": "1.0",
        "operation": "xlsx.render",
        "input": str(qa_xlsx),
        "output": str(output),
        "arguments": {},
    })

    assert result["status"] == "failed"
    assert output.read_bytes() == b"sentinel"
    assert any(capability.operation == "xlsx.render" for capability in definition.capabilities)


def test_provider_only_operations_are_reported_unavailable_without_callability(
    project_root: Path,
    qa_xlsx: Path,
    tmp_path: Path,
) -> None:
    registry = ProviderCatalog()
    dotnet, _provider = build_dotnet_provider(
        project_root,
        detector=_AbsentDetector(),
        runner=_DotnetRunner({}),
    )
    libreoffice, _provider = build_libreoffice_provider(
        project_root,
        detector=_AbsentDetector(),
        runner=_LibreOfficeRunner(b""),
    )
    registry.register_provider(dotnet)
    registry.register_provider(libreoffice)
    report = build_capabilities(project_root, "xlsx", registry)
    operations = {item["operation"]: item for item in report["operations"]}
    assert operations["xlsx.validate.schema"]["available"] is False
    assert operations["xlsx.render"]["available"] is False

    output = tmp_path / "not-created.pdf"
    with pytest.raises(DocumentSkillsError) as captured:
        registry.execute({
            "operation": "xlsx.render",
            "input": str(qa_xlsx),
            "output": str(output),
            "arguments": {},
        })
    assert captured.value.status == "unavailable"
    assert not output.exists()


def test_public_worker_grants_provider_operations_bounded_time(
    project_root: Path,
    tmp_path: Path,
) -> None:
    supervisor = PublicCommandSupervisor(project_root, timeout_seconds=8.0)
    observed = {}
    for operation in ("xlsx.render", "xlsx.validate.schema"):
        request = tmp_path / f"{operation}.json"
        request.write_text(json.dumps({"operation": operation}), encoding="utf-8")
        observed[operation] = supervisor._command_limits(
            PublicCommand("run", ("run", "--request", str(request))),
            tmp_path,
        )
    assert observed == {
        "xlsx.render": (45.0, 2_097_152),
        "xlsx.validate.schema": (240.0, 2_097_152),
    }


def test_public_worker_distinguishes_auto_and_required_recalculation_time(
    project_root: Path,
    tmp_path: Path,
) -> None:
    supervisor = PublicCommandSupervisor(project_root, timeout_seconds=8.0)
    requests = {
        "auto": {
            "operation": "xlsx.create",
            "arguments": {"recalculation": "auto"},
        },
        "required": {
            "operation": "xlsx.create",
            "arguments": {"recalculation": "required"},
        },
        "required-edit": {
            "operation": "xlsx.edit",
            "arguments": {"recalculation": "required"},
        },
        "explicit": {"operation": "xlsx.recalculate", "arguments": {}},
        "legacy": {"operation": "xlsx.convert", "arguments": {}},
    }
    observed = {}
    for name, payload in requests.items():
        request = tmp_path / f"{name}.json"
        request.write_text(json.dumps(payload), encoding="utf-8")
        observed[name] = supervisor._command_limits(
            PublicCommand("run", ("run", "--request", str(request))),
            tmp_path,
        )

    assert observed == {
        "auto": (8.0, 2_097_152),
        "required": (90.0, 2_097_152),
        "required-edit": (90.0, 2_097_152),
        "explicit": (90.0, 2_097_152),
        "legacy": (45.0, 2_097_152),
    }


def test_libreoffice_pdf_filter_name_reopens_through_bounded_flate_decoder() -> None:
    content = b"BT /F1 12 Tf (LibreOffice) Tj ET"
    assert decode_stream(zlib.compress(content), ["/FlateDecode"]) == content


@pytest.mark.parametrize(
    ("formula", "expected_token"),
    [
        ('WEBSERVICE("http://127.0.0.1:9/secret")', "WEBSERVICE"),
        (
            '_xlfn._xlws.WEBSERVICE("http://127.0.0.1:9/secret")',
            "WEBSERVICE",
        ),
        ("cmd|" + ("A" * 1024) + "!A1", "DDE_LINK"),
    ],
)
def test_active_formula_is_rejected_before_libreoffice_render(
    tmp_path: Path,
    formula: str,
    expected_token: str,
) -> None:
    source = tmp_path / "active.xlsx"
    create_xlsx(
        source,
        {
            "metadata": {},
            "sheets": [{
                "name": "Data",
                "rows": [{"cells": [{
                    "ref": "A1",
                    "formula": formula,
                    "type": "n",
                }]}],
                "number_formats": [],
            }],
            "defined_names": [],
            "tables": [],
        },
    )
    with pytest.raises(DocumentSkillsError) as caught:
        assert_provider_formula_safe(source)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details["tokens"] == [expected_token]


def test_render_provider_receives_the_screened_private_snapshot(
    project_root: Path,
    qa_xlsx: Path,
    qa_pdf: Path,
    tmp_path: Path,
) -> None:
    original_bytes = qa_xlsx.read_bytes()
    received: list[Path] = []

    def mutate_original_after_preflight(provider_source: Path) -> bytes:
        received.append(provider_source)
        qa_xlsx.write_bytes(b"unscreened replacement")
        assert provider_source != qa_xlsx
        assert provider_source.read_bytes() == original_bytes
        return qa_pdf.read_bytes()

    with pytest.raises(DocumentSkillsError) as caught:
        execute_render(
            {
                "operation": "xlsx.render",
                "input": str(qa_xlsx),
                "output": str(tmp_path / "not-promoted.pdf"),
                "arguments": {},
            },
            project_root=project_root,
            converter=mutate_original_after_preflight,
        )
    assert received
    source_preservation = caught.value.details["source_preservation"]
    assert source_preservation["status"] == "fail"
    assert source_preservation["error"]["code"] == "DS_VALIDATION_FAILED"


def test_render_sampling_streams_only_selected_sheets_and_bounded_cells(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from document_skills_core.formats.xlsx import render_sampling

    source = tmp_path / "dense-many-sheet.xlsx"
    create_xlsx(
        source,
        {
            "metadata": {},
            "sheets": [
                {
                    "name": f"Sheet{sheet_index + 1}",
                    "rows": [
                        {
                            "cells": [
                                {
                                    "ref": f"{chr(ord('A') + column_index)}{row_index}",
                                    "value": f"dense-{sheet_index}-{row_index}-{column_index}",
                                    "type": "s",
                                }
                                for column_index in range(10)
                            ]
                        }
                        for row_index in range(1, 21)
                    ],
                    "number_formats": [],
                }
                for sheet_index in range(12)
            ],
            "defined_names": [],
            "tables": [],
            "charts": [],
        },
    )

    parsed_worksheets: list[str] = []
    unselected_read_sizes: list[int] = []
    sampled_cells = 0
    original_read = zipfile.ZipFile.read
    original_stream_read = zipfile.ZipExtFile.read
    original_iterparse = render_sampling.iterparse
    original_sample_cell = render_sampling._sample_cell

    def reject_materialized_xml(archive, member, *args, **kwargs):
        name = member.filename if isinstance(member, zipfile.ZipInfo) else member
        if name.startswith("xl/") and name.endswith(".xml"):
            raise AssertionError(f"SpreadsheetML XML was materialized: {name}")
        return original_read(archive, member, *args, **kwargs)

    def track_iterparse(source_stream, *args, **kwargs):
        name = str(getattr(source_stream, "name", ""))
        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
            parsed_worksheets.append(name)
        return original_iterparse(source_stream, *args, **kwargs)

    def track_stream_read(source_stream, size=-1):
        name = str(getattr(source_stream, "name", ""))
        if (
            name.startswith("xl/worksheets/")
            and name.endswith(".xml")
            and name != "xl/worksheets/sheet1.xml"
        ):
            unselected_read_sizes.append(size)
        return original_stream_read(source_stream, size)

    def count_sampled_cell(element, row_style_index: int):
        nonlocal sampled_cells
        sampled_cells += 1
        return original_sample_cell(element, row_style_index)

    monkeypatch.setattr(zipfile.ZipFile, "read", reject_materialized_xml)
    monkeypatch.setattr(zipfile.ZipExtFile, "read", track_stream_read)
    monkeypatch.setattr(render_sampling, "iterparse", track_iterparse)
    monkeypatch.setattr(render_sampling, "_sample_cell", count_sampled_cell)

    evidence, _warnings = render_sampling.sample_render_source(
        source,
        {"max_sheets": 1, "max_cells_per_sheet": 7, "max_findings": 10},
    )

    assert parsed_worksheets == ["xl/worksheets/sheet1.xml"]
    assert unselected_read_sizes
    assert all(0 < size <= 64 * 1024 for size in unselected_read_sizes)
    assert sampled_cells == 7
    assert evidence["sheet_count"] == 12
    assert evidence["sampled_sheet_count"] == 1
    assert evidence["truncated_sheet_count"] == 11
    sample = evidence["samples"][0]
    assert sample["sampled_cells"] == 7
    assert sample["total_cells"] == 8
    assert sample["total_cells_is_lower_bound"] is True
    assert sample["cell_sampling_truncated"] is True

    index = render_sampling.RenderPackageIndex.open(source)
    assert not hasattr(index, "parts")
    assert not hasattr(index, "relationships")
    assert all(not isinstance(value, bytes) for value in vars(index).values())


def test_render_sampling_inventories_print_setup_tables_and_charts(
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.table import Table
    from document_skills_core.formats.xlsx import render_sampling

    source = tmp_path / "render-inventory.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(["Label", "Value"])
    sheet.append(["One", 1])
    sheet.append(["Two", 2])
    sheet.add_table(Table(displayName="RenderInventory", ref="A1:B3"))
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, "D2")
    sheet.print_area = "A1:H20"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.oddHeader.center.text = "Render inventory"
    workbook.save(source)
    workbook.close()

    evidence, _warnings = render_sampling.sample_render_source(
        source,
        {"max_cells_per_sheet": 50, "max_findings": 20, "max_sheets": 5},
    )
    sample = evidence["samples"][0]

    assert sample["table_count"] == 1
    assert sample["chart_count"] == 1
    assert sample["print_area"] == "A1:H20"
    assert sample["print_titles"] == {"rows": "1:1", "columns": None}
    assert sample["page_setup"]["orientation"] == "landscape"
    assert sample["header_footer_present"] is True


def test_formula_preflight_streams_workbook_xml_without_opc_parts(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from document_skills_core.formats.xlsx import formula_security, package

    source = tmp_path / "active-streamed.xlsx"
    create_xlsx(
        source,
        {
            "metadata": {},
            "sheets": [{
                "name": "Data",
                "rows": [{"cells": [{
                    "ref": "A1",
                    "formula": 'WEBSERVICE("http://127.0.0.1:9/secret")',
                    "type": "n",
                }]}],
                "number_formats": [],
            }],
            "defined_names": [],
            "tables": [],
        },
    )
    parsed_worksheets: list[str] = []
    original_read = zipfile.ZipFile.read
    original_iterparse = formula_security.iterparse

    def reject_materialized_xml(archive, member, *args, **kwargs):
        name = member.filename if isinstance(member, zipfile.ZipInfo) else member
        if name.startswith("xl/") and name.endswith(".xml"):
            raise AssertionError(f"SpreadsheetML XML was materialized: {name}")
        return original_read(archive, member, *args, **kwargs)

    def track_iterparse(source_stream, *args, **kwargs):
        name = str(getattr(source_stream, "name", ""))
        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
            parsed_worksheets.append(name)
        return original_iterparse(source_stream, *args, **kwargs)

    def reject_opc_open(*_args, **_kwargs):
        raise AssertionError("formula preflight used OpcPackage.open")

    monkeypatch.setattr(zipfile.ZipFile, "read", reject_materialized_xml)
    monkeypatch.setattr(formula_security, "iterparse", track_iterparse)
    monkeypatch.setattr(package.OpcPackage, "open", reject_opc_open)

    with pytest.raises(DocumentSkillsError) as caught:
        assert_provider_formula_safe(source)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details["tokens"] == ["WEBSERVICE"]
    assert parsed_worksheets == ["xl/worksheets/sheet1.xml"]


def test_render_rejects_disguised_xlm_root_before_converter(
    project_root: Path,
    qa_xlsx: Path,
    tmp_path: Path,
) -> None:
    _add_neutral_xml_part(
        qa_xlsx,
        part="xl/custom/provider.xml",
        payload=(
            f'<macroSheet xmlns="{NS["main"]}"><sheetData /></macroSheet>'
        ).encode(),
        relationship_part=WORKBOOK_RELS,
        relationship_target="custom/provider.xml",
    )
    converter_calls: list[Path] = []

    with pytest.raises(DocumentSkillsError) as caught:
        execute_render(
            {
                "operation": "xlsx.render",
                "input": str(qa_xlsx),
                "output": str(tmp_path / "not-rendered.pdf"),
                "arguments": {},
            },
            project_root=project_root,
            converter=lambda path: converter_calls.append(path) or b"not used",
        )
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    inventory = caught.value.details["security_inventory"]
    assert inventory["counts"]["xlm"] == 1
    assert converter_calls == []


def test_render_rejects_malformed_non_xl_xml_before_converter(
    project_root: Path,
    qa_xlsx: Path,
    tmp_path: Path,
) -> None:
    _add_neutral_xml_part(
        qa_xlsx,
        part="customXml/item1.xml",
        payload=b"<neutral>",
        relationship_part=PACKAGE_RELS,
        relationship_target="customXml/item1.xml",
    )
    converter_calls: list[Path] = []

    with pytest.raises(DocumentSkillsError) as caught:
        execute_render(
            {
                "operation": "xlsx.render",
                "input": str(qa_xlsx),
                "output": str(tmp_path / "not-rendered.pdf"),
                "arguments": {},
            },
            project_root=project_root,
            converter=lambda path: converter_calls.append(path) or b"not used",
        )
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert converter_calls == []


def test_render_relationship_count_is_bounded_and_not_retained(
    qa_xlsx: Path,
) -> None:
    from document_skills_core.formats.xlsx import render_package

    with zipfile.ZipFile(qa_xlsx) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    relationships = fromstring(members[WORKBOOK_RELS][1])
    existing = len(relationships)
    relationship_limit = MAX_PARTS * 2
    for index in range(relationship_limit + 1 - existing):
        SubElement(
            relationships,
            f"{{{NS['rels']}}}Relationship",
            {
                "Id": f"rOverflow{index}",
                "Type": (
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships/customXml"
                ),
                "Target": f"https://example.invalid/{index}",
                "TargetMode": "External",
            },
        )
    _replace_member(members, WORKBOOK_RELS, tostring(relationships))
    with zipfile.ZipFile(
        qa_xlsx,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)

    with pytest.raises(DocumentSkillsError) as caught:
        render_package.RenderPackageIndex.open(qa_xlsx)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details["relationship_limit"] == relationship_limit
    assert "security_inventory" not in caught.value.details


def test_render_content_type_count_is_bounded(
    qa_xlsx: Path,
) -> None:
    from document_skills_core.formats.xlsx import render_package

    with zipfile.ZipFile(qa_xlsx) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    content_types = fromstring(members[CONTENT_TYPES][1])
    for index in range(MAX_PARTS + 1 - len(content_types)):
        SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
            {
                "Extension": f"safe{index}",
                "ContentType": f"application/x-safe-{index}",
            },
        )
    _replace_member(members, CONTENT_TYPES, tostring(content_types))
    with zipfile.ZipFile(
        qa_xlsx,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)

    with pytest.raises(DocumentSkillsError) as caught:
        render_package.RenderPackageIndex.open(qa_xlsx)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details["content_type_count"] == MAX_PARTS + 1
    assert caught.value.details["content_type_limit"] == MAX_PARTS


def test_render_content_type_ceiling_precedes_dangerous_inventory(
    qa_xlsx: Path,
) -> None:
    from document_skills_core.formats.xlsx import render_package

    with zipfile.ZipFile(qa_xlsx) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    content_types = fromstring(members[CONTENT_TYPES][1])
    for index in range(MAX_PARTS + 1 - len(content_types)):
        SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
            {
                "Extension": f"macro{index}",
                "ContentType": (
                    "application/vnd.ms-excel.sheet.macroEnabled."
                    f"overflow-{index}"
                ),
            },
        )
    _replace_member(members, CONTENT_TYPES, tostring(content_types))
    with zipfile.ZipFile(
        qa_xlsx,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)

    with pytest.raises(DocumentSkillsError) as caught:
        render_package.RenderPackageIndex.open(qa_xlsx)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details == {
        "content_type_count": MAX_PARTS + 1,
        "content_type_limit": MAX_PARTS,
    }


def test_render_rejects_nested_content_type_declarations_before_inventory(
    qa_xlsx: Path,
) -> None:
    from document_skills_core.formats.xlsx import render_package

    with zipfile.ZipFile(qa_xlsx) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    content_types = fromstring(members[CONTENT_TYPES][1])
    wrapper = SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
        {
            "Extension": "nestedwrapper",
            "ContentType": "application/x-safe-wrapper",
        },
    )
    for index in range(MAX_PARTS + 1):
        SubElement(
            wrapper,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
            {
                "Extension": f"nestedmacro{index}",
                "ContentType": (
                    "application/vnd.ms-excel.sheet.macroEnabled."
                    f"nested-{index}"
                ),
            },
        )
    _replace_member(members, CONTENT_TYPES, tostring(content_types))
    with zipfile.ZipFile(
        qa_xlsx,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)

    with pytest.raises(DocumentSkillsError) as caught:
        render_package.RenderPackageIndex.open(qa_xlsx)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details == {}


def test_render_rejects_nonempty_content_type_declaration(
    qa_xlsx: Path,
) -> None:
    from document_skills_core.formats.xlsx import render_package

    with zipfile.ZipFile(qa_xlsx) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    content_types = fromstring(members[CONTENT_TYPES][1])
    declaration = SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
        {
            "Extension": "nonempty",
            "ContentType": "application/x-safe-nonempty",
        },
    )
    declaration.text = "attacker-controlled-content"
    _replace_member(members, CONTENT_TYPES, tostring(content_types))
    with zipfile.ZipFile(
        qa_xlsx,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)

    with pytest.raises(DocumentSkillsError) as caught:
        render_package.RenderPackageIndex.open(qa_xlsx)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details == {}


def test_render_content_type_errors_do_not_retain_attacker_input(
    qa_xlsx: Path,
) -> None:
    from document_skills_core.formats.xlsx import render_package

    with zipfile.ZipFile(qa_xlsx) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    content_types = fromstring(members[CONTENT_TYPES][1])
    SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
        {
            "Extension": ("attacker" * 8192) + "/",
            "ContentType": "application/x-invalid",
        },
    )
    _replace_member(members, CONTENT_TYPES, tostring(content_types))
    with zipfile.ZipFile(
        qa_xlsx,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)

    with pytest.raises(DocumentSkillsError) as caught:
        render_package.RenderPackageIndex.open(qa_xlsx)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details == {}


@pytest.mark.parametrize(
    "malformation",
    ["unknown-child", "root-text", "declaration-tail"],
)
def test_render_rejects_invalid_content_type_root_content(
    qa_xlsx: Path,
    malformation: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from document_skills_core.formats.xlsx import render_package

    with zipfile.ZipFile(qa_xlsx) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    content_types = fromstring(members[CONTENT_TYPES][1])
    if malformation == "unknown-child":
        SubElement(content_types, "{urn:attacker}Unexpected")
    elif malformation == "root-text":
        content_types.text = "attacker-controlled-content"
    else:
        declaration = SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
            {
                "Extension": "tail",
                "ContentType": "application/x-safe-tail",
            },
        )
        declaration.tail = "attacker-controlled-content"
    _replace_member(members, CONTENT_TYPES, tostring(content_types))
    with zipfile.ZipFile(
        qa_xlsx,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)

    def reject_inventory(*_args, **_kwargs):
        raise AssertionError("malformed content types reached security inventory")

    monkeypatch.setattr(
        render_package,
        "spreadsheet_security_inventory",
        reject_inventory,
    )

    with pytest.raises(DocumentSkillsError) as caught:
        render_package.RenderPackageIndex.open(qa_xlsx)
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert caught.value.details == {}


def test_render_ignores_relationship_elements_outside_rels_parts(
    qa_xlsx: Path,
) -> None:
    from document_skills_core.formats.xlsx import render_package

    root = fromstring(b"<container />")
    for index in range(MAX_PARTS * 2 + 1):
        SubElement(
            root,
            f"{{{NS['rels']}}}Relationship",
            {
                "Id": f"rFake{index}",
                "Type": (
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships/customXml"
                ),
                "Target": f"https://example.invalid/{index}",
                "TargetMode": "External",
            },
        )
    _add_neutral_xml_part(
        qa_xlsx,
        part="customXml/fake-relationships.xml",
        payload=tostring(root),
        relationship_part=PACKAGE_RELS,
        relationship_target="customXml/fake-relationships.xml",
    )

    index = render_package.RenderPackageIndex.open(qa_xlsx)

    assert index.security["counts"]["external_targets"] == 0
    assert index.relationship_count < MAX_PARTS * 2


@pytest.mark.parametrize(
    "malformation",
    [
        "missing-sheet-rid",
        "missing-relationship",
        "wrong-relationship-type",
        "missing-target",
        "wrong-content-type",
    ],
)
def test_render_rejects_malformed_sheet_relationship_before_converter(
    project_root: Path,
    qa_xlsx: Path,
    tmp_path: Path,
    malformation: str,
) -> None:
    _malform_sheet_binding(qa_xlsx, malformation)
    converter_calls: list[Path] = []

    def converter(path: Path) -> bytes:
        converter_calls.append(path)
        return b"converter must not run"

    with pytest.raises(DocumentSkillsError) as caught:
        execute_render(
            {
                "operation": "xlsx.render",
                "input": str(qa_xlsx),
                "output": str(tmp_path / "not-rendered.pdf"),
                "arguments": {"max_sheets": 1},
            },
            project_root=project_root,
            converter=converter,
        )
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert converter_calls == []


def test_render_validates_unselected_sheet_binding_before_sampling_limit(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "two-sheets.xlsx"
    create_xlsx(
        source,
        {
            "metadata": {},
            "sheets": [
                {
                    "name": name,
                    "rows": [{"cells": [{"ref": "A1", "value": name, "type": "s"}]}],
                    "number_formats": [],
                }
                for name in ("Selected", "NotSelected")
            ],
            "defined_names": [],
            "tables": [],
            "charts": [],
        },
    )
    _malform_sheet_binding(source, "duplicate-target")
    converter_calls: list[Path] = []

    with pytest.raises(DocumentSkillsError) as caught:
        execute_render(
            {
                "operation": "xlsx.render",
                "input": str(source),
                "output": str(tmp_path / "not-rendered.pdf"),
                "arguments": {"max_sheets": 1},
            },
            project_root=project_root,
            converter=lambda path: converter_calls.append(path) or b"not used",
        )
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert converter_calls == []


@pytest.mark.parametrize(
    "malformation",
    [
        "invalid-column-width",
        "nonfinite-column-width",
        "invalid-row-style",
        "malformed-worksheet-xml",
    ],
)
def test_render_wraps_streamed_worksheet_parse_errors_as_archive_unsafe(
    project_root: Path,
    qa_xlsx: Path,
    tmp_path: Path,
    malformation: str,
) -> None:
    _malform_worksheet_xml(qa_xlsx, malformation)
    converter_calls: list[Path] = []

    with pytest.raises(DocumentSkillsError) as caught:
        execute_render(
            {
                "operation": "xlsx.render",
                "input": str(qa_xlsx),
                "output": str(tmp_path / "not-rendered.pdf"),
                "arguments": {},
            },
            project_root=project_root,
            converter=lambda path: converter_calls.append(path) or b"not used",
        )
    assert caught.value.code == ErrorCode.ARCHIVE_UNSAFE
    assert converter_calls == []


def _malform_sheet_binding(path: Path, malformation: str) -> None:
    with zipfile.ZipFile(path) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    workbook = fromstring(members[WORKBOOK_MAIN][1])
    sheets = workbook.findall(f"{{{NS['main']}}}sheets/{{{NS['main']}}}sheet")
    assert sheets
    sheet = sheets[0]
    relationship_id = sheet.attrib[f"{{{NS['r']}}}id"]
    relationships = fromstring(members[WORKBOOK_RELS][1])
    relationships_by_id = {
        node.attrib.get("Id"): node for node in relationships
    }
    relationship = relationships_by_id[relationship_id]
    content_types = fromstring(members[CONTENT_TYPES][1])

    if malformation == "missing-sheet-rid":
        del sheet.attrib[f"{{{NS['r']}}}id"]
        _replace_member(members, WORKBOOK_MAIN, tostring(workbook))
    elif malformation == "missing-relationship":
        relationships.remove(relationship)
        _replace_member(members, WORKBOOK_RELS, tostring(relationships))
    elif malformation == "wrong-relationship-type":
        relationship.attrib["Type"] = REL_TABLE
        _replace_member(members, WORKBOOK_RELS, tostring(relationships))
    elif malformation == "missing-target":
        relationship.attrib["Target"] = "worksheets/missing.xml"
        _replace_member(members, WORKBOOK_RELS, tostring(relationships))
    elif malformation == "wrong-content-type":
        override = next(
            node
            for node in content_types
            if node.attrib.get("PartName") == "/xl/worksheets/sheet1.xml"
        )
        override.attrib["ContentType"] = "application/xml"
        _replace_member(members, CONTENT_TYPES, tostring(content_types))
    elif malformation == "duplicate-target":
        assert len(sheets) == 2
        second_id = sheets[1].attrib[f"{{{NS['r']}}}id"]
        relationships_by_id[second_id].attrib["Target"] = relationship.attrib["Target"]
        _replace_member(members, WORKBOOK_RELS, tostring(relationships))
    else:
        raise AssertionError(f"unknown malformation: {malformation}")

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)


def _malform_worksheet_xml(path: Path, malformation: str) -> None:
    sheet_part = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(path) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    if malformation == "malformed-worksheet-xml":
        _replace_member(members, sheet_part, b"<worksheet>")
    else:
        worksheet = fromstring(members[sheet_part][1])
        if malformation in {"invalid-column-width", "nonfinite-column-width"}:
            node = worksheet.find(f"{{{NS['main']}}}cols/{{{NS['main']}}}col")
            assert node is not None
            node.attrib["width"] = (
                "not-a-float"
                if malformation == "invalid-column-width"
                else "NaN"
            )
        elif malformation == "invalid-row-style":
            node = worksheet.find(
                f"{{{NS['main']}}}sheetData/{{{NS['main']}}}row"
            )
            assert node is not None
            node.attrib["s"] = "not-an-integer"
        else:
            raise AssertionError(f"unknown malformation: {malformation}")
        _replace_member(members, sheet_part, tostring(worksheet))
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for info, payload in members.values():
            archive.writestr(info, payload)


def _add_neutral_xml_part(
    path: Path,
    *,
    part: str,
    payload: bytes,
    relationship_part: str,
    relationship_target: str,
) -> None:
    with zipfile.ZipFile(path) as archive:
        members = {
            info.filename: (info, archive.read(info))
            for info in archive.infolist()
        }
    content_types = fromstring(members[CONTENT_TYPES][1])
    SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        {"PartName": f"/{part}", "ContentType": "application/xml"},
    )
    _replace_member(members, CONTENT_TYPES, tostring(content_types))
    relationships = fromstring(members[relationship_part][1])
    SubElement(
        relationships,
        f"{{{NS['rels']}}}Relationship",
        {
            "Id": "rNeutralPayload",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/customXml"
            ),
            "Target": relationship_target,
        },
    )
    _replace_member(members, relationship_part, tostring(relationships))
    info = zipfile.ZipInfo(part)
    info.compress_type = zipfile.ZIP_DEFLATED
    members[part] = (info, payload)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for member_info, member_payload in members.values():
            archive.writestr(member_info, member_payload)


def _replace_member(
    members: dict[str, tuple[zipfile.ZipInfo, bytes]],
    name: str,
    payload: bytes,
) -> None:
    info, _old_payload = members[name]
    members[name] = (info, payload)
