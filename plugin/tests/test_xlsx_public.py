"""XLSX public command surface tests — frozen uv subprocess boundary."""

import hashlib
import json
from pathlib import Path
import subprocess
import zipfile
from xml.etree.ElementTree import fromstring, tostring

import pytest

from document_skills_core.core.contracts.schemas import SchemaCatalog
from document_skills_core.formats.xlsx.constants import NS
from document_skills_core.formats.xlsx.create import create_xlsx
from document_skills_core.formats.xlsx.mapping import map_workbook
from document_skills_core.formats.xlsx.package import OpcPackage
from tests.support.xlsx_macro_fixture import create_package_fixture


def _strip_style_children(path: Path) -> None:
    """Replace styles.xml with an empty-container form that fails the style gate."""

    with zipfile.ZipFile(path, "r") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/styles.xml"] = (
        b'<?xml version="1.0" encoding="UTF-8"?>\n'
        b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<fonts count="1"/><fills count="2"/><borders count="1"/>'
        b'<cellStyleXfs count="1"/><cellXfs count="1"/>'
        b'<cellStyles count="1"/><dxfs count="0"/></styleSheet>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in sorted(parts.items()):
            archive.writestr(name, data)


def _set_sheet_numeric_attribute(
    path: Path,
    element_name: str,
    attribute: str,
    value: str,
    *,
    part: str = "xl/worksheets/sheet1.xml",
) -> None:
    """Replace one SpreadsheetML numeric attribute without normalizing its value."""

    with zipfile.ZipFile(path, "r") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = fromstring(parts[part])
    element = root.find(f".//{{{NS['main']}}}{element_name}")
    assert element is not None
    element.attrib[attribute] = value
    parts[part] = tostring(root, encoding="UTF-8", xml_declaration=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in sorted(parts.items()):
            archive.writestr(name, data)


def _public(
    project_root: Path,
    *arguments: str,
    check: bool = True,
    cwd: Path | None = None,
) -> dict[str, object]:
    process = subprocess.run(
        [
            "uv",
            "run",
            "--project",
            str(project_root),
            "--frozen",
            "python",
            str(project_root / "skills/document-xlsx/scripts/run.py"),
            *arguments,
        ],
        cwd=cwd or project_root,
        check=False,
        capture_output=True,
        text=False,
        shell=False,
        timeout=300,
    )
    if check:
        assert process.returncode == 0, (
            process.stderr.decode("utf-8", errors="replace")
            or process.stdout.decode("utf-8", errors="replace")
        )
    assert process.stderr == b""
    text = process.stdout.decode("utf-8", errors="strict")
    decoder = json.JSONDecoder()
    payload, end = decoder.raw_decode(text)
    assert text[end:].strip() == ""
    assert type(payload) is dict
    return payload


def _request(tmp_path: Path, name: str, payload: dict[str, object]) -> Path:
    path = tmp_path / name
    path.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
        newline="\n",
    )
    return path


def _workbook() -> dict[str, object]:
    return {
        "metadata": {"title": "Public XLSX", "creator": "Test", "subject": ""},
        "sheets": [
            {
                "name": "Sheet1",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Name", "type": "s"},
                            {"ref": "B1", "value": "10", "type": "n"},
                            {"ref": "B2", "formula": "B1*2", "type": "n"},
                        ],
                    }
                ],
                "number_formats": [],
            }
        ],
        "defined_names": [],
        "tables": [],
        "chart_reference": None,
        "page_setup": None,
    }


def _claimed_create_features_workbook() -> dict[str, object]:
    return {
        "metadata": {"title": "Claimed XLSX features", "creator": "Test", "subject": ""},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Name", "type": "s"},
                            {"ref": "B1", "value": "12", "type": "n"},
                            {
                                "ref": "B2",
                                "formula": "B1*2",
                                "cached_value": "24",
                                "type": "n",
                            },
                            {"ref": "D1", "value": "Item", "type": "s"},
                            {"ref": "E1", "value": "Amount", "type": "s"},
                            {"ref": "D2", "value": "Alpha", "type": "s"},
                            {"ref": "E2", "value": "12", "type": "n"},
                        ]
                    }
                ],
                "number_formats": [],
                "data_validations": [
                    {
                        "ref": "C2:C20",
                        "type": "list",
                        "formula1": "\"Low,High\"",
                        "allow_blank": True,
                        "show_input_message": True,
                        "show_error_message": True,
                        "prompt_title": "Level",
                        "prompt": "Choose a level",
                        "error_title": "Invalid",
                        "error": "Choose Low or High",
                        "error_style": "stop",
                    }
                ],
                "conditional_formats": [
                    {
                        "ref": "E2:E20",
                        "type": "cellIs",
                        "operator": "greaterThan",
                        "formulas": ["10"],
                        "style": {
                            "fill": {"pattern": "solid", "color": "#FFCC00"}
                        },
                    }
                ],
            },
            {
                "name": "Summary",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Total", "type": "s"},
                        ]
                    }
                ],
                "number_formats": [],
            },
        ],
        "defined_names": [
            {"name": "TotalValue", "ref": "Data!$B$2", "scope": "workbook"},
        ],
        "tables": [
            {
                "name": "PublicData",
                "ref": "D1:E2",
                "sheet": "Data",
                "style": "TableStyleMedium2",
            }
        ],
        "charts": [
            {
                "name": "PublicChart",
                "sheet": "Data",
                "type": "column",
                "title": "Public native chart",
                "anchor": "G2:N16",
                "series": [
                    {
                        "name": "Amount",
                        "categories": "Data!$D$2:$D$2",
                        "values": "Data!$E$2:$E$2",
                        "color": "#4472C4",
                    }
                ],
                "show_legend": True,
                "legend_position": "r",
                "x_axis_title": "Item",
                "y_axis_title": "Amount",
                "y_axis_number_format": "#,##0",
                "data_labels": {"show_value": True},
            }
        ],
        "chart_reference": None,
        "page_setup": None,
    }


@pytest.fixture
def public_created(project_root: Path, tmp_path: Path) -> Path:
    from openpyxl import Workbook

    output = tmp_path / "public-created.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "Name"
    sheet["B1"] = 10
    sheet["B2"] = "=B1*2"
    workbook.save(output)
    return output


def test_public_create_is_truthful_and_promotes_bounded_artifact(
    project_root: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "bounded.xlsx"
    request = _request(
        tmp_path,
        "create-bounded.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": _workbook()},
        },
    )
    result = _public(project_root, "run", "--request", str(request))
    SchemaCatalog(project_root).validate("operation-result", result)
    assert result["status"] in {"success", "degraded"}
    assert output.is_file()
    assert any(
        item["path"] == str(output.resolve())
        for item in result["artifacts"]
    )
    assert any(
        gate["id"] == "operation.consumer-package-conformance"
        and gate["outcome"] == "pass"
        for gate in result["validation"]["gates"]
    )


def test_public_feature_truth_table_is_asserted(project_root: Path) -> None:
    truth_table_path = project_root / "skills/document-xlsx/references/feature-truth-table.json"
    truth_table = json.loads(truth_table_path.read_text(encoding="utf-8"))

    assert truth_table == {
        "schema_version": "1.0",
        "operations": {
            "xlsx.read": {
                "available": [
                    "static_formula_syntax_and_reference_report",
                    "formula_type_classification",
                    "xlsm_inert_vba_read",
                    "native_pivot_readback",
                ],
                "limitations": [
                    "static_analysis_is_not_a_calculation_engine",
                ],
            },
            "xlsx.inspect.structure": {
                "available": [
                    "inert_formula_type_classification",
                    "external_formula_inventory",
                    "xlsm_vba_and_signature_inventory",
                    "native_pivot_relationship_inventory",
                ],
            },
            "xlsx.create": {
                "available": [
                    "multiple_sheets",
                    "typed_cell_values",
                    "formulas_with_optional_cached_values",
                    "defined_names",
                    "cell_style",
                    "row_style",
                    "column_style",
                    "custom_number_format",
                    "row_height_and_hidden",
                    "column_width_and_hidden",
                    "native_table",
                    "data_validation",
                    "conditional_formatting",
                    "native_chart",
                    "native_area_radar_bubble_chart",
                    "native_combo_chart_secondary_axis",
                    "native_chart_trendline_error_bars",
                    "chart_style_and_structural_reference_maintenance",
                    "native_sparklines",
                    "page_setup",
                    "header_footer",
                    "sheet_view",
                    "print_area",
                    "print_titles",
                    "internal_hyperlink",
                    "cell_comments",
                    "workbook_properties",
                    "recalculation_policy_auto_required_skip",
                    "static_formula_syntax_and_reference_validation",
                    "formula_type_classification",
                ],
                "enhancement_required": [
                    "external_hyperlink_authoring",
                ],
            },
            "xlsx.edit": {
                "available": [
                    "cell_value",
                    "cell_formula",
                    "sheet_rename",
                    "cell_style",
                    "row_style",
                    "column_style",
                    "custom_number_format",
                    "row_insert",
                    "row_delete",
                    "column_insert",
                    "column_delete",
                    "row_height",
                    "row_hidden",
                    "column_width",
                    "column_hidden",
                    "sheet_add",
                    "sheet_delete",
                    "sheet_copy_plain_worksheet",
                    "sheet_reorder",
                    "cells_merge",
                    "cells_unmerge",
                    "range_clear",
                    "freeze_panes",
                    "auto_filter",
                    "print_area",
                    "manual_page_breaks",
                    "defined_name_crud",
                    "native_table_add",
                    "native_table_resize",
                    "native_table_rename",
                    "native_table_style",
                    "native_table_delete",
                    "data_validation_crud",
                    "conditional_formatting_crud",
                    "native_chart_crud",
                    "advanced_native_chart_crud",
                    "chart_style_and_structural_reference_maintenance",
                    "native_sparkline_crud",
                    "sparkline_structural_reference_maintenance",
                    "page_setup",
                    "header_footer",
                    "sheet_view",
                    "print_titles",
                    "internal_hyperlink_crud",
                    "cell_comment_crud",
                    "workbook_properties",
                    "recalculation_policy_auto_required_skip",
                    "static_formula_syntax_and_reference_validation",
                    "formula_type_classification",
                    "xlsm_keep_vba_copy_through",
                    "vba_relationship_and_signature_preservation",
                    "macro_signature_invalidation_disclosure",
                ],
                "enhancement_required": [
                    "sheet_copy_with_related_objects",
                    "sheet_delete_with_related_objects",
                    "sheet_delete_with_inbound_references",
                    "shared_array_or_data_table_formula_structural_edit",
                    "external_workbook_reference_structural_edit",
                    "pivot_structural_edit",
                    "external_hyperlink_authoring",
                ],
            },
            "xlsx.recalculate": {
                "available": [
                    "formula_identity_validation",
                    "formula_error_token_scan",
                    "cached_value_harvest",
                    "static_formula_syntax_and_reference_validation",
                    "special_formula_fail_closed",
                    "unknown_part_copy_through",
                    "source_preservation",
                    "atomic_promotion",
                ],
                "unavailable_without": [
                    "libreoffice_for_formula_workbooks",
                ],
            },
            "xlsx.convert": {
                "available": [
                    "xlsx_csv_tsv_json",
                    "encoding_bom_delimiter_quote_line_ending",
                    "typed_null_empty_boolean_number_date_time_timezone",
                    "leading_zero_and_large_integer_policy",
                    "formula_preserve_text_evaluated_reject_policy",
                    "stable_multisheet_typed_json",
                    "csv_injection_default_escape",
                    "bounded_streaming_and_resource_limits",
                    "semantic_loss_reporting",
                    "source_preservation",
                    "atomic_promotion",
                    "legacy_xls_to_xlsx_via_libreoffice",
                ],
                "limitations": [
                    "delimited_output_is_single_sheet",
                    "xlsx_objects_and_styles_are_not_tabular",
                    "evaluated_formula_cache_is_not_recalculated",
                    "shared_array_data_table_formula_conversion_fails_closed",
                ],
                "unavailable_without": [
                    "libreoffice_for_legacy_xls_conversion",
                ],
            },
            "xlsx.template.instantiate": {
                "available": [
                    "xltx_to_xlsx",
                    "xltm_to_xlsm_keep_vba",
                    "optional_bounded_edits",
                    "vba_relationship_and_signature_preservation",
                    "macro_signature_invalidation_disclosure",
                    "source_preservation",
                    "atomic_promotion",
                ],
                "limitations": [
                    "macro_enabled_template_recalculation_is_skipped",
                    "template_code_and_macros_are_never_executed",
                ],
            },
            "xlsx.summary.aggregate": {
                "available": [
                    "ordinary_summary_sheet",
                    "group_by",
                    "sum_average_min_max",
                    "count_nonblank_and_distinct",
                    "multi_key_sort",
                    "top_n",
                    "native_table_output",
                    "stored_value_type_preservation",
                    "explicit_cached_formula_policy",
                    "xlsm_keep_vba_copy_through",
                    "source_preservation",
                    "atomic_promotion",
                ],
                "limitations": [
                    "ordinary_summary_is_not_native_pivot",
                    "formula_caches_are_not_recalculated",
                    "formatted_dates_group_as_stored_values",
                ],
            },
            "xlsx.pivot.create": {
                "available": [
                    "native_pivot_table_and_cache",
                    "worksheet_source_range",
                    "single_row_axis",
                    "optional_single_column_axis",
                    "optional_single_page_filter",
                    "single_value_sum_average_min_max_count",
                    "saved_pivot_cache_records",
                    "relationship_and_content_type_validation",
                    "independent_public_readback",
                    "xlsm_keep_vba_copy_through",
                    "source_preservation",
                    "atomic_promotion",
                ],
                "limitations": [
                    "one_row_one_column_one_filter_one_value_maximum",
                    "target_sheet_must_be_new",
                    "formula_caches_are_not_recalculated",
                    "pivot_structural_edit_is_unavailable",
                ],
            },
            "xlsx.validate.schema": {
                "available": [
                    "spreadsheetdocument_openxml_schema_validation",
                    "bounded_per_part_error_report",
                    "xlsx_and_inert_xlsm_read_only_validation",
                    "source_preservation",
                ],
                "limitations": [
                    "schema_validation_does_not_calculate_formulas_or_prove_visual_fidelity",
                ],
                "unavailable_without": [
                    "dotnet_openxml_provider",
                ],
            },
            "xlsx.render": {
                "available": [
                    "libreoffice_xlsx_to_pdf",
                    "pdf_reopen_validation",
                    "bounded_per_sheet_risk_sampling",
                    "hidden_data_wide_column_and_text_truncation_inventory",
                    "print_setup_table_and_chart_inventory",
                    "source_preservation",
                    "atomic_promotion",
                ],
                "limitations": [
                    "metadata_sampling_does_not_prove_per_object_visual_parity",
                    "render_output_is_pdf_only",
                ],
                "unavailable_without": [
                    "libreoffice_provider",
                ],
            },
        },
    }


def test_public_create_claimed_features_reopen(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "claimed-create-features.xlsx"
    request = _request(
        tmp_path,
        "claimed-create-features.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": _claimed_create_features_workbook()},
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] in {"success", "degraded"}
    reopened = load_workbook(output, data_only=False)
    assert reopened.sheetnames == ["Data", "Summary"]
    assert reopened["Data"]["A1"].value == "Name"
    assert reopened["Data"]["B1"].value == 12
    assert reopened["Data"]["B2"].value == "=B1*2"
    assert reopened.defined_names["TotalValue"].attr_text == "Data!$B$2"
    assert reopened["Data"].tables["PublicData"].ref == "D1:E2"
    validations = list(reopened["Data"].data_validations.dataValidation)
    assert len(validations) == 1
    assert str(validations[0].sqref) == "C2:C20"
    conditional_rules = [
        rule
        for conditional_format in reopened["Data"].conditional_formatting
        for rule in reopened["Data"].conditional_formatting[conditional_format]
    ]
    assert [rule.type for rule in conditional_rules] == ["cellIs"]
    assert [type(chart).__name__ for chart in reopened["Data"]._charts] == ["BarChart"]
    assert reopened["Data"]._charts[0].type == "col"


def test_public_native_object_edits_reopen(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    source = tmp_path / "native-objects-source.xlsx"
    first_output = tmp_path / "native-objects-first.xlsx"
    renamed_output = tmp_path / "native-objects-renamed.xlsx"
    deleted_output = tmp_path / "native-objects-deleted.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in [
        ["Name", "Amount", "Category", "Code", "Score"],
        ["Alpha", 10, "A", "X", 1],
        ["Beta", 20, "B", "Y", 2],
        ["Gamma", 30, "C", "Z", 3],
    ]:
        sheet.append(row)
    table = Table(displayName="DataTable", ref="A1:C3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    for ref, validation_type, formula in (
        ("A2:A4", "whole", "1"),
        ("C2:C4", "list", '"A,B,C"'),
    ):
        validation = DataValidation(
            type=validation_type,
            operator="greaterThanOrEqual" if validation_type == "whole" else None,
            formula1=formula,
        )
        validation.add(ref)
        sheet.add_data_validation(validation)
    red = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    sheet.conditional_formatting.add(
        "A2:A4", CellIsRule(operator="greaterThan", formula=["1"], fill=red)
    )
    sheet.conditional_formatting.add(
        "B2:B4", FormulaRule(formula=["MOD(B2,2)=0"], fill=red)
    )
    workbook.save(source)

    edits = [
        {"sheet": "Data", "type": "table_resize", "name": "DataTable", "ref": "A1:C4"},
        {
            "sheet": "Data",
            "type": "table_style",
            "name": "DataTable",
            "table_style": "TableStyleLight1",
        },
        {
            "sheet": "Data",
            "type": "table_add",
            "name": "ScoreTable",
            "ref": "D1:E4",
            "table_style": "TableStyleDark1",
        },
        {
            "sheet": "Data",
            "type": "data_validation_update",
            "ref": "A2:A4",
            "validation": {
                "ref": "A2:A4",
                "type": "decimal",
                "operator": "greaterThanOrEqual",
                "formula1": "0",
            },
        },
        {
            "sheet": "Data",
            "type": "data_validation_add",
            "validation": {"ref": "B2:B4", "type": "list", "formula1": '"Low,High"'},
        },
        {"sheet": "Data", "type": "data_validation_delete", "ref": "C2:C4"},
        {
            "sheet": "Data",
            "type": "conditional_format_update",
            "ref": "A2:A4",
            "priority": 1,
            "rule": {
                "ref": "A2:A4",
                "type": "cellIs",
                "operator": "lessThan",
                "formulas": ["4"],
                "style": {"font": {"bold": True, "color": "#008000"}},
            },
        },
        {
            "sheet": "Data",
            "type": "conditional_format_delete",
            "ref": "B2:B4",
            "priority": 2,
        },
        {
            "sheet": "Data",
            "type": "conditional_format_add",
            "rule": {
                "ref": "C2:C4",
                "type": "dataBar",
                "thresholds": [{"type": "min"}, {"type": "max"}],
                "color": "#638EC6",
            },
        },
    ]
    request = _request(
        tmp_path,
        "native-object-edits.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(first_output),
            "arguments": {"edits": edits, "expected_edits": len(edits)},
        },
    )
    assert _public(project_root, "run", "--request", str(request))["status"] == "success"
    reopened = load_workbook(first_output)
    assert reopened["Data"].tables["DataTable"].ref == "A1:C4"
    assert reopened["Data"].tables["DataTable"].tableStyleInfo.name == "TableStyleLight1"
    assert reopened["Data"].tables["ScoreTable"].ref == "D1:E4"
    assert {str(item.sqref) for item in reopened["Data"].data_validations.dataValidation} == {
        "A2:A4",
        "B2:B4",
    }
    rules = [
        item
        for conditional_format in reopened["Data"].conditional_formatting
        for item in reopened["Data"].conditional_formatting[conditional_format]
    ]
    assert {(item.type, item.priority) for item in rules} == {("cellIs", 1), ("dataBar", 2)}

    rename_request = _request(
        tmp_path,
        "native-object-rename.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(first_output),
            "output": str(renamed_output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "table_rename",
                        "name": "DataTable",
                        "value": "SalesTable",
                    }
                ]
            },
        },
    )
    assert _public(project_root, "run", "--request", str(rename_request))["status"] == "success"
    assert "SalesTable" in load_workbook(renamed_output)["Data"].tables

    delete_request = _request(
        tmp_path,
        "native-object-delete.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(renamed_output),
            "output": str(deleted_output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "table_delete", "name": "SalesTable"},
                    {"sheet": "Data", "type": "table_delete", "name": "ScoreTable"},
                ]
            },
        },
    )
    assert _public(project_root, "run", "--request", str(delete_request))["status"] == "success"
    assert not load_workbook(deleted_output)["Data"].tables


def test_public_native_chart_edits_reopen(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "public-chart-source.xlsx"
    edited = tmp_path / "public-chart-edited.xlsx"
    deleted = tmp_path / "public-chart-deleted.xlsx"
    workbook = _workbook()
    cells = workbook["sheets"][0]["rows"][0]["cells"]
    cells[2] = {"ref": "B2", "value": "20", "type": "n"}
    cells.append({"ref": "A2", "value": "Beta", "type": "s"})

    def chart(name: str, chart_type: str, anchor: str) -> dict[str, object]:
        return {
            "name": name,
            "sheet": "Sheet1",
            "type": chart_type,
            "title": name,
            "anchor": anchor,
            "series": [
                {
                    "name": "Amount",
                    "categories": "Sheet1!$A$1:$A$2",
                    "values": "Sheet1!$B$1:$B$2",
                    "color": "#4472C4",
                }
            ],
            "show_legend": True,
            "legend_position": "r",
            "data_labels": {"show_value": True},
        }

    workbook["charts"] = [chart("OriginalChart", "column", "D2:K16")]
    create_request = _request(
        tmp_path,
        "public-chart-create.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(source),
            "arguments": {"workbook": workbook},
        },
    )
    assert _public(project_root, "run", "--request", str(create_request))["status"] == "success"

    update = chart("UpdatedChart", "line", "E3:L17")
    added = chart("AddedPie", "pie", "M3:T17")
    edit_request = _request(
        tmp_path,
        "public-chart-edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(edited),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Sheet1",
                        "type": "chart_update",
                        "name": "OriginalChart",
                        "chart": update,
                    },
                    {"sheet": "Sheet1", "type": "chart_add", "chart": added},
                ],
                "expected_edits": 2,
            },
        },
    )
    assert _public(project_root, "run", "--request", str(edit_request))["status"] == "success"
    assert [type(item).__name__ for item in load_workbook(edited)["Sheet1"]._charts] == [
        "LineChart",
        "PieChart",
    ]

    delete_request = _request(
        tmp_path,
        "public-chart-delete.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(edited),
            "output": str(deleted),
            "arguments": {
                "edits": [
                    {"sheet": "Sheet1", "type": "chart_delete", "name": "UpdatedChart"},
                    {"sheet": "Sheet1", "type": "chart_delete", "name": "AddedPie"},
                ],
                "expected_edits": 2,
            },
        },
    )
    assert _public(project_root, "run", "--request", str(delete_request))["status"] == "success"
    assert not load_workbook(deleted)["Sheet1"]._charts


def test_public_edit_claimed_features_reopen(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "claimed-edit-features.xlsx"
    request = _request(
        tmp_path,
        "claimed-edit-features.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(public_created),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Sheet1", "type": "cell_value", "ref": "B1", "value": "42"},
                    {
                        "sheet": "Sheet1",
                        "type": "cell_formula",
                        "ref": "B2",
                        "value": "B1*3",
                    },
                    {
                        "sheet": "Sheet1",
                        "type": "sheet_rename",
                        "ref": "A1",
                        "value": "Renamed",
                    },
                ],
                "expected_edits": 3,
            },
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "degraded"
    reopened = load_workbook(output, data_only=False)
    assert reopened.sheetnames == ["Renamed"]
    assert reopened["Renamed"]["B1"].value == 42
    assert reopened["Renamed"]["B2"].value == "=B1*3"


def test_public_edit_structural_sheet_and_range_features_reopen(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.workbook.defined_name import DefinedName

    source = tmp_path / "structural-public-source.xlsx"
    output = tmp_path / "structural-public-output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["A1"] = "Name"
    sheet["B1"] = 10
    sheet["B2"] = "=B1*2"
    sheet["E1"] = "Merged"
    sheet.merge_cells("E1:F1")
    workbook.create_sheet("DeleteMe")["A1"] = "delete"
    workbook.defined_names.add(
        DefinedName("ExistingName", attr_text="Main!$A$1")
    )
    workbook.defined_names.add(
        DefinedName("DeleteName", attr_text="Main!$B$1")
    )
    workbook.save(source)
    edits = [
        {"sheet": "Main", "type": "row_insert", "ref": "2", "count": 1},
        {"sheet": "Main", "type": "row_delete", "ref": "4", "count": 1},
        {"sheet": "Main", "type": "column_insert", "ref": "C", "count": 1},
        {"sheet": "Main", "type": "column_delete", "ref": "D", "count": 1},
        {"sheet": "Main", "type": "row_height", "ref": "2:3", "height": 25},
        {"sheet": "Main", "type": "row_hidden", "ref": "2", "hidden": True},
        {"sheet": "Main", "type": "column_width", "ref": "B", "width": 18},
        {"sheet": "Main", "type": "column_hidden", "ref": "C", "hidden": True},
        {"sheet": "Main", "type": "cells_unmerge", "ref": "E1:F1"},
        {"sheet": "Main", "type": "cells_merge", "ref": "C1:D1"},
        {"sheet": "Main", "type": "range_clear", "ref": "A1", "clear": "contents"},
        {"sheet": "Main", "type": "freeze_panes", "ref": "C3"},
        {"sheet": "Main", "type": "auto_filter", "ref": "A1:D3"},
        {"sheet": "Main", "type": "print_area", "ref": "A1:G5"},
        {"sheet": "Main", "type": "row_page_break", "ref": "4"},
        {"sheet": "Main", "type": "column_page_break", "ref": "E"},
        {
            "sheet": "Main",
            "type": "defined_name_add",
            "name": "NewName",
            "ref": "Main!$A$1",
        },
        {
            "sheet": "Main",
            "type": "defined_name_update",
            "name": "ExistingName",
            "ref": "Main!$B$1",
        },
        {
            "sheet": "Main",
            "type": "defined_name_delete",
            "name": "DeleteName",
        },
        {"sheet": "Main", "type": "sheet_copy", "name": "Clone", "position": 1},
        {"sheet": "Added", "type": "sheet_add", "position": 2},
        {"sheet": "Added", "type": "cell_value", "ref": "A1", "value": "7"},
        {"sheet": "Added", "type": "sheet_reorder", "position": 0},
        {"sheet": "DeleteMe", "type": "sheet_delete"},
    ]
    request = _request(
        tmp_path,
        "structural-edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {"edits": edits, "expected_edits": len(edits)},
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "degraded"
    reopened = load_workbook(output, data_only=False)
    assert reopened.sheetnames == ["Added", "Main", "Clone"]
    main = reopened["Main"]
    assert main["B3"].value == "=B1*2"
    assert main["A1"].value is None
    assert main.row_dimensions[2].height == 25
    assert main.row_dimensions[2].hidden is True
    assert main.column_dimensions["B"].width == 18
    assert main.column_dimensions["C"].hidden is True
    assert [str(item) for item in main.merged_cells.ranges] == ["C1:D1"]
    assert main.freeze_panes == "C3"
    assert main.auto_filter.ref == "A1:D3"
    assert main.print_area == "'Main'!$A$1:$G$5"
    assert [item.id for item in main.row_breaks.brk] == [4]
    assert [item.id for item in main.col_breaks.brk] == [5]
    assert reopened["Clone"]["B3"].value == "=B1*2"
    assert reopened["Added"]["A1"].value == 7
    assert reopened.defined_names["ExistingName"].attr_text == "Main!$B$1"
    assert reopened.defined_names["NewName"].attr_text == "Main!$A$1"
    assert "DeleteName" not in reopened.defined_names


def test_public_create_styles_and_number_formats_reopen(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "styled.xlsx"
    workbook = {
        "metadata": {"title": "Styled XLSX", "creator": "Test", "subject": ""},
        "sheets": [
            {
                "name": "Styled",
                "columns": [
                    {
                        "ref": "A",
                        "width": 20,
                        "hidden": False,
                        "style": {"font": {"color": "#008800"}},
                    },
                    {"ref": "C", "hidden": True, "style": None},
                ],
                "rows": [
                    {
                        "height": 24,
                        "hidden": True,
                        "style": {
                            "font": {"name": "Aptos", "size": 12, "bold": True},
                            "fill": {"pattern": "solid", "color": "#DDEEFF"},
                            "border": {
                                "bottom": {"style": "thin", "color": "#112233"}
                            },
                            "alignment": {
                                "horizontal": "center",
                                "vertical": "top",
                                "wrap": True,
                            },
                            "protection": {"locked": True, "hidden": False},
                        },
                        "cells": [
                            {
                                "ref": "A1",
                                "value": "Header",
                                "type": "s",
                                "style": {
                                    "font": {
                                        "italic": True,
                                        "underline": "single",
                                        "color": "#FF0000",
                                    },
                                    "alignment": {"rotation": 45},
                                },
                            },
                            {"ref": "B1", "value": "Amount", "type": "s"},
                        ],
                    },
                    {
                        "cells": [
                            {
                                "ref": "B2",
                                "value": "1234.5",
                                "type": "n",
                                "style": {
                                    "number_format": {"id": 165},
                                    "alignment": {"horizontal": "right"},
                                    "protection": {"locked": False},
                                },
                            },
                            {
                                "ref": "C2",
                                "value": "8.25",
                                "type": "n",
                                "style": {
                                    "number_format": {"id": 165},
                                    "alignment": {"horizontal": "right"},
                                    "protection": {"locked": False},
                                },
                            },
                        ]
                    },
                ],
                "number_formats": [{"id": 165, "code": "$#,##0.00"}],
            }
        ],
        "defined_names": [],
        "tables": [],
        "chart_reference": None,
        "page_setup": None,
    }
    request = _request(
        tmp_path,
        "styled.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": workbook},
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "success"
    reopened = load_workbook(output, data_only=False)
    sheet = reopened["Styled"]
    header = sheet["A1"]
    assert header.font.name == "Aptos"
    assert header.font.sz == 12
    assert header.font.bold is True
    assert header.font.italic is True
    assert header.font.underline == "single"
    assert header.font.color.rgb == "FFFF0000"
    assert header.fill.fill_type == "solid"
    assert header.fill.fgColor.rgb == "FFDDEEFF"
    assert header.border.bottom.style == "thin"
    assert header.border.bottom.color.rgb == "FF112233"
    assert header.alignment.horizontal == "center"
    assert header.alignment.vertical == "top"
    assert header.alignment.wrap_text is True
    assert header.alignment.text_rotation == 45
    assert header.protection.locked is True
    assert sheet["B2"].number_format == "$#,##0.00"
    assert sheet["B2"].protection.locked is False
    assert sheet["B2"].style_id == sheet["C2"].style_id
    assert sheet.row_dimensions[1].height == 24
    assert sheet.row_dimensions[1].hidden is True
    assert sheet.column_dimensions["A"].width == 20
    assert sheet.column_dimensions["C"].hidden is True

    read_request = _request(
        tmp_path,
        "styled-read.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(output),
            "arguments": {},
        },
    )
    read_result = _public(project_root, "run", "--request", str(read_request))
    cells = {
        cell["ref"]: cell
        for row in read_result["diagnostics"]["operation_result"]["sheets"][0]["rows"]
        for cell in row["cells"]
    }
    assert cells["A1"]["style_index"] > 0
    assert cells["A1"]["style"]["font"]["bold"] is True
    assert cells["A1"]["style"]["font"]["italic"] is True
    assert cells["A1"]["style"]["alignment"]["rotation"] == 45
    assert cells["B2"]["number_format"] == "$#,##0.00"


def test_public_edit_styles_and_number_format_reopen(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "styled-edit.xlsx"
    request = _request(
        tmp_path,
        "styled-edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(public_created),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Sheet1",
                        "type": "cell_style",
                        "ref": "A1",
                        "style": {
                            "font": {"bold": True, "color": "#AA0000"},
                            "fill": {"color": "#FFFF00"},
                            "border": {
                                "right": {"style": "double", "color": "#0000AA"}
                            },
                            "alignment": {
                                "horizontal": "center",
                                "wrap": True,
                                "rotation": -30,
                            },
                            "protection": {"hidden": True},
                        },
                    },
                    {
                        "sheet": "Sheet1",
                        "type": "row_style",
                        "ref": "2",
                        "style": {"font": {"italic": True}},
                    },
                    {
                        "sheet": "Sheet1",
                        "type": "column_style",
                        "ref": "C",
                        "style": {"number_format": {"code": "0.000"}},
                    },
                ],
                "expected_edits": 3,
            },
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] in {"success", "degraded"}
    reopened = load_workbook(output, data_only=False)
    sheet = reopened["Sheet1"]
    assert sheet["A1"].font.bold is True
    assert sheet["A1"].font.color.rgb == "FFAA0000"
    assert sheet["A1"].fill.fgColor.rgb == "FFFFFF00"
    assert sheet["A1"].border.right.style == "double"
    assert sheet["A1"].border.right.color.rgb == "FF0000AA"
    assert sheet["A1"].alignment.horizontal == "center"
    assert sheet["A1"].alignment.wrap_text is True
    assert sheet["A1"].alignment.text_rotation == 120
    assert sheet["A1"].protection.hidden is True
    assert sheet.row_dimensions[2].style_id > 0
    assert sheet.column_dimensions["C"].style_id > 0
    assert sheet.column_dimensions["C"].number_format == "0.000"


def test_public_capabilities_list_xlsx_operations(project_root: Path) -> None:
    report = _public(project_root, "capabilities", "--json")
    operations = {item["operation"]: item for item in report["operations"]}
    assert "xlsx.read" in operations
    assert "xlsx.inspect.structure" in operations
    assert "xlsx.create" in operations
    assert "xlsx.edit" in operations
    assert "xlsx.recalculate" in operations
    assert "xlsx.convert" in operations
    assert "xlsx.template.instantiate" in operations
    assert "xlsx.summary.aggregate" in operations
    assert "xlsx.pivot.create" in operations
    assert "xlsx.validate.schema" in operations
    assert "xlsx.render" in operations
    core_operations = set(operations) - {"xlsx.validate.schema", "xlsx.render"}
    assert all(operations[name]["available"] for name in core_operations)
    assert operations["xlsx.validate.schema"]["providers"] in (
        [],
        ["dotnet-openxml"],
    )
    assert operations["xlsx.render"]["providers"] in ([], ["libreoffice"])


@pytest.mark.parametrize(
    ("operation", "output_name"),
    [
        ("xlsx.validate.schema", None),
        ("xlsx.render", "provider-render.pdf"),
    ],
)
def test_public_provider_operation_is_honestly_unavailable_when_not_callable(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
    operation: str,
    output_name: str | None,
) -> None:
    capabilities = _public(project_root, "capabilities", "--json")
    capability = next(
        item for item in capabilities["operations"]
        if item["operation"] == operation
    )
    if capability["available"]:
        pytest.skip(f"{operation} is callable on this test machine")
    payload = {
        "schema_version": "1.0",
        "operation": operation,
        "input": str(public_created),
        "arguments": {},
    }
    output = None if output_name is None else tmp_path / output_name
    if output is not None:
        payload["output"] = str(output)
    request = _request(tmp_path, f"{operation}.json", payload)

    result = _public(
        project_root,
        "run",
        "--request",
        str(request),
        check=False,
    )

    assert result["status"] == "unavailable"
    assert result["provider_chain"] == []
    assert result["errors"][0]["code"] == "DS_PROVIDER_UNAVAILABLE"
    if output is not None:
        assert not output.exists()


def test_public_macro_read_and_edit_preserve_inert_vba(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = create_package_fixture(tmp_path / "public-macro.xlsm", "xlsm", signed=True)
    read_request = _request(
        tmp_path,
        "public-macro-read.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(source),
            "arguments": {},
        },
    )
    output = tmp_path / "public-macro-edited.xlsm"
    edit_request = _request(
        tmp_path,
        "public-macro-edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "keep_vba": True,
                "recalculation": "skip",
                "edits": [
                    {
                        "sheet": "Sheet1",
                        "type": "cell_value",
                        "ref": "A1",
                        "value": "15",
                    }
                ],
            },
        },
    )

    read_result = _public(project_root, "run", "--request", str(read_request))
    edit_result = _public(project_root, "run", "--request", str(edit_request))

    assert read_result["status"] in {"success", "degraded"}
    assert read_result["diagnostics"]["operation_result"]["macro"][
        "vba_execution"
    ] == "not_executed"
    assert edit_result["status"] == "degraded"
    macro = edit_result["diagnostics"]["operation_result"]["macro"]
    assert macro["vba_payload"] == "preserved"
    assert macro["signature_state"] == "invalidated_by_package_mutation"
    assert output.is_file()
    reopened = load_workbook(output, data_only=False, keep_vba=True)
    assert reopened["Sheet1"]["A1"].value == 15
    assert reopened.vba_archive is not None
    reopened.vba_archive.close()
    reopened.close()


def test_public_template_instantiation_reopens_output(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = create_package_fixture(tmp_path / "public-template.xltx", "xltx")
    output = tmp_path / "public-template-output.xlsx"
    request = _request(
        tmp_path,
        "public-template.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.template.instantiate",
            "input": str(source),
            "output": str(output),
            "arguments": {"recalculation": "skip"},
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] in {"success", "degraded"}
    assert result["provider_chain"] == ["core-python"]
    assert OpcPackage.open(output).workbook_format == "xlsx"
    assert any(
        gate["id"] == "operation.template-content-type-transition"
        and gate["outcome"] == "pass"
        for gate in result["validation"]["gates"]
    )
    reopened = load_workbook(output, data_only=False)
    assert reopened.sheetnames == ["Sheet1"]
    reopened.close()


def test_public_summary_aggregate_reopens_ordinary_native_table(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "public-summary-source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Region", "Category", "Revenue", "Units"])
    sheet.append(["001", "A", 10, 2])
    sheet.append(["010", "A", 20, 4])
    sheet.append(["001", "B", 15, 3])
    sheet.append(["010", "B", 5, 1])
    sheet.append(["100", "C", 1, 1])
    workbook.save(source)
    workbook.close()
    output = tmp_path / "public-summary.xlsx"
    request = _request(
        tmp_path,
        "public-summary.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.summary.aggregate",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "source": {"sheet": "Data", "range": "A1:D6"},
                "group_by": ["Region"],
                "aggregates": [
                    {"column": "Revenue", "function": "sum", "as": "Revenue Total"},
                    {"column": "Units", "function": "average", "as": "Units Average"},
                    {"column": "Revenue", "function": "min", "as": "Revenue Minimum"},
                    {"column": "Revenue", "function": "max", "as": "Revenue Maximum"},
                    {"function": "count", "as": "Rows"},
                    {
                        "column": "Category",
                        "function": "count_nonblank",
                        "as": "Categories Nonblank",
                    },
                    {
                        "column": "Category",
                        "function": "count_distinct",
                        "as": "Categories Distinct",
                    },
                ],
                "sort": [
                    {"column": "Revenue Total", "direction": "desc"},
                    {"column": "Region", "direction": "asc"},
                ],
                "top_n": 2,
                "target": {
                    "sheet": "Summary",
                    "start_cell": "A1",
                    "table_name": "PublicSummary",
                },
            },
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "success"
    summary = result["diagnostics"]["operation_result"]["summary"]
    assert summary["kind"] == "ordinary_table"
    assert summary["native_pivot"] is False
    reopened = load_workbook(output, data_only=False)
    assert reopened["Summary"]["A2"].value == "001"
    assert reopened["Summary"]["B2"].value == 25
    assert reopened["Summary"]["C2"].value == 2.5
    assert reopened["Summary"]["D2"].value == 10
    assert reopened["Summary"]["E2"].value == 15
    assert reopened["Summary"]["F2"].value == 2
    assert reopened["Summary"]["G2"].value == 2
    assert reopened["Summary"]["H2"].value == 2
    assert reopened["Summary"].tables["PublicSummary"].ref == "A1:H3"
    assert not reopened["Summary"]._pivots
    reopened.close()


def test_public_summary_cached_formula_xlsm_preserves_vba(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    workbook = {
        "metadata": {},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Region", "type": "s"},
                            {"ref": "B1", "value": "Revenue", "type": "s"},
                            {"ref": "A2", "value": "East", "type": "s"},
                            {
                                "ref": "B2",
                                "formula": "5+5",
                                "cached_value": "10",
                                "type": "n",
                            },
                        ]
                    }
                ],
                "number_formats": [],
            }
        ],
        "defined_names": [],
        "tables": [],
    }
    source = create_package_fixture(
        tmp_path / "public-summary-macro.xlsm",
        "xlsm",
        signed=True,
        workbook=workbook,
    )
    output = tmp_path / "public-summary-macro-output.xlsm"
    request = _request(
        tmp_path,
        "public-summary-macro.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.summary.aggregate",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "source": {"sheet": "Data", "range": "A1:B2"},
                "group_by": ["Region"],
                "aggregates": [
                    {"column": "Revenue", "function": "sum", "as": "Revenue Total"}
                ],
                "sort": [{"column": "Revenue Total", "direction": "desc"}],
                "top_n": 1,
                "target": {"sheet": "Summary", "table_name": "MacroSummary"},
                "formula_policy": "cached",
                "keep_vba": True,
            },
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "degraded"
    operation_result = result["diagnostics"]["operation_result"]
    assert operation_result["summary"]["source"]["formula_cells_used"] == ["Data!B2"]
    assert operation_result["macro"]["vba_payload"] == "preserved"
    reopened = load_workbook(output, data_only=False, keep_vba=True)
    assert reopened["Summary"]["A2"].value == "East"
    assert reopened["Summary"]["B2"].value == 10
    assert reopened.vba_archive is not None
    reopened.vba_archive.close()
    reopened.close()


def test_public_convert_csv_to_canonical_json(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "public.csv"
    source.write_text("code,value\n00123,12.5\n", encoding="utf-8", newline="\n")
    output = tmp_path / "public.json"
    request = _request(
        tmp_path,
        "public-convert.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.convert",
            "input": str(source),
            "output": str(output),
            "arguments": {"source_format": "csv", "target_format": "json"},
        },
    )

    result = _public(project_root, "run", "--request", str(request))
    document = json.loads(output.read_text(encoding="utf-8"))

    assert result["status"] == "degraded"
    assert result["provider_chain"] == ["core-python"]
    assert document["format"] == "document-skills-tabular"
    assert document["sheets"][0]["rows"][1] == [
        {"type": "string", "value": "00123"},
        {"type": "number", "value": "12.5"},
    ]
    assert result["diagnostics"]["promotion"]["filesystem_state"].startswith(
        "committed"
    )


def test_public_convert_tsv_to_csv_uses_default_dialects(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "public.tsv"
    source.write_text("a\tb\r\n1\t2\r\n", encoding="utf-8", newline="")
    output = tmp_path / "public.csv"
    request = _request(
        tmp_path,
        "public-tsv-convert.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.convert",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "source_format": "tsv",
                "target_format": "csv",
                "values": {"infer_types": False},
            },
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "success"
    assert output.read_text(encoding="utf-8") == "a,b\n1,2\n"


def test_public_convert_json_to_xlsx_reopens_typed_output(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "public-source.json"
    source.write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "format": "document-skills-tabular",
                "sheets": [
                    {
                        "name": "Data",
                        "rows": [[
                            {"type": "string", "value": "Name"},
                            {"type": "number", "value": "12.5"},
                        ]],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    output = tmp_path / "public-output.xlsx"
    request = _request(
        tmp_path,
        "public-json-xlsx-convert.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.convert",
            "input": str(source),
            "output": str(output),
            "arguments": {"source_format": "json", "target_format": "xlsx"},
        },
    )

    result = _public(project_root, "run", "--request", str(request))
    mapped = map_workbook(OpcPackage.open(output))

    assert result["status"] == "success"
    assert mapped["sheets"][0]["rows"][0]["cells"][0]["value"] == "Name"
    assert mapped["sheets"][0]["rows"][0]["cells"][1]["value"] == "12.5"


def test_public_convert_xlsx_to_json_reports_formula_text_loss(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "public-source.xlsx"
    create_xlsx(source, _workbook())
    output = tmp_path / "public-output.json"
    request = _request(
        tmp_path,
        "public-xlsx-json-convert.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.convert",
            "input": str(source),
            "output": str(output),
            "arguments": {"source_format": "xlsx", "target_format": "json"},
        },
    )

    result = _public(project_root, "run", "--request", str(request))
    document = json.loads(output.read_text(encoding="utf-8"))
    losses = {
        item["code"]
        for item in result["diagnostics"]["operation_result"]["semantic_losses"]
    }

    assert result["status"] == "degraded"
    assert document["sheets"][0]["rows"][1][1] == {
        "type": "string",
        "value": "=B1*2",
    }
    assert "formulas-preserved-as-text" in losses


def test_public_recalculate_without_formulas_is_not_applicable_success(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "plain.xlsx"
    workbook = _workbook()
    workbook["sheets"][0]["rows"][0]["cells"] = [
        {"ref": "A1", "value": "Name", "type": "s"},
        {"ref": "B1", "value": "10", "type": "n"},
    ]
    create_xlsx(source, workbook)
    output = tmp_path / "plain-recalculated.xlsx"
    request = _request(
        tmp_path,
        "recalculate-no-formulas.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.recalculate",
            "input": str(source),
            "output": str(output),
            "arguments": {},
        },
    )

    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "success"
    assert result["provider_chain"] == ["core-python"]
    assert result["diagnostics"]["operation_result"]["recalculation"][
        "outcome"
    ] == "not_applicable"
    assert output.is_file()


def test_public_static_formula_validation_blocks_invalid_create(
    project_root: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "invalid-formula.xlsx"
    workbook = _workbook()
    workbook["sheets"][0]["rows"][0]["cells"][2]["formula"] = "Missing!A1"
    request = _request(
        tmp_path,
        "invalid-formula.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": workbook},
        },
    )

    result = _public(
        project_root,
        "run",
        "--request",
        str(request),
        check=False,
    )

    assert result["status"] == "failed"
    assert not output.exists()
    assert any(
        gate["id"] == "operation.formula-static-analysis"
        and gate["outcome"] == "fail"
        for gate in result["validation"]["gates"]
    )


def test_public_external_formula_read_fails_closed_and_inspect_reports(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "external-formula.xlsx"
    workbook = _workbook()
    workbook["sheets"][0]["rows"][0]["cells"][2]["formula"] = (
        "'[Book.xlsx]Sheet 1'!A1"
    )
    create_xlsx(source, workbook)
    read_request = _request(
        tmp_path,
        "external-read.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(source),
            "arguments": {},
        },
    )
    inspect_request = _request(
        tmp_path,
        "external-inspect.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.inspect.structure",
            "input": str(source),
            "arguments": {},
        },
    )

    read_result = _public(
        project_root,
        "run",
        "--request",
        str(read_request),
        check=False,
    )
    inspect_result = _public(
        project_root,
        "run",
        "--request",
        str(inspect_request),
    )

    assert read_result["status"] == "failed"
    assert read_result["errors"][0]["code"] == "DS_ARCHIVE_UNSAFE"
    assert inspect_result["diagnostics"]["operation_result"]["formula_analysis"][
        "categories"
    ]["external_reference"] == 1


def test_public_recalculate_special_formula_fails_closed(
    project_root: Path,
    tmp_path: Path,
) -> None:
    core_source = tmp_path / "normal-formula.xlsx"
    create_xlsx(core_source, _workbook())
    package = OpcPackage.open(core_source)
    root = package.xml("xl/worksheets/sheet1.xml")
    formula = next(
        item
        for item in root.findall(f".//{{{NS['main']}}}f")
        if item.text == "B1*2"
    )
    formula.attrib.update({"t": "array", "ref": "B2"})
    source = tmp_path / "array-formula.xlsx"
    package.write_copy(
        source,
        changed_parts={
            "xl/worksheets/sheet1.xml": tostring(
                root,
                encoding="UTF-8",
                xml_declaration=True,
            )
        },
    )
    output = tmp_path / "must-not-exist.xlsx"
    request = _request(
        tmp_path,
        "recalculate-array.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.recalculate",
            "input": str(source),
            "output": str(output),
            "arguments": {},
        },
    )

    result = _public(
        project_root,
        "run",
        "--request",
        str(request),
        check=False,
    )

    assert result["status"] == "enhancement_required"
    assert not output.exists()


def test_public_doctor_succeeds(project_root: Path) -> None:
    report = _public(project_root, "doctor", "--json")
    assert report["status"] == "healthy"


def test_public_create_and_read(project_root: Path, public_created: Path, tmp_path: Path) -> None:
    read_request = _request(
        tmp_path,
        "read.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(public_created),
            "arguments": {},
        },
    )
    result = _public(project_root, "run", "--request", str(read_request))
    assert result["status"] in {"success", "degraded"}
    assert result["diagnostics"]["operation_result"]["sheet_count"] == 1


def test_public_inspect_inert(project_root: Path, public_created: Path, tmp_path: Path) -> None:
    inspect_request = _request(
        tmp_path,
        "inspect.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.inspect.structure",
            "input": str(public_created),
            "arguments": {},
        },
    )
    result = _public(project_root, "run", "--request", str(inspect_request))
    assert result["status"] == "success"
    assert result["diagnostics"]["operation_result"]["mutation_authorized"] is False


def test_public_edit_produces_distinct_output(project_root: Path, public_created: Path, tmp_path: Path) -> None:
    output = tmp_path / "public-edited.xlsx"
    edit_request = _request(
        tmp_path,
        "edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(public_created),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Sheet1", "type": "cell_value", "ref": "B1", "value": "42"},
                ],
            },
        },
    )
    result = _public(project_root, "run", "--request", str(edit_request))
    assert result["status"] in {"success", "degraded"}
    assert output.is_file()
    assert public_created.is_file()
    assert any(item["path"] == str(output.resolve()) for item in result["artifacts"])
    assert any(
        gate["id"] == "operation.consumer-package-conformance"
        and gate["outcome"] == "pass"
        for gate in result["validation"]["gates"]
    )


def test_public_edit_rejects_consumer_invalid_package_and_preserves_paths(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "core-invalid.xlsx"
    create_xlsx(source, _workbook())
    _strip_style_children(source)
    source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "existing.xlsx"
    existing = b"existing-public-destination"
    output.write_bytes(existing)
    edit_request = _request(
        tmp_path,
        "invalid-edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Sheet1", "type": "cell_value", "ref": "B1", "value": "42"}
                ]
            },
        },
    )

    result = _public(project_root, "run", "--request", str(edit_request), check=False)

    SchemaCatalog(project_root).validate("operation-result", result)
    assert result["status"] == "failed"
    assert result["validation"]["status"] == "fail"
    assert any(
        gate["id"] == "operation.consumer-package-conformance"
        and gate["required"] is True
        and gate["outcome"] == "fail"
        for gate in result["validation"]["gates"]
    )
    assert result["artifacts"] == []
    assert output.read_bytes() == existing
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_sha256


def test_public_validate_reopens_valid_xlsx(project_root: Path, public_created: Path) -> None:
    result = _public(
        project_root,
        "validate",
        "--input",
        str(public_created),
        "--json",
    )
    outcomes = {item["id"]: item["outcome"] for item in result["gates"]}
    assert result["status"] == "pass"
    assert outcomes["provider.reopen"] == "pass"
    assert outcomes["visual.render"] == "unavailable"
    assert outcomes["schema.full"] == "unavailable"


def test_public_unknown_operation_rejected(project_root: Path, tmp_path: Path) -> None:
    request = _request(
        tmp_path,
        "unknown.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.nonexistent",
            "input": str(tmp_path / "input.xlsx"),
            "arguments": {},
        },
    )
    result = _public(project_root, "run", "--request", str(request), check=False)
    assert result["status"] == "invalid_request"
    assert result["errors"][0]["code"] == "DS_OPERATION_UNKNOWN"


def test_public_formula_state_never_claims_unverified_recalculation(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
) -> None:
    """Recalculated formulas must be provider-backed rather than unverified claims."""
    read_request = _request(
        tmp_path,
        "read.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(public_created),
            "arguments": {"include_formulas": True},
        },
    )
    result = _public(project_root, "run", "--request", str(read_request))
    formula_state = result["diagnostics"]["operation_result"]["formula_state"]
    summary = formula_state["summary"]
    if summary["recalculation_provider"] == "unavailable":
        for ref, cell in formula_state["cells"].items():
            assert cell["state"] != "recalculated", f"Cell {ref} falsely reports recalculated"
    assert summary["no_unverified_claimed_recalculated"] is True


def test_public_edit_invalidates_dependents(project_root: Path, public_created: Path, tmp_path: Path) -> None:
    """Editing a precedent marks dependent formulas as recalculation_required."""
    output = tmp_path / "invalidated.xlsx"
    edit_request = _request(
        tmp_path,
        "edit.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(public_created),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Sheet1", "type": "cell_value", "ref": "B1", "value": "99"},
                ],
            },
        },
    )
    result = _public(project_root, "run", "--request", str(edit_request))
    formula_cells = result["diagnostics"]["operation_result"].get("formula_state", {}).get("cells", {})
    # The formula B2=B1*2 should be invalidated because B1 was edited
    invalidated = [
        ref for ref, c in formula_cells.items()
        if c.get("state") == "recalculation_required"
    ]
    assert len(invalidated) > 0, "Expected dependent invalidation"


@pytest.mark.parametrize("operation", ["xlsx.read", "xlsx.inspect.structure"])
@pytest.mark.parametrize(
    "unsafe_priority",
    ["not-a-number", "-1", "9" * 256],
    ids=["malformed", "negative", "overlong"],
)
def test_public_projection_rejects_unsafe_conditional_format_priority(
    project_root: Path,
    tmp_path: Path,
    operation: str,
    unsafe_priority: str,
) -> None:
    source = tmp_path / f"unsafe-priority-{operation.replace('.', '-')}.xlsx"
    workbook = _workbook()
    workbook["sheets"][0]["conditional_formats"] = [
        {
            "ref": "B1:B2",
            "type": "cellIs",
            "priority": 1,
            "operator": "greaterThan",
            "formulas": ["0"],
            "style": {"font": {"color": "#FF0000"}},
        }
    ]
    create_xlsx(source, workbook)
    _set_sheet_numeric_attribute(source, "cfRule", "priority", unsafe_priority)
    request = _request(
        tmp_path,
        f"unsafe-priority-{operation.replace('.', '-')}.json",
        {
            "schema_version": "1.0",
            "operation": operation,
            "input": str(source),
            "arguments": {},
        },
    )

    result = _public(project_root, "run", "--request", str(request), check=False)

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_ARCHIVE_UNSAFE"
    assert unsafe_priority not in json.dumps(result["errors"][0], ensure_ascii=False)


@pytest.mark.parametrize("operation", ["xlsx.read", "xlsx.inspect.structure"])
@pytest.mark.parametrize(
    ("element_name", "attribute", "unsafe_value"),
    [
        ("sheetView", "zoomScale", "-31415926"),
        ("pageMargins", "left", "1e999"),
    ],
    ids=["out-of-range-integer", "non-finite-number"],
)
def test_public_projection_rejects_other_unsafe_numeric_attributes(
    project_root: Path,
    tmp_path: Path,
    operation: str,
    element_name: str,
    attribute: str,
    unsafe_value: str,
) -> None:
    source = tmp_path / f"unsafe-metadata-{operation.replace('.', '-')}.xlsx"
    workbook = _workbook()
    workbook["sheets"][0]["view"] = {
        "show_grid_lines": True,
        "zoom_scale": 100,
        "selected_cell": "A1",
    }
    workbook["sheets"][0]["page_setup"] = {
        "orientation": "portrait",
        "paper_size": "letter",
        "margins": {
            "left": 0.7,
            "right": 0.7,
            "top": 0.75,
            "bottom": 0.75,
            "header": 0.3,
            "footer": 0.3,
        },
        "fit_to_width": 1,
        "fit_to_height": 0,
        "horizontal_centered": False,
        "vertical_centered": False,
    }
    create_xlsx(source, workbook)
    _set_sheet_numeric_attribute(source, element_name, attribute, unsafe_value)
    request = _request(
        tmp_path,
        f"unsafe-metadata-{operation.replace('.', '-')}.json",
        {
            "schema_version": "1.0",
            "operation": operation,
            "input": str(source),
            "arguments": {},
        },
    )

    result = _public(project_root, "run", "--request", str(request), check=False)

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_ARCHIVE_UNSAFE"
    assert unsafe_value not in json.dumps(result["errors"][0], ensure_ascii=False)


@pytest.mark.parametrize("operation", ["xlsx.read", "xlsx.inspect.structure"])
def test_public_projection_rejects_overlong_comment_author_id(
    project_root: Path,
    tmp_path: Path,
    operation: str,
) -> None:
    unsafe_author_id = "7" * 256
    source = tmp_path / f"unsafe-author-{operation.replace('.', '-')}.xlsx"
    workbook = _workbook()
    workbook["sheets"][0]["comments"] = [
        {"ref": "A1", "text": "Review", "author": "Alice"}
    ]
    create_xlsx(source, workbook)
    _set_sheet_numeric_attribute(
        source,
        "comment",
        "authorId",
        unsafe_author_id,
        part="xl/comments1.xml",
    )
    request = _request(
        tmp_path,
        f"unsafe-author-{operation.replace('.', '-')}.json",
        {
            "schema_version": "1.0",
            "operation": operation,
            "input": str(source),
            "arguments": {},
        },
    )

    result = _public(project_root, "run", "--request", str(request), check=False)

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_ARCHIVE_UNSAFE"
    assert unsafe_author_id not in json.dumps(result["errors"][0], ensure_ascii=False)
