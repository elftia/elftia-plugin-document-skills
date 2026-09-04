"""Regression coverage for reviewed DOCX template-pack remediation findings."""

import json
import os
from pathlib import Path
import shutil
import subprocess
from xml.etree.ElementTree import Element, SubElement

from docx import Document
import pytest

import document_skills_core.formats.docx.template_pack_operation as pack_operation
from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.core.io.paths import sha256_file
from document_skills_core.formats.docx.authoring_validation import assert_authoring_parts
from document_skills_core.formats.docx.constants import qn
from document_skills_core.formats.docx.create import create_docx
from document_skills_core.formats.docx.document_spec import parse_document_spec
from document_skills_core.formats.docx.package import OpcPackage, write_deterministic_zip
from document_skills_core.formats.docx.service import DocxService


def _request(operation: str, **values: object) -> dict[str, object]:
    request: dict[str, object] = {
        "schema_version": "1.0",
        "operation": operation,
        "arguments": {},
    }
    request.update(values)
    return request


def _raw_report(text: str) -> dict[str, object]:
    return {
        "blocks": [{"type": "paragraph", "text": text, "style": "Normal"}],
        "image": None,
        "header": None,
        "footer": None,
        "sections": [{"orientation": "portrait", "title": None}],
        "metadata": {"title": "Synthetic", "subject": "", "creator": "", "keywords": ""},
        "style_profile": None,
    }


def test_template_authoring_uses_snapshot_after_original_pack_mutates(
    project_root: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "source.docx"
    create_docx(source, _raw_report("ORIGINAL {name}"))
    pack_root = tmp_path / "pack"
    imported = DocxService(project_root).execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(source),
            output=str(pack_root),
            arguments={
                "expected_source_sha256": sha256_file(source),
                "manifest": {
                    "id": "snapshot-race",
                    "version": "1.0.0",
                    "display_name": "Snapshot Race",
                    "description": "Synthetic non-private substitution regression.",
                    "modes": ["template"],
                    "variables": [{"name": "name", "required": True}],
                    "regions": [],
                    "role_styles": {},
                    "compatibility": {
                        "document_spec_versions": [],
                        "domain_profiles": [],
                        "template_engine_versions": ["1.0"],
                        "style_profile_versions": [],
                    },
                },
                "provenance": {
                    "source_kind": "local",
                    "imported_at": "2026-08-30T00:00:00Z",
                    "redistributable": False,
                },
            },
        ),
    )
    assert imported["status"] == "success", imported
    original_payload_sha256 = sha256_file(pack_root / "template.docx")
    substitute = tmp_path / "substitute.docx"
    create_docx(substitute, _raw_report("MUTATED {name}"))
    original_snapshot = pack_operation.snapshot_pack_reference

    def snapshot_then_mutate(*args: object, **kwargs: object):
        snapshot = original_snapshot(*args, **kwargs)
        shutil.copyfile(substitute, pack_root / "template.docx")
        return snapshot

    monkeypatch.setattr(
        pack_operation,
        "snapshot_pack_reference",
        snapshot_then_mutate,
    )
    output = tmp_path / "authored.docx"
    result = DocxService(project_root).execute(
        "docx.template.pack.instantiate",
        _request(
            "docx.template.pack.instantiate",
            output=str(output),
            arguments={
                "pack": {
                    "kind": "local",
                    "path": str(pack_root),
                    "expected_manifest_sha256": imported["diagnostics"]["operation_result"]["manifest_sha256"],
                },
                "mode": "template",
                "variables": {"name": "Agent"},
            },
        ),
    )
    assert result["status"] == "success", json.dumps(result, default=str)
    document_xml = OpcPackage.open(output).parts["word/document.xml"]
    assert b"ORIGINAL Agent" in document_xml
    assert b"MUTATED Agent" not in document_xml
    pack_result = result["diagnostics"]["operation_result"]["pack"]
    assert pack_result["authoring_source"] == "operation-owned-verified-snapshot"
    assert pack_result["payload_sha256"] == original_payload_sha256
    assert sha256_file(pack_root / "template.docx") != original_payload_sha256


def test_document_spec_authoring_uses_snapshot_after_original_pack_mutates(
    project_root: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    builtin = (
        project_root
        / "skills/document-docx/assets/template-packs/general-academic-paper/1.0.0"
    )
    pack_root = tmp_path / "document-spec-pack"
    shutil.copytree(builtin, pack_root)
    manifest_sha256 = sha256_file(pack_root / "manifest.json")
    original_payload_sha256 = sha256_file(pack_root / "template.docx")
    substitute = tmp_path / "substitute.docx"
    create_docx(substitute, _raw_report("MUTATED DOCUMENT SPEC SOURCE"))
    assert sha256_file(substitute) != original_payload_sha256
    original_snapshot = pack_operation.snapshot_pack_reference

    def snapshot_then_mutate(*args: object, **kwargs: object):
        snapshot = original_snapshot(*args, **kwargs)
        shutil.copyfile(substitute, pack_root / "template.docx")
        return snapshot

    monkeypatch.setattr(
        pack_operation,
        "snapshot_pack_reference",
        snapshot_then_mutate,
    )
    output = tmp_path / "document-spec-authored.docx"
    result = DocxService(project_root).execute(
        "docx.template.pack.instantiate",
        _request(
            "docx.template.pack.instantiate",
            output=str(output),
            arguments={
                "pack": {
                    "kind": "local",
                    "path": str(pack_root),
                    "expected_manifest_sha256": manifest_sha256,
                },
                "mode": "document-spec",
                "document_spec": _academic_spec(),
            },
        ),
    )
    assert result["status"] == "success", result
    styles = OpcPackage.open(output).xml("word/styles.xml")
    body = _style(styles, "ElftiaBody")
    fonts = body.find(f"./{qn('w', 'rPr')}/{qn('w', 'rFonts')}")
    assert fonts is not None
    assert fonts.attrib[qn("w", "ascii")] == "Times New Roman"
    assert fonts.attrib[qn("w", "eastAsia")] == "SimSun"
    pack_result = result["diagnostics"]["operation_result"]["pack"]
    assert pack_result == {
        "id": "general-academic-paper",
        "version": "1.0.0",
        "manifest_sha256": manifest_sha256,
        "payload_sha256": original_payload_sha256,
        "mode": "document-spec",
        "authoring_source": "operation-owned-verified-snapshot",
    }
    assert sha256_file(pack_root / "template.docx") != original_payload_sha256


def test_import_regions_match_direct_body_executor_eligibility(
    project_root: Path,
    tmp_path: Path,
) -> None:
    reference = tmp_path / "regions.docx"
    create_docx(
        reference,
        {
            **_raw_report("Direct {direct}"),
            "blocks": [
                {"type": "paragraph", "text": "Direct {direct}", "style": "Normal"},
                {
                    "type": "table",
                    "style": "TableGrid",
                    "rows": [["Nested {nested}"]],
                },
            ],
        },
    )
    inspected = DocxService(project_root).execute(
        "docx.template.import.inspect",
        _request(
            "docx.template.import.inspect",
            input=str(reference),
            arguments={"expected_source_sha256": sha256_file(reference)},
        ),
    )
    assert inspected["status"] == "success", inspected
    operation = inspected["diagnostics"]["operation_result"]
    assert {item["name"] for item in operation["variables"]} == {"direct", "nested"}
    assert {item["expected_text"] for item in operation["regions"]} == {"Direct {direct}"}
    assert operation["mode_compatibility"]["document-spec"]["status"] == "degraded"


def test_import_inventory_reports_material_direct_format_and_mode_loss(
    project_root: Path,
    tmp_path: Path,
) -> None:
    reference = tmp_path / "direct-format.docx"
    create_docx(
        reference,
        {
            **_raw_report("Synthetic direct-format paragraph"),
            "blocks": [
                {
                    "type": "paragraph",
                    "text": "Synthetic direct-format paragraph",
                    "style": "Normal",
                },
                {
                    "type": "table",
                    "style": "TableGrid",
                    "rows": [["Metric", "Value"], ["A", "1"]],
                    "borders": {
                        "top": {"style": "single", "size_eighth_points": 8, "color": "000000"},
                        "header_bottom": {"style": "single", "size_eighth_points": 8, "color": "000000"},
                        "bottom": {"style": "single", "size_eighth_points": 8, "color": "000000"},
                    },
                },
            ],
            "footer": {
                "paragraphs": [
                    {"alignment": "center", "runs": [{"field": "PAGE"}]}
                ]
            },
        },
    )
    package = OpcPackage.open(reference)
    document = package.xml("word/document.xml")
    paragraph = document.find(f"./{qn('w', 'body')}/{qn('w', 'p')}")
    assert paragraph is not None
    properties = paragraph.find(qn("w", "pPr"))
    if properties is None:
        properties = Element(qn("w", "pPr"))
        paragraph.insert(0, properties)
    SubElement(properties, qn("w", "ind"), {qn("w", "firstLineChars"): "200"})
    SubElement(
        properties,
        qn("w", "spacing"),
        {qn("w", "after"): "120", qn("w", "line"): "276", qn("w", "lineRule"): "auto"},
    )
    run = paragraph.find(qn("w", "r"))
    assert run is not None
    run_properties = SubElement(run, qn("w", "rPr"))
    SubElement(
        run_properties,
        qn("w", "rFonts"),
        {
            qn("w", "ascii"): "Noto Serif",
            qn("w", "hAnsi"): "Noto Serif",
            qn("w", "eastAsia"): "Noto Serif CJK SC",
            qn("w", "cs"): "Noto Serif",
        },
    )
    SubElement(run_properties, qn("w", "sz"), {qn("w", "val"): "22"})
    parts = dict(package.parts)
    from document_skills_core.formats.docx.xml_utils import xml_bytes

    parts["word/document.xml"] = xml_bytes(document)
    write_deterministic_zip(reference, parts)
    inspected = DocxService(project_root).execute(
        "docx.template.import.inspect",
        _request(
            "docx.template.import.inspect",
            input=str(reference),
            arguments={"expected_source_sha256": sha256_file(reference)},
        ),
    )
    assert inspected["status"] == "success", inspected
    operation = inspected["diagnostics"]["operation_result"]
    inventory = operation["format_inventory"]
    assert any(
        item["format"].get("font_eastAsia") == "Noto Serif CJK SC"
        for item in inventory["direct_run_formats"]
    )
    assert any(
        item["format"].get("indent_firstLineChars") == "200"
        for item in inventory["direct_paragraph_formats"]
    )
    assert inventory["tables"][0]["grid_positive"] is True
    assert inventory["tables"][0]["borders"]["top"]["val"] == "single"
    assert inventory["stories"][0]["fields"] == [
        {"instruction": "PAGE", "allowlisted": True}
    ]
    assert operation["mode_compatibility"]["template"]["status"] == "compatible"
    document_spec = operation["mode_compatibility"]["document-spec"]
    assert document_spec["status"] == "degraded"
    assert "direct-run-format-requires-explicit-authoring-map" in document_spec["reasons"]
    assert "story-fields-require-explicit-authoring-map" in document_spec["reasons"]


def test_explicit_authoring_format_import_read_and_instantiate_round_trip(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "explicit-authoring-format.docx"
    profile = _create_formatted_reference(source)
    builtin_manifest = json.loads(
        (
            project_root
            / "skills/document-docx/assets/template-packs/general-academic-paper/1.0.0/manifest.json"
        ).read_text(encoding="utf-8")
    )
    role_styles = builtin_manifest["capabilities"]["role_styles"]
    service = DocxService(project_root)

    missing = service.execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(source),
            output=str(tmp_path / "missing-profile-pack"),
            arguments=_formatted_import_arguments(source, role_styles, None),
        ),
    )
    assert missing["status"] == "enhancement_required", missing
    assert missing["errors"][0]["details"]["diagnostic"] == (
        "DS_DOCX_TEMPLATE_AUTHORING_FORMAT_REQUIRED"
    )

    lossy_profile = json.loads(json.dumps(profile))
    lossy_profile["role_formats"]["paragraph"]["latin_font"] = "Invented Font"
    lossy = service.execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(source),
            output=str(tmp_path / "lossy-profile-pack"),
            arguments=_formatted_import_arguments(
                source,
                role_styles,
                lossy_profile,
            ),
        ),
    )
    assert lossy["status"] == "enhancement_required", lossy
    assert lossy["errors"][0]["details"]["diagnostic"] == (
        "DS_DOCX_TEMPLATE_AUTHORING_FORMAT_LOSSY"
    )

    pack_root = tmp_path / "explicit-profile-pack"
    imported = service.execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(source),
            output=str(pack_root),
            arguments=_formatted_import_arguments(source, role_styles, profile),
        ),
    )
    assert imported["status"] == "success", json.dumps(imported, ensure_ascii=False)
    pack = imported["diagnostics"]["operation_result"]["pack"]
    local = {
        "kind": "local",
        "path": str(pack_root),
        "expected_manifest_sha256": pack["manifest_sha256"],
    }
    read = service.execute(
        "docx.template.pack.read",
        _request("docx.template.pack.read", arguments={"pack": local}),
    )
    assert read["status"] == "success", read
    read_capabilities = read["diagnostics"]["operation_result"]["pack"]["capabilities"]
    assert read_capabilities["authoring_format"] == profile
    assert read_capabilities["mode_compatibility"]["document-spec"] == {
        "status": "compatible",
        "preservation": "reconstruction",
        "reasons": [],
        "unsupported_fields": [],
        "mapped_by": "explicit-authoring-format/v1",
    }

    output = tmp_path / "explicit-profile.docx"
    instantiated = service.execute(
        "docx.template.pack.instantiate",
        _request(
            "docx.template.pack.instantiate",
            output=str(output),
            arguments={
                "pack": local,
                "mode": "document-spec",
                "document_spec": _academic_spec(),
            },
        ),
    )
    assert instantiated["status"] == "success", instantiated
    package = OpcPackage.open(output)
    styles = package.xml("word/styles.xml")
    body = _style(styles, "ElftiaBody")
    body_fonts = body.find(f"./{qn('w', 'rPr')}/{qn('w', 'rFonts')}")
    body_size = body.find(f"./{qn('w', 'rPr')}/{qn('w', 'sz')}")
    body_alignment = body.find(f"./{qn('w', 'pPr')}/{qn('w', 'jc')}")
    body_indent = body.find(f"./{qn('w', 'pPr')}/{qn('w', 'ind')}")
    body_spacing = body.find(f"./{qn('w', 'pPr')}/{qn('w', 'spacing')}")
    assert body_fonts is not None
    assert body_fonts.attrib[qn("w", "ascii")] == "Noto Serif"
    assert body_fonts.attrib[qn("w", "eastAsia")] == "Noto Serif CJK SC"
    assert body_size is not None and body_size.attrib[qn("w", "val")] == "22"
    assert body_alignment is not None and body_alignment.attrib[qn("w", "val")] == "both"
    assert body_indent is not None and body_indent.attrib[qn("w", "firstLineChars")] == "200"
    assert body_spacing is not None
    assert body_spacing.attrib[qn("w", "after")] == "120"
    assert body_spacing.attrib[qn("w", "line")] == "276"

    bibliography = _style(styles, "ElftiaBibliography")
    bibliography_fonts = bibliography.find(f"./{qn('w', 'rPr')}/{qn('w', 'rFonts')}")
    bibliography_size = bibliography.find(f"./{qn('w', 'rPr')}/{qn('w', 'sz')}")
    bibliography_indent = bibliography.find(f"./{qn('w', 'pPr')}/{qn('w', 'ind')}")
    assert bibliography_fonts is not None
    assert bibliography_fonts.attrib[qn("w", "ascii")] == "Liberation Serif"
    assert bibliography_fonts.attrib[qn("w", "eastAsia")] == "Noto Serif CJK SC"
    assert bibliography_size is not None
    assert bibliography_size.attrib[qn("w", "val")] == "20"
    assert bibliography_indent is not None
    assert bibliography_indent.attrib[qn("w", "hanging")] == "360"
    assert bibliography_indent.attrib[qn("w", "left")] == "360"

    table = next(package.xml("word/document.xml").iter(qn("w", "tbl")))
    grid = table.find(qn("w", "tblGrid"))
    assert grid is not None
    assert [int(item.attrib[qn("w", "w")]) for item in grid] == [4_680, 4_680]
    borders = table.find(f"./{qn('w', 'tblPr')}/{qn('w', 'tblBorders')}")
    assert borders is not None
    assert borders.find(qn("w", "top")).attrib[qn("w", "val")] == "single"
    assert borders.find(qn("w", "bottom")).attrib[qn("w", "val")] == "single"
    footer = package.xml("word/footer1.xml")
    assert [
        " ".join((item.text or "").split())
        for item in footer.iter(qn("w", "instrText"))
    ] == ["PAGE", "NUMPAGES"]


def test_reference_import_rejects_footer_paragraph_inventory_overflow(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "footer-overflow.docx"
    profile = _create_formatted_reference(source)
    _replace_footer_paragraphs(source, count=17)
    bounded_profile = json.loads(json.dumps(profile))
    bounded_profile["default_footer"] = {
        "paragraphs": [
            {"alignment": "center", "runs": [{"text": "Synthetic footer"}]}
            for _paragraph_index in range(16)
        ]
    }
    builtin_manifest = json.loads(
        (
            project_root
            / "skills/document-docx/assets/template-packs/general-academic-paper/1.0.0/manifest.json"
        ).read_text(encoding="utf-8")
    )
    service = DocxService(project_root)

    inspected = service.execute(
        "docx.template.import.inspect",
        _request(
            "docx.template.import.inspect",
            input=str(source),
            arguments={"expected_source_sha256": sha256_file(source)},
        ),
    )
    assert inspected["status"] == "success", inspected
    operation = inspected["diagnostics"]["operation_result"]
    footer = next(
        story
        for story in operation["format_inventory"]["stories"]
        if story["kind"] == "footer"
    )
    assert footer["paragraph_count"] == 17
    assert footer["paragraphs_truncated"] is True
    assert len(footer["paragraphs"]) == 16
    document_spec = operation["mode_compatibility"]["document-spec"]
    assert document_spec["status"] == "incompatible"
    assert "footer-paragraph-inventory-overflow" in document_spec["reasons"]
    assert any(
        item["code"] == "DS_DOCX_TEMPLATE_FORMAT_INVENTORY_OVERFLOW"
        for item in operation["unsupported_features"]
    )

    template_source = tmp_path / "footer-overflow-template-source.docx"
    create_docx(
        template_source,
        {**_raw_report("Template base copy"), "footer": "Synthetic footer"},
    )
    _replace_footer_paragraphs(template_source, count=17)
    template_inspection = service.execute(
        "docx.template.import.inspect",
        _request(
            "docx.template.import.inspect",
            input=str(template_source),
            arguments={"expected_source_sha256": sha256_file(template_source)},
        ),
    )
    assert template_inspection["status"] == "success", template_inspection
    assert template_inspection["diagnostics"]["operation_result"][
        "mode_compatibility"
    ]["template"]["status"] == "compatible"
    template_pack = tmp_path / "footer-overflow-template-pack"
    template_import = service.execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(template_source),
            output=str(template_pack),
            arguments={
                "expected_source_sha256": sha256_file(template_source),
                "manifest": {
                    "id": "footer-overflow-template",
                    "version": "1.0.0",
                    "display_name": "Footer Overflow Template",
                    "description": "Synthetic base-copy fidelity regression.",
                    "modes": ["template"],
                    "variables": [],
                    "regions": [],
                    "role_styles": {},
                    "compatibility": {
                        "document_spec_versions": [],
                        "domain_profiles": [],
                        "template_engine_versions": ["1.0"],
                        "style_profile_versions": [],
                    },
                },
                "provenance": {
                    "source_kind": "local",
                    "imported_at": "2026-08-30T00:00:00Z",
                    "redistributable": False,
                },
            },
        ),
    )
    assert template_import["status"] == "success", template_import
    template_output = tmp_path / "footer-overflow-template.docx"
    template_instantiation = service.execute(
        "docx.template.pack.instantiate",
        _request(
            "docx.template.pack.instantiate",
            output=str(template_output),
            arguments={
                "pack": {
                    "kind": "local",
                    "path": str(template_pack),
                    "expected_manifest_sha256": template_import["diagnostics"][
                        "operation_result"
                    ]["manifest_sha256"],
                },
                "mode": "template",
                "variables": {},
            },
        ),
    )
    assert template_instantiation["status"] == "success", template_instantiation
    assert len(OpcPackage.open(template_output).xml("word/footer1.xml").findall(qn("w", "p"))) == 17

    pack_root = tmp_path / "footer-overflow-pack"
    imported = service.execute(
        "docx.template.import.create",
        _request(
            "docx.template.import.create",
            input=str(source),
            output=str(pack_root),
            arguments=_formatted_import_arguments(
                source,
                builtin_manifest["capabilities"]["role_styles"],
                bounded_profile,
            ),
        ),
    )
    assert imported["status"] == "enhancement_required", imported
    details = imported["errors"][0]["details"]
    assert details["diagnostic"] == "DS_DOCX_TEMPLATE_FORMAT_INVENTORY_OVERFLOW"
    assert "footer-paragraph-inventory-overflow" in details["reasons"]
    assert not pack_root.exists()


def test_every_capped_format_inventory_reports_limit_plus_one_overflow(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "all-format-overflows.docx"
    _create_inventory_overflow_reference(source)
    inspected = DocxService(project_root).execute(
        "docx.template.import.inspect",
        _request(
            "docx.template.import.inspect",
            input=str(source),
            arguments={"expected_source_sha256": sha256_file(source)},
        ),
    )
    assert inspected["status"] == "success", inspected
    operation = inspected["diagnostics"]["operation_result"]
    overflows = operation["format_inventory"]["overflows"]
    assert {item["reason"] for item in overflows} == {
        "direct-run-format-inventory-overflow",
        "direct-paragraph-format-inventory-overflow",
        "story-field-inventory-overflow",
        "table-cell-width-inventory-overflow",
        "table-inventory-overflow",
    }
    assert all(item["observed_at_least"] == item["limit"] + 1 for item in overflows)
    compatibility = operation["mode_compatibility"]
    assert compatibility["template"]["status"] == "compatible"
    assert compatibility["document-spec"]["status"] == "incompatible"


@pytest.mark.slow
def test_academic_pack_emits_strict_geometry_typography_layout_and_fields(
    project_root: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "academic.docx"
    result = DocxService(project_root).execute(
        "docx.template.pack.instantiate",
        _request(
            "docx.template.pack.instantiate",
            output=str(output),
            arguments={
                "pack": {
                    "kind": "builtin",
                    "id": "general-academic-paper",
                    "version": "1.0.0",
                },
                "mode": "document-spec",
                "document_spec": _academic_spec(),
            },
        ),
    )
    assert result["status"] == "success", json.dumps(result, default=str)
    _strict_traverse(output)
    package = OpcPackage.open(output)
    document = package.xml("word/document.xml")
    table = next(document.iter(qn("w", "tbl")))
    grid = table.find(qn("w", "tblGrid"))
    assert grid is not None
    grid_widths = [int(item.attrib[qn("w", "w")]) for item in grid]
    assert grid_widths == [4_680, 4_680]
    table_width = table.find(f"./{qn('w', 'tblPr')}/{qn('w', 'tblW')}")
    assert table_width is not None and table_width.attrib[qn("w", "w")] == "9360"
    assert all(
        int(width.attrib[qn("w", "w")]) > 0
        for width in table.iter(qn("w", "tcW"))
    )
    borders = table.find(f"./{qn('w', 'tblPr')}/{qn('w', 'tblBorders')}")
    assert borders is not None
    assert borders.find(qn("w", "top")).attrib[qn("w", "val")] == "single"
    assert borders.find(qn("w", "bottom")).attrib[qn("w", "val")] == "single"
    assert borders.find(qn("w", "insideV")).attrib[qn("w", "val")] == "nil"

    styles = package.xml("word/styles.xml")
    body_style = _style(styles, "ElftiaBody")
    body_fonts = body_style.find(f"./{qn('w', 'rPr')}/{qn('w', 'rFonts')}")
    assert body_fonts is not None
    assert body_fonts.attrib[qn("w", "ascii")] == "Times New Roman"
    assert body_fonts.attrib[qn("w", "eastAsia")] == "SimSun"
    body_indent = body_style.find(f"./{qn('w', 'pPr')}/{qn('w', 'ind')}")
    assert body_indent is not None and body_indent.attrib[qn("w", "firstLineChars")] == "200"
    bibliography = _style(styles, "ElftiaBibliography")
    bibliography_indent = bibliography.find(f"./{qn('w', 'pPr')}/{qn('w', 'ind')}")
    assert bibliography_indent is not None
    assert bibliography_indent.attrib[qn("w", "hanging")] == "360"
    assert bibliography_indent.attrib[qn("w", "left")] == "360"

    footer = package.xml("word/footer1.xml")
    alignment = footer.find(f"./{qn('w', 'p')}/{qn('w', 'pPr')}/{qn('w', 'jc')}")
    assert alignment is not None and alignment.attrib[qn("w", "val")] == "center"
    instructions = [" ".join((item.text or "").split()) for item in footer.iter(qn("w", "instrText"))]
    assert instructions == ["PAGE", "NUMPAGES"]
    field_types = [item.attrib[qn("w", "fldCharType")] for item in footer.iter(qn("w", "fldChar"))]
    assert field_types == ["begin", "separate", "end", "begin", "separate", "end"]
    _libreoffice_open_and_render(output, tmp_path / "libreoffice")


def test_legacy_story_and_row_only_table_remain_strict_consumer_valid(
    project_root: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "legacy.docx"
    result = DocxService(project_root).execute(
        "docx.create",
        _request(
            "docx.create",
            output=str(output),
            arguments={
                "report": {
                    "metadata": {"title": "Legacy"},
                    "blocks": [
                        {"type": "table", "rows": [["A", "B"], ["1", "2"]]},
                    ],
                    "footer": "Legacy footer",
                }
            },
        ),
    )
    assert result["status"] == "success", result
    _strict_traverse(output)
    package = OpcPackage.open(output)
    footer = package.xml("word/footer1.xml")
    assert next(footer.iter(qn("w", "t"))).text == "Legacy footer"
    assert list(footer.iter(qn("w", "instrText"))) == []
    grid = package.xml("word/document.xml").find(f".//{qn('w', 'tblGrid')}")
    assert grid is not None
    assert [int(item.attrib[qn("w", "w")]) for item in grid] == [4_680, 4_680]


def test_invalid_story_field_and_table_geometry_fail_before_publication(
    tmp_path: Path,
) -> None:
    invalid_field = _academic_spec()
    invalid_field["footer"] = {
        "paragraphs": [{"alignment": "center", "runs": [{"field": "INCLUDETEXT"}]}]
    }
    with pytest.raises(DocumentSkillsError):
        parse_document_spec(invalid_field)
    invalid_table = _academic_spec()
    invalid_table["nodes"][6].update(
        {"width_twips": 9_360, "column_widths_twips": [4_000, 4_000]}
    )
    with pytest.raises(DocumentSkillsError):
        parse_document_spec(invalid_table)
    assert not (tmp_path / "invalid.docx").exists()


def test_empty_grid_synthetic_failure_shape_is_rejected(
    tmp_path: Path,
) -> None:
    valid = tmp_path / "valid.docx"
    create_docx(
        valid,
        {
            **_raw_report("Synthetic failure shape"),
            "blocks": [
                {"type": "table", "style": "TableGrid", "rows": [["A", "B"]]},
            ],
        },
    )
    package = OpcPackage.open(valid)
    document = package.xml("word/document.xml")
    grid = document.find(f".//{qn('w', 'tblGrid')}")
    assert grid is not None
    original_grid = list(grid)
    for child in list(grid):
        grid.remove(child)
    parts = dict(package.parts)
    from document_skills_core.formats.docx.xml_utils import xml_bytes

    parts["word/document.xml"] = xml_bytes(document)
    for name in ("_rels/.rels", "word/_rels/document.xml.rels"):
        parts[name] = _prefix_relationship_namespace(parts[name])
    with pytest.raises(DocumentSkillsError):
        assert_authoring_parts(parts)
    broken = tmp_path / "empty-grid.docx"
    write_deterministic_zip(broken, parts)
    # A basic consumer may still traverse this shape, which is why an empty
    # tblGrid cannot be accepted as proof of strict compatibility by itself.
    opened = Document(broken)
    assert len(opened.tables[0].rows[0].cells) == 2
    for child in original_grid:
        grid.append(child)
    parts["word/document.xml"] = xml_bytes(document)
    with pytest.raises(DocumentSkillsError) as captured:
        assert_authoring_parts(parts)
    assert captured.value.details["reason"] == "relationship-namespace-prefix"


def _create_formatted_reference(path: Path) -> dict[str, object]:
    rule = {"style": "single", "size_eighth_points": 8, "color": "000000"}
    footer = {
        "paragraphs": [
            {
                "alignment": "center",
                "runs": [
                    {"field": "PAGE"},
                    {"text": " / "},
                    {"field": "NUMPAGES"},
                ],
            }
        ]
    }
    profile: dict[str, object] = {
        "version": "1.0",
        "role_formats": {
            "paragraph": {
                "latin_font": "Noto Serif",
                "east_asia_font": "Noto Serif CJK SC",
                "size_half_points": 22,
                "alignment": "both",
                "first_line_chars": 200,
                "space_after_twips": 120,
                "line_twips": 276,
            },
            "bibliography": {
                "latin_font": "Liberation Serif",
                "east_asia_font": "Noto Serif CJK SC",
                "size_half_points": 20,
                "alignment": "left",
                "hanging_twips": 360,
                "left_twips": 360,
                "space_after_twips": 80,
            },
        },
        "default_footer": footer,
        "table": {
            "width_twips": 9_360,
            "borders": {"top": rule, "header_bottom": rule, "bottom": rule},
        },
    }
    create_docx(
        path,
        {
            **_raw_report("Body"),
            "blocks": [
                {"type": "paragraph", "text": "正文 Body", "style": "ElftiaBody"},
                {
                    "type": "paragraph",
                    "text": "Reference entry",
                    "style": "ElftiaBibliography",
                },
                {
                    "type": "table",
                    "style": "ElftiaTable",
                    "rows": [["Metric", "Value"], ["A", "1"]],
                    "width_twips": 9_360,
                    "column_widths_twips": [4_680, 4_680],
                    "borders": {"top": rule, "header_bottom": rule, "bottom": rule},
                },
            ],
            "footer": footer,
            "style_profile": {"id": "academic-pack", "version": "1.0"},
        },
    )
    package = OpcPackage.open(path)
    document = package.xml("word/document.xml")
    paragraphs = document.findall(f"./{qn('w', 'body')}/{qn('w', 'p')}")
    _set_direct_format(
        paragraphs[0],
        latin="Noto Serif",
        east_asia="Noto Serif CJK SC",
        size=22,
        alignment="both",
        indentation={"firstLineChars": 200},
        spacing={"after": 120, "line": 276, "lineRule": "auto"},
    )
    _set_direct_format(
        paragraphs[1],
        latin="Liberation Serif",
        east_asia="Noto Serif CJK SC",
        size=20,
        alignment="left",
        indentation={"hanging": 360, "left": 360},
        spacing={"after": 80},
    )
    parts = dict(package.parts)
    from document_skills_core.formats.docx.xml_utils import xml_bytes

    parts["word/document.xml"] = xml_bytes(document)
    write_deterministic_zip(path, parts)
    return profile


def _replace_footer_paragraphs(path: Path, *, count: int) -> None:
    package = OpcPackage.open(path)
    footer = package.xml("word/footer1.xml")
    for child in list(footer):
        footer.remove(child)
    for _paragraph_index in range(count):
        paragraph = SubElement(footer, qn("w", "p"))
        properties = SubElement(paragraph, qn("w", "pPr"))
        SubElement(properties, qn("w", "jc"), {qn("w", "val"): "center"})
        run = SubElement(paragraph, qn("w", "r"))
        SubElement(run, qn("w", "t")).text = "Synthetic footer"
    parts = dict(package.parts)
    from document_skills_core.formats.docx.xml_utils import xml_bytes

    parts["word/footer1.xml"] = xml_bytes(footer)
    write_deterministic_zip(path, parts)


def _create_inventory_overflow_reference(path: Path) -> None:
    footer = {
        "paragraphs": [
            {
                "alignment": "center",
                "runs": [{"field": "PAGE"} for _run_index in range(17)],
            }
            for _paragraph_index in range(16)
        ]
    }
    tables = [
        {
            "type": "table",
            "style": "TableGrid",
            "rows": [[f"row-{row_index}"] for row_index in range(257)],
        }
    ]
    tables.extend(
        {
            "type": "table",
            "style": "TableGrid",
            "rows": [[f"table-{table_index}"]],
        }
        for table_index in range(1, 257)
    )
    create_docx(
        path,
        {
            **_raw_report("Overflow inventory"),
            "blocks": [
                {
                    "type": "paragraph",
                    "text": f"Paragraph {paragraph_index}",
                    "style": "Normal",
                }
                for paragraph_index in range(257)
            ]
            + tables,
            "footer": footer,
        },
    )
    package = OpcPackage.open(path)
    document = package.xml("word/document.xml")
    paragraphs = document.findall(f"./{qn('w', 'body')}/{qn('w', 'p')}")
    assert len(paragraphs) == 257
    for paragraph_index, paragraph in enumerate(paragraphs):
        _set_direct_format(
            paragraph,
            latin=f"Synthetic Font {paragraph_index:03d}",
            east_asia="Synthetic CJK",
            size=22,
            alignment="left",
            indentation={"firstLineChars": paragraph_index + 1},
            spacing={"after": 20},
        )
    parts = dict(package.parts)
    from document_skills_core.formats.docx.xml_utils import xml_bytes

    parts["word/document.xml"] = xml_bytes(document)
    write_deterministic_zip(path, parts)


def _set_direct_format(
    paragraph: object,
    *,
    latin: str,
    east_asia: str,
    size: int,
    alignment: str,
    indentation: dict[str, int],
    spacing: dict[str, int | str],
) -> None:
    properties = paragraph.find(qn("w", "pPr"))
    assert properties is not None
    SubElement(properties, qn("w", "jc"), {qn("w", "val"): alignment})
    SubElement(
        properties,
        qn("w", "ind"),
        {qn("w", key): str(value) for key, value in indentation.items()},
    )
    SubElement(
        properties,
        qn("w", "spacing"),
        {qn("w", key): str(value) for key, value in spacing.items()},
    )
    run = paragraph.find(qn("w", "r"))
    assert run is not None
    run_properties = SubElement(run, qn("w", "rPr"))
    SubElement(
        run_properties,
        qn("w", "rFonts"),
        {
            qn("w", "ascii"): latin,
            qn("w", "hAnsi"): latin,
            qn("w", "eastAsia"): east_asia,
            qn("w", "cs"): latin,
        },
    )
    SubElement(run_properties, qn("w", "sz"), {qn("w", "val"): str(size)})
    SubElement(run_properties, qn("w", "szCs"), {qn("w", "val"): str(size)})


def _formatted_import_arguments(
    source: Path,
    role_styles: dict[str, str],
    authoring_format: dict[str, object] | None,
) -> dict[str, object]:
    manifest: dict[str, object] = {
        "id": "explicit-authoring-format",
        "version": "1.0.0",
        "display_name": "Explicit Authoring Format",
        "description": "Synthetic non-private direct-format import regression.",
        "modes": ["document-spec"],
        "variables": [],
        "regions": [],
        "role_styles": role_styles,
        "compatibility": {
            "document_spec_versions": ["1.0"],
            "domain_profiles": [{"id": "academic-paper", "version": "1.0"}],
            "template_engine_versions": [],
            "style_profile_versions": ["template-mapped/1.0"],
        },
    }
    if authoring_format is not None:
        manifest["authoring_format"] = authoring_format
    return {
        "expected_source_sha256": sha256_file(source),
        "manifest": manifest,
        "provenance": {
            "source_kind": "local",
            "imported_at": "2026-08-30T00:00:00Z",
            "redistributable": False,
        },
    }


def _academic_spec() -> dict[str, object]:
    return {
        "version": "1.0",
        "domain_profile": {
            "id": "academic-paper",
            "version": "1.0",
            "locale": "zh-CN",
        },
        "metadata": {"title": "Synthetic Academic Fixture"},
        "nodes": [
            {"id": "title", "type": "title", "text": "合成学术文档 Synthetic Paper"},
            {"id": "authors", "type": "authors", "items": ["Example Author"]},
            {"id": "abstract", "type": "abstract", "text": "仅含非私有测试内容。"},
            {"id": "keywords", "type": "keywords", "items": ["DOCX", "OOXML"]},
            {"id": "heading", "type": "heading", "level": 1, "text": "结果 Results"},
            {"id": "body", "type": "paragraph", "text": "中文与 Latin text 保持调用方内容。"},
            {"id": "table", "type": "table", "rows": [["指标", "Value"], ["A", "1"]]},
            {"id": "citation", "type": "citation", "keys": ["ref1"]},
            {"id": "bibliography", "type": "bibliography"},
            {
                "id": "reference_entry",
                "type": "bibliography_entry",
                "key": "ref1",
                "authors": ["Example Author"],
                "year": 2026,
                "title": "Synthetic Reference",
                "container": "Fixture Journal",
            },
        ],
    }


def _style(styles: object, style_id: str):
    return next(
        item
        for item in styles.findall(qn("w", "style"))
        if item.attrib.get(qn("w", "styleId")) == style_id
    )


def _strict_traverse(path: Path) -> None:
    document = Document(path)
    _ = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                _ = [paragraph.text for paragraph in cell.paragraphs]
    for section in document.sections:
        _ = [paragraph.text for paragraph in section.header.paragraphs]
        _ = [paragraph.text for paragraph in section.footer.paragraphs]


def _prefix_relationship_namespace(payload: bytes) -> bytes:
    namespace = b"http://schemas.openxmlformats.org/package/2006/relationships"
    return (
        payload.replace(
            b'<Relationships xmlns="' + namespace + b'">',
            b'<rels:Relationships xmlns:rels="' + namespace + b'">',
        )
        .replace(b"<Relationship ", b"<rels:Relationship ")
        .replace(b"</Relationships>", b"</rels:Relationships>")
    )


def _libreoffice_open_and_render(path: Path, output: Path) -> None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None and os.name == "nt":
        known = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
        executable = str(known) if known.is_file() else None
    assert executable is not None, "LibreOffice is mandatory for the DOCX pack release gate"
    output.mkdir()
    profile = output / "profile"
    completed = subprocess.run(
        [
            executable,
            "--headless",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output),
            str(path),
        ],
        check=False,
        capture_output=True,
        timeout=120,
    )
    assert completed.returncode == 0, completed.stderr.decode(errors="replace")[-2000:]
    rendered = output / f"{path.stem}.pdf"
    assert rendered.is_file() and rendered.stat().st_size > 0
