"""Worksheet view, page setup, headers/footers, and print-title tests."""

from hashlib import sha256
from pathlib import Path

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.service import XlsxService


def _page_setup(*, fit: bool = True) -> dict:
    result = {
        "orientation": "landscape",
        "paper_size": "a4",
        "margins": {
            "left": 0.5,
            "right": 0.5,
            "top": 0.6,
            "bottom": 0.6,
            "header": 0.2,
            "footer": 0.2,
        },
        "horizontal_centered": True,
        "vertical_centered": False,
    }
    if fit:
        result.update({"fit_to_width": 1, "fit_to_height": 0})
    else:
        result["scale"] = 85
    return result


def _header_footer() -> dict:
    return {
        "odd_header": "&CQuarterly report",
        "odd_footer": "&RPage &P of &N",
        "even_header": "&LEven page",
        "even_footer": "&CEven footer",
        "first_header": "&CFirst page",
        "first_footer": "&CPrepared by Elftia",
        "different_first": True,
        "different_odd_even": True,
        "scale_with_doc": False,
        "align_with_margins": True,
    }


def _workbook() -> dict:
    return {
        "metadata": {"title": "Page setup", "creator": "Test", "subject": ""},
        "sheets": [
            {
                "name": "Report",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Name", "type": "s"},
                            {"ref": "B1", "value": "Amount", "type": "s"},
                            {"ref": "A2", "value": "Alpha", "type": "s"},
                            {"ref": "B2", "value": "10", "type": "n"},
                        ]
                    }
                ],
                "columns": [],
                "number_formats": [],
                "view": {
                    "show_grid_lines": False,
                    "zoom_scale": 125,
                    "selected_cell": "B2",
                },
                "page_setup": _page_setup(),
                "header_footer": _header_footer(),
                "print_area": "A1:B20",
                "print_titles": {"rows": "1:2", "columns": "A:B"},
            },
            {
                "name": "Scaled",
                "rows": [{"cells": [{"ref": "A1", "value": "Scaled", "type": "s"}]}],
                "columns": [],
                "number_formats": [],
                "page_setup": _page_setup(fit=False),
            },
        ],
        "defined_names": [],
        "tables": [],
        "charts": [],
        "chart_reference": None,
        "page_setup": None,
    }


def test_page_setup_create_reopens_and_projects_full_details(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "page-setup.xlsx"
    result = XlsxService(project_root).execute(
        "xlsx.create",
        {
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": _workbook()},
        },
    )

    assert result["status"] == "success", result
    reopened = load_workbook(output)
    report = reopened["Report"]
    assert report.sheet_view.showGridLines is False
    assert report.sheet_view.zoomScale == 125
    assert report.sheet_view.selection[0].activeCell == "B2"
    assert report.page_setup.orientation == "landscape"
    assert report.page_setup.paperSize == 9
    assert report.page_setup.fitToWidth == 1
    assert report.page_setup.fitToHeight == 0
    assert report.sheet_properties.pageSetUpPr.fitToPage is True
    assert report.print_options.horizontalCentered is True
    assert report.oddHeader.center.text == "Quarterly report"
    assert report.firstHeader.center.text == "First page"
    assert report.HeaderFooter.differentFirst is True
    assert report.HeaderFooter.differentOddEven is True
    assert reopened["Scaled"].page_setup.scale == 85

    read_result = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(output), "arguments": {}},
    )
    metadata = {
        item["sheet"]: item
        for item in read_result["diagnostics"]["operation_result"]["worksheet_metadata"]
    }
    assert metadata["Report"]["view"] == _workbook()["sheets"][0]["view"]
    assert metadata["Report"]["page_setup"] == {
        **parse_xlsx_request(
            {
                "operation": "xlsx.create",
                "output": "normalized.xlsx",
                "arguments": {"workbook": _workbook()},
            }
        ).arguments["workbook"]["sheets"][0]["page_setup"],
        "fit_to_page": True,
    }
    assert metadata["Report"]["header_footer"] == _header_footer()
    assert metadata["Report"]["print_area"] == "A1:B20"
    assert metadata["Report"]["print_titles"] == {"rows": "1:2", "columns": "A:B"}


def test_page_setup_create_is_deterministic(
    project_root: Path,
    tmp_path: Path,
) -> None:
    outputs = [tmp_path / "first.xlsx", tmp_path / "second.xlsx"]
    for output in outputs:
        result = XlsxService(project_root).execute(
            "xlsx.create",
            {
                "operation": "xlsx.create",
                "output": str(output),
                "arguments": {"workbook": _workbook()},
            },
        )
        assert result["status"] == "success"
    assert outputs[0].read_bytes() == outputs[1].read_bytes()


@pytest.mark.parametrize(
    "mutate",
    [
        lambda sheet: sheet["page_setup"].update({"orientation": "diagonal"}),
        lambda sheet: sheet["page_setup"].update({"paper_size": "unknown"}),
        lambda sheet: sheet["page_setup"].update({"scale": 85}),
        lambda sheet: sheet["view"].update({"zoom_scale": 500}),
        lambda sheet: sheet.update({"print_titles": {}}),
    ],
)
def test_page_setup_contract_rejects_invalid_boundaries(mutate) -> None:
    workbook = _workbook()
    mutate(workbook["sheets"][0])
    with pytest.raises(DocumentSkillsError) as exc:
        parse_xlsx_request(
            {
                "operation": "xlsx.create",
                "output": "page-setup.xlsx",
                "arguments": {"workbook": workbook},
            }
        )
    assert exc.value.code.value == "DS_REQUEST_INVALID"


def _source(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "Name"
    sheet["B1"] = "Amount"
    sheet.print_title_rows = "1:1"
    workbook.save(path)


def test_page_setup_edit_reopens_and_preserves_source(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "edited.xlsx"
    _source(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    edits = [
        {"sheet": "Report", "type": "page_setup", "page_setup": _page_setup()},
        {"sheet": "Report", "type": "header_footer", "header_footer": _header_footer()},
        {
            "sheet": "Report",
            "type": "sheet_view",
            "view": {"show_grid_lines": False, "zoom_scale": 140, "selected_cell": "B2"},
        },
        {
            "sheet": "Report",
            "type": "print_titles",
            "print_titles": {"rows": "1:2", "columns": "A:B"},
        },
    ]
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {"edits": edits, "expected_edits": 4},
        },
    )

    assert result["status"] == "success", result
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    reopened = load_workbook(output)
    assert reopened["Report"].page_setup.orientation == "landscape"
    assert reopened["Report"].sheet_view.zoomScale == 140
    assert reopened["Report"].print_title_rows == "$1:$2"
    assert reopened["Report"].print_title_cols == "$A:$B"


def test_print_titles_clear_and_failure_destination_preservation(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    cleared = tmp_path / "cleared.xlsx"
    existing = tmp_path / "existing.xlsx"
    _source(source)
    clear_result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(cleared),
            "arguments": {"edits": [{"sheet": "Report", "type": "print_titles_clear"}]},
        },
    )
    assert clear_result["status"] == "success"
    assert load_workbook(cleared)["Report"].print_title_rows is None

    existing.write_bytes(b"existing-page-setup-destination")
    failure = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(existing),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Report",
                        "type": "page_setup",
                        "page_setup": {"fit_to_width": 1, "scale": 90},
                    }
                ]
            },
        },
    )
    assert failure["status"] == "invalid_request"
    assert existing.read_bytes() == b"existing-page-setup-destination"
