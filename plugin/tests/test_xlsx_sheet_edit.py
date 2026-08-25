"""Sheet CRUD, reorder, copy, and reference-aware rename tests."""

import hashlib
from pathlib import Path

from document_skills_core.formats.xlsx.service import XlsxService


def _rename_source(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.workbook.defined_name import DefinedName

    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Amount", 10, 20])
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "=Data!B1"
    summary["B1"] = "Jump"
    summary["B1"].hyperlink = Hyperlink(ref="B1", location="Data!C1")
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, max_col=3, min_row=1, max_row=1))
    summary.add_chart(chart, "D2")
    workbook.defined_names.add(
        DefinedName("DataValue", attr_text="Data!$B$1")
    )
    workbook.save(path)


def _plain_source(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    plain = workbook.active
    plain.title = "Plain"
    plain["A1"] = "copy me"
    plain.print_area = "A1:A1"
    second = workbook.create_sheet("Second")
    second["A1"] = "keep me"
    second.print_area = "A1:A1"
    workbook.save(path)


def test_sheet_rename_migrates_package_wide_references(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "rename-source.xlsx"
    output = tmp_path / "renamed.xlsx"
    _rename_source(source)
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "sheet_rename",
                        "value": "Renamed Data",
                    }
                ]
            },
        },
    )

    assert result["status"] == "degraded"
    reopened = load_workbook(output, data_only=False)
    assert reopened.sheetnames == ["Renamed Data", "Summary"]
    assert reopened["Summary"]["A1"].value == "='Renamed Data'!B1"
    assert reopened["Summary"]["B1"].hyperlink.location == "'Renamed Data'!C1"
    assert reopened.defined_names["DataValue"].attr_text == "'Renamed Data'!$B$1"
    chart_formulas = [
        series.val.numRef.f for series in reopened["Summary"]._charts[0].ser
    ]
    assert chart_formulas == ["'Renamed Data'!$B$1", "'Renamed Data'!$C$1"]
    formula_cells = result["diagnostics"]["operation_result"]["formula_state"]["cells"]
    assert formula_cells["Summary!A1"]["formula"] == "'Renamed Data'!B1"
    assert formula_cells["Summary!A1"]["state"] == "recalculation_required"


def test_sheet_add_copy_reorder_and_followup_cell_edit(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "plain-source.xlsx"
    output = tmp_path / "sheet-edited.xlsx"
    _plain_source(source)
    edits = [
        {
            "sheet": "Plain",
            "type": "sheet_copy",
            "name": "Clone",
            "position": 1,
        },
        {"sheet": "Added", "type": "sheet_add", "position": 0},
        {"sheet": "Added", "type": "cell_value", "ref": "A1", "value": "new sheet"},
        {"sheet": "Second", "type": "sheet_reorder", "position": 1},
    ]
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {"edits": edits, "expected_edits": len(edits)},
        },
    )

    assert result["status"] == "success"
    reopened = load_workbook(output, data_only=False)
    assert reopened.sheetnames == ["Added", "Second", "Plain", "Clone"]
    assert reopened["Added"]["A1"].value == "new sheet"
    assert reopened["Clone"]["A1"].value == "copy me"
    assert reopened["Plain"].print_area == "'Plain'!$A$1"
    assert reopened["Clone"].print_area == "'Clone'!$A$1"
    assert reopened["Second"].print_area == "'Second'!$A$1"


def test_sheet_delete_removes_only_declared_parts(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "delete-source.xlsx"
    output = tmp_path / "deleted.xlsx"
    _plain_source(source)
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {"edits": [{"sheet": "Plain", "type": "sheet_delete"}]},
        },
    )

    assert result["status"] == "success"
    assert load_workbook(output).sheetnames == ["Second"]
    preservation = result["diagnostics"]["operation_result"]["preservation"]
    assert preservation["removed_parts"] == ["xl/worksheets/sheet1.xml"]


def test_sheet_delete_with_inbound_formula_fails_without_promotion(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "referenced-source.xlsx"
    output = tmp_path / "existing.xlsx"
    _rename_source(source)
    output.write_bytes(b"existing-sheet-delete-destination")
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {"edits": [{"sheet": "Data", "type": "sheet_delete"}]},
        },
    )

    assert result["status"] == "enhancement_required"
    assert output.read_bytes() == b"existing-sheet-delete-destination"


def test_sheet_copy_and_delete_with_related_objects_fail_without_promotion(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "related-object-source.xlsx"
    _rename_source(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    cases = [
        (
            {
                "sheet": "Summary",
                "type": "sheet_copy",
                "name": "Summary Copy",
                "position": 2,
            },
            "xlsx.related-object-sheet-copy",
        ),
        (
            {"sheet": "Summary", "type": "sheet_delete"},
            "xlsx.related-object-sheet-delete",
        ),
    ]
    for index, (edit, capability) in enumerate(cases):
        expected = f"existing-related-object-{index}".encode()
        output = tmp_path / f"existing-related-object-{index}.xlsx"
        output.write_bytes(expected)
        result = XlsxService(project_root).execute(
            "xlsx.edit",
            {
                "operation": "xlsx.edit",
                "input": str(source),
                "output": str(output),
                "arguments": {"edits": [edit]},
            },
        )

        assert result["status"] == "enhancement_required"
        assert result["errors"][0]["details"]["capability"] == capability
        assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
        assert output.read_bytes() == expected


def test_sheet_edit_output_is_deterministic(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "deterministic-source.xlsx"
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    _plain_source(source)
    edits = [
        {"sheet": "Plain", "type": "sheet_copy", "name": "Clone", "position": 1},
        {"sheet": "Added", "type": "sheet_add", "position": 0},
        {"sheet": "Added", "type": "cell_value", "ref": "A1", "value": "stable"},
    ]
    for output in (first, second):
        result = XlsxService(project_root).execute(
            "xlsx.edit",
            {
                "schema_version": "1.0",
                "operation": "xlsx.edit",
                "input": str(source),
                "output": str(output),
                "arguments": {"edits": edits},
            },
        )
        assert result["status"] == "success"

    assert first.read_bytes() == second.read_bytes()
