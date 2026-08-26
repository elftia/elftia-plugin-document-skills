"""Native pivot contract, package, consumer, and transaction coverage."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.create import create_xlsx
from document_skills_core.formats.xlsx.package import OpcPackage
from document_skills_core.formats.xlsx.pivot_projection import project_pivot_tables
from document_skills_core.formats.xlsx.service import XlsxService
from tests.support.xlsx_macro_fixture import create_package_fixture


def _arguments(**overrides: object) -> dict[str, object]:
    result: dict[str, object] = {
        "source": {"sheet": "Data", "range": "A1:D7"},
        "rows": [{"column": "Region", "sort": "asc"}],
        "columns": [{"column": "Category", "sort": "asc"}],
        "values": [
            {"column": "Revenue", "function": "sum", "as": "Total Revenue"}
        ],
        "filters": [{"column": "Segment", "value": "Retail"}],
        "target": {
            "sheet": "Pivot",
            "start_cell": "A1",
            "name": "SalesPivot",
            "style": "PivotStyleMedium9",
        },
        "formula_policy": "reject",
        "numeric_policy": "strict",
        "recalculation": "skip",
    }
    result.update(overrides)
    return result


def _request(input_path: Path, output: Path, **arguments: object) -> dict[str, object]:
    return {
        "operation": "xlsx.pivot.create",
        "input": str(input_path),
        "output": str(output),
        "arguments": _arguments(**arguments),
    }


def _write_source(path: Path, *, revenue_as_text: bool = False) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Region", "Category", "Segment", "Revenue"])
    rows = [
        ("East", "A", "Retail", 10),
        ("West", "A", "Retail", 20),
        ("East", "B", "Wholesale", 15),
        ("West", "B", "Wholesale", 10),
        ("East", "A", "Retail", 5),
        ("North", "C", "Retail", 1),
    ]
    for region, category, segment, revenue in rows:
        sheet.append(
            [
                region,
                category,
                segment,
                str(revenue) if revenue_as_text else revenue,
            ]
        )
    workbook.save(path)
    workbook.close()
    return path


def test_pivot_contract_is_bounded_and_macro_explicit() -> None:
    parsed = parse_xlsx_request(
        {
            "operation": "xlsx.pivot.create",
            "input": "source.xlsx",
            "output": "pivot.xlsx",
            "arguments": _arguments(),
        }
    )
    assert parsed.arguments["rows"] == [{"column": "Region", "sort": "asc"}]
    assert parsed.arguments["limits"]["max_axis_items"] == 1_000

    invalid_arguments = [
        _arguments(
            rows=[
                {"column": "Region", "sort": "asc"},
                {"column": "Category", "sort": "asc"},
            ]
        ),
        _arguments(values=[]),
        _arguments(columns=[{"column": "Region", "sort": "asc"}]),
        _arguments(source={"sheet": "Data", "range": "A1:D1"}),
    ]
    for arguments in invalid_arguments:
        with pytest.raises(DocumentSkillsError):
            parse_xlsx_request(
                {
                    "operation": "xlsx.pivot.create",
                    "input": "source.xlsx",
                    "output": "pivot.xlsx",
                    "arguments": arguments,
                }
            )

    with pytest.raises(DocumentSkillsError):
        parse_xlsx_request(
            {
                "operation": "xlsx.pivot.create",
                "input": "source.xlsm",
                "output": "pivot.xlsm",
                "arguments": _arguments(),
            }
        )
    macro = parse_xlsx_request(
        {
            "operation": "xlsx.pivot.create",
            "input": "source.xlsm",
            "output": "pivot.xlsm",
            "arguments": _arguments(keep_vba=True),
        }
    )
    assert macro.arguments["keep_vba"] is True


def test_pivot_writes_complete_native_relationship_chain_and_reopens(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = _write_source(tmp_path / "source.xlsx")
    output = tmp_path / "pivot.xlsx"
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    result = XlsxService(project_root).execute(
        "xlsx.pivot.create",
        _request(source, output),
    )

    assert result["status"] == "success", result["errors"]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    pivot = result["diagnostics"]["operation_result"]["pivot"]
    assert pivot["native"] is True
    assert pivot["location"] == "A4:D9"
    assert pivot["source"]["records"] == 6
    assert pivot["cache"]["records_written"] == 6
    assert pivot["cache"]["workbook_relationship_matches"] is True
    assert any(
        gate["id"] == "operation.native-pivot-table"
        and gate["outcome"] == "pass"
        and gate["evidence"]["native_pivot"] is True
        for gate in result["validation"]["gates"]
    )

    package = OpcPackage.open(output)
    target_sheet = package.xml("xl/worksheets/sheet2.xml")
    assert not target_sheet.findall(
        ".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}pivotTableParts"
    )
    projected = project_pivot_tables(package)
    assert projected == [{key: value for key, value in pivot.items() if key != "display_range" and key != "source"} | {
        "source": {"sheet": "Data", "range": "A1:D7"}
    }]
    expected_parts = {
        "xl/pivotTables/pivotTable1.xml",
        "xl/pivotTables/_rels/pivotTable1.xml.rels",
        "xl/pivotCache/pivotCacheDefinition1.xml",
        "xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels",
        "xl/pivotCache/pivotCacheRecords1.xml",
    }
    assert expected_parts.issubset(package.parts)

    reopened = load_workbook(output, data_only=False)
    sheet = reopened["Pivot"]
    assert len(sheet._pivots) == 1
    native = sheet._pivots[0]
    assert native.name == "SalesPivot"
    assert native.location.ref == "A4:D9"
    assert native.cache.cacheSource.worksheetSource.sheet == "Data"
    assert native.cache.cacheSource.worksheetSource.ref == "A1:D7"
    assert native.cache.recordCount == 6
    assert len(native.cache.records.r) == 6
    assert [(field.name, field.fld, field.subtotal) for field in native.dataFields] == [
        ("Total Revenue", 3, "sum")
    ]
    assert [sheet.cell(5, column).value for column in range(1, 5)] == [
        "Region",
        "A",
        "C",
        "Grand Total",
    ]
    assert [sheet.cell(6, column).value for column in range(1, 5)] == [
        "East",
        15,
        None,
        15,
    ]
    assert [sheet.cell(9, column).value for column in range(1, 5)] == [
        "Grand Total",
        35,
        1,
        36,
    ]
    reopened.close()


def test_pivot_supports_row_only_count_and_descending_order(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = _write_source(tmp_path / "count-source.xlsx")
    output = tmp_path / "count-pivot.xlsx"
    result = XlsxService(project_root).execute(
        "xlsx.pivot.create",
        _request(
            source,
            output,
            rows=[{"column": "Region", "sort": "desc"}],
            columns=[],
            filters=[],
            values=[{"column": "Revenue", "function": "count", "as": "Rows"}],
            target={
                "sheet": "Counts",
                "start_cell": "B2",
                "name": "CountPivot",
                "style": "PivotStyleLight16",
            },
        ),
    )

    assert result["status"] == "success"
    reopened = load_workbook(output, data_only=False)
    sheet = reopened["Counts"]
    assert [sheet.cell(row, 2).value for row in range(2, 7)] == [
        "Region",
        "West",
        "North",
        "East",
        "Grand Total",
    ]
    assert [sheet.cell(row, 3).value for row in range(2, 7)] == [
        "Rows",
        2,
        1,
        3,
        6,
    ]
    reopened.close()


def test_pivot_failure_preserves_existing_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _write_source(tmp_path / "text-source.xlsx", revenue_as_text=True)
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing-pivot-destination")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    result = XlsxService(project_root).execute(
        "xlsx.pivot.create",
        _request(source, output),
    )

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_VALIDATION_FAILED"
    assert output.read_bytes() == b"existing-pivot-destination"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash


def test_pivot_formula_policy_is_explicit_and_cached_use_is_degraded(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "formula-source.xlsx"
    workbook = {
        "metadata": {},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Region", "type": "s"},
                            {"ref": "B1", "value": "Revenue", "type": "s"},
                            {"ref": "A2", "value": "East", "type": "s"},
                            {
                                "ref": "B2",
                                "formula": "5+5",
                                "cached_value": "10",
                                "type": "n",
                            },
                        ]
                    }
                ],
                "number_formats": [],
            }
        ],
        "defined_names": [],
        "tables": [],
    }
    parsed = parse_xlsx_request(
        {
            "operation": "xlsx.create",
            "output": str(source),
            "arguments": {"workbook": workbook, "recalculation": "skip"},
        }
    )
    create_xlsx(source, parsed.arguments["workbook"])
    overrides = {
        "source": {"sheet": "Data", "range": "A1:B2"},
        "columns": [],
        "filters": [],
        "values": [{"column": "Revenue", "function": "sum", "as": "Revenue"}],
    }
    rejected = XlsxService(project_root).execute(
        "xlsx.pivot.create",
        _request(source, tmp_path / "rejected.xlsx", **overrides),
    )
    assert rejected["status"] == "enhancement_required"

    result = XlsxService(project_root).execute(
        "xlsx.pivot.create",
        _request(
            source,
            tmp_path / "cached.xlsx",
            **overrides,
            formula_policy="cached",
        ),
    )
    assert result["status"] == "degraded", result["errors"]
    pivot = result["diagnostics"]["operation_result"]["pivot"]
    assert pivot["source"]["formula_cells_used"] == ["Data!B2"]
    assert any(item["code"] == "pivot-cached-formula-values" for item in result["degradations"])


def test_pivot_xlsm_preserves_vba_and_discloses_signature_invalidation(
    project_root: Path,
    tmp_path: Path,
) -> None:
    workbook = {
        "metadata": {},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Region", "type": "s"},
                            {"ref": "B1", "value": "Category", "type": "s"},
                            {"ref": "C1", "value": "Segment", "type": "s"},
                            {"ref": "D1", "value": "Revenue", "type": "s"},
                            {"ref": "A2", "value": "East", "type": "s"},
                            {"ref": "B2", "value": "A", "type": "s"},
                            {"ref": "C2", "value": "Retail", "type": "s"},
                            {"ref": "D2", "value": "10", "type": "n"},
                        ]
                    }
                ],
                "number_formats": [],
            }
        ],
        "defined_names": [],
        "tables": [],
    }
    source = create_package_fixture(
        tmp_path / "pivot-source.xlsm",
        "xlsm",
        signed=True,
        workbook=workbook,
    )
    output = tmp_path / "pivot-output.xlsm"
    result = XlsxService(project_root).execute(
        "xlsx.pivot.create",
        _request(
            source,
            output,
            source={"sheet": "Data", "range": "A1:D2"},
            keep_vba=True,
        ),
    )

    assert result["status"] == "degraded", result["errors"]
    macro = result["diagnostics"]["operation_result"]["macro"]
    assert macro["vba_payload"] == "preserved"
    assert macro["signature_state"] == "invalidated_by_package_mutation"
    package = OpcPackage.open(output, allowed_inert_categories=frozenset({"vba"}))
    assert package.parts["xl/vbaProject.bin"] == b"inert-vba-project-payload\x00"
    assert project_pivot_tables(package)[0]["native"] is True
