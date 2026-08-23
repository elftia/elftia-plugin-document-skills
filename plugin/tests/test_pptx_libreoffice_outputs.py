"""LibreOffice-backed public PPTX render and conversion tests."""

from hashlib import sha256
import json
from pathlib import Path
import struct
import zipfile
import zlib

from document_skills_core.core.capabilities import (
    DetectionEvidence,
    ProviderCatalog,
)
from document_skills_core.core.capabilities.reports import build_capabilities
from document_skills_core.core.contracts.errors import DocumentSkillsError, ErrorCode
from document_skills_core.formats.pdf.validation import reopen_pdf
from document_skills_core.formats.pptx.legacy_conversion import (
    MAX_LEGACY_PPT_BYTES,
    preflight_legacy_ppt,
)
from document_skills_core.formats.pptx.render_validation import validate_render_bundle
from document_skills_core.formats.pptx.validation import reopen_pptx
from document_skills_core.providers.libreoffice import build_libreoffice_provider


class _CallableDetector:
    def detect(self) -> DetectionEvidence:
        return DetectionEvidence(True, version="25.8.0", path="/fake/soffice")


class _AbsentDetector:
    def detect(self) -> DetectionEvidence:
        return DetectionEvidence(False, reason="LibreOffice is absent")


class _RenderRunner:
    def __init__(
        self,
        *,
        legacy_slides: int = 2,
        pdf_pages: int | None = None,
    ) -> None:
        self.legacy_slides = legacy_slides
        self.pdf_pages = pdf_pages
        self.calls: list[dict[str, object]] = []

    def set_executable(self, _executable: str) -> None:
        pass

    def convert(
        self,
        input_path: Path,
        target_format: str,
        output_dir: Path,
        *,
        timeout_seconds: float | None = None,
    ) -> Path:
        self.calls.append({
            "format": target_format,
            "input": str(input_path),
            "timeout": timeout_seconds,
        })
        output = output_dir / f"{input_path.stem}.{target_format}"
        if target_format == "pdf":
            from pptx import Presentation

            pages = self.pdf_pages
            if pages is None:
                pages = len(Presentation(input_path).slides)
            _write_pdf(output, pages)
        elif target_format == "png":
            from pptx import Presentation

            presentation = Presentation(input_path)
            assert len(presentation.slides) == 1
            self.calls[-1]["title"] = presentation.slides[0].shapes.title.text
            output.write_bytes(_png(4, 3, (40, 80, 120, 255)))
        elif target_format == "pptx":
            assert input_path.suffix.casefold() == ".ppt"
            _deck(output, slides=self.legacy_slides)
        else:  # pragma: no cover - operation contract
            raise AssertionError(target_format)
        return output


class _MalformedLegacyRunner:
    def set_executable(self, _executable: str) -> None:
        pass

    def convert(
        self,
        input_path: Path,
        target_format: str,
        output_dir: Path,
        *,
        timeout_seconds: float | None = None,
    ) -> Path:
        assert input_path.suffix.casefold() == ".ppt"
        assert target_format == "pptx"
        output = output_dir / f"{input_path.stem}.pptx"
        output.write_bytes(b"not an OOXML presentation")
        return output


class _TimeoutRunner:
    def set_executable(self, _executable: str) -> None:
        pass

    def convert(self, *_args, **_kwargs) -> Path:
        raise DocumentSkillsError(
            ErrorCode.PROCESS_TIMEOUT,
            "LibreOffice timed out.",
        )


def test_pptx_convert_pdf_promotes_reopened_page_matched_output(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _deck(tmp_path / "source.pptx", slides=2)
    source_hash = sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "converted.pdf"
    definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector(),
        runner=_RenderRunner(),
    )

    result = provider.execute("pptx.convert.pdf", _request(
        "pptx.convert.pdf",
        source,
        output,
    ))

    assert result["status"] == "success", result
    assert reopen_pdf(output)["pages"] == 2
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    outcomes = {gate["id"]: gate["outcome"] for gate in result["validation"]["gates"]}
    assert outcomes["operation.pptx-source-deep-validation"] == "pass"
    assert outcomes["operation.pptx-pdf-page-count"] == "pass"
    assert {item.operation for item in definition.capabilities} >= {
        "pptx.convert.legacy",
        "pptx.convert.pdf",
        "pptx.render",
    }
    assert callable(definition.validators["visual"])


def test_pptx_libreoffice_profile_capability_truth(
    project_root: Path,
) -> None:
    for detector, expected in [(_CallableDetector(), True), (_AbsentDetector(), False)]:
        definition, _provider = build_libreoffice_provider(
            project_root,
            detector=detector,
            runner=_RenderRunner(),
        )
        registry = ProviderCatalog()
        registry.register_provider(definition)
        report = build_capabilities(project_root, "pptx", registry)
        operations = {
            item["operation"]: item
            for item in report["operations"]
        }
        assert operations["pptx.convert.pdf"]["available"] is expected
        assert operations["pptx.convert.legacy"]["available"] is expected
        assert operations["pptx.render"]["available"] is expected
        assert report["validation"]["visual"] == (
            "available" if expected else "unavailable"
        )


def test_pptx_convert_legacy_promotes_deep_reopened_semantic_conversion(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _legacy_ppt(tmp_path / "source.ppt")
    source_hash = sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "converted.pptx"
    runner = _RenderRunner(legacy_slides=2)
    _definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector(),
        runner=runner,
    )

    result = provider.execute(
        "pptx.convert.legacy",
        _request("pptx.convert.legacy", source, output),
    )

    assert result["status"] == "degraded", result
    assert result["degraded"] is True
    assert reopen_pptx(output)["slides"] == 2
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    assert len(runner.calls) == 1
    assert runner.calls[0]["format"] == "pptx"
    assert Path(str(runner.calls[0]["input"])).suffix.casefold() == ".ppt"
    assert runner.calls[0]["timeout"] == 30.0
    conversion = result["diagnostics"]["operation_result"]["conversion"]
    assert conversion == {
        "provider": "libreoffice",
        "provider_version": "25.8.0",
        "semantic_conversion": True,
        "source_format": "ppt",
        "source_visual_preservation_claimed": False,
        "target_format": "pptx",
        "slides": 2,
    }
    assert result["warnings"][0]["code"] == "PPT_LEGACY_SEMANTIC_CONVERSION"
    assert "exact source visual fidelity is not claimed" in result["warnings"][0]["message"]
    outcomes = {
        gate["id"]: gate["outcome"]
        for gate in result["validation"]["gates"]
    }
    assert outcomes["provider.reopen"] == "pass"
    assert outcomes["operation.legacy-cfb-preflight"] == "pass"
    assert outcomes["operation.legacy-pptx-deep-validation"] == "pass"
    assert outcomes["visual.render"] == "unavailable"


def test_legacy_ppt_preflight_rejects_bad_magic_header_and_size(
    tmp_path: Path,
) -> None:
    valid = _legacy_ppt(tmp_path / "valid.ppt")
    assert preflight_legacy_ppt(valid) == {
        "bytes": 1536,
        "container": "cfb",
        "major_version": 3,
        "sector_size": 512,
    }

    bad_magic = tmp_path / "bad-magic.ppt"
    bad_magic.write_bytes(b"x" * 512)
    bad_header = tmp_path / "bad-header.ppt"
    payload = bytearray(valid.read_bytes())
    payload[26:28] = (4).to_bytes(2, "little")
    bad_header.write_bytes(payload)
    truncated = tmp_path / "truncated.ppt"
    truncated.write_bytes(valid.read_bytes()[:511])
    oversized = tmp_path / "oversized.ppt"
    with oversized.open("wb") as handle:
        handle.truncate(MAX_LEGACY_PPT_BYTES + 1)

    for candidate in (bad_magic, bad_header, truncated, oversized):
        try:
            preflight_legacy_ppt(candidate)
        except DocumentSkillsError as error:
            assert error.code == ErrorCode.ARCHIVE_UNSAFE
            assert error.status == "invalid_request"
        else:  # pragma: no cover - negative assertion
            raise AssertionError(f"unsafe legacy PPT was accepted: {candidate.name}")


def test_pptx_convert_legacy_timeout_preserves_source_and_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _legacy_ppt(tmp_path / "source.ppt")
    source_payload = source.read_bytes()
    output = tmp_path / "prior.pptx"
    prior_output = b"prior destination"
    output.write_bytes(prior_output)
    _definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector(),
        runner=_TimeoutRunner(),
    )

    result = provider.execute(
        "pptx.convert.legacy",
        _request("pptx.convert.legacy", source, output),
    )

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == ErrorCode.PROCESS_TIMEOUT.value
    assert source.read_bytes() == source_payload
    assert output.read_bytes() == prior_output


def test_pptx_convert_legacy_rejects_malformed_generated_pptx_without_promotion(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _legacy_ppt(tmp_path / "source.ppt")
    source_payload = source.read_bytes()
    output = tmp_path / "prior.pptx"
    prior_output = b"prior destination"
    output.write_bytes(prior_output)
    _definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector(),
        runner=_MalformedLegacyRunner(),
    )

    result = provider.execute(
        "pptx.convert.legacy",
        _request("pptx.convert.legacy", source, output),
    )

    assert result["status"] == "failed"
    assert source.read_bytes() == source_payload
    assert output.read_bytes() == prior_output


def test_pptx_render_promotes_pdf_and_every_slide_png_bundle(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _deck(tmp_path / "source.pptx", slides=2)
    output = tmp_path / "render.zip"
    runner = _RenderRunner()
    _definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector(),
        runner=runner,
    )

    result = provider.execute("pptx.render", _request("pptx.render", source, output))

    assert result["status"] == "success", result
    assert [item["format"] for item in runner.calls] == ["pdf", "png", "png"]
    assert [item.get("title") for item in runner.calls[1:]] == ["Slide 1", "Slide 2"]
    with zipfile.ZipFile(output) as archive:
        assert set(archive.namelist()) == {
            "deck.pdf",
            "manifest.json",
            "slides/slide-001.png",
            "slides/slide-002.png",
        }
        manifest = json.loads(archive.read("manifest.json"))
    assert manifest["slides"] == 2
    assert manifest["pdf"]["pages"] == 2
    assert [item["slide"] for item in manifest["images"]] == [1, 2]
    visual = next(
        gate for gate in result["validation"]["gates"]
        if gate["id"] == "visual.render"
    )
    assert visual["required"] is True
    assert visual["outcome"] == "pass"
    assert visual["evidence"]["slides_rendered"] == 2


def test_pptx_convert_pdf_rejects_page_count_mismatch_without_output(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _deck(tmp_path / "source.pptx", slides=2)
    source_hash = sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "prior-output.pdf"
    prior_output = b"prior destination"
    output.write_bytes(prior_output)
    _definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector(),
        runner=_RenderRunner(pdf_pages=1),
    )

    result = provider.execute("pptx.convert.pdf", _request(
        "pptx.convert.pdf",
        source,
        output,
    ))

    assert result["status"] == "failed"
    assert output.read_bytes() == prior_output
    assert sha256(source.read_bytes()).hexdigest() == source_hash


def test_pptx_render_timeout_preserves_source_and_publishes_nothing(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _deck(tmp_path / "source.pptx", slides=1)
    source_hash = sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "must-not-exist.zip"
    _definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector(),
        runner=_TimeoutRunner(),
    )

    result = provider.execute("pptx.render", _request("pptx.render", source, output))

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == ErrorCode.PROCESS_TIMEOUT.value
    assert not output.exists()
    assert sha256(source.read_bytes()).hexdigest() == source_hash


def test_render_bundle_validator_rejects_manifest_byte_mismatch(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _deck(tmp_path / "source.pptx", slides=1)
    output = tmp_path / "render.zip"
    _definition, provider = build_libreoffice_provider(
        project_root,
        detector=_CallableDetector(),
        runner=_RenderRunner(),
    )
    result = provider.execute("pptx.render", _request("pptx.render", source, output))
    assert result["status"] == "success"
    with zipfile.ZipFile(output) as archive:
        manifest = json.loads(archive.read("manifest.json"))
    manifest["images"][0]["bytes"] += 1

    try:
        validate_render_bundle(output, manifest)
    except DocumentSkillsError as error:
        assert error.code == ErrorCode.VALIDATION_FAILED
    else:  # pragma: no cover - negative assertion
        raise AssertionError("tampered render manifest was accepted")


def _request(operation: str, source: Path, output: Path) -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "operation": operation,
        "input": str(source),
        "output": str(output),
        "arguments": {},
        "options": {"fidelity": "enhanced"},
    }


def _deck(path: Path, *, slides: int) -> Path:
    from pptx import Presentation

    presentation = Presentation()
    for number in range(1, slides + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Slide {number}"
        slide.placeholders[1].text = "Body"
    presentation.save(path)
    return path


def _legacy_ppt(path: Path) -> Path:
    """Write a minimal valid CFB v3 container for bounded preflight tests."""

    free_sector = (0xFFFFFFFF).to_bytes(4, "little")
    end_of_chain = (0xFFFFFFFE).to_bytes(4, "little")
    fat_sector = (0xFFFFFFFD).to_bytes(4, "little")
    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    header[24:26] = (0x003E).to_bytes(2, "little")
    header[26:28] = (3).to_bytes(2, "little")
    header[28:30] = b"\xfe\xff"
    header[30:32] = (9).to_bytes(2, "little")
    header[32:34] = (6).to_bytes(2, "little")
    header[44:48] = (1).to_bytes(4, "little")
    header[48:52] = (0).to_bytes(4, "little")
    header[56:60] = (4096).to_bytes(4, "little")
    header[60:64] = end_of_chain
    header[68:72] = end_of_chain
    header[76:80] = (1).to_bytes(4, "little")
    header[80:512] = free_sector * 108

    directory = bytearray(512)
    root_name = "Root Entry\x00".encode("utf-16le")
    directory[:len(root_name)] = root_name
    directory[64:66] = len(root_name).to_bytes(2, "little")
    directory[66] = 5
    directory[67] = 1
    directory[68:80] = free_sector * 3
    directory[116:120] = end_of_chain

    fat = bytearray(free_sector * 128)
    fat[:4] = end_of_chain
    fat[4:8] = fat_sector
    path.write_bytes(bytes(header + directory + fat))
    return path


def _write_pdf(path: Path, pages: int) -> None:
    import fitz

    document = fitz.open()
    for number in range(1, pages + 1):
        page = document.new_page(width=800, height=600)
        page.insert_text((72, 72), f"Slide {number}")
    document.save(path)
    document.close()


def _png(width: int, height: int, rgba: tuple[int, int, int, int]) -> bytes:
    raw = b"".join(b"\x00" + bytes(rgba) * width for _row in range(height))
    return b"\x89PNG\r\n\x1a\n" + b"".join([
        _chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 6, 0, 0, 0)),
        _chunk(b"IDAT", zlib.compress(raw)),
        _chunk(b"IEND", b""),
    ])


def _chunk(kind: bytes, payload: bytes) -> bytes:
    return (
        struct.pack(">I", len(payload))
        + kind
        + payload
        + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)
    )
