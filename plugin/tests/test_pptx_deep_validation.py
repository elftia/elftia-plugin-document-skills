"""Deep PPTX graph, chart/workbook, and static-layout validation tests."""

from copy import deepcopy
from pathlib import Path
from typing import Callable
from xml.etree.ElementTree import Element, tostring
import zipfile

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError, ErrorCode
from document_skills_core.formats.pptx.constants import NS, local_name
from document_skills_core.formats.pptx.deep_validation import validate_deep_package
from document_skills_core.formats.pptx.package import write_deterministic_zip
from document_skills_core.formats.pptx.service import PptxService
import document_skills_core.formats.pptx.service as pptx_service_module


@pytest.fixture
def chart_deck(tmp_path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    output = tmp_path / "chart-deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Deep validation"
    data = ChartData()
    data.categories = ["North", "South", "West"]
    data.add_series("Sales", (3.0, 5.0, 8.0))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1.5),
        Inches(7),
        Inches(4),
        data,
    )
    presentation.save(output)
    return output


def test_deep_validator_checks_chart_axis_cache_and_embedded_workbook(
    chart_deck: Path,
) -> None:
    report = validate_deep_package(chart_deck)

    assert report["valid"] is True
    assert report["inventory"]["slides"] == 1
    assert report["inventory"]["charts"] == 1
    assert report["charts"]["axes"] == 2
    assert report["charts"]["series"] == 1
    assert report["charts"]["workbook_ranges_checked"] >= 3
    assert report["graph"]["orphan_parts"] == 0


def test_deep_validator_rejects_orphan_part(chart_deck: Path, tmp_path: Path) -> None:
    parts = _parts(chart_deck)
    parts["ppt/media/orphan.png"] = b"\x89PNG\r\n\x1a\n"
    output = tmp_path / "orphan.pptx"
    write_deterministic_zip(output, parts)

    failures = _failures(output)
    assert "orphan-part:ppt/media/orphan.png" in failures


def test_deep_validator_rejects_dangling_xml_relationship(
    chart_deck: Path,
    tmp_path: Path,
) -> None:
    output = _mutate_xml(
        chart_deck,
        tmp_path / "dangling.pptx",
        "ppt/slides/slide1.xml",
        lambda root: _set_chart_relationship(root, "rIdMissing"),
    )

    failures = _failures(output)
    assert "xml-relationship-missing:ppt/slides/slide1.xml:rIdMissing" in failures


def test_deep_validator_rejects_unlinked_chart_axis(
    chart_deck: Path,
    tmp_path: Path,
) -> None:
    output = _mutate_xml(
        chart_deck,
        tmp_path / "axis.pptx",
        "ppt/charts/chart1.xml",
        lambda root: _set_first_value(root, "crossAx", "999999"),
    )

    failures = _failures(output)
    assert "chart-cross-axis:ppt/charts/chart1.xml:999999" in failures


def test_deep_validator_rejects_chart_cache_workbook_mismatch(
    chart_deck: Path,
    tmp_path: Path,
) -> None:
    output = _mutate_xml(
        chart_deck,
        tmp_path / "cache.pptx",
        "ppt/charts/chart1.xml",
        _tamper_numeric_cache,
    )

    failures = _failures(output)
    assert any(value.startswith("chart-workbook-cache-mismatch:") for value in failures)


def test_deep_validator_rejects_wrong_embedded_workbook_content_type(
    chart_deck: Path,
    tmp_path: Path,
) -> None:
    output = _mutate_xml(
        chart_deck,
        tmp_path / "workbook-content-type.pptx",
        "[Content_Types].xml",
        _replace_xlsx_content_type,
    )

    failures = _failures(output)
    assert "chart-workbook-content-type:ppt/charts/chart1.xml" in failures


@pytest.mark.parametrize("branch", ["fallback", "later-choice"])
@pytest.mark.parametrize("fault", ["zero", "non-numeric", "external-collision"])
def test_equation_alternate_branches_block_invalid_candidate_promotion(
    project_root: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    branch: str,
    fault: str,
) -> None:
    output = tmp_path / f"{branch}-{fault}.pptx"
    real_create = pptx_service_module.create_pptx

    def tampered_create(
        destination: Path,
        deck: dict[str, object],
        template: Path | None = None,
    ) -> dict[str, object]:
        creation = real_create(destination, deck, template=template)
        _tamper_equation_alternate(destination, branch=branch, fault=fault)
        return creation

    monkeypatch.setattr(pptx_service_module, "create_pptx", tampered_create)
    result = PptxService(project_root).execute(
        "pptx.create",
        {
            "schema_version": "1.0",
            "operation": "pptx.create",
            "output": str(output),
            "arguments": {
                "deck": {
                    "metadata": {
                        "title": "Deep equation branches",
                        "creator": "Elftia",
                        "subject": "B6",
                    },
                    "slides": [
                        {
                            "layout": "content",
                            "title": "Alternate branches",
                            "shapes": [
                                {
                                    "type": "equation",
                                    "id": "eq-branch",
                                    "bbox": {
                                        "x": 0.5,
                                        "y": 1.5,
                                        "w": 4.0,
                                        "h": 0.7,
                                    },
                                    "source": {
                                        "kind": "latex",
                                        "value": "x+1",
                                    },
                                    "fallback": "reject",
                                }
                            ],
                            "table": None,
                            "chart_reference": None,
                            "image_reference": None,
                            "notes": None,
                        }
                    ],
                }
            },
        },
    )

    deep_gate = next(
        gate
        for gate in result["validation"]["gates"]
        if gate["id"] == "operation.pptx-deep-validation"
    )
    expected = (
        "drawing-id-duplicate:ppt/slides/slide1.xml"
        if fault == "external-collision"
        else "drawing-id-invalid:ppt/slides/slide1.xml"
    )
    assert result["errors"][0]["code"] == "DS_VALIDATION_FAILED"
    assert deep_gate["outcome"] == "fail"
    assert expected in deep_gate["evidence"]["failures"]
    assert not output.exists()


def test_static_layout_reports_tokens_overflow_contrast_and_geometry(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    output = tmp_path / "lint.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    body = slide.placeholders[1]
    body.text = "Body"
    for left in (1, 1.1):
        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left),
            Inches(2),
            Inches(3),
            Inches(2),
        )
    text_box = slide.shapes.add_textbox(
        Inches(-1),
        Inches(5),
        Inches(1),
        Inches(0.3),
    )
    text_box.fill.solid()
    text_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    run = text_box.text_frame.paragraphs[0].add_run()
    run.text = "{{customer}} TODO debug " + "overflow " * 20
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(255, 255, 255)
    presentation.save(output)

    findings = validate_deep_package(output)["static_layout"]["by_code"]
    assert {
        "debug-text",
        "empty-title",
        "low-contrast",
        "minimum-font-size",
        "out-of-bounds",
        "overlap-risk",
        "placeholder-token",
        "text-overflow-risk",
        "todo-text",
    } <= set(findings)


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {
            info.filename: archive.read(info)
            for info in archive.infolist()
            if not info.is_dir()
        }


def _mutate_xml(
    source: Path,
    output: Path,
    part: str,
    mutate: Callable[[Element], None],
) -> Path:
    from defusedxml.ElementTree import fromstring

    parts = _parts(source)
    root = fromstring(parts[part])
    mutate(root)
    parts[part] = tostring(root, encoding="utf-8", xml_declaration=True)
    write_deterministic_zip(output, parts)
    return output


def _set_chart_relationship(root: Element, value: str) -> None:
    chart = next(node for node in root.iter() if local_name(node.tag) == "chart")
    chart.attrib[f"{{{NS['r']}}}id"] = value


def _set_first_value(root: Element, local: str, value: str) -> None:
    node = next(node for node in root.iter() if local_name(node.tag) == local)
    node.attrib["val"] = value


def _tamper_numeric_cache(root: Element) -> None:
    cache = next(node for node in root.iter() if local_name(node.tag) == "numCache")
    value = next(node for node in cache.iter() if local_name(node.tag) == "v")
    value.text = "999"


def _replace_xlsx_content_type(root: Element) -> None:
    declaration = next(
        node
        for node in root
        if node.attrib.get("Extension", "").casefold() == "xlsx"
    )
    declaration.attrib["ContentType"] = "application/octet-stream"


def _tamper_equation_alternate(
    path: Path,
    *,
    branch: str,
    fault: str,
) -> None:
    from defusedxml.ElementTree import fromstring

    parts = _parts(path)
    part = "ppt/slides/slide1.xml"
    root = fromstring(parts[part])
    alternate = next(
        node for node in root.iter() if local_name(node.tag) == "AlternateContent"
    )
    if branch == "fallback":
        selected = next(
            child for child in alternate if local_name(child.tag) == "Fallback"
        )
    else:
        first_choice = next(
            child for child in alternate if local_name(child.tag) == "Choice"
        )
        selected = deepcopy(first_choice)
        selected.attrib["Requires"] = "a15"
        fallback_index = next(
            index
            for index, child in enumerate(alternate)
            if local_name(child.tag) == "Fallback"
        )
        alternate.insert(fallback_index, selected)
    properties = next(
        node for node in selected.iter() if local_name(node.tag) == "cNvPr"
    )
    if fault == "zero":
        properties.attrib["id"] = "0"
    elif fault == "non-numeric":
        properties.attrib["id"] = "invalid"
    else:
        selected_nodes = set(selected.iter())
        external = next(
            node
            for node in root.iter()
            if local_name(node.tag) == "cNvPr" and node not in selected_nodes
            and node.attrib.get("id") != properties.attrib.get("id")
        )
        properties.attrib["id"] = external.attrib["id"]
    parts[part] = tostring(root, encoding="utf-8", xml_declaration=True)
    write_deterministic_zip(path, parts)


def _failures(path: Path) -> list[str]:
    with pytest.raises(DocumentSkillsError) as captured:
        validate_deep_package(path)
    assert captured.value.code == ErrorCode.VALIDATION_FAILED
    return captured.value.details["failures"]
