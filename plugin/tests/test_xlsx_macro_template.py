"""Macro-enabled, template-as-base, and legacy XLS coverage."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.constants import CONTENT_TYPES, WORKBOOK_CONTENT_TYPES
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.format_policy import allowed_inert_categories
from document_skills_core.formats.xlsx.package import OpcPackage
from document_skills_core.formats.xlsx.service import XlsxService
from tests.support.xlsx_macro_fixture import create_package_fixture


def test_contract_accepts_macro_read_and_requires_explicit_keep_vba() -> None:
    parsed = parse_xlsx_request(
        {
            "operation": "xlsx.read",
            "input": "book.xlsm",
            "arguments": {},
        }
    )
    assert parsed.input_path is not None and parsed.input_path.suffix == ".xlsm"
    with pytest.raises(DocumentSkillsError):
        parse_xlsx_request(
            {
                "operation": "xlsx.edit",
                "input": "book.xlsm",
                "output": "edited.xlsm",
                "arguments": {
                    "edits": [
                        {
                            "sheet": "Sheet1",
                            "type": "cell_value",
                            "ref": "A1",
                            "value": "11",
                        }
                    ]
                },
            }
        )
    parsed_edit = parse_xlsx_request(
        {
            "operation": "xlsx.edit",
            "input": "book.xlsm",
            "output": "edited.xlsm",
            "arguments": {
                "keep_vba": True,
                "edits": [
                    {
                        "sheet": "Sheet1",
                        "type": "cell_value",
                        "ref": "A1",
                        "value": "11",
                    }
                ],
            },
        }
    )
    assert parsed_edit.arguments["keep_vba"] is True
    assert parsed_edit.arguments["recalculation"] == "skip"
    for input_name, output_name, arguments in (
        (
            "book.xlsx",
            "edited.xlsx",
            {
                "keep_vba": False,
                "edits": [
                    {"sheet": "Sheet1", "type": "cell_value", "ref": "A1", "value": "1"}
                ],
            },
        ),
        (
            "book.xlsm",
            "edited.xlsm",
            {
                "keep_vba": True,
                "recalculation": "auto",
                "edits": [
                    {"sheet": "Sheet1", "type": "cell_value", "ref": "A1", "value": "1"}
                ],
            },
        ),
    ):
        with pytest.raises(DocumentSkillsError):
            parse_xlsx_request(
                {
                    "operation": "xlsx.edit",
                    "input": input_name,
                    "output": output_name,
                    "arguments": arguments,
                }
            )


@pytest.mark.parametrize(
    ("source", "output", "keep_vba"),
    [
        ("base.xltx", "created.xlsx", None),
        ("base.xltm", "created.xlsm", True),
    ],
)
def test_contract_accepts_only_exact_template_transitions(
    source: str,
    output: str,
    keep_vba: bool | None,
) -> None:
    arguments = {} if keep_vba is None else {"keep_vba": keep_vba}
    parsed = parse_xlsx_request(
        {
            "operation": "xlsx.template.instantiate",
            "input": source,
            "output": output,
            "arguments": arguments,
        }
    )
    assert parsed.arguments["recalculation"] == "skip"
    with pytest.raises(DocumentSkillsError):
        parse_xlsx_request(
            {
                "operation": "xlsx.template.instantiate",
                "input": source,
                "output": "wrong.xlsx" if output.endswith(".xlsm") else "wrong.xlsm",
                "arguments": arguments,
            }
        )


def test_macro_read_and_inspect_are_inert(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = create_package_fixture(tmp_path / "signed.xlsm", "xlsm", signed=True)
    service = XlsxService(project_root, libreoffice=_MustNotRunProvider())

    read = service.execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(source), "arguments": {}},
    )
    inspected = service.execute(
        "xlsx.inspect.structure",
        {
            "operation": "xlsx.inspect.structure",
            "input": str(source),
            "arguments": {},
        },
    )

    assert read["status"] in {"success", "degraded"}
    read_result = read["diagnostics"]["operation_result"]
    assert read_result["macro"]["vba_execution"] == "not_executed"
    assert read_result["recalculation"]["reason"] == "macro-workbook-provider-disabled"
    inspect_result = inspected["diagnostics"]["operation_result"]
    assert inspect_result["mutation_authorized"] is False
    assert inspect_result["dangerous_content"]["categories"]["vba"]
    assert inspect_result["macro"]["signature_state"] == (
        "present_not_cryptographically_verified"
    )


def test_macro_read_still_rejects_external_targets(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = create_package_fixture(
        tmp_path / "external.xlsm",
        "xlsm",
        external_target=True,
    )

    result = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(source), "arguments": {}},
    )

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_ARCHIVE_UNSAFE"


def test_macro_edit_preserves_vba_and_reports_signature_invalidation(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = create_package_fixture(tmp_path / "source.xlsm", "xlsm", signed=True)
    output = tmp_path / "edited.xlsm"
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    result = XlsxService(project_root, libreoffice=_MustNotRunProvider()).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "keep_vba": True,
                "recalculation": "skip",
                "edits": [
                    {
                        "sheet": "Sheet1",
                        "type": "cell_value",
                        "ref": "A1",
                        "value": "11",
                    }
                ],
            },
        },
    )

    assert result["status"] == "degraded"
    evidence = result["diagnostics"]["operation_result"]["macro"]
    assert evidence["vba_payload"] == "preserved"
    assert evidence["signature_state"] == "invalidated_by_package_mutation"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    before = OpcPackage.open(
        source,
        allowed_inert_categories=allowed_inert_categories("xlsm"),
    )
    after = OpcPackage.open(
        output,
        allowed_inert_categories=allowed_inert_categories("xlsm"),
    )
    assert before.part_hashes["xl/vbaProject.bin"] == after.part_hashes["xl/vbaProject.bin"]
    assert before.part_hashes["xl/vbaProjectSignature.bin"] == after.part_hashes[
        "xl/vbaProjectSignature.bin"
    ]
    reopened = load_workbook(output, data_only=False, keep_vba=True)
    assert reopened["Sheet1"]["A1"].value == 11
    assert reopened.vba_archive is not None
    assert reopened.vba_archive.read("xl/vbaProject.bin") == before.parts[
        "xl/vbaProject.bin"
    ]
    reopened.vba_archive.close()
    reopened.close()


def test_xltx_template_instantiates_and_applies_optional_edit(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = create_package_fixture(tmp_path / "base.xltx", "xltx")
    output = tmp_path / "instance.xlsx"

    result = XlsxService(project_root).execute(
        "xlsx.template.instantiate",
        {
            "operation": "xlsx.template.instantiate",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "recalculation": "skip",
                "edits": [
                    {
                        "sheet": "Sheet1",
                        "type": "cell_value",
                        "ref": "A1",
                        "value": "12",
                    }
                ],
            },
        },
    )

    assert result["status"] == "degraded"
    output_package = OpcPackage.open(output)
    assert output_package.workbook_format == "xlsx"
    assert output_package.content_type_for("xl/workbook.xml") == WORKBOOK_CONTENT_TYPES[
        "xlsx"
    ]
    assert CONTENT_TYPES in result["diagnostics"]["operation_result"]["preservation"][
        "changed_parts"
    ]
    reopened = load_workbook(output, data_only=False)
    assert reopened["Sheet1"]["A1"].value == 12
    reopened.close()


def test_xltm_template_preserves_vba_into_xlsm(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = create_package_fixture(tmp_path / "base.xltm", "xltm", signed=True)
    output = tmp_path / "instance.xlsm"

    result = XlsxService(project_root).execute(
        "xlsx.template.instantiate",
        {
            "operation": "xlsx.template.instantiate",
            "input": str(source),
            "output": str(output),
            "arguments": {"keep_vba": True, "recalculation": "skip"},
        },
    )

    assert result["status"] == "degraded"
    macro = result["diagnostics"]["operation_result"]["macro"]
    assert macro["vba_payload"] == "preserved"
    assert macro["signature_state"] == "invalidated_by_package_mutation"
    package = OpcPackage.open(
        output,
        allowed_inert_categories=allowed_inert_categories("xlsm"),
    )
    assert package.workbook_format == "xlsm"


def test_template_security_failure_preserves_existing_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = create_package_fixture(
        tmp_path / "unsafe.xltm",
        "xltm",
        external_target=True,
    )
    output = tmp_path / "existing.xlsm"
    output.write_bytes(b"existing-template-destination")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    result = XlsxService(project_root).execute(
        "xlsx.template.instantiate",
        {
            "operation": "xlsx.template.instantiate",
            "input": str(source),
            "output": str(output),
            "arguments": {"keep_vba": True, "recalculation": "skip"},
        },
    )

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_ARCHIVE_UNSAFE"
    assert output.read_bytes() == b"existing-template-destination"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash


def test_legacy_xls_conversion_requires_provider_and_validates_output(
    project_root: Path,
    tmp_path: Path,
) -> None:
    canned = create_package_fixture(tmp_path / "canned.xlsx", "xlsx").read_bytes()
    source = tmp_path / "legacy.xls"
    source.write_bytes(b"legacy-binary-workbook")
    output = tmp_path / "converted.xlsx"
    provider = _LegacyProvider(canned)

    result = XlsxService(project_root, libreoffice=provider).execute(
        "xlsx.convert",
        {
            "operation": "xlsx.convert",
            "input": str(source),
            "output": str(output),
            "arguments": {"source_format": "xls", "target_format": "xlsx"},
        },
    )

    assert result["status"] == "degraded"
    assert result["provider_chain"] == ["libreoffice"]
    assert provider.calls == [(source.resolve(), "xlsx")]
    assert OpcPackage.open(output).workbook_format == "xlsx"
    losses = result["diagnostics"]["operation_result"]["semantic_losses"]
    assert [item["code"] for item in losses] == ["legacy-provider-conversion"]


def test_legacy_contract_rejects_reverse_or_tabular_options() -> None:
    for source_format, target_format, input_name, output_name, extra in (
        ("xlsx", "xls", "input.xlsx", "output.xls", {}),
        ("xls", "json", "input.xls", "output.json", {}),
        ("xls", "xlsx", "input.xls", "output.xlsx", {"values": {}}),
    ):
        with pytest.raises(DocumentSkillsError):
            parse_xlsx_request(
                {
                    "operation": "xlsx.convert",
                    "input": input_name,
                    "output": output_name,
                    "arguments": {
                        "source_format": source_format,
                        "target_format": target_format,
                        **extra,
                    },
                }
            )


def test_legacy_xls_conversion_without_provider_preserves_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "legacy.xls"
    source.write_bytes(b"legacy")
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing")

    result = XlsxService(project_root).execute(
        "xlsx.convert",
        {
            "operation": "xlsx.convert",
            "input": str(source),
            "output": str(output),
            "arguments": {"source_format": "xls", "target_format": "xlsx"},
        },
    )

    assert result["status"] == "unavailable"
    assert output.read_bytes() == b"existing"


def test_legacy_xls_conversion_rejects_macro_enabled_provider_output(
    project_root: Path,
    tmp_path: Path,
) -> None:
    dangerous = create_package_fixture(
        tmp_path / "provider-output.xlsm",
        "xlsm",
    ).read_bytes()
    source = tmp_path / "legacy.xls"
    source.write_bytes(b"legacy")
    output = tmp_path / "existing-dangerous.xlsx"
    output.write_bytes(b"existing-legacy-destination")

    result = XlsxService(
        project_root,
        libreoffice=_LegacyProvider(dangerous),
    ).execute(
        "xlsx.convert",
        {
            "operation": "xlsx.convert",
            "input": str(source),
            "output": str(output),
            "arguments": {"source_format": "xls", "target_format": "xlsx"},
        },
    )

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_ARCHIVE_UNSAFE"
    assert output.read_bytes() == b"existing-legacy-destination"


class _MustNotRunProvider:
    def recalculate_xlsx_artifact(self, _path: Path) -> None:
        raise AssertionError("Macro-enabled reads must not invoke LibreOffice.")


class _LegacyProvider:
    def __init__(self, payload: bytes) -> None:
        self.payload = payload
        self.calls: list[tuple[Path, str]] = []

    def convert_legacy_required(self, path: Path, *, target_format: str) -> bytes:
        self.calls.append((path.resolve(), target_format))
        return self.payload

    def diagnostics(self) -> dict[str, str]:
        return {"path": "fake-soffice", "version": "test"}
