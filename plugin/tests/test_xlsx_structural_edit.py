"""Row/column structural edit integration and consumer-reopen tests."""

import hashlib
from pathlib import Path
from xml.etree.ElementTree import tostring

from document_skills_core.formats.xlsx.constants import NS
from document_skills_core.formats.xlsx.package import OpcPackage
from document_skills_core.formats.xlsx.service import XlsxService
from test_xlsx_pivot import _request as _pivot_request
from test_xlsx_pivot import _write_source as _write_pivot_source


def _service(project_root: Path) -> XlsxService:
    return XlsxService(project_root)


def _complex_source(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.defined_name import DefinedName

    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Name", "Amount", "Note", "Merged", None, "Link"])
    for row_number, amount in enumerate((10, 20, 30, 40), start=2):
        data.cell(row_number, 1, f"Item {row_number - 1}")
        data.cell(row_number, 2, amount)
    data["B6"] = "=SUM(B2:B5)"
    data["F2"] = "Jump"
    data["F2"].hyperlink = Hyperlink(
        ref="F2",
        location="Data!A4",
        display="Jump",
    )
    data.merge_cells("D2:E2")
    data["D2"] = "Merged"
    data.freeze_panes = "A4"
    data.print_area = "A1:H6"
    data.print_title_rows = "1:1"
    table = Table(displayName="DataTable", ref="A1:B5")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    data.add_table(table)
    validation = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2="100",
    )
    validation.add("B2:B5")
    data.add_data_validation(validation)
    data.conditional_formatting.add(
        "B2:B5",
        CellIsRule(
            operator="greaterThan",
            formula=["10"],
            fill=PatternFill(fill_type="solid", fgColor="FFFF00"),
        ),
    )
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    data.add_chart(chart, "H4")
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "=Data!B6"
    summary_validation = DataValidation(type="custom", formula1="Data!B4>0")
    summary_validation.add("A2")
    summary.add_data_validation(summary_validation)
    workbook.defined_names.add(
        DefinedName("DataRange", attr_text="Data!$A$2:$B$5")
    )
    workbook.save(path)


def _simple_column_source(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName

    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data["A1"] = 2
    data["B1"] = 3
    data["C1"] = "=A1+B1"
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "=Data!C1"
    workbook.defined_names.add(
        DefinedName("InputRow", attr_text="Data!$A$1:$C$1")
    )
    workbook.save(path)


def test_row_insert_migrates_package_wide_references(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "complex-source.xlsx"
    _complex_source(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "row-inserted.xlsx"

    result = _service(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "row_insert",
                        "ref": "3",
                        "count": 2,
                    }
                ],
                "expected_edits": 1,
            },
        },
    )

    assert result["status"] == "degraded"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    reopened = load_workbook(output, data_only=False)
    data = reopened["Data"]
    assert data["B8"].value == "=SUM(B2:B7)"
    assert reopened["Summary"]["A1"].value == "=Data!B8"
    assert reopened["Summary"].data_validations.dataValidation[0].formula1 == "Data!B6>0"
    assert data.tables["DataTable"].ref == "A1:B7"
    assert str(data.data_validations.dataValidation[0].sqref) == "B2:B7"
    assert str(next(iter(data.conditional_formatting)).sqref) == "B2:B7"
    assert [str(cell_range) for cell_range in data.merged_cells.ranges] == ["D2:E2"]
    assert data.freeze_panes == "A6"
    assert data["F2"].hyperlink.location == "Data!A6"
    assert reopened.defined_names["DataRange"].attr_text == "Data!$A$2:$B$7"
    assert data.print_area == "'Data'!$A$1:$H$8"
    assert data._charts[0].ser[0].val.numRef.f == "'Data'!$B$2:$B$7"
    assert data._charts[0].anchor._from.row == 5
    formula_cells = result["diagnostics"]["operation_result"]["formula_state"]["cells"]
    assert formula_cells["Data!B8"]["formula"] == "SUM(B2:B7)"
    assert formula_cells["Data!B8"]["state"] == "recalculation_required"
    assert formula_cells["Summary!A1"]["formula"] == "Data!B8"
    assert "Data!B6" not in formula_cells
    assert any(
        gate["id"] == "operation.mutation-semantics"
        and gate["outcome"] == "pass"
        for gate in result["validation"]["gates"]
    )


def test_column_insert_migrates_formulas_and_defined_names(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "column-source.xlsx"
    _simple_column_source(source)
    output = tmp_path / "column-inserted.xlsx"
    result = _service(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "column_insert",
                        "ref": "B",
                        "count": 1,
                    }
                ]
            },
        },
    )

    assert result["status"] == "degraded"
    reopened = load_workbook(output, data_only=False)
    assert reopened["Data"]["C1"].value == 3
    assert reopened["Data"]["D1"].value == "=A1+C1"
    assert reopened["Summary"]["A1"].value == "=Data!D1"
    assert reopened.defined_names["InputRow"].attr_text == "Data!$A$1:$D$1"


def test_structural_delete_that_would_create_ref_error_fails_without_promotion(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "delete-source.xlsx"
    _simple_column_source(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing-structural-destination")

    result = _service(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "column_delete",
                        "ref": "B",
                        "count": 1,
                    }
                ]
            },
        },
    )

    assert result["status"] == "enhancement_required"
    assert output.read_bytes() == b"existing-structural-destination"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash


def test_special_formula_structural_edit_fails_without_promotion(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook

    ordinary = tmp_path / "ordinary-formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = 1
    sheet["B1"] = "=A1"
    workbook.save(ordinary)
    workbook.close()

    package = OpcPackage.open(ordinary)
    root = package.xml("xl/worksheets/sheet1.xml")
    formula = root.find(f".//{{{NS['main']}}}f")
    assert formula is not None
    formula.attrib.update({"t": "array", "ref": "B1"})
    source = tmp_path / "array-formula.xlsx"
    package.write_copy(
        source,
        changed_parts={
            "xl/worksheets/sheet1.xml": tostring(
                root,
                encoding="UTF-8",
                xml_declaration=True,
            )
        },
    )
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "existing-special-formula.xlsx"
    output.write_bytes(b"existing-special-formula-destination")

    result = _service(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "row_insert", "ref": "1", "count": 1}
                ]
            },
        },
    )

    assert result["status"] == "enhancement_required"
    assert result["errors"][0]["details"]["capability"] == (
        "xlsx.special-formula-structural-edit"
    )
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    assert output.read_bytes() == b"existing-special-formula-destination"


def test_pivot_structural_edit_fails_without_promotion(
    project_root: Path,
    tmp_path: Path,
) -> None:
    pivot_source = _write_pivot_source(tmp_path / "pivot-source.xlsx")
    source = tmp_path / "native-pivot.xlsx"
    created = _service(project_root).execute(
        "xlsx.pivot.create",
        _pivot_request(pivot_source, source),
    )
    assert created["status"] == "success", created
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "existing-pivot-structural.xlsx"
    output.write_bytes(b"existing-pivot-structural-destination")

    result = _service(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "row_insert", "ref": "2", "count": 1}
                ]
            },
        },
    )

    assert result["status"] == "enhancement_required"
    assert result["errors"][0]["details"]["capability"] == (
        "xlsx.pivot-structural-edit"
    )
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    assert output.read_bytes() == b"existing-pivot-structural-destination"


def test_table_column_mutation_that_needs_table_schema_edit_fails_closed(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "table-source.xlsx"
    _complex_source(source)
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing-table-destination")

    result = _service(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "column_insert",
                        "ref": "B",
                        "count": 1,
                    }
                ]
            },
        },
    )

    assert result["status"] == "enhancement_required"
    assert output.read_bytes() == b"existing-table-destination"


def test_structural_insert_that_would_overflow_sheet_bounds_fails_closed(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook

    source = tmp_path / "boundary-source.xlsx"
    workbook = Workbook()
    workbook.active["XFD1"] = "edge"
    workbook.save(source)
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing-boundary-destination")

    result = _service(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Sheet",
                        "type": "column_insert",
                        "ref": "XFD",
                        "count": 1,
                    }
                ]
            },
        },
    )

    assert result["status"] == "enhancement_required"
    assert output.read_bytes() == b"existing-boundary-destination"
