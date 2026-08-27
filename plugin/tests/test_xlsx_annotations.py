"""Internal hyperlink, cell-note, and workbook-property tests."""

from hashlib import sha256
from pathlib import Path
import zipfile

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.service import XlsxService


def _metadata() -> dict:
    return {
        "title": "Annotated workbook",
        "creator": "Elftia Tester",
        "subject": "P1.5",
        "description": "Workbook metadata round trip",
        "keywords": "xlsx,metadata",
        "category": "Test",
        "last_modified_by": "Review Bot",
        "created": "2026-08-24T08:30:00Z",
        "modified": "2026-08-24T09:45:00+08:00",
        "company": "Elftia",
        "manager": "Quality Lead",
    }


def _workbook() -> dict:
    return {
        "metadata": _metadata(),
        "sheets": [
            {
                "name": "Report",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Jump", "type": "s"},
                            {"ref": "B2", "value": "Reviewed", "type": "s"},
                            {"ref": "C2", "value": "More", "type": "s"},
                        ]
                    }
                ],
                "columns": [],
                "number_formats": [],
                "hyperlinks": [
                    {
                        "ref": "A1",
                        "location": "Details!A1",
                        "display": "Open details",
                        "tooltip": "Internal workbook link",
                    }
                ],
                "comments": [
                    {"ref": "B2", "text": "Initial note", "author": "Alice"}
                ],
            },
            {
                "name": "Details",
                "rows": [{"cells": [{"ref": "A1", "value": "Target", "type": "s"}]}],
                "columns": [],
                "number_formats": [],
            },
        ],
        "defined_names": [],
        "tables": [],
        "charts": [],
        "chart_reference": None,
        "page_setup": None,
    }


def _create(project_root: Path, path: Path) -> dict:
    return XlsxService(project_root).execute(
        "xlsx.create",
        {
            "operation": "xlsx.create",
            "output": str(path),
            "arguments": {"workbook": _workbook()},
        },
    )


def test_annotations_create_reopens_and_projects_details(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "annotated.xlsx"
    result = _create(project_root, output)
    assert result["status"] == "success", result

    reopened = load_workbook(output)
    assert reopened["Report"]["A1"].hyperlink.location == "Details!A1"
    assert reopened["Report"]["A1"].hyperlink.target is None
    assert reopened["Report"]["B2"].comment.text == "Initial note"
    assert reopened["Report"]["B2"].comment.author == "Alice"
    assert reopened.properties.title == "Annotated workbook"
    assert reopened.properties.description == "Workbook metadata round trip"
    assert reopened.properties.lastModifiedBy == "Review Bot"

    read_result = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(output), "arguments": {}},
    )
    operation_result = read_result["diagnostics"]["operation_result"]
    assert operation_result["hyperlinks"] == [
        {
            "sheet": "Report",
            "ref": "A1",
            "location": "Details!A1",
            "display": "Open details",
            "tooltip": "Internal workbook link",
            "relationship_id": "",
            "target": None,
            "external": False,
        }
    ]
    assert operation_result["comments"][0] == {
        "sheet": "Report",
        "ref": "B2",
        "text": "Initial note",
        "author": "Alice",
        "part": "xl/comments1.xml",
        "vml_part": "xl/drawings/commentsDrawing1.vml",
    }
    assert operation_result["workbook_properties"] == _metadata()
    inspect_result = XlsxService(project_root).execute(
        "xlsx.inspect.structure",
        {
            "operation": "xlsx.inspect.structure",
            "input": str(output),
            "arguments": {},
        },
    )["diagnostics"]["operation_result"]
    assert inspect_result["hyperlinks"] == operation_result["hyperlinks"]
    assert inspect_result["comments"] == operation_result["comments"]
    assert inspect_result["workbook_properties"] == _metadata()


def test_annotations_create_is_deterministic(
    project_root: Path,
    tmp_path: Path,
) -> None:
    outputs = [tmp_path / "first.xlsx", tmp_path / "second.xlsx"]
    for output in outputs:
        result = _create(project_root, output)
        assert result["status"] == "success", result
    assert outputs[0].read_bytes() == outputs[1].read_bytes()


def test_comments_coexist_with_table_and_chart_relationships(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    from test_xlsx_chart_create import _workbook as chart_workbook

    workbook = chart_workbook()
    workbook["charts"] = workbook["charts"][:1]
    workbook["tables"] = [
        {
            "name": "AnnotatedData",
            "sheet": "Data",
            "ref": "A1:C5",
            "style": "TableStyleMedium2",
        }
    ]
    workbook["sheets"][0]["comments"] = [
        {"ref": "A2", "text": "Coexists", "author": "Elftia"}
    ]
    output = tmp_path / "related-objects.xlsx"
    result = XlsxService(project_root).execute(
        "xlsx.create",
        {
            "operation": "xlsx.create",
            "output": str(output),
            "arguments": {"workbook": workbook},
        },
    )
    assert result["status"] == "success", result
    reopened = load_workbook(output)
    assert "AnnotatedData" in reopened["Data"].tables
    assert len(reopened["Data"]._charts) == 1
    assert reopened["Data"]["A2"].comment.text == "Coexists"


def test_annotation_edits_are_deterministic(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "deterministic-source.xlsx"
    outputs = [tmp_path / "first-edited.xlsx", tmp_path / "second-edited.xlsx"]
    assert _create(project_root, source)["status"] == "success"
    edits = [
        {
            "sheet": "Report",
            "type": "hyperlink_update",
            "ref": "A1",
            "hyperlink": {"ref": "A1", "location": "Details!A1", "display": "Updated"},
        },
        {
            "sheet": "Report",
            "type": "comment_update",
            "ref": "B2",
            "comment": {"ref": "B2", "text": "Updated", "author": "Bob"},
        },
        {
            "sheet": "",
            "type": "workbook_properties",
            "properties": {"title": "Updated"},
        },
    ]
    for output in outputs:
        result = XlsxService(project_root).execute(
            "xlsx.edit",
            {
                "operation": "xlsx.edit",
                "input": str(source),
                "output": str(output),
                "arguments": {"edits": edits},
            },
        )
        assert result["status"] == "success", result
    assert outputs[0].read_bytes() == outputs[1].read_bytes()


def test_annotations_edit_crud_reopens_and_preserves_source(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    edited = tmp_path / "edited.xlsx"
    deleted = tmp_path / "deleted.xlsx"
    assert _create(project_root, source)["status"] == "success"
    source_hash = sha256(source.read_bytes()).hexdigest()
    edits = [
        {
            "sheet": "Report",
            "type": "hyperlink_update",
            "ref": "A1",
            "hyperlink": {
                "ref": "B1",
                "location": "Details!A1",
                "display": "Moved link",
                "tooltip": None,
            },
        },
        {
            "sheet": "Report",
            "type": "hyperlink_add",
            "hyperlink": {
                "ref": "C1",
                "location": "Details!A1",
                "display": None,
                "tooltip": "Second link",
            },
        },
        {
            "sheet": "Report",
            "type": "comment_update",
            "ref": "B2",
            "comment": {"ref": "B2", "text": "Updated note", "author": "Bob"},
        },
        {
            "sheet": "Report",
            "type": "comment_add",
            "comment": {"ref": "C2", "text": "Second note", "author": "Alice"},
        },
        {
            "sheet": "",
            "type": "workbook_properties",
            "properties": {
                "title": "Edited title",
                "company": "Updated Company",
                "modified": "2026-08-24T12:00:00Z",
            },
        },
    ]
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(edited),
            "arguments": {"edits": edits, "expected_edits": len(edits)},
        },
    )
    assert result["status"] == "success", result
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    reopened = load_workbook(edited)
    assert reopened["Report"]["A1"].hyperlink is None
    assert reopened["Report"]["B1"].hyperlink.location == "Details!A1"
    assert reopened["Report"]["C1"].hyperlink.location == "Details!A1"
    assert reopened["Report"]["B2"].comment.text == "Updated note"
    assert reopened["Report"]["B2"].comment.author == "Bob"
    assert reopened["Report"]["C2"].comment.text == "Second note"
    assert reopened.properties.title == "Edited title"
    projected_properties = XlsxService(project_root).execute(
        "xlsx.read",
        {"operation": "xlsx.read", "input": str(edited), "arguments": {}},
    )["diagnostics"]["operation_result"]["workbook_properties"]
    assert projected_properties["creator"] == "Elftia Tester"
    assert projected_properties["manager"] == "Quality Lead"
    assert projected_properties["company"] == "Updated Company"

    delete_edits = [
        {"sheet": "Report", "type": "hyperlink_delete", "ref": "B1"},
        {"sheet": "Report", "type": "hyperlink_delete", "ref": "C1"},
        {"sheet": "Report", "type": "comment_delete", "ref": "B2"},
        {"sheet": "Report", "type": "comment_delete", "ref": "C2"},
    ]
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(edited),
            "output": str(deleted),
            "arguments": {"edits": delete_edits, "expected_edits": 4},
        },
    )
    assert result["status"] == "success", result
    reopened = load_workbook(deleted)
    assert reopened["Report"]["B1"].hyperlink is None
    assert reopened["Report"]["C1"].hyperlink is None
    assert reopened["Report"]["B2"].comment is None
    assert reopened["Report"]["C2"].comment is None
    with zipfile.ZipFile(deleted) as archive:
        names = set(archive.namelist())
    assert "xl/comments1.xml" not in names
    assert "xl/drawings/commentsDrawing1.vml" not in names


def test_comment_adds_parts_to_unannotated_existing_workbook(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "plain.xlsx"
    output = tmp_path / "commented.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Value"
    workbook.save(source)
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Sheet",
                        "type": "comment_add",
                        "comment": {"ref": "A1", "text": "Added", "author": "Elftia"},
                    }
                ]
            },
        },
    )
    assert result["status"] == "success", result
    assert load_workbook(output)["Sheet"]["A1"].comment.text == "Added"


def test_comment_update_and_delete_accept_consumer_authored_parts(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.comments import Comment

    source = tmp_path / "consumer-comment.xlsx"
    updated = tmp_path / "consumer-comment-updated.xlsx"
    deleted = tmp_path / "consumer-comment-deleted.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Value"
    workbook.active["A1"].comment = Comment("Consumer note", "Consumer")
    workbook.save(source)
    update = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(updated),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Sheet",
                        "type": "comment_update",
                        "ref": "A1",
                        "comment": {"ref": "A1", "text": "Updated", "author": "Elftia"},
                    }
                ]
            },
        },
    )
    assert update["status"] == "success", update
    assert load_workbook(updated)["Sheet"]["A1"].comment.text == "Updated"
    delete = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(updated),
            "output": str(deleted),
            "arguments": {
                "edits": [{"sheet": "Sheet", "type": "comment_delete", "ref": "A1"}]
            },
        },
    )
    assert delete["status"] == "success", delete
    assert load_workbook(deleted)["Sheet"]["A1"].comment is None


@pytest.mark.parametrize(
    "location",
    [
        "https://example.com",
        "mailto:test@example.com",
        "[Other.xlsx]Sheet1!A1",
        "\\\\server\\share",
    ],
)
def test_external_hyperlink_authoring_remains_fail_closed(location: str) -> None:
    workbook = _workbook()
    workbook["sheets"][0]["hyperlinks"][0]["location"] = location
    with pytest.raises(DocumentSkillsError) as exc:
        parse_xlsx_request(
            {
                "operation": "xlsx.create",
                "output": "external.xlsx",
                "arguments": {"workbook": workbook},
            }
        )
    assert exc.value.code.value == "DS_ENHANCEMENT_REQUIRED"


@pytest.mark.parametrize(
    "timestamp",
    ["2026-08-24T08:30:00", "2026-99-24T08:30:00Z"],
)
def test_workbook_property_timestamp_rejects_invalid_value(timestamp: str) -> None:
    workbook = _workbook()
    workbook["metadata"]["created"] = timestamp
    with pytest.raises(DocumentSkillsError) as exc:
        parse_xlsx_request(
            {
                "operation": "xlsx.create",
                "output": "invalid.xlsx",
                "arguments": {"workbook": workbook},
            }
        )
    assert exc.value.code.value == "DS_REQUEST_INVALID"


def test_annotation_edit_failure_preserves_existing_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "existing.xlsx"
    assert _create(project_root, source)["status"] == "success"
    destination.write_bytes(b"existing-annotation-destination")
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(destination),
            "arguments": {
                "edits": [
                    {"sheet": "Report", "type": "comment_delete", "ref": "Z99"}
                ]
            },
        },
    )
    assert result["status"] == "invalid_request"
    assert destination.read_bytes() == b"existing-annotation-destination"
