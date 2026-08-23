"""Transactional slide delete, duplicate, same-deck copy, and cross-deck copy."""

from hashlib import sha256
from pathlib import Path
from typing import Any

from pptx import Presentation
import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError, ErrorCode
from document_skills_core.formats.pptx.contracts import parse_pptx_request
from document_skills_core.formats.pptx.create import create_pptx
from document_skills_core.formats.pptx.edit import edit_pptx
from document_skills_core.formats.pptx.mapping import map_slides
from document_skills_core.formats.pptx.package import OpcPackage
from document_skills_core.formats.pptx.read import read_pptx
from document_skills_core.formats.pptx.service import PptxService
from tests.fixtures.recipes.docx_fixture_support import PNG_1X1


def _slide(title: str, *, image: Path | None = None, rich: bool = False) -> dict[str, Any]:
    return {
        "layout": "content",
        "title": title,
        "shapes": [{"text": f"{title} body", "runs": []}],
        "table": {"rows": [{"cells": ["A", "B"]}]} if rich else None,
        "chart_reference": (
            {
                "title": f"{title} chart",
                "chart_type": "line",
                "categories": ["Q1", "Q2"],
                "series": [{"name": "Value", "values": [2.0, 5.0]}],
            }
            if rich
            else None
        ),
        "image_reference": (
            {
                "path": str(image),
                "content_type": "image/png",
                "alt_text": f"{title} image",
            }
            if image is not None
            else None
        ),
        "notes": f"{title} notes" if rich else None,
    }


def _deck(slides: list[dict[str, Any]], title: str = "Lifecycle") -> dict[str, Any]:
    return {
        "metadata": {"title": title, "creator": "Test", "subject": ""},
        "slide_size": {"cx": "9144000", "cy": "6858000", "type": "screen4x3"},
        "slides": slides,
    }


def _created(tmp_path: Path, name: str, slides: list[dict[str, Any]]) -> Path:
    output = tmp_path / name
    create_pptx(output, _deck(slides, name))
    return output


def test_contract_accepts_slide_lifecycle_primitives(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    parsed = parse_pptx_request({
        "schema_version": "1.0",
        "operation": "pptx.edit",
        "input": str(tmp_path / "target.pptx"),
        "output": str(tmp_path / "out.pptx"),
        "arguments": {
            "edits": [
                {"type": "slide_add", "position": 1, "slide": _slide("Added")},
                {"type": "slide_delete", "slide": 2},
                {"type": "slide_duplicate", "slide": 1, "position": 2},
                {
                    "type": "slide_copy",
                    "source": str(source),
                    "source_slide": 3,
                    "position": 1,
                },
                {"type": "slide_move", "slide": 2, "position": 1},
            ]
        },
    })
    assert [edit["type"] for edit in parsed.arguments["edits"]] == [
        "slide_add",
        "slide_delete",
        "slide_duplicate",
        "slide_copy",
        "slide_reorder",
    ]
    assert parsed.arguments["edits"][3]["source"] == source.resolve()


def test_slide_add_uses_target_layout_and_creates_native_objects(tmp_path: Path) -> None:
    image = tmp_path / "added.png"
    image.write_bytes(PNG_1X1)
    source = _created(tmp_path, "add-source.pptx", [_slide("Existing")])
    added = _slide("Added", image=image, rich=True)
    output = tmp_path / "added.pptx"
    result, manifest = edit_pptx(
        source,
        output,
        {"edits": [{"type": "slide_add", "position": 1, "slide": added}]},
    )
    package = OpcPackage.open(output)
    slides = map_slides(package)

    assert len(slides) == 2
    assert len(package.slide_layout_parts()) == 2
    assert len(package.slide_master_parts()) == 1
    assert len(package.theme_parts()) == 1
    assert len(package.media_parts()) == 1
    assert len(package.chart_parts()) == 1
    assert len(package.notes_slide_parts()) == 1
    assert len(package.notes_master_parts()) == 1
    assert slides[0]["layout"]["part"] in package.slide_layout_parts()
    assert any(name.startswith("ppt/slides/slide") for name in manifest.added)
    evidence = result["slide_lifecycle"][0]
    assert evidence["type"] == "slide_add"
    assert evidence["image"]["fallback"] == "native"
    assert evidence["chart"]["editable"] is True
    presentation = Presentation(output)
    assert any(getattr(shape, "has_chart", False) for shape in presentation.slides[0].shapes)


def test_slide_delete_removes_only_unreachable_dependency_graph(tmp_path: Path) -> None:
    image = tmp_path / "delete.png"
    image.write_bytes(PNG_1X1)
    source = _created(
        tmp_path,
        "delete-source.pptx",
        [_slide("Keep"), _slide("Delete", image=image, rich=True)],
    )
    source_hash = sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "deleted.pptx"
    result, manifest = edit_pptx(
        source,
        output,
        {"edits": [{"type": "slide_delete", "slide": 2}]},
    )
    package = OpcPackage.open(output)
    operation_result, _warnings = read_pptx(output, {})

    assert operation_result["slide_count"] == 1
    assert "ppt/slides/slide2.xml" in manifest.removed
    assert "ppt/charts/chart1.xml" in manifest.removed
    assert "ppt/media/image1.png" in manifest.removed
    assert "ppt/notesSlides/notesSlide2.xml" in manifest.removed
    assert package.slide_layout_parts()
    assert package.slide_master_parts()
    assert package.theme_parts()
    assert result["slide_lifecycle"][0]["type"] == "slide_delete"
    assert sha256(source.read_bytes()).hexdigest() == source_hash


def test_slide_duplicate_copies_media_chart_notes_and_reuses_layout(tmp_path: Path) -> None:
    image = tmp_path / "duplicate.png"
    image.write_bytes(PNG_1X1)
    source = _created(tmp_path, "duplicate-source.pptx", [_slide("Rich", image=image, rich=True)])
    output = tmp_path / "duplicated.pptx"
    result, manifest = edit_pptx(
        source,
        output,
        {"edits": [{"type": "slide_duplicate", "slide": 1, "position": 2}]},
    )
    package = OpcPackage.open(output)
    slides = map_slides(package)

    assert len(slides) == 2
    assert len(package.media_parts()) == 2
    assert len(package.chart_parts()) == 2
    assert len(package.notes_slide_parts()) == 2
    assert len(package.slide_layout_parts()) == 2
    assert package.parts[slides[0]["part"]] == package.parts[slides[1]["part"]]
    assert slides[0]["layout"]["part"] == slides[1]["layout"]["part"]
    assert any(name.startswith("ppt/media/") for name in manifest.added)
    evidence = result["slide_lifecycle"][0]
    assert evidence["cross_deck"] is False
    assert any(item["source_part"].startswith("ppt/charts/") for item in evidence["dependencies"])
    assert len(Presentation(output).slides) == 2


def test_same_deck_slide_copy_is_a_deep_copy(tmp_path: Path) -> None:
    image = tmp_path / "same-copy.png"
    image.write_bytes(PNG_1X1)
    source = _created(tmp_path, "same-source.pptx", [_slide("Rich", image=image, rich=True)])
    output = tmp_path / "same-copy.pptx"
    result, _manifest = edit_pptx(
        source,
        output,
        {"edits": [{"type": "slide_copy", "source_slide": 1, "position": 2}]},
    )
    package = OpcPackage.open(output)
    assert len(package.slide_parts()) == 2
    assert len(package.media_parts()) == 2
    assert len(package.chart_parts()) == 2
    assert result["slide_lifecycle"][0]["cross_deck"] is False


def test_cross_deck_slide_copy_imports_layout_master_theme_and_dependencies(tmp_path: Path) -> None:
    source_image = tmp_path / "cross.png"
    source_image.write_bytes(PNG_1X1)
    source = _created(tmp_path, "external-source.pptx", [_slide("Imported", image=source_image, rich=True)])
    target = _created(tmp_path, "target.pptx", [_slide("Existing")])
    source_hash = sha256(source.read_bytes()).hexdigest()
    target_hash = sha256(target.read_bytes()).hexdigest()
    output = tmp_path / "cross-copy.pptx"
    result, _manifest = edit_pptx(
        target,
        output,
        {
            "edits": [{
                "type": "slide_copy",
                "source": source,
                "source_slide": 1,
                "position": 2,
            }]
        },
    )
    package = OpcPackage.open(output)
    slides = map_slides(package)

    assert len(slides) == 2
    assert len(package.slide_master_parts()) == 2
    assert len(package.slide_layout_parts()) == 4
    assert len(package.theme_parts()) == 2
    assert len(package.media_parts()) == 1
    assert len(package.chart_parts()) == 1
    assert len(package.notes_slide_parts()) == 1
    assert len(package.notes_master_parts()) == 1
    assert slides[1]["layout"]["part"] != slides[0]["layout"]["part"]
    assert slides[1]["master"]["part"] != slides[0]["master"]["part"]
    evidence = result["slide_lifecycle"][0]
    assert evidence["cross_deck"] is True
    copied_targets = {item["target_part"] for item in evidence["dependencies"]}
    assert slides[1]["part"] in copied_targets
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    assert sha256(target.read_bytes()).hexdigest() == target_hash
    presentation = Presentation(output)
    assert len(presentation.slides) == 2
    assert any(getattr(shape, "has_chart", False) for shape in presentation.slides[1].shapes)


def test_precondition_failure_publishes_no_partial_output(tmp_path: Path) -> None:
    source = _created(tmp_path, "precondition-source.pptx", [_slide("One"), _slide("Two")])
    output = tmp_path / "must-not-exist.pptx"
    with pytest.raises(DocumentSkillsError) as captured:
        edit_pptx(
            source,
            output,
            {
                "edits": [
                    {"type": "slide_delete", "slide": 2},
                    {
                        "type": "slide_duplicate",
                        "slide": 1,
                        "precondition_sha256": "0" * 64,
                    },
                ]
            },
        )
    assert captured.value.code == ErrorCode.VALIDATION_FAILED
    assert not output.exists()


def test_service_validates_and_promotes_slide_lifecycle(
    project_root: Path,
    tmp_path: Path,
) -> None:
    image = tmp_path / "service.png"
    image.write_bytes(PNG_1X1)
    source = _created(tmp_path, "service-source.pptx", [_slide("Existing")])
    output = tmp_path / "service-output.pptx"
    service = PptxService(project_root)
    result = service.execute("pptx.edit", {
        "schema_version": "1.0",
        "operation": "pptx.edit",
        "input": str(source),
        "output": str(output),
        "arguments": {
            "edits": [{
                "type": "slide_add",
                "position": 2,
                "slide": _slide("Added", image=image, rich=True),
            }]
        },
    })

    assert result["status"] == "success", result
    assert output.is_file()
    outcomes = {gate["id"]: gate["outcome"] for gate in result["validation"]["gates"]}
    assert outcomes["operation.mutation-semantics"] == "pass"
    assert len(Presentation(output).slides) == 2
