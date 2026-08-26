"""PPTX public command surface tests — frozen uv subprocess boundary."""

from hashlib import sha256
import json
from pathlib import Path
import subprocess

import pytest

from document_skills_core.formats.pptx.package import OpcPackage
from tests.fixtures.recipes.docx_fixture_support import PNG_1X1


def _public(
    project_root: Path,
    *arguments: str,
    check: bool = True,
    cwd: Path | None = None,
) -> dict[str, object]:
    process = subprocess.run(
        [
            "uv",
            "run",
            "--project",
            str(project_root),
            "--frozen",
            "python",
            str(project_root / "skills/document-pptx/scripts/run.py"),
            *arguments,
        ],
        cwd=cwd or project_root,
        check=False,
        capture_output=True,
        text=False,
        shell=False,
        timeout=60,
    )
    if check:
        assert process.returncode == 0, (
            process.stderr.decode("utf-8", errors="replace")
            or process.stdout.decode("utf-8", errors="replace")
        )
    assert process.stderr == b""
    text = process.stdout.decode("utf-8", errors="strict")
    decoder = json.JSONDecoder()
    payload, end = decoder.raw_decode(text)
    assert text[end:].strip() == ""
    assert type(payload) is dict
    return payload


def _request(tmp_path: Path, name: str, payload: dict[str, object]) -> Path:
    path = tmp_path / name
    path.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
        newline="\n",
    )
    return path


def _deck(image_path: Path) -> dict[str, object]:
    return {
        "metadata": {"title": "Public PPTX", "creator": "Test", "subject": ""},
        "slide_size": {"cx": "9144000", "cy": "6858000", "type": "screen4x3"},
        "slides": [
            {
                "layout": "title",
                "title": "Title",
                "shapes": [
                    {"text": "Body", "runs": [{"text": "Body", "style": None}]}
                ],
                "table": None,
                "chart_reference": None,
                "image_reference": None,
                "notes": "Speaker notes",
            },
            {
                "layout": "content",
                "title": "Content",
                "shapes": [
                    {"text": "Content body", "runs": [{"text": "Content", "style": None}]}
                ],
                "table": {"rows": [{"cells": ["A", "B"]}]},
                "chart_reference": {
                    "title": "Chart",
                    "chart_type": "column",
                    "categories": ["North", "South"],
                    "series": [{"name": "Sales", "values": [4, 7]}],
                    "legend": {"show": True, "position": "bottom"},
                    "axes": {
                        "category": {"title": "Region", "number_format": "General"},
                        "value": {"title": "Units", "number_format": "0"},
                    },
                    "data_labels": {"show_value": True},
                    "colors": ["3366CC"],
                },
                "image_reference": {
                    "path": str(image_path),
                    "content_type": "image/png",
                    "fit": "contain",
                    "alt_text": "Public embedded image",
                },
                "notes": None,
            },
        ],
    }


@pytest.fixture
def public_created(project_root: Path, tmp_path: Path) -> Path:
    from pptx import Presentation

    output = tmp_path / "public-created.pptx"
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Title"
    first.placeholders[1].text = "Body"
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Content"
    second.placeholders[1].text = "Content body"
    presentation.save(output)
    return output


def test_public_capabilities_list_pptx_operations(project_root: Path) -> None:
    report = _public(project_root, "capabilities", "--json")
    operations = {item["operation"]: item for item in report["operations"]}
    core = (
        "pptx.read",
        "pptx.inspect.structure",
        "pptx.outline.create",
        "pptx.create",
        "pptx.create.from-markdown",
        "pptx.edit",
    )
    for operation in core:
        assert operation in operations
        assert operations[operation]["available"], operations[operation]

    # Provider-backed operations are environment-dependent by design:
    # `pptx.create.from-html` is served solely by `html-browser`, which is
    # optional (`required=False`) and needs Node, the locked Playwright library
    # and a launchable Chromium-family browser. It is genuinely unavailable
    # without one. `tests/test_runtime.py` asserts that sole-provider binding.
    #
    # The blanket "every pptx operation is available" assertion this replaces
    # asserted the opposite of what capability discovery exists for. Scope of
    # that failure, from run 31529508680: the two ubuntu legs failed here
    # (1 failed, 1026 passed, 26 skipped); the windows legs failed a different
    # test and the macos legs failed 86. This repair addresses the ubuntu
    # failure only.
    #
    # The check below is a regression guard over `reports.py`, not an exercised
    # negative path: that module sets a constant non-empty reason whenever
    # `available` is false, so today's production output cannot reach the
    # failure branch. It exists so a later change cannot publish a missing
    # provider as a silent, reasonless absence.
    for operation, item in operations.items():
        if "pptx" not in operation or operation in core:
            continue
        assert item["available"] or item["reason"], item
    schema = operations["pptx.validate.schema"]
    assert schema["providers"] in ([], ["dotnet-openxml"])
    assert report["validation"]["schema"] == (
        "available" if schema["available"] else "unavailable"
    )
    for operation in ("pptx.convert.legacy", "pptx.convert.pdf", "pptx.render"):
        item = operations[operation]
        assert item["providers"] in ([], ["libreoffice"])
        assert item["fidelity"] == ("enhanced" if item["available"] else "none")
    libreoffice_available = operations["pptx.render"]["available"]
    assert operations["pptx.convert.pdf"]["available"] is libreoffice_available
    assert operations["pptx.convert.legacy"]["available"] is libreoffice_available
    assert report["validation"]["visual"] == (
        "available" if libreoffice_available else "unavailable"
    )


def test_public_doctor_succeeds(project_root: Path) -> None:
    report = _public(project_root, "doctor", "--json")
    assert report["status"] == "healthy"


def test_public_outline_and_markdown_content_entry(
    project_root: Path,
    tmp_path: Path,
) -> None:
    outline_output = tmp_path / "public-outline.json"
    outline_request = _request(tmp_path, "outline-request.json", {
        "schema_version": "1.0",
        "operation": "pptx.outline.create",
        "output": str(outline_output),
        "arguments": {
            "title": "Public plan",
            "slides": [{"title": "Opening", "bullets": ["Context"]}],
        },
    })
    outline_result = _public(
        project_root,
        "run",
        "--request",
        str(outline_request),
    )
    assert outline_result["status"] == "success", outline_result
    outline = json.loads(outline_output.read_text(encoding="utf-8"))
    assert outline["presentation_generated"] is False

    image = tmp_path / "markdown-image.png"
    image.write_bytes(PNG_1X1)
    source = tmp_path / "public-deck.md"
    source.write_text(
        "# Public deck\n\nNative paragraph.\n\n"
        "## Evidence\n\n| Item | Value |\n| --- | --- |\n| A | 1 |\n\n"
        "![Evidence](markdown-image.png)\n",
        encoding="utf-8",
        newline="\n",
    )
    output = tmp_path / "public-markdown.pptx"
    markdown_request = _request(tmp_path, "markdown-request.json", {
        "schema_version": "1.0",
        "operation": "pptx.create.from-markdown",
        "input": str(source),
        "output": str(output),
        "arguments": {},
    })
    markdown_result = _public(
        project_root,
        "run",
        "--request",
        str(markdown_request),
    )
    assert markdown_result["status"] == "success", markdown_result
    package = OpcPackage.open(output)
    assert len(package.slide_parts()) == 2
    assert len(package.media_parts()) == 1
    assert markdown_result["diagnostics"]["operation_result"]["creation"]["has_table"] is True


def test_public_create_and_read(project_root: Path, public_created: Path, tmp_path: Path) -> None:
    read_request = _request(
        tmp_path,
        "read.json",
        {
            "schema_version": "1.0",
            "operation": "pptx.read",
            "input": str(public_created),
            "arguments": {},
        },
    )
    result = _public(project_root, "run", "--request", str(read_request))
    assert result["status"] == "success"
    assert result["diagnostics"]["operation_result"]["slide_count"] == 2


def test_public_inspect_inert(project_root: Path, public_created: Path, tmp_path: Path) -> None:
    inspect_request = _request(
        tmp_path,
        "inspect.json",
        {
            "schema_version": "1.0",
            "operation": "pptx.inspect.structure",
            "input": str(public_created),
            "arguments": {},
        },
    )
    result = _public(project_root, "run", "--request", str(inspect_request))
    assert result["status"] == "success"
    assert result["diagnostics"]["operation_result"]["mutation_authorized"] is False


def test_public_edit_produces_distinct_output(project_root: Path, public_created: Path, tmp_path: Path) -> None:
    output = tmp_path / "public-edited.pptx"
    edit_request = _request(
        tmp_path,
        "edit.json",
        {
            "schema_version": "1.0",
            "operation": "pptx.edit",
            "input": str(public_created),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"slide": 1, "type": "slide_text", "ref": "", "value": "Modified"},
                ],
            },
        },
    )
    result = _public(project_root, "run", "--request", str(edit_request))
    assert result["status"] == "success"
    assert output.is_file()
    assert public_created.is_file()
    assert any(item["path"] == str(output.resolve()) for item in result["artifacts"])


def test_public_reorder_preserves_slides(project_root: Path, public_created: Path, tmp_path: Path) -> None:
    output = tmp_path / "public-reordered.pptx"
    edit_request = _request(
        tmp_path,
        "reorder.json",
        {
            "schema_version": "1.0",
            "operation": "pptx.edit",
            "input": str(public_created),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"slide": 2, "type": "slide_reorder", "ref": "", "value": "1"},
                ],
            },
        },
    )
    result = _public(project_root, "run", "--request", str(edit_request))
    assert result["status"] == "success"
    op_result = result["diagnostics"]["operation_result"]
    assert "reorder" in op_result
    assert op_result["reorder"][0]["shape_ids_preserved"] is True


def test_public_validate_reopens_valid_pptx(project_root: Path, public_created: Path) -> None:
    result = _public(
        project_root,
        "validate",
        "--input",
        str(public_created),
        "--json",
    )
    outcomes = {item["id"]: item["outcome"] for item in result["gates"]}
    assert result["status"] == "pass"
    assert outcomes["provider.reopen"] == "pass"
    assert outcomes["operation.pptx-deep-validation"] == "pass"
    assert outcomes["visual.render"] == "unavailable"
    assert outcomes["schema.full"] == "unavailable"


def test_public_schema_validation_is_honestly_provider_gated(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
) -> None:
    before = sha256(public_created.read_bytes()).hexdigest()
    request = _request(tmp_path, "schema.json", {
        "schema_version": "1.0",
        "operation": "pptx.validate.schema",
        "input": str(public_created),
        "arguments": {},
        "options": {"fidelity": "enhanced"},
    })
    result = _public(project_root, "run", "--request", str(request), check=False)

    assert sha256(public_created.read_bytes()).hexdigest() == before
    assert result["status"] in {"success", "unavailable"}, result
    if result["status"] == "success":
        assert result["provider_chain"] == ["dotnet-openxml"]
        schema_gate = next(
            gate
            for gate in result["validation"]["gates"]
            if gate["id"] == "schema.full"
        )
        assert schema_gate["outcome"] == "pass", schema_gate
    else:
        assert result["errors"], result
        reason = result["errors"][0]["message"]
        assert isinstance(reason, str) and reason, result


def test_public_libreoffice_outputs_are_honestly_provider_gated(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
) -> None:
    capabilities = _public(project_root, "capabilities", "--json")
    operations = {
        item["operation"]: item
        for item in capabilities["operations"]
    }
    source_hash = sha256(public_created.read_bytes()).hexdigest()
    cases = [
        ("pptx.convert.pdf", tmp_path / "public-converted.pdf"),
        ("pptx.render", tmp_path / "public-render.zip"),
    ]
    for operation, output in cases:
        request = _request(tmp_path, f"{output.stem}.json", {
            "schema_version": "1.0",
            "operation": operation,
            "input": str(public_created),
            "output": str(output),
            "arguments": {},
            "options": {"fidelity": "enhanced"},
        })
        result = _public(
            project_root,
            "run",
            "--request",
            str(request),
            check=False,
        )
        assert sha256(public_created.read_bytes()).hexdigest() == source_hash
        if operations[operation]["available"]:
            assert result["status"] == "success", result
            assert result["provider_chain"] == ["libreoffice"]
            assert output.is_file()
        else:
            assert result["status"] == "unavailable"
            assert not output.exists()


def test_public_unknown_operation_rejected(project_root: Path, tmp_path: Path) -> None:
    request = _request(
        tmp_path,
        "unknown.json",
        {
            "schema_version": "1.0",
            "operation": "pptx.nonexistent",
            "input": str(tmp_path / "input.pptx"),
            "arguments": {},
        },
    )
    result = _public(project_root, "run", "--request", str(request), check=False)
    assert result["status"] == "invalid_request"


def test_public_create_is_deterministic(project_root: Path, tmp_path: Path) -> None:
    output1 = tmp_path / "d1.pptx"
    output2 = tmp_path / "d2.pptx"
    image = tmp_path / "img.png"
    image.write_bytes(PNG_1X1)
    results = []
    for out in [output1, output2]:
        req = _request(tmp_path, f"create_{out.stem}.json", {
            "schema_version": "1.0",
            "operation": "pptx.create",
            "output": str(out),
            "arguments": {"deck": _deck(image)},
        })
        results.append(
            _public(project_root, "run", "--request", str(req), check=False)
        )
    assert all(result["status"] == "success" for result in results)
    assert all(result["artifacts"] for result in results)
    assert output1.exists() and output2.exists()
    candidate_hashes = [
        next(
            gate["evidence"]["sha256"]
            for gate in result["validation"]["gates"]
            if gate["id"] == "artifact.exists-size"
        )
        for result in results
    ]
    assert candidate_hashes[0] == candidate_hashes[1]


def test_public_create_reopens_real_native_objects(project_root: Path, tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    image = tmp_path / "native.png"
    image.write_bytes(PNG_1X1)
    output = tmp_path / "native-objects.pptx"
    request = _request(tmp_path, "native-create.json", {
        "schema_version": "1.0",
        "operation": "pptx.create",
        "output": str(output),
        "arguments": {"deck": _deck(image)},
    })
    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "success"
    creation = result["diagnostics"]["operation_result"]["creation"]
    assert creation["images"][0]["embedded_media_part"] == "ppt/media/image1.png"
    assert creation["images"][0]["fallback"] == "native"
    assert creation["charts"][0]["editable"] is True
    gates = {gate["id"]: gate for gate in result["validation"]["gates"]}
    outcomes = {gate_id: gate["outcome"] for gate_id, gate in gates.items()}
    assert outcomes["visual.render"] == "unavailable"
    schema_gate = gates["schema.full"]
    assert schema_gate["outcome"] in {"pass", "unavailable"}, schema_gate
    assert schema_gate["validator"] == "dotnet-openxml"
    if schema_gate["outcome"] == "pass":
        assert schema_gate["required"] is True
        assert schema_gate["evidence"]["valid"] is True
        assert schema_gate["evidence"]["error_count"] == 0
    else:
        assert schema_gate["required"] is False
        reason = schema_gate["evidence"].get("reason")
        assert isinstance(reason, str) and reason, schema_gate
    presentation = Presentation(output)
    assert len(presentation.slides) == 2
    shapes = list(presentation.slides[1].shapes)
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
    chart_shapes = [shape for shape in shapes if getattr(shape, "has_chart", False)]
    assert len(chart_shapes) == 1
    assert list(chart_shapes[0].chart.series[0].values) == [4.0, 7.0]
    assert any(getattr(shape, "has_table", False) for shape in shapes)


def test_public_create_applies_design_tokens_and_reuses_template_graph(
    project_root: Path,
    tmp_path: Path,
) -> None:
    theme = {
        "name": "Public Design",
        "palette": {"accent1": "006D77", "accent2": "E29578"},
        "fonts": {"major": "Aptos Display", "minor": "Aptos"},
        "background": "F8F5EC",
        "default_text": {
            "title_color": "101820",
            "body_color": "203040",
            "title_size": 32,
            "body_size": 17,
            "bold_titles": True,
        },
        "default_shape": {"fill": "FFFFFF", "line": "006D77", "opacity": 0.95},
        "default_chart": {"colors": ["006D77", "E29578"]},
    }
    slide = {
        "layout": "content",
        "recipe": "comparison",
        "title": "Designed publicly",
        "shapes": [
            {"text": "Left", "runs": []},
            {"text": "Right", "runs": []},
        ],
        "table": None,
        "chart_reference": None,
        "image_reference": None,
        "notes": None,
    }
    deck = {
        "metadata": {"title": "Public design", "creator": "Test", "subject": ""},
        "slide_size": {"cx": "9144000", "cy": "6858000", "type": "screen4x3"},
        "theme": theme,
        "layout_tokens": {
            "safe_margins": {
                "top": 300_000,
                "right": 400_000,
                "bottom": 300_000,
                "left": 400_000,
            },
            "grid": {"columns": 12, "gutter": 200_000},
            "spacing": {"sm": 150_000, "md": 250_000, "lg": 400_000},
            "typography_scale": {"title": 32, "section": 25, "body": 17, "caption": 11},
        },
        "slides": [slide],
    }
    template = tmp_path / "public-design-template.pptx"
    template_request = _request(tmp_path, "public-design-template.json", {
        "schema_version": "1.0",
        "operation": "pptx.create",
        "output": str(template),
        "arguments": {"deck": deck},
    })
    template_result = _public(project_root, "run", "--request", str(template_request))
    assert template_result["status"] == "success", template_result
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = OpcPackage.open(template)
    design_parts = source.slide_master_parts() + source.slide_layout_parts() + source.theme_parts()
    design_hashes = {part: source.part_hashes[part] for part in design_parts}

    based_deck = {**deck, "slides": [{**slide, "recipe": "summary", "title": "Based publicly"}]}
    based_deck.pop("theme")
    output = tmp_path / "public-template-based.pptx"
    based_request = _request(tmp_path, "public-template-based.json", {
        "schema_version": "1.0",
        "operation": "pptx.create",
        "output": str(output),
        "arguments": {"deck": based_deck, "template": str(template)},
    })
    result = _public(project_root, "run", "--request", str(based_request))

    assert result["status"] == "success", result
    assert sha256(template.read_bytes()).hexdigest() == template_hash
    candidate = OpcPackage.open(output)
    assert {part: candidate.part_hashes[part] for part in design_parts} == design_hashes
    creation = result["diagnostics"]["operation_result"]["creation"]
    assert creation["template_reuse"]["source_sha256"] == template_hash
    assert creation["layout_recipes"][0]["name"] == "summary"
    outcomes = {gate["id"]: gate["outcome"] for gate in result["validation"]["gates"]}
    assert outcomes["operation.typed-design-correspondence"] == "pass"


def test_public_create_missing_image_fails_closed(project_root: Path, tmp_path: Path) -> None:
    output = tmp_path / "must-not-exist.pptx"
    request = _request(tmp_path, "missing-image.json", {
        "schema_version": "1.0",
        "operation": "pptx.create",
        "output": str(output),
        "arguments": {"deck": _deck(tmp_path / "missing.png")},
    })
    result = _public(project_root, "run", "--request", str(request), check=False)
    assert result["status"] == "invalid_request"
    assert not output.exists()


def test_public_slide_lifecycle_reopens_after_add_duplicate_delete(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    output = tmp_path / "public-lifecycle.pptx"
    request = _request(tmp_path, "public-lifecycle.json", {
        "schema_version": "1.0",
        "operation": "pptx.edit",
        "input": str(public_created),
        "output": str(output),
        "arguments": {
            "edits": [
                {
                    "type": "slide_add",
                    "position": 2,
                    "slide": {
                        "layout": "content",
                        "title": "Added publicly",
                        "shapes": [{"text": "Native body", "runs": []}],
                        "table": None,
                        "chart_reference": None,
                        "image_reference": None,
                        "notes": "Added notes",
                    },
                },
                {"type": "slide_duplicate", "slide": 1, "position": 4},
                {"type": "slide_delete", "slide": 3},
            ]
        },
    })
    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "success", result
    assert len(Presentation(output).slides) == 3
    evidence = result["diagnostics"]["operation_result"]["slide_lifecycle"]
    assert [item["type"] for item in evidence] == [
        "slide_add",
        "slide_duplicate",
        "slide_delete",
    ]
    outcomes = {gate["id"]: gate["outcome"] for gate in result["validation"]["gates"]}
    assert outcomes["operation.mutation-semantics"] == "pass"


def test_public_object_edit_batch_covers_every_primitive(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    image_a = tmp_path / "public-object-a.png"
    image_b = tmp_path / "public-object-b.png"
    image_a.write_bytes(PNG_1X1)
    image_b.write_bytes(PNG_1X1)
    frame = {"x": 500_000, "y": 1_500_000, "cx": 2_000_000, "cy": 1_200_000}
    shape = {
        "name": "Public shape",
        "geometry": "rect",
        "frame": frame,
        "fill": "336699",
        "text": "Public text",
        "z_order": 2,
    }
    image = {
        "name": "Public image",
        "path": str(image_a),
        "content_type": "image/png",
        "frame": frame,
        "fit": "cover",
        "alt_text": "Public image",
    }
    replacement_image = {
        **image,
        "name": "Replaced public image",
        "path": str(image_b),
    }
    table = {
        "name": "Public table",
        "frame": frame,
        "rows": [["A", "B"], ["C", "D"]],
        "z_order": 3,
    }
    chart = {
        "name": "Public chart",
        "frame": {**frame, "cx": 4_000_000, "cy": 2_400_000},
        "chart_type": "column",
        "title": "Public chart",
        "categories": ["A", "B"],
        "series": [{"name": "Values", "values": [1, 2]}],
        "z_order": 4,
    }
    updated_chart = {
        **chart,
        "name": "Updated public chart",
        "chart_type": "line",
        "series": [{"name": "Values", "values": [3, 4]}],
    }
    shape_selector = {"name": "Public shape", "type": "shape"}
    edits = [
        {"type": "shape_add", "slide": 1, "object": shape},
        {
            "type": "shape_update",
            "slide": 1,
            "selector": shape_selector,
            "properties": {"fill": "AA5500", "rotation": 10},
        },
        {
            "type": "text_update",
            "slide": 1,
            "selector": shape_selector,
            "properties": {"text": "Updated public text"},
        },
        {
            "type": "text_style",
            "slide": 1,
            "selector": shape_selector,
            "properties": {"bold": True, "font_size": 18},
        },
        {"type": "action_add", "slide": 1, "selector": shape_selector, "action": "next"},
        {"type": "action_update", "slide": 1, "selector": shape_selector, "action": "last"},
        {"type": "action_remove", "slide": 1, "selector": shape_selector},
        {"type": "hyperlink_add", "slide": 1, "selector": shape_selector, "target_slide": 2},
        {"type": "hyperlink_update", "slide": 1, "selector": shape_selector, "target_slide": 1},
        {"type": "hyperlink_remove", "slide": 1, "selector": shape_selector},
        {"type": "shape_delete", "slide": 1, "selector": shape_selector},
        {"type": "image_add", "slide": 1, "object": image},
        {
            "type": "image_crop",
            "slide": 1,
            "selector": {"name": "Public image", "type": "image"},
            "properties": {"fit": "contain", "opacity": 0.8},
        },
        {
            "type": "image_replace",
            "slide": 1,
            "selector": {"name": "Public image", "type": "image"},
            "object": replacement_image,
        },
        {
            "type": "image_delete",
            "slide": 1,
            "selector": {"name": "Replaced public image", "type": "image"},
        },
        {"type": "table_add", "slide": 1, "object": table},
        {
            "type": "table_update",
            "slide": 1,
            "selector": {"name": "Public table", "type": "table"},
            "properties": {"name": "Updated public table", "rows": [["E", "F"]]},
        },
        {
            "type": "table_delete",
            "slide": 1,
            "selector": {"name": "Updated public table", "type": "table"},
        },
        {"type": "chart_add", "slide": 1, "object": chart},
        {
            "type": "chart_update",
            "slide": 1,
            "selector": {"name": "Public chart", "type": "chart"},
            "object": updated_chart,
        },
        {
            "type": "chart_delete",
            "slide": 1,
            "selector": {"name": "Updated public chart", "type": "chart"},
        },
        {"type": "notes_update", "slide": 2, "value": "Public notes update"},
    ]
    output = tmp_path / "public-object-edits.pptx"
    request = _request(tmp_path, "public-object-edits.json", {
        "schema_version": "1.0",
        "operation": "pptx.edit",
        "input": str(public_created),
        "output": str(output),
        "arguments": {"edits": edits},
    })
    result = _public(project_root, "run", "--request", str(request))

    assert result["status"] == "success", result
    assert len(Presentation(output).slides) == 2
    counts = result["diagnostics"]["operation_result"]["edit_counts"]
    assert set(counts) == {edit["type"] for edit in edits}
    assert len(result["diagnostics"]["operation_result"]["object_edits"]) == len(edits)


def test_public_external_hyperlink_is_rejected(
    project_root: Path,
    public_created: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "external-link-must-not-exist.pptx"
    request = _request(tmp_path, "external-link.json", {
        "schema_version": "1.0",
        "operation": "pptx.edit",
        "input": str(public_created),
        "output": str(output),
        "arguments": {"edits": [{
            "type": "hyperlink_add",
            "slide": 1,
            "selector": {"name": "Title"},
            "target": "https://example.com",
        }]},
    })
    result = _public(
        project_root,
        "run",
        "--request",
        str(request),
        check=False,
    )
    assert result["status"] == "invalid_request"
    assert not output.exists()
