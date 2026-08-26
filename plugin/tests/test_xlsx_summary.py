"""Ordinary grouped summary operation contract and transaction coverage."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.create import create_xlsx
from document_skills_core.formats.xlsx.package import OpcPackage
from document_skills_core.formats.xlsx.service import XlsxService
from tests.support.xlsx_macro_fixture import create_package_fixture


def _arguments(**overrides: object) -> dict[str, object]:
    result: dict[str, object] = {
        "source": {"sheet": "Data", "range": "A1:D7"},
        "group_by": ["Region"],
        "aggregates": [
            {"column": "Revenue", "function": "sum", "as": "Total Revenue"},
            {"column": "Units", "function": "average", "as": "Average Units"},
            {"function": "count", "as": "Rows"},
            {"column": "Category", "function": "count_distinct", "as": "Categories"},
        ],
        "sort": [
            {"column": "Total Revenue", "direction": "desc"},
            {"column": "Region", "direction": "asc"},
        ],
        "top_n": 2,
        "target": {
            "sheet": "Summary",
            "start_cell": "B2",
            "table_name": "RegionSummary",
            "table_style": "TableStyleMedium2",
        },
        "formula_policy": "reject",
        "numeric_policy": "strict",
        "recalculation": "skip",
    }
    result.update(overrides)
    return result


def _request(input_path: Path, output: Path, **arguments: object) -> dict[str, object]:
    return {
        "operation": "xlsx.summary.aggregate",
        "input": str(input_path),
        "output": str(output),
        "arguments": _arguments(**arguments),
    }


def _write_source(path: Path, *, revenue_as_text: bool = False) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Region", "Category", "Revenue", "Units"])
    rows = [
        ("East", "A", 10, 2),
        ("West", "A", 20, 4),
        ("East", "B", 15, 3),
        ("West", "B", 10, 1),
        ("East", "A", 5, 1),
        ("North", "C", 1, 1),
    ]
    for region, category, revenue, units in rows:
        sheet.append(
            [region, category, str(revenue) if revenue_as_text else revenue, units]
        )
    workbook.save(path)
    workbook.close()
    return path


def test_summary_contract_is_bounded_and_macro_explicit() -> None:
    parsed = parse_xlsx_request(
        {
            "operation": "xlsx.summary.aggregate",
            "input": "source.xlsx",
            "output": "summary.xlsx",
            "arguments": _arguments(),
        }
    )
    assert parsed.arguments["top_n"] == 2
    assert parsed.arguments["limits"]["max_source_columns"] == 256
    assert parsed.arguments["target"]["start_cell"] == "B2"

    with pytest.raises(DocumentSkillsError):
        parse_xlsx_request(
            {
                "operation": "xlsx.summary.aggregate",
                "input": "source.xlsx",
                "output": "summary.xlsx",
                "arguments": _arguments(sort=[]),
            }
        )
    with pytest.raises(DocumentSkillsError):
        parse_xlsx_request(
            {
                "operation": "xlsx.summary.aggregate",
                "input": "source.xlsm",
                "output": "summary.xlsm",
                "arguments": _arguments(),
            }
        )
    macro = parse_xlsx_request(
        {
            "operation": "xlsx.summary.aggregate",
            "input": "source.xlsm",
            "output": "summary.xlsm",
            "arguments": _arguments(keep_vba=True),
        }
    )
    assert macro.arguments["keep_vba"] is True


def test_summary_groups_sorts_top_n_and_reopens_native_table(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = _write_source(tmp_path / "source.xlsx")
    output = tmp_path / "summary.xlsx"
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    result = XlsxService(project_root).execute(
        "xlsx.summary.aggregate",
        _request(source, output),
    )

    assert result["status"] == "success"
    summary = result["diagnostics"]["operation_result"]["summary"]
    assert summary["kind"] == "ordinary_table"
    assert summary["native_pivot"] is False
    assert summary["groups_before_top_n"] == 3
    assert summary["groups_written"] == 2
    assert summary["output"]["range"] == "B2:F4"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    assert any(
        gate["id"] == "operation.ordinary-summary-not-pivot"
        and gate["outcome"] == "pass"
        and gate["evidence"]["native_pivot"] is False
        for gate in result["validation"]["gates"]
    )

    reopened = load_workbook(output, data_only=False)
    sheet = reopened["Summary"]
    assert [sheet.cell(2, column).value for column in range(2, 7)] == [
        "Region",
        "Total Revenue",
        "Average Units",
        "Rows",
        "Categories",
    ]
    assert [sheet.cell(3, column).value for column in range(2, 7)] == [
        "East",
        30,
        2,
        3,
        2,
    ]
    assert [sheet.cell(4, column).value for column in range(2, 7)] == [
        "West",
        30,
        2.5,
        2,
        2,
    ]
    assert sheet.tables["RegionSummary"].ref == "B2:F4"
    assert not sheet._pivots
    reopened.close()


def test_summary_multi_key_sort_breaks_primary_ties(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "multi-key-source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Region", "Category", "Revenue", "Units"])
    for row in [
        ("West", "A", 20, 4),
        ("East", "A", 10, 2),
        ("West", "B", 10, 1),
        ("East", "B", 15, 3),
        ("East", "A", 5, 1),
        ("North", "C", 1, 1),
    ]:
        sheet.append(row)
    workbook.save(source)
    workbook.close()
    output = tmp_path / "multi-key-summary.xlsx"

    result = XlsxService(project_root).execute(
        "xlsx.summary.aggregate",
        _request(source, output, top_n=3),
    )

    assert result["status"] == "success", result["errors"]
    reopened = load_workbook(output, data_only=False)
    assert [reopened["Summary"].cell(row, 2).value for row in range(3, 6)] == [
        "East",
        "West",
        "North",
    ]
    reopened.close()


def test_summary_preserves_text_group_keys_and_can_coerce_numeric_text(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = _write_source(tmp_path / "text-source.xlsx", revenue_as_text=True)
    workbook = load_workbook(source)
    sheet = workbook["Data"]
    sheet["A2"] = "001"
    sheet["A3"] = "010"
    sheet["A4"] = "001"
    sheet["A5"] = "010"
    sheet["A6"] = "001"
    workbook.save(source)
    workbook.close()
    output = tmp_path / "coerced.xlsx"

    result = XlsxService(project_root).execute(
        "xlsx.summary.aggregate",
        _request(
            source,
            output,
            group_by=["Region"],
            numeric_policy="coerce-text",
            sort=[{"column": "Region", "direction": "asc"}],
            top_n=2,
        ),
    )

    assert result["status"] == "success"
    reopened = load_workbook(output, data_only=False)
    assert reopened["Summary"]["B3"].value == "001"
    assert reopened["Summary"]["B4"].value == "010"
    assert reopened["Summary"]["C3"].value == 30
    assert reopened["Summary"]["C4"].value == 30
    reopened.close()


def test_summary_supports_min_max_and_nonblank_count(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = _write_source(tmp_path / "aggregate-functions.xlsx")
    output = tmp_path / "aggregate-functions-summary.xlsx"
    result = XlsxService(project_root).execute(
        "xlsx.summary.aggregate",
        _request(
            source,
            output,
            aggregates=[
                {"column": "Revenue", "function": "min", "as": "Minimum"},
                {"column": "Revenue", "function": "max", "as": "Maximum"},
                {
                    "column": "Category",
                    "function": "count_nonblank",
                    "as": "Nonblank",
                },
            ],
            sort=[{"column": "Maximum", "direction": "desc"}],
            top_n=3,
        ),
    )

    assert result["status"] == "success"
    reopened = load_workbook(output, data_only=False)
    sheet = reopened["Summary"]
    assert [sheet.cell(3, column).value for column in range(2, 6)] == [
        "West",
        10,
        20,
        2,
    ]
    assert [sheet.cell(4, column).value for column in range(2, 6)] == [
        "East",
        5,
        15,
        3,
    ]
    reopened.close()


def test_summary_failure_preserves_existing_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _write_source(tmp_path / "text-numbers.xlsx", revenue_as_text=True)
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing-summary-destination")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    result = XlsxService(project_root).execute(
        "xlsx.summary.aggregate",
        _request(source, output),
    )

    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == "DS_VALIDATION_FAILED"
    assert output.read_bytes() == b"existing-summary-destination"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash


def test_summary_formula_policy_rejects_by_default_and_discloses_cached_values(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "formula-source.xlsx"
    workbook = {
        "metadata": {},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Region", "type": "s"},
                            {"ref": "B1", "value": "Revenue", "type": "s"},
                            {"ref": "A2", "value": "East", "type": "s"},
                            {
                                "ref": "B2",
                                "formula": "5+5",
                                "cached_value": "10",
                                "type": "n",
                            },
                        ]
                    }
                ],
                "number_formats": [],
            }
        ],
        "defined_names": [],
        "tables": [],
    }
    parsed_create = parse_xlsx_request(
        {
            "operation": "xlsx.create",
            "output": str(source),
            "arguments": {"workbook": workbook, "recalculation": "skip"},
        }
    )
    create_xlsx(source, parsed_create.arguments["workbook"])
    rejected = XlsxService(project_root).execute(
        "xlsx.summary.aggregate",
        _request(
            source,
            tmp_path / "rejected.xlsx",
            source={"sheet": "Data", "range": "A1:B2"},
            aggregates=[
                {"column": "Revenue", "function": "sum", "as": "Total Revenue"}
            ],
            sort=[{"column": "Total Revenue", "direction": "desc"}],
        ),
    )
    assert rejected["status"] == "enhancement_required"

    output = tmp_path / "cached.xlsx"
    result = XlsxService(project_root).execute(
        "xlsx.summary.aggregate",
        _request(
            source,
            output,
            source={"sheet": "Data", "range": "A1:B2"},
            aggregates=[
                {"column": "Revenue", "function": "sum", "as": "Total Revenue"}
            ],
            sort=[{"column": "Total Revenue", "direction": "desc"}],
            formula_policy="cached",
        ),
    )
    assert result["status"] == "degraded", result["errors"]
    assert result["diagnostics"]["operation_result"]["summary"]["source"][
        "formula_cells_used"
    ] == ["Data!B2"]
    assert any(
        item["code"] == "summary-cached-formula-values"
        for item in result["degradations"]
    )
    assert OpcPackage.open(output).workbook_format == "xlsx"


def test_summary_xlsm_preserves_vba_and_discloses_signature_invalidation(
    project_root: Path,
    tmp_path: Path,
) -> None:
    workbook = {
        "metadata": {},
        "sheets": [
            {
                "name": "Data",
                "rows": [
                    {
                        "cells": [
                            {"ref": "A1", "value": "Region", "type": "s"},
                            {"ref": "B1", "value": "Category", "type": "s"},
                            {"ref": "C1", "value": "Revenue", "type": "s"},
                            {"ref": "D1", "value": "Units", "type": "s"},
                            {"ref": "A2", "value": "East", "type": "s"},
                            {"ref": "B2", "value": "A", "type": "s"},
                            {"ref": "C2", "value": "10", "type": "n"},
                            {"ref": "D2", "value": "2", "type": "n"},
                        ]
                    }
                ],
                "number_formats": [],
            }
        ],
        "defined_names": [],
        "tables": [],
    }
    source = create_package_fixture(
        tmp_path / "summary-source.xlsm",
        "xlsm",
        signed=True,
        workbook=workbook,
    )
    output = tmp_path / "summary-output.xlsm"

    result = XlsxService(project_root).execute(
        "xlsx.summary.aggregate",
        _request(
            source,
            output,
            source={"sheet": "Data", "range": "A1:D2"},
            keep_vba=True,
        ),
    )

    assert result["status"] == "degraded"
    macro = result["diagnostics"]["operation_result"]["macro"]
    assert macro["vba_payload"] == "preserved"
    assert macro["signature_state"] == "invalidated_by_package_mutation"
    package = OpcPackage.open(
        output,
        allowed_inert_categories=frozenset({"vba"}),
    )
    assert package.parts["xl/vbaProject.bin"] == b"inert-vba-project-payload\x00"
