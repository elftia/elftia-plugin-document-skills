"""Public subprocess coverage for native pivot creation and readback."""

from __future__ import annotations

import json
from pathlib import Path
import subprocess


def _public(project_root: Path, *arguments: str) -> dict[str, object]:
    process = subprocess.run(
        [
            "uv",
            "run",
            "--project",
            str(project_root),
            "--frozen",
            "python",
            str(project_root / "skills/document-xlsx/scripts/run.py"),
            *arguments,
        ],
        cwd=project_root,
        check=False,
        capture_output=True,
        text=False,
        shell=False,
        timeout=30,
    )
    assert process.returncode == 0, (
        process.stderr.decode("utf-8", errors="replace")
        or process.stdout.decode("utf-8", errors="replace")
    )
    assert process.stderr == b""
    text = process.stdout.decode("utf-8", errors="strict")
    payload, end = json.JSONDecoder().raw_decode(text)
    assert text[end:].strip() == ""
    assert type(payload) is dict
    return payload


def _request(tmp_path: Path, name: str, payload: dict[str, object]) -> Path:
    path = tmp_path / name
    path.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
        newline="\n",
    )
    return path


def _source(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Region", "Category", "Segment", "Revenue"])
    for row in [
        ("East", "A", "Retail", 10),
        ("West", "A", "Retail", 20),
        ("East", "B", "Wholesale", 15),
        ("West", "B", "Wholesale", 10),
    ]:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def test_public_pivot_create_and_structured_readback(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = _source(tmp_path / "public-pivot-source.xlsx")
    output = tmp_path / "public-pivot.xlsx"
    create_request = _request(
        tmp_path,
        "pivot-create.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.pivot.create",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "source": {"sheet": "Data", "range": "A1:D5"},
                "rows": [{"column": "Region", "sort": "asc"}],
                "columns": [{"column": "Category", "sort": "asc"}],
                "values": [
                    {
                        "column": "Revenue",
                        "function": "sum",
                        "as": "Total Revenue",
                    }
                ],
                "filters": [{"column": "Segment", "value": None}],
                "target": {
                    "sheet": "Pivot",
                    "start_cell": "A1",
                    "name": "PublicPivot",
                    "style": "PivotStyleMedium9",
                },
                "recalculation": "skip",
            },
        },
    )
    created = _public(project_root, "run", "--request", str(create_request))

    assert created["status"] == "success"
    pivot = created["diagnostics"]["operation_result"]["pivot"]
    assert pivot["native"] is True
    assert pivot["cache"]["records_written"] == 4
    reopened = load_workbook(output, data_only=False)
    assert len(reopened["Pivot"]._pivots) == 1
    assert reopened["Pivot"]._pivots[0].name == "PublicPivot"
    reopened.close()

    read_request = _request(
        tmp_path,
        "pivot-read.json",
        {
            "schema_version": "1.0",
            "operation": "xlsx.read",
            "input": str(output),
            "arguments": {"include_formulas": False},
        },
    )
    read = _public(project_root, "run", "--request", str(read_request))
    projected = read["diagnostics"]["operation_result"]["pivot_tables"]
    assert len(projected) == 1
    assert projected[0]["name"] == "PublicPivot"
    assert projected[0]["source"] == {"sheet": "Data", "range": "A1:D5"}
    assert projected[0]["values"] == [
        {
            "column": "Revenue",
            "field_index": 3,
            "function": "sum",
            "as": "Total Revenue",
        }
    ]


def test_public_capabilities_advertise_callable_native_pivot(project_root: Path) -> None:
    report = _public(project_root, "capabilities", "--json")
    operation = next(
        item for item in report["operations"] if item["operation"] == "xlsx.pivot.create"
    )
    assert operation["available"] is True
    assert operation["providers"] == ["core-python"]
