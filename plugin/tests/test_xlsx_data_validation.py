"""Native XLSX data-validation creation, editing, and readback tests."""

from hashlib import sha256
from pathlib import Path

import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError
from document_skills_core.formats.xlsx.contracts import parse_xlsx_request
from document_skills_core.formats.xlsx.service import XlsxService


def _validation(
    ref: str,
    validation_type: str,
    formula1: str,
    *,
    operator: str | None = None,
    formula2: str | None = None,
) -> dict:
    result = {
        "ref": ref,
        "type": validation_type,
        "formula1": formula1,
        "allow_blank": True,
        "show_input_message": True,
        "show_error_message": True,
        "prompt_title": "Input",
        "prompt": "Enter an allowed value",
        "error_title": "Invalid",
        "error": "Value is outside the allowed range",
        "error_style": "stop",
    }
    if operator is not None:
        result["operator"] = operator
    if formula2 is not None:
        result["formula2"] = formula2
    return result


def _workbook() -> dict:
    validations = [
        _validation("B2:B10", "list", '"Low,Medium,High"'),
        _validation("C2:C10", "whole", "1", operator="between", formula2="100"),
        _validation("D2:D10", "decimal", "0", operator="greaterThanOrEqual"),
        _validation("E2:E10", "date", "DATE(2020,1,1)", operator="greaterThanOrEqual"),
        _validation("F2:F10", "time", "TIME(9,0,0)", operator="greaterThanOrEqual"),
        _validation("G2:G10", "textLength", "12", operator="lessThanOrEqual"),
        _validation("H2:H10", "custom", "MOD(H2,2)=0"),
    ]
    return {
        "metadata": {"title": "Validation", "creator": "Test", "subject": ""},
        "sheets": [
            {
                "name": "Data",
                "rows": [{"cells": [{"ref": "A1", "value": "Value", "type": "s"}]}],
                "columns": [],
                "number_formats": [],
                "data_validations": validations,
            }
        ],
        "defined_names": [],
        "tables": [],
        "chart_reference": None,
        "page_setup": None,
    }


def _request(operation: str, **paths: Path) -> dict:
    return {
        "schema_version": "1.0",
        "operation": operation,
        **{name: str(path) for name, path in paths.items()},
    }


def test_data_validation_create_supports_all_public_types_and_readback(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "validations.xlsx"
    request = _request("xlsx.create", output=output)
    request["arguments"] = {"workbook": _workbook()}
    result = XlsxService(project_root).execute("xlsx.create", request)

    assert result["status"] == "success"
    rules = list(load_workbook(output)["Data"].data_validations.dataValidation)
    assert [rule.type for rule in rules] == [
        "list",
        "whole",
        "decimal",
        "date",
        "time",
        "textLength",
        "custom",
    ]
    assert str(rules[0].sqref) == "B2:B10"
    assert rules[1].operator == "between"
    assert rules[1].formula2 == "100"
    assert rules[0].promptTitle == "Input"
    assert rules[0].errorTitle == "Invalid"

    read_request = _request("xlsx.read", input=output)
    read_request["arguments"] = {}
    read_result = XlsxService(project_root).execute("xlsx.read", read_request)
    projected = read_result["diagnostics"]["operation_result"]["data_validations"]
    assert len(projected) == 7
    assert projected[0] == {"sheet": "Data", **_workbook()["sheets"][0]["data_validations"][0], "operator": None, "formula2": None}


def test_data_validation_create_is_deterministic(
    project_root: Path,
    tmp_path: Path,
) -> None:
    outputs = [tmp_path / "first.xlsx", tmp_path / "second.xlsx"]
    for output in outputs:
        request = _request("xlsx.create", output=output)
        request["arguments"] = {"workbook": _workbook()}
        assert XlsxService(project_root).execute("xlsx.create", request)["status"] == "success"
    assert outputs[0].read_bytes() == outputs[1].read_bytes()


@pytest.mark.parametrize(
    "mutate",
    [
        lambda item: item.pop("formula1"),
        lambda item: item.update({"type": "whole", "operator": "between"}),
        lambda item: item.update({"operator": "equal"}),
        lambda item: item.update({"formula1": "[Book.xlsx]Data!$A$1"}),
        lambda item: item.update({"type": "unsupported"}),
    ],
)
def test_data_validation_contract_rejects_invalid_boundaries(mutate) -> None:
    workbook = _workbook()
    validation = workbook["sheets"][0]["data_validations"][0]
    mutate(validation)
    with pytest.raises(DocumentSkillsError) as exc:
        parse_xlsx_request(
            {
                "operation": "xlsx.create",
                "output": "validation.xlsx",
                "arguments": {"workbook": workbook},
            }
        )
    assert exc.value.code.value == "DS_REQUEST_INVALID"


def test_data_validation_contract_rejects_overlapping_create_ranges() -> None:
    workbook = _workbook()
    workbook["sheets"][0]["data_validations"].append(
        _validation("B5:B12", "custom", "B5<>0")
    )
    with pytest.raises(DocumentSkillsError):
        parse_xlsx_request(
            {
                "operation": "xlsx.create",
                "output": "validation.xlsx",
                "arguments": {"workbook": workbook},
            }
        )


def _source(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Value"
    first = DataValidation(type="whole", operator="between", formula1="1", formula2="10")
    first.add("A2:A10")
    sheet.add_data_validation(first)
    obsolete = DataValidation(type="list", formula1='"Yes,No"')
    obsolete.add("D2:D10")
    sheet.add_data_validation(obsolete)
    workbook.save(path)


def test_data_validation_edit_add_update_delete_is_transactional(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "edited.xlsx"
    _source(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    edits = [
        {
            "sheet": "Data",
            "type": "data_validation_update",
            "ref": "A2:A10",
            "validation": _validation(
                "A2:A10", "decimal", "0", operator="greaterThanOrEqual"
            ),
        },
        {
            "sheet": "Data",
            "type": "data_validation_add",
            "validation": _validation("B2:B10", "list", '"Red,Green,Blue"'),
        },
        {"sheet": "Data", "type": "data_validation_delete", "ref": "D2:D10"},
    ]
    request = _request("xlsx.edit", input=source, output=output)
    request["arguments"] = {"edits": edits, "expected_edits": 3}
    result = XlsxService(project_root).execute("xlsx.edit", request)

    assert result["status"] == "success"
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    rules = list(load_workbook(output)["Data"].data_validations.dataValidation)
    by_ref = {str(rule.sqref): rule for rule in rules}
    assert set(by_ref) == {"A2:A10", "B2:B10"}
    assert by_ref["A2:A10"].type == "decimal"
    assert by_ref["B2:B10"].formula1 == '"Red,Green,Blue"'


def test_data_validation_edit_failure_preserves_existing_destination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "existing.xlsx"
    _source(source)
    output.write_bytes(b"existing-data-validation-destination")
    request = _request("xlsx.edit", input=source, output=output)
    request["arguments"] = {
        "edits": [
            {
                "sheet": "Data",
                "type": "data_validation_add",
                "validation": _validation("A5:B5", "custom", "A5<>0"),
            }
        ]
    }
    result = XlsxService(project_root).execute("xlsx.edit", request)

    assert result["status"] == "invalid_request"
    assert output.read_bytes() == b"existing-data-validation-destination"
