"""Native master/layout/theme authoring, copying, deletion, and inheritance."""

from pathlib import Path
from typing import Any
from xml.etree.ElementTree import SubElement

from pptx import Presentation
import pytest

from document_skills_core.core.contracts.errors import DocumentSkillsError, ErrorCode
from document_skills_core.formats.pptx.constants import NS, local_name
from document_skills_core.formats.pptx.contracts import parse_pptx_request
from document_skills_core.formats.pptx.create import create_pptx
from document_skills_core.formats.pptx.mutation import MutablePptxPackage
from document_skills_core.formats.pptx.object_xml import drawable_elements, slide_shape_tree
from document_skills_core.formats.pptx.package import OpcPackage
from document_skills_core.formats.pptx.scaffold import _to_xml_bytes
from document_skills_core.formats.pptx.service import PptxService

_MASTER = "ppt/slideMasters/slideMaster1.xml"


def _deck(*, notes: bool = False) -> dict[str, Any]:
    return {
        "metadata": {"title": "Design graph", "creator": "Test", "subject": ""},
        "slide_size": {"cx": "9144000", "cy": "6858000", "type": "screen4x3"},
        "slides": [{
            "layout": "content",
            "title": "Existing",
            "shapes": [{"text": "Body", "runs": []}],
            "table": None,
            "chart_reference": None,
            "image_reference": None,
            "notes": "Speaker notes" if notes else None,
        }],
    }


def _created(tmp_path: Path, name: str = "source.pptx", *, notes: bool = False) -> Path:
    source = tmp_path / name
    create_pptx(source, _deck(notes=notes))
    return source


def _request(source: Path, output: Path, edits: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "schema_version": "1.0",
        "operation": "pptx.edit",
        "input": str(source),
        "output": str(output),
        "arguments": {"edits": edits},
    }


def _layout(name: str = "Brand Content") -> dict[str, Any]:
    return {
        "name": name,
        "type": "obj",
        "background": "F7FAFC",
        "show_master_shapes": True,
        "placeholders": [
            {
                "type": "title",
                "idx": 0,
                "name": "Brand title",
                "frame": {"x": 457200, "y": 365760, "cx": 8229600, "cy": 914400},
            },
            {
                "type": "body",
                "idx": 1,
                "name": "Brand body",
                "frame": {"x": 457200, "y": 1600200, "cx": 8229600, "cy": 4389120},
            },
        ],
        "footer": {"enabled": True, "text": "Confidential"},
        "date": {"enabled": True, "text": "2026-08-24"},
        "slide_number": {"enabled": True},
    }


def _theme() -> dict[str, Any]:
    return {
        "name": "Brand Theme",
        "palette": {"accent1": "006D77", "accent2": "E29578"},
        "fonts": {"major": "Aptos Display", "minor": "Aptos"},
        "effects": {
            "shadow": {
                "enabled": True,
                "blur": 60000,
                "distance": 40000,
                "direction": 315,
                "color": "112233",
                "opacity": 0.35,
            }
        },
    }


def test_design_edit_contracts_are_closed_and_part_selected(tmp_path: Path) -> None:
    parsed = parse_pptx_request(_request(
        tmp_path / "source.pptx",
        tmp_path / "output.pptx",
        [
            {"type": "master_add", "master": {"name": "Brand", "theme": _theme(), "layout": _layout()}},
            {"type": "layout_add", "master": _MASTER, "layout": _layout("Extra")},
            {"type": "layout_copy", "layout": "ppt/slideLayouts/slideLayout1.xml", "master": _MASTER},
            {"type": "theme_update", "master": _MASTER, "theme": _theme()},
        ],
    ))
    assert [item["type"] for item in parsed.arguments["edits"]] == [
        "master_add", "layout_add", "layout_copy", "theme_update",
    ]
    assert parsed.arguments["edits"][0]["master"]["theme"]["effects"]["shadow"]["enabled"] is True

    with pytest.raises(DocumentSkillsError) as captured:
        parse_pptx_request(_request(
            tmp_path / "source.pptx",
            tmp_path / "output.pptx",
            [{"type": "layout_delete", "layout": "../slideLayout1.xml"}],
        ))
    assert captured.value.code == ErrorCode.REQUEST_INVALID


def test_master_add_authors_theme_layout_placeholders_and_footer_slots(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _created(tmp_path)
    output = tmp_path / "master-added.pptx"
    result = PptxService(project_root).execute("pptx.edit", _request(
        source,
        output,
        [{
            "type": "master_add",
            "master": {"name": "Brand Master", "theme": _theme(), "layout": _layout()},
        }],
    ))

    assert result["status"] == "success", result
    package = OpcPackage.open(output)
    assert len(package.slide_master_parts()) == 2
    assert len(package.slide_layout_parts()) == 8
    assert len(package.theme_parts()) == 2
    evidence = result["diagnostics"]["operation_result"]["design_edits"][0]
    layout = package.xml(evidence["layout_part"])
    placeholder_types = {
        node.attrib.get("type")
        for node in layout.iter()
        if local_name(node.tag) == "ph"
    }
    assert placeholder_types == {"body", "dt", "ftr", "sldNum", "title"}
    header_footer = next(node for node in layout if local_name(node.tag) == "hf")
    assert header_footer.attrib == {"dt": "1", "ftr": "1", "sldNum": "1"}
    common = layout.find(f"{{{NS['p']}}}cSld")
    assert common is not None and common.find(f"{{{NS['p']}}}bg") is not None
    theme = package.xml(evidence["theme_part"])
    assert sum(local_name(node.tag) == "outerShdw" for node in theme.iter()) == 2
    assert len(Presentation(output).slide_masters) == 2


def test_layout_add_copy_and_delete_preserve_master_inheritance(
    project_root: Path,
    tmp_path: Path,
) -> None:
    service = PptxService(project_root)
    source = _created(tmp_path)
    added_output = tmp_path / "layout-added.pptx"
    added = service.execute("pptx.edit", _request(
        source,
        added_output,
        [{"type": "layout_add", "master": _MASTER, "layout": _layout("Added Layout")}],
    ))
    added_layout = added["diagnostics"]["operation_result"]["design_edits"][0]["layout_part"]
    assert len(OpcPackage.open(added_output).slide_layout_parts()) == 8

    copied_output = tmp_path / "layout-copied.pptx"
    copied = service.execute("pptx.edit", _request(
        added_output,
        copied_output,
        [{"type": "layout_copy", "layout": added_layout, "master": _MASTER}],
    ))
    copied_layout = copied["diagnostics"]["operation_result"]["design_edits"][0]["layout_part"]
    copied_package = OpcPackage.open(copied_output)
    assert len(copied_package.slide_layout_parts()) == 9
    inherited_master = next(
        item.resolved_target
        for item in copied_package.part_rels(copied_layout)
        if item.relationship_type.endswith("/slideMaster")
    )
    assert inherited_master == _MASTER

    deleted_output = tmp_path / "layout-deleted.pptx"
    deleted = service.execute("pptx.edit", _request(
        copied_output,
        deleted_output,
        [{"type": "layout_delete", "layout": copied_layout}],
    ))
    assert deleted["status"] == "success", deleted
    assert copied_layout not in OpcPackage.open(deleted_output).parts


def test_master_copy_then_delete_round_trips_dependency_graph(
    project_root: Path,
    tmp_path: Path,
) -> None:
    service = PptxService(project_root)
    source = _created(tmp_path)
    copied_output = tmp_path / "master-copied.pptx"
    copied = service.execute("pptx.edit", _request(
        source,
        copied_output,
        [{"type": "master_copy", "master": _MASTER}],
    ))
    evidence = copied["diagnostics"]["operation_result"]["design_edits"][0]
    copied_master = evidence["master_part"]
    copied_package = OpcPackage.open(copied_output)
    assert len(copied_package.slide_master_parts()) == 2
    assert len(copied_package.slide_layout_parts()) == 14
    assert len(copied_package.theme_parts()) == 2
    assert copied_master != _MASTER

    deleted_output = tmp_path / "master-deleted.pptx"
    deleted = service.execute("pptx.edit", _request(
        copied_output,
        deleted_output,
        [{"type": "master_delete", "master": copied_master}],
    ))
    assert deleted["status"] == "success", deleted
    package = OpcPackage.open(deleted_output)
    assert len(package.slide_master_parts()) == 1
    assert len(package.slide_layout_parts()) == 7
    assert len(package.theme_parts()) == 1


def test_theme_update_clones_shared_theme_and_writes_effects(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _created(tmp_path, notes=True)
    output = tmp_path / "theme-updated.pptx"
    result = PptxService(project_root).execute("pptx.edit", _request(
        source,
        output,
        [{"type": "theme_update", "master": _MASTER, "theme": _theme()}],
    ))

    assert result["status"] == "success", result
    evidence = result["diagnostics"]["operation_result"]["design_edits"][0]
    assert evidence["clone_on_write"] is True
    package = OpcPackage.open(output)
    assert len(package.theme_parts()) == 2
    theme = package.xml(evidence["theme_part"])
    assert theme.attrib["name"] == "Brand Theme"
    assert any(local_name(node.tag) == "outerShdw" for node in theme.iter())


def test_template_lint_reports_inheritance_and_direct_format_contamination(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _created(tmp_path)
    clean = PptxService(project_root).execute("pptx.inspect.structure", {
        "schema_version": "1.0",
        "operation": "pptx.inspect.structure",
        "input": str(source),
        "arguments": {},
    })
    clean_lint = clean["diagnostics"]["operation_result"]["template_lint"]
    assert clean_lint["inheritance"]["valid"] is True
    assert clean_lint["direct_format_contamination"]["present"] is False
    assert len(clean_lint["masters"]) == 1
    assert len(clean_lint["layouts"]) == 7

    package = OpcPackage.open(source)
    target = MutablePptxPackage(package)
    slide = target.xml("ppt/slides/slide1.xml")
    shape = drawable_elements(slide_shape_tree(slide))[0]
    application = next(
        node for node in shape.iter() if local_name(node.tag) == "nvPr"
    )
    SubElement(application, f"{{{NS['p']}}}ph", {"idx": "0", "type": "title"})
    target.set_part("ppt/slides/slide1.xml", _to_xml_bytes(slide))
    contaminated = tmp_path / "contaminated-template.pptx"
    target.emit(contaminated)

    report = PptxService(project_root).execute("pptx.inspect.structure", {
        "schema_version": "1.0",
        "operation": "pptx.inspect.structure",
        "input": str(contaminated),
        "arguments": {},
    })
    lint = report["diagnostics"]["operation_result"]["template_lint"]
    assert lint["mutation_authorized"] is False
    assert lint["direct_format_contamination"]["present"] is True
    finding = lint["direct_format_contamination"]["findings"][0]
    assert finding["slide"] == 1
    assert "solidFill" in finding["properties"]


def test_used_design_parts_fail_closed_without_output(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = _created(tmp_path)
    package = OpcPackage.open(source)
    used_layout = next(
        item.resolved_target
        for item in package.part_rels("ppt/slides/slide1.xml")
        if item.relationship_type.endswith("/slideLayout")
    )
    for edit in (
        {"type": "layout_delete", "layout": used_layout},
        {"type": "master_delete", "master": _MASTER},
    ):
        output = tmp_path / f"{edit['type']}-must-not-exist.pptx"
        result = PptxService(project_root).execute(
            "pptx.edit",
            _request(source, output, [edit]),
        )
        assert result["status"] == "invalid_request", result
        assert result["errors"][0]["code"] == ErrorCode.REQUEST_INVALID.value
        assert not output.exists()
