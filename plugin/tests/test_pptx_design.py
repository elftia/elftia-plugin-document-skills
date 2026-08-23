"""Typed theme, layout recipe, and template-as-base public service tests."""

from hashlib import sha256
from pathlib import Path
from typing import Any

from pptx import Presentation

from document_skills_core.formats.pptx.mutation import MutablePptxPackage
from document_skills_core.formats.pptx.package import OpcPackage
from document_skills_core.formats.pptx.read import read_pptx
from document_skills_core.formats.pptx.scaffold import _to_xml_bytes
from document_skills_core.formats.pptx.service import PptxService

_TEMPLATE_MAIN = (
    "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
)
_PRESENTATION_MAIN = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
)


def _slide(name: str, recipe: str) -> dict[str, Any]:
    return {
        "layout": "content",
        "recipe": recipe,
        "title": name,
        "shapes": [
            {"text": f"{name} primary", "runs": []},
            {"text": f"{name} secondary", "runs": []},
        ],
        "table": None,
        "chart_reference": None,
        "image_reference": None,
        "notes": None,
    }


def _theme() -> dict[str, Any]:
    return {
        "name": "Research Theme",
        "palette": {
            "dk1": "101820",
            "lt1": "F8F5EC",
            "accent1": "006D77",
            "accent2": "E29578",
        },
        "fonts": {"major": "Aptos Display", "minor": "Aptos"},
        "background": "F8F5EC",
        "default_text": {
            "title_color": "101820",
            "body_color": "203040",
            "title_size": 32,
            "body_size": 17,
            "bold_titles": True,
        },
        "default_shape": {"fill": "FFFFFF", "line": "006D77", "opacity": 0.95},
        "default_chart": {"colors": ["006D77", "E29578"]},
    }


def _deck(*, themed: bool, all_recipes: bool = False) -> dict[str, Any]:
    recipes = (
        ["cover", "section", "content", "two-column", "image-focus", "comparison", "summary"]
        if all_recipes
        else ["content"]
    )
    deck: dict[str, Any] = {
        "metadata": {"title": "Designed deck", "creator": "Test", "subject": "Design"},
        "slide_size": {"cx": "9144000", "cy": "6858000", "type": "screen4x3"},
        "layout_tokens": {
            "safe_margins": {"top": 300_000, "right": 400_000, "bottom": 300_000, "left": 400_000},
            "grid": {"columns": 12, "gutter": 200_000},
            "spacing": {"sm": 150_000, "md": 250_000, "lg": 400_000},
            "typography_scale": {"title": 32, "section": 25, "body": 17, "caption": 11},
        },
        "slides": [_slide(recipe.title(), recipe) for recipe in recipes],
    }
    if themed:
        deck["theme"] = _theme()
        deck["slides"][0]["chart_reference"] = {
            "title": "Theme chart",
            "chart_type": "column",
            "categories": ["A", "B"],
            "series": [{"name": "Value", "values": [2, 3]}],
        }
    return deck


def _request(output: Path, deck: dict[str, Any], template: Path | None = None) -> dict[str, Any]:
    arguments: dict[str, Any] = {"deck": deck}
    if template is not None:
        arguments["template"] = str(template)
    return {
        "schema_version": "1.0",
        "operation": "pptx.create",
        "output": str(output),
        "arguments": arguments,
    }


def _make_potx(source_path: Path, destination: Path) -> None:
    source = OpcPackage.open(source_path)
    target = MutablePptxPackage(source)
    root = target.xml("[Content_Types].xml")
    matches = [
        node
        for node in root
        if node.attrib.get("PartName") == "/ppt/presentation.xml"
    ]
    assert len(matches) == 1
    matches[0].set("ContentType", _TEMPLATE_MAIN)
    target.set_part("[Content_Types].xml", _to_xml_bytes(root))
    target.emit(destination)


def test_typed_theme_and_all_layout_recipes_are_native(
    project_root: Path,
    tmp_path: Path,
) -> None:
    output = tmp_path / "designed.pptx"
    service = PptxService(project_root)
    result = service.execute("pptx.create", _request(output, _deck(themed=True, all_recipes=True)))

    assert result["status"] == "success", result
    assert len(Presentation(output).slides) == 7
    creation = result["diagnostics"]["operation_result"]["creation"]
    assert creation["layouts"] == 7
    assert {item["name"] for item in creation["layout_recipes"]} == {
        "cover", "section", "content", "two-column", "image-focus", "comparison", "summary"
    }
    assert creation["theme"]["fonts"] == {
        "major": "Aptos Display",
        "minor": "Aptos",
    }
    projected, _warnings = read_pptx(output, {})
    assert projected["theme"]["name"] == "Research Theme"
    assert projected["theme"]["palette"]["accent1"] == "006D77"
    assert projected["theme"]["fonts"]["major"] == "Aptos Display"
    assert {item["name"] for item in projected["layout_recipes"]} == {
        "Cover", "Section", "Content", "Two Column", "Image Focus", "Comparison", "Summary"
    }
    chart_xml = OpcPackage.open(output).parts["ppt/charts/chart1.xml"]
    assert b'val="006D77"' in chart_xml
    outcomes = {gate["id"]: gate["outcome"] for gate in result["validation"]["gates"]}
    assert outcomes["operation.typed-design-correspondence"] == "pass"


def test_template_as_base_reuses_master_layout_theme_byte_for_byte(
    project_root: Path,
    tmp_path: Path,
) -> None:
    service = PptxService(project_root)
    template = tmp_path / "template.pptx"
    template_result = service.execute(
        "pptx.create",
        _request(template, _deck(themed=True)),
    )
    assert template_result["status"] == "success", template_result
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = OpcPackage.open(template)
    design_parts = source.slide_master_parts() + source.slide_layout_parts() + source.theme_parts()
    design_hashes = {part: source.part_hashes[part] for part in design_parts}

    output = tmp_path / "from-template.pptx"
    result = service.execute(
        "pptx.create",
        _request(output, _deck(themed=False), template),
    )

    assert result["status"] == "success", result
    assert sha256(template.read_bytes()).hexdigest() == template_hash
    assert len(Presentation(output).slides) == 1
    candidate = OpcPackage.open(output)
    assert {part: candidate.part_hashes[part] for part in design_parts} == design_hashes
    creation = result["diagnostics"]["operation_result"]["creation"]
    assert creation["template_reuse"]["source_sha256"] == template_hash
    assert creation["template_reuse"]["master_parts"] == source.slide_master_parts()
    outcomes = {gate["id"]: gate["outcome"] for gate in result["validation"]["gates"]}
    assert outcomes["operation.typed-design-correspondence"] == "pass"


def test_potx_template_base_normalizes_output_identity_and_preserves_design(
    project_root: Path,
    tmp_path: Path,
) -> None:
    service = PptxService(project_root)
    source_pptx = tmp_path / "source-template.pptx"
    assert service.execute(
        "pptx.create",
        _request(source_pptx, _deck(themed=True)),
    )["status"] == "success"
    template = tmp_path / "template.potx"
    _make_potx(source_pptx, template)
    template_hash = sha256(template.read_bytes()).hexdigest()
    source = OpcPackage.open(template, allow_dangerous_inventory=True)
    assert source.content_type_for("ppt/presentation.xml") == _TEMPLATE_MAIN
    design_parts = source.slide_master_parts() + source.slide_layout_parts() + source.theme_parts()

    output = tmp_path / "from-potx.pptx"
    result = service.execute(
        "pptx.create",
        _request(output, _deck(themed=False), template),
    )

    assert result["status"] == "success", result
    assert sha256(template.read_bytes()).hexdigest() == template_hash
    assert len(Presentation(output).slides) == 1
    candidate = OpcPackage.open(output)
    assert candidate.content_type_for("ppt/presentation.xml") == _PRESENTATION_MAIN
    assert {
        part: candidate.part_hashes[part] for part in design_parts
    } == {
        part: source.part_hashes[part] for part in design_parts
    }
    reuse = result["diagnostics"]["operation_result"]["creation"]["template_reuse"]
    assert reuse["source_extension"] == ".potx"
    assert reuse["presentation_content_type"] == _PRESENTATION_MAIN
    assert reuse["template_main_type_normalized"] is True


def test_potx_admission_rejects_renamed_pptx_and_other_dangerous_inventory(
    project_root: Path,
    tmp_path: Path,
) -> None:
    service = PptxService(project_root)
    source_pptx = tmp_path / "source.pptx"
    assert service.execute(
        "pptx.create",
        _request(source_pptx, _deck(themed=True)),
    )["status"] == "success"

    renamed = tmp_path / "renamed.potx"
    renamed.write_bytes(source_pptx.read_bytes())
    renamed_output = tmp_path / "renamed-output.pptx"
    renamed_result = service.execute(
        "pptx.create",
        _request(renamed_output, _deck(themed=False), renamed),
    )
    assert renamed_result["status"] == "invalid_request"
    assert not renamed_output.exists()

    clean_potx = tmp_path / "clean.potx"
    _make_potx(source_pptx, clean_potx)
    active_package = OpcPackage.open(clean_potx, allow_dangerous_inventory=True)
    active_target = MutablePptxPackage(active_package)
    active_target.set_part("ppt/vbaProject.bin", b"inert-test-vba")
    active_potx = tmp_path / "active.potx"
    active_target.emit(active_potx)
    active_output = tmp_path / "active-output.pptx"
    active_result = service.execute(
        "pptx.create",
        _request(active_output, _deck(themed=False), active_potx),
    )
    assert active_result["status"] == "invalid_request"
    assert not active_output.exists()


def test_design_contract_rejects_unknown_tokens_and_template_theme(
    project_root: Path,
    tmp_path: Path,
) -> None:
    service = PptxService(project_root)
    invalid = _deck(themed=True)
    invalid["theme"]["unsupported_effect"] = "glow"
    output = tmp_path / "invalid-theme.pptx"
    result = service.execute("pptx.create", _request(output, invalid))
    assert result["status"] == "invalid_request"
    assert not output.exists()

    template = tmp_path / "base.pptx"
    assert service.execute(
        "pptx.create",
        _request(template, _deck(themed=False)),
    )["status"] == "success"
    rejected = tmp_path / "must-not-use-theme.pptx"
    result = service.execute(
        "pptx.create",
        _request(rejected, _deck(themed=True), template),
    )
    assert result["status"] == "invalid_request"
    assert not rejected.exists()
