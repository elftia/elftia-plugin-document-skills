"""Native XLSX table mutation tests."""

from pathlib import Path

from document_skills_core.formats.xlsx.service import XlsxService


def _source(path: Path, *, with_formula: bool = False) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    rows = [
        ["Name", "Amount", "Category", "Code", "Score"],
        ["Alpha", 10, "A", "X", 1],
        ["Beta", 20, "B", "Y", 2],
        ["Gamma", 30, "C", "Z", 3],
    ]
    for row in rows:
        sheet.append(row)
    table = Table(displayName="DataTable", ref="A1:C3")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    if with_formula:
        sheet["G1"] = "=SUM(DataTable[Amount])"
    workbook.save(path)


def test_table_add_resize_and_style_reopen(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "edited.xlsx"
    _source(source)
    edits = [
        {
            "sheet": "Data",
            "type": "table_resize",
            "name": "DataTable",
            "ref": "A1:C4",
        },
        {
            "sheet": "Data",
            "type": "table_style",
            "name": "DataTable",
            "table_style": "TableStyleLight1",
        },
        {
            "sheet": "Data",
            "type": "table_add",
            "name": "ScoreTable",
            "ref": "D1:E4",
            "table_style": "TableStyleDark1",
        },
    ]
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {"edits": edits, "expected_edits": len(edits)},
        },
    )

    assert result["status"] == "success"
    reopened = load_workbook(output, data_only=False)
    data_table = reopened["Data"].tables["DataTable"]
    score_table = reopened["Data"].tables["ScoreTable"]
    assert data_table.ref == "A1:C4"
    assert data_table.tableStyleInfo.name == "TableStyleLight1"
    assert [column.name for column in data_table.tableColumns] == [
        "Name",
        "Amount",
        "Category",
    ]
    assert score_table.ref == "D1:E4"
    assert score_table.tableStyleInfo.name == "TableStyleDark1"
    assert [column.name for column in score_table.tableColumns] == ["Code", "Score"]


def test_table_rename_migrates_structured_formula_and_formula_state(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "renamed.xlsx"
    _source(source, with_formula=True)
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "table_rename",
                        "name": "DataTable",
                        "value": "SalesTable",
                    }
                ]
            },
        },
    )

    assert result["status"] == "degraded"
    reopened = load_workbook(output, data_only=False)
    assert "SalesTable" in reopened["Data"].tables
    assert "DataTable" not in reopened["Data"].tables
    assert reopened["Data"]["G1"].value == "=SUM(SalesTable[Amount])"
    formula_cells = result["diagnostics"]["operation_result"]["formula_state"]["cells"]
    assert formula_cells["Data!G1"]["formula"] == "SUM(SalesTable[Amount])"
    assert formula_cells["Data!G1"]["state"] == "recalculation_required"


def test_table_delete_removes_declared_parts(
    project_root: Path,
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "source.xlsx"
    output = tmp_path / "deleted.xlsx"
    _source(source)
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "table_delete", "name": "DataTable"}
                ]
            },
        },
    )

    assert result["status"] == "success"
    assert not load_workbook(output)["Data"].tables
    removed = result["diagnostics"]["operation_result"]["preservation"]["removed_parts"]
    assert removed == [
        "xl/tables/table1.xml",
        "xl/worksheets/_rels/sheet1.xml.rels",
    ]


def test_table_delete_with_structured_reference_fails_without_promotion(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "existing.xlsx"
    _source(source, with_formula=True)
    output.write_bytes(b"existing-table-delete-destination")
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {"sheet": "Data", "type": "table_delete", "name": "DataTable"}
                ]
            },
        },
    )

    assert result["status"] == "enhancement_required"
    assert output.read_bytes() == b"existing-table-delete-destination"


def test_table_shrink_that_removes_referenced_column_fails_closed(
    project_root: Path,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "existing.xlsx"
    _source(source)
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    workbook["Data"]["G1"] = "=COUNTA(DataTable[Category])"
    workbook.save(source)
    output.write_bytes(b"existing-table-resize-destination")
    result = XlsxService(project_root).execute(
        "xlsx.edit",
        {
            "schema_version": "1.0",
            "operation": "xlsx.edit",
            "input": str(source),
            "output": str(output),
            "arguments": {
                "edits": [
                    {
                        "sheet": "Data",
                        "type": "table_resize",
                        "name": "DataTable",
                        "ref": "A1:B3",
                    }
                ]
            },
        },
    )

    assert result["status"] == "enhancement_required"
    assert output.read_bytes() == b"existing-table-resize-destination"
